"""Schema-detected TIOM Despatch source adapters.

Supported .xls/.xlsx/.csv providers are detected by headers, not filenames:
- Full trip ledgers: Trip Summary or Date Wise / weighment with E-Pass +
  tare/gross/net. These produce new, validated E-Pass trips.
- Transporter Wise / dispatch manifests: E-Pass + net weight + vehicle
  grouped under transporter name. These enrich matching trips only.

Never infer missing weights or create E-Passes from aggregate totals. Additional
providers may be accepted by extending the column aliases below, without
hardcoding a date, filename, sheet name or column number.
"""
from __future__ import annotations

import csv
from datetime import date, datetime, timedelta
from html import unescape
from io import BytesIO, StringIO
import re
from typing import Any

from openpyxl import load_workbook

COLUMN_ALIASES = {
    "epass_no": ("epassno", "epassnumber", "epass", "passno", "passnumber",
                 "pass", "passid", "gatepassno", "ewaypass"),
    "permit_no": ("permitno", "permitnumber", "permit", "permitid"),
    "permit_qty": ("permitqty", "permitquantity", "permitmt", "permittonnage"),
    "buyer": ("buyer", "buyername", "purchaser", "purchasername"),
    "vehicle_no": ("vehicleregno", "vehicleregistrationno", "vehicleno",
                   "vehiclenumber", "vehicle", "truckno", "trucknumber", "regno"),
    "transporter": ("transporter", "transportername", "carrier",
                    "carriernam", "transportagency", "transportingagency"),
    "destination": ("destination", "destinationname", "unloadingpoint"),
    "material_grade": ("material", "materialname", "mineralname", "grade",
                       "materialgrade", "product"),
    "stack_no": ("stack", "stackno", "stacknumber", "stackid"),
    "consignee": ("consignee", "consigneename"),
    "tare_mt": ("tarewt", "tareweight", "tare", "tarem", "taremt"),
    "gross_mt": ("grosswt", "grossweight", "gross", "grossmt"),
    "net_mt": ("netwt", "netweight", "netmt", "qtytons", "qtyton",
               "qtymt", "quantitymt", "netquantitymt", "nettons",
               "weightmt", "dispatchqty"),
    "report_date": ("date", "reportdate", "dispatchdate", "despatchdate",
                    "transportingdate", "dateofdispatch", "dateofdespatch"),
    "dispatch_time": ("time", "dispatcheddate", "dispatchtime", "despatchtime"),
}
_ALIAS_TO_FIELD = {alias: field for field, aliases in COLUMN_ALIASES.items()
                   for alias in aliases}
_PASS = re.compile(r"^[A-Za-z0-9][A-Za-z0-9/_.-]{2,119}$")


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", unescape(str(value or "")).casefold())


def _day(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 30000 < value < 80000:
        return (datetime(1899, 12, 30) + timedelta(days=int(value))).date().isoformat()
    s = str(value or "").strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%y", "%d.%m.%Y", "%d/%m/%Y",
                "%d-%m-%Y", "%d-%b-%Y", "%b %d %Y", "%d %b %Y",
                "%d/%m/%y"):
        try:
            return datetime.strptime(s[:len(datetime.now().strftime(fmt))], fmt).date().isoformat()
        except ValueError:
            continue
    # Legacy Date Wise strings contain a timestamp like
    # "Oct 27 2026 11:15:31:910AM"; the date is its first three tokens.
    m = re.search(r"\b([A-Za-z]{3})\s+(\d{1,2})\s+(20\d{2})\b", s)
    if m:
        try:
            return datetime.strptime(" ".join(m.groups()), "%b %d %Y").date().isoformat()
        except ValueError:
            pass
    return None


def _read_sheets(data: bytes, filename: str):
    name = filename.casefold()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig")
        return [("CSV", [list(row) for row in csv.reader(StringIO(text))])]
    if not name.endswith((".xls", ".xlsx")):
        raise ValueError("Supported sources: .xls, .xlsx, .csv")
    if data.startswith(bytes.fromhex("504B0304")):
        book = load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            return [(ws.title, [list(row) for row in ws.iter_rows(values_only=True)])
                    for ws in book]
        finally:
            book.close()
    if data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError("Legacy .xls reader unavailable; install xlrd") from exc
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            return [(ws.name, [ws.row_values(i) for i in range(ws.nrows)])
                    for ws in book.sheets()]
        finally:
            book.release_resources()
    raise ValueError("Unsupported Excel binary signature")


def _header_mapping(rows):
    # Real MIS files include decorative title/from-to blocks and the header
    # may be on row 2, 6, 8 or further down. Locate by semantic aliases.
    for index, row in enumerate(rows[:45]):
        mapping = {}
        for col, cell in enumerate(row):
            key = _ALIAS_TO_FIELD.get(normalize_header(cell))
            if key and key not in mapping:
                mapping[key] = col
        if "epass_no" not in mapping:
            continue
        has_full = all(field in mapping for field in
                       ("tare_mt", "gross_mt", "net_mt", "vehicle_no"))
        has_aux = all(field in mapping for field in ("net_mt", "vehicle_no"))
        if has_full or has_aux:
            return index, mapping, ("detailed" if has_full else "enrichment")
    return None


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return unescape(str(value)).strip()


def parse_despatch_sources(data: bytes, filename: str) -> list[dict]:
    if len(data) > 12_000_000:
        raise ValueError("File exceeds 12 MB")
    extracted: list[dict] = []
    recognized = False
    for sheet, rows in _read_sheets(data, filename):
        if len(rows) > 100_000:
            raise ValueError("Too many source rows")
        detected = _header_mapping(rows)
        if not detected:
            continue
        recognized = True
        header_index, mapping, role = detected
        sheet_day = _day(sheet)
        group_transporter = ""
        for row in rows[header_index + 1:]:
            values = [_text(c) for c in row if c is not None]
            # Legacy Transporter Wise groups are visible before their rows.
            for value in values:
                group = re.match(r"(?i)^\s*transporter\s*name\s*:\s*(.+)", value)
                if group:
                    group_transporter = unescape(group.group(1)).strip()
                    break
            col = mapping["epass_no"]
            epass = _text(row[col]) if col < len(row) else ""
            if not _PASS.fullmatch(epass):
                continue  # skip subtotal / repeated headings / empty footers
            fields = {field: _text(row[col]) if col < len(row) else ""
                      for field, col in mapping.items()}
            day = _day(fields.get("report_date")) or sheet_day
            if not day:
                raise ValueError("Missing/invalid report date for E-Pass " + epass)
            if fields.get("report_date") and not _day(fields["report_date"]):
                raise ValueError("Invalid explicit source date for E-Pass " + epass)
            fields["report_date"] = day
            if role == "enrichment" and group_transporter:
                fields["transporter"] = group_transporter
            if role == "enrichment" and not fields.get("transporter"):
                raise ValueError("Transporter missing for E-Pass " + epass)
            if role == "detailed" and not fields.get("dispatch_time"):
                fields["dispatch_time"] = fields.get("dispatch_date", "")
            fields["_role"] = role
            fields["source"] = (filename + " / " + sheet)[:200]
            extracted.append(fields)
    if not recognized or not extracted:
        raise ValueError("No E-Pass-level source detected. Check the date/vehicle/pass/weight column headings")
    if len(extracted) > 100_000:
        raise ValueError("Too many E-Pass rows")
    return extracted


def detect_source_type(data: bytes, filename: str) -> str:
    """Preflight for multi-attachment ordering; never changes the database."""
    rows = parse_despatch_sources(data, filename)
    return rows[0]["_role"]
