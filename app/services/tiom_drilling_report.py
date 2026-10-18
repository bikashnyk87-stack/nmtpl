"""Excel reports for TIOM drilling, independent of live database writes."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GREEN = "A9D18E"
PEACH = "F4B183"
BLUE = "234B65"
WHITE = "FFFFFF"
LINE = Side(style="thin", color="687785")


def drilling_workbook(report, detail_rows, period):
    """Build the printed day/MTD form plus every selected-period shift entry.

    report: materialized day and month-to-date data; no SQL or side effects here.
    period: {"fromDate": str, "toDate": str, "shift": str}
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Drilling Report"
    ws.merge_cells("A1:J1")
    ws["A1"] = "Drilling Report"
    ws["A1"].font = Font(size=19, bold=True, color="000000")
    ws["A1"].fill = PatternFill("solid", fgColor=GREEN)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:J2")
    ws["A2"] = f"TIOM  |  Day: {report['date']} (all shifts)  |  Month-to-date: {report['monthFrom']} to {report['date']}"
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["A2"].font = Font(bold=True, size=10, color=BLUE)
    ws.row_dimensions[2].height = 22

    for region, label in (("A3:A4", "Sl.No"), ("B3:B4", "Equipment name"),
                          ("C3:F3", "For The Day"), ("G3:J3", "For The Month")):
        ws.merge_cells(region)
        cell = ws[region.split(":")[0]]
        cell.value = label
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for idx, label in enumerate(("HMR", "Holes", "Meterage", "Meters/hr.",
                                  "HMR", "Holes", "Meterage", "Meters/hr."), 3):
        ws.cell(4, idx, label)
    for row in ws.iter_rows(min_row=3, max_row=4, min_col=1, max_col=10):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=PEACH)
            cell.font = Font(bold=True, color="000000")
            cell.border = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 24

    def put_metric(row_no, start, d):
        ws.cell(row_no, start, d.get("hmr"))
        ws.cell(row_no, start + 1, d.get("holes"))
        ws.cell(row_no, start + 2, d.get("meterage"))
        if d.get("metersPerHour") is not None:
            ws.cell(row_no, start + 3, f'=IFERROR({get_column_letter(start+2)}{row_no}/{get_column_letter(start)}{row_no},"")')

    for pos, item in enumerate(report.get("rows", []), start=1):
        row_no = 4 + pos
        ws.cell(row_no, 1, pos)
        ws.cell(row_no, 2, item["label"])
        put_metric(row_no, 3, item["day"])
        put_metric(row_no, 7, item["month"])
        for cell in ws[row_no][:10]:
            cell.border = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
            cell.alignment = Alignment(vertical="center", horizontal="left" if cell.column == 2 else "center")
        for col in (3, 5, 6, 7, 9, 10):
            ws.cell(row_no, col).number_format = "0.0"
        ws.row_dimensions[row_no].height = 22

    total_row = 5 + len(report.get("rows", []))
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
    ws.cell(total_row, 1, "Total")
    for start, label in ((3, "dayTotals"), (7, "monthTotals")):
        tot = report.get(label, {})
        if report.get("rows"):
            for offset, field in ((0, "hmr"), (1, "holes"), (2, "meterage")):
                if tot.get(field) is not None:
                    col = get_column_letter(start + offset)
                    ws.cell(total_row, start + offset, f"=SUM({col}5:{col}{total_row-1})")
        if tot.get("metersPerHour") is not None:
            ws.cell(total_row, start + 3, f'=IFERROR({get_column_letter(start+2)}{total_row}/{get_column_letter(start)}{total_row},"")')
    for cell in ws[total_row][:10]:
        cell.fill = PatternFill("solid", fgColor="E8F0E5")
        cell.font = Font(bold=True)
        cell.border = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if cell.column >= 3 and cell.column != 4 and cell.column != 8:
            cell.number_format = "0.0"
    ws.row_dimensions[total_row].height = 23

    notes_start = total_row + 1
    notes_end = notes_start + max(1, len(report.get("rows", []))) - 1
    ws.merge_cells(start_row=notes_start, start_column=1, end_row=notes_end, end_column=1)
    ws.cell(notes_start, 1, "Remarks:-")
    ws.cell(notes_start, 1).alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    for pos, item in enumerate(report.get("rows", [])):
        r = notes_start + pos
        ws.cell(r, 2, item["label"])
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=10)
        comments = item.get("remarks", [])
        ws.cell(r, 3, " | ".join(f"{x['shift']}: {x['text']}" for x in comments) if comments else "—")
        ws.cell(r, 3).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = min(68, max(24, 18 * (1 + len(str(ws.cell(r, 3).value)) // 95)))
    if not report.get("rows"):
        ws.cell(notes_start, 2, "No drilling entries for this date")
        ws.merge_cells(start_row=notes_start, start_column=3, end_row=notes_start, end_column=10)
    for row in ws.iter_rows(min_row=notes_start, max_row=notes_end, min_col=1, max_col=10):
        for cell in row:
            cell.border = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
    for idx, width in enumerate((11, 35, 13, 11, 15, 15, 13, 11, 15, 15), 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "C5"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "1:4"
    ws.print_area = f"A1:J{notes_end}"

    def add_sheet(name, headers, records, widths):
        sheet = wb.create_sheet(name)
        sheet.append(headers)
        for col, cell in enumerate(sheet[1], 1):
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.column_dimensions[get_column_letter(col)].width = widths[col-1]
        sheet.row_dimensions[1].height = 34
        for record in records:
            sheet.append(record)
            for cell in sheet[sheet.max_row]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        return sheet

    cols = ["Date", "Shift", "Drill Set", "Drill Machine", "Compressor", "Source",
            "Drill Open HMR", "Drill Close HMR", "Drill Hours", "Comp Open HMR",
            "Comp Close HMR", "Comp Hours", "Total Holes", "ROM/OB Holes",
            "BHJ/BHQ Holes", "ROM/OB m", "BHJ/BHQ m", "Total m",
            "Meters/hr", "Breakdown Hours", "Remarks", "Entered By", "Entered At"]
    records = []
    for item in detail_rows:
        records.append([item.get(k) for k in
            ("date", "shift", "label", "drillMachine", "compressor", "source",
             "drillOpenHmr", "drillCloseHmr", "drillHours", "compressorOpenHmr",
             "compressorCloseHmr", "compressorHours", "holes", "romObHoles",
             "bhjBhqHoles", "romObMeterage", "bhjBhqMeterage", "totalMeterage",
             "metersPerHour", "breakdownHours", "remarks", "enteredBy", "enteredAt")])
    add_sheet("Shift Details", cols, records,
              [13, 8, 32, 32, 32, 28, 15, 15, 12, 15, 15, 12, 13, 14, 14, 14, 14, 14, 12, 15, 65, 18, 21])

    # Split summary deliberately omits historic holes not assigned to a material.
    for title, holes_key, metres_key in (
        ("ROM-OB Summary", "romObHoles", "romObMeterage"),
        ("BHJ-BHQ Summary", "bhjBhqHoles", "bhjBhqMeterage"),
    ):
        grouped = {}
        for item in detail_rows:
            key = (item["drillSetId"], item["label"])
            group = grouped.setdefault(key, {"holes": 0, "metres": 0, "unclassified": False})
            h = item.get(holes_key)
            if h is None:
                group["unclassified"] = True
            else:
                group["holes"] += h
            group["metres"] += item.get(metres_key) or 0
        add_sheet(title, ["Drill Set ID", "Equipment", "Assigned Holes", "Meterage (m)",
                          "Unclassified Hole Records", "From", "To", "Shift"],
            [[key[0], key[1], val["holes"], round(val["metres"], 3),
              "YES" if val["unclassified"] else "NO", period["fromDate"],
              period["toDate"], period["shift"]] for key, val in sorted(grouped.items())],
            [22, 40, 18, 18, 30, 15, 15, 14])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
