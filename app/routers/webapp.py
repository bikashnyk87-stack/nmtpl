from datetime import date, datetime, timedelta, time as dtime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from openpyxl import Workbook
from pathlib import Path
from uuid import uuid4
import hashlib
import json
import re
import logging
import base64
from time import perf_counter
from email.message import EmailMessage

from fastapi import APIRouter, Depends, HTTPException, Request, Response, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import select, delete, text, func
from sqlalchemy.orm import Session
from sqlalchemy.exc import SQLAlchemyError

from app.db import get_db
from app.maintenance_models import MaintenanceBreakdown
from app.auth import (WebUser, WebSession, MODULES, COOKIE, csrf, get_user, require,
                      hash_password, check_password, issue_session, digest_token, utcnow, aware)
from app.models import (Person, Equipment, Location, Product, ShiftMaster, ShiftRotation, ActivityMaster, PersonAttendance,
                        EquipmentAttendance, ShiftCrew, ShiftDeployment, ShiftState, LoadTrip, LoadWbMatch,
                        HsdTanker, HsdPurchaseLot, HsdIssue, HsdIssueAllocation, AuditLog, WbImportBatch, WbMovement,
                        VehicleAlias, LocationAlias, MaterialAlias, MasterOption, WbHeaderAlias)
from app.services.time_context import now_local, operating_context, TZ, week_monday
from app.services.shift_rotation import resolve_rotation_shift
from app.services.hsd_fifo import apply_fifo_issue
from app.services.reconcile import auto_reconcile
from app.site_auth import has_permission
from app.services.wb_mapping import (
    ensure_wb_header_mapping, find_wb_sheet, row_value, build_location_resolver, norm_vehicle,
    expected_movement_date
)
from app.services.attendance_automation import auto_close_shift
from app.services.tiom_erp import (
    authoritative_wb as tiom_authoritative_wb,
    vehicle_matches as tiom_wb_vehicle_matches,
    movement_payload as tiom_wb_payload,
    wb_report_contributions as tiom_wb_report_contributions,
    allocate_blended_crusher_output,
)
from app.services.tiom_location_erp import canonicalize_wb_rows, location_options, ensure_location_master_roles
from app.site_models import (
    TiomSourceDeployment, TiomSourceContext, TiomRouteMaster, TiomDrillSet, TiomDrillingShift,
    TiomMisReport, TiomMisTripRow, TiomMisTripDetail, TiomTripFactor,
    TiomHsdReceiptDetail, TiomHsdIssueDetail, TiomHsdStockCheck, SiteAssetMeter,
    TiomShiftProductionReport, TiomShiftProductionMovement, TiomShiftReportBaseline, TiomMisReconciliation,
    TiomLeadDistance, TiomMisTripLead, TiomWbCanonical
)

router = APIRouter(prefix='/api/web', tags=['webapp'])
log = logging.getLogger('nmtpl.webapp')


def lock(db):
    # Serialize pilot writes with a transaction-scoped lock, including empty-row creation.
    # Legacy write routes are disabled; all pilot writes use this transaction boundary.
    if db.bind.dialect.name == 'postgresql':
        db.execute(text('SELECT pg_advisory_xact_lock(7304001)'))


def audit(db, user, action, entity, key, detail):
    db.add(AuditLog(created_at=now_local(), actor=user.login_id, action=action,
                    entity=entity, entity_id=str(key), detail=json.dumps(detail, default=str)))


class LoginIn(BaseModel):
    loginId: str = Field(min_length=1, max_length=60, pattern=r'^[A-Za-z0-9_.-]+$')
    password: str = Field(min_length=1, max_length=128)
    name: str = Field(default='Administrator', min_length=1, max_length=120)


@router.get('/setup-status')
def setup_status(db: Session = Depends(get_db)):
    return {'setupRequired': db.scalar(select(func.count()).select_from(WebUser)) == 0}


@router.post('/setup')
def setup(p: LoginIn, request: Request, response: Response, db: Session = Depends(get_db)):
    csrf(request)
    if request.url.hostname not in {'127.0.0.1', 'localhost', '::1'}:
        raise HTTPException(403, 'Create the first administrator from http://127.0.0.1:8000 on the server laptop before enabling remote access.')
    lock(db)
    if db.scalar(select(func.count()).select_from(WebUser)):
        raise HTTPException(409, 'Administrator already configured. Sign in.')
    user = WebUser(login_id=p.loginId.lower(), name=p.name, password_hash=hash_password(p.password),
                   modules=','.join(sorted(MODULES)), shifts='ALL', admin=True, active=True)
    db.add(user); db.flush()
    audit(db, user, 'CREATE_ADMIN', 'web_user', user.login_id, {})
    issue_session(db, user, response, request)
    db.commit()
    return {'ok': True}


@router.post('/login')
def login(p: LoginIn, request: Request, response: Response, db: Session = Depends(get_db)):
    csrf(request); lock(db)
    user = db.get(WebUser, p.loginId.lower())
    if user and user.locked_until and aware(user.locked_until) > utcnow():
        raise HTTPException(429, 'Too many attempts. Try again after 15 minutes.')
    if not user or not user.active or not check_password(p.password, user.password_hash):
        if user:
            user.failed_attempts += 1
            if user.failed_attempts >= 5:
                user.locked_until = utcnow() + timedelta(minutes=15)
                user.failed_attempts = 0
            db.commit()
        raise HTTPException(401, 'Invalid user ID or password.')
    user.failed_attempts = 0; user.locked_until = None
    old = request.cookies.get(COOKIE)
    if old:
        db.execute(delete(WebSession).where(WebSession.token_hash == digest_token(old)))
    db.execute(delete(WebSession).where(WebSession.expires_at < utcnow()))
    issue_session(db, user, response, request); db.commit()
    return {'ok': True}


@router.post('/logout')
def logout(request: Request, response: Response, db: Session = Depends(get_db)):
    csrf(request)
    db.execute(delete(WebSession).where(WebSession.token_hash == digest_token(request.cookies.get(COOKIE, ''))))
    db.commit(); response.delete_cookie(COOKIE, path='/')
    return {'ok': True}



# NMTPL_FLEX_UI_DATE_V2_START
def _parse_ui_date_v2(value, label='date'):
    """Accept ISO plus Indian browser/display date formats without changing stored SQL dates."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or '').strip()
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%d.%m.%Y'):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    raise HTTPException(422, f'Choose a valid {label}.')
# NMTPL_FLEX_UI_DATE_V2_END

def context(db, user, day, shift):
    day = _parse_ui_date_v2(day, 'operating date')
    sh = str(shift).strip().upper()
    definition = db.get(ShiftMaster, sh)
    if not definition or not definition.active:
        raise HTTPException(422, 'Choose an active shift.')
    require(user, shift=sh)
    return day, sh, definition


def rows(db, model, day, shift):
    return db.scalars(select(model).where(model.operating_date == day, model.shift == shift)).all()


def open_shift(db, day, shift):
    state = db.scalar(select(ShiftState).where(ShiftState.operating_date == day, ShiftState.shift == shift))
    if state and state.status == 'CLOSED':
        raise HTTPException(409, 'This shift is closed.')


def shift_bounds(day, definition):
    start = datetime.combine(day, definition.start_time, TZ)
    end = datetime.combine(day, definition.end_time, TZ)
    if end <= start:
        end += timedelta(days=1)
    return start, end


def active_resource(db, model, key):
    value = db.get(model, key)
    if not value or not value.active:
        raise HTTPException(422, f'Unknown or inactive resource: {key}')
    return value


def busy(db, machine):
    return db.scalar(select(LoadTrip.trip_id).where(
        ((LoadTrip.machine_id == machine) & (LoadTrip.status == 'LOADING')) |
        ((LoadTrip.vehicle_id == machine) & LoadTrip.status.in_(['LOADING', 'HAULING']))).limit(1))


def time_text(value):
    return value.astimezone(TZ).strftime('%H:%M:%S') if value else ''


def person_row(person, att=None):
    return {'id': person.employee_id, 'name': person.name, 'role': person.role,
            'rotationTeam': person.rotation_team or '', 'attendance': att.status if att else '',
            'inTime': time_text(att.in_at) if att else '', 'outTime': time_text(att.out_at) if att else '',
            'workedHours': att.worked_hours if att and att.worked_hours is not None else '',
            'outSource': (att.out_source or '') if att else '',
            'reviewStatus': (att.review_status or '') if att else '', 'remarks': (att.remarks or '') if att else ''}


def auto_close(db, user, day, sh, definition):
    return auto_close_shift(db, day, sh, actor=user.login_id)


# NMTPL_SIMPLE_UI_V4_START
def _attendance_roster_shift_v4(rotation_team, team, operating_date):
    key = str(rotation_team or '').strip().upper()
    direct = {'TEAM-A': 'A', 'TEAM-B': 'B', 'TEAM-C': 'C', 'GENERAL': 'GENERAL'}
    if key in direct:
        return direct[key]
    return resolve_rotation_shift(team, operating_date) if team else ''
# NMTPL_SIMPLE_UI_V4_END

def attendance(db, user, day, sh, definition):
    can_p = user.admin or 'PERSON_ATTENDANCE' in user.modules.split(',')
    can_e = user.admin or 'EQUIPMENT_ATTENDANCE' in user.modules.split(',')
    if not can_p and not can_e:
        raise HTTPException(403, 'Attendance access required.')
    # Auto-close is a separate explicit maintenance operation, never a side effect of a read.
    persons = db.scalars(select(Person).where(Person.active).order_by(Person.employee_id)).all()
    teams = {r.team_id: r for r in db.scalars(select(ShiftRotation).where(ShiftRotation.active))}
    pmap = {r.employee_id: r for r in rows(db, PersonAttendance, day, sh)}
    emap = {r.machine_id: r for r in rows(db, EquipmentAttendance, day, sh)}
    prows = []
    for p in persons:
        r = person_row(p, pmap.get(p.employee_id))
        team = teams.get(p.rotation_team)
        resolved = _attendance_roster_shift_v4(p.rotation_team, team, day)
        r.update(rosterShift=resolved, rostered=(resolved == sh or resolved == 'GENERAL'))
        prows.append(r)
    # Suggest the most recent prior closing meter without overwriting a saved opening.
    # This is calculated by operating date/shift, not by browser time.
    shift_order = {'A': 0, 'B': 1, 'C': 2, 'GENERAL': 3}
    current_rank = shift_order.get(sh, 99)
    prior_close = {}
    prior_rows = db.scalars(select(EquipmentAttendance).where(
        EquipmentAttendance.closing_meter.is_not(None),
        EquipmentAttendance.operating_date <= day
    ).order_by(EquipmentAttendance.operating_date.desc(), EquipmentAttendance.id.desc()).limit(10000)).all()
    for old in prior_rows:
        if old.operating_date == day and shift_order.get(old.shift, 99) >= current_rank:
            continue
        if old.machine_id not in prior_close:
            prior_close[old.machine_id] = old.closing_meter
    erows = []
    for e in db.scalars(select(Equipment).where(Equipment.active).order_by(Equipment.machine_id)):
        r = emap.get(e.machine_id)
        erows.append({'id': e.machine_id, 'type': e.type, 'group': e.group,
                      'meterType': 'KMR' if e.group in {'TRANSPORT', 'HSD_TANKER'} else 'HMR',
                      'attendance': r.status if r else '', 'condition': (r.condition or '') if r else '',
                      'opening': r.opening_meter if r and r.opening_meter is not None else '',
                      'closing': r.closing_meter if r and r.closing_meter is not None else '',
                      'delta': r.run_meter if r and r.run_meter is not None else '',
                      'suggestedOpening': prior_close.get(e.machine_id, ''), 'remarks': (r.remarks or '') if r else ''})
    assigned = sum(bool(p.rotation_team in teams or p.rotation_team == 'GENERAL') for p in persons)
    return {'ok': True, 'date': day, 'shift': sh, 'canPersons': can_p, 'canEquipment': can_e,
            'shiftStart': str(definition.start_time), 'shiftEnd': str(definition.end_time),
            'persons': prows if can_p else [], 'equipment': erows if can_e else [],
            'roster': {'configured': bool(teams), 'ready': assigned == len(persons),
                       'weekMonday': week_monday(day), 'assignedCount': assigned,
                       'unassignedCount': len(persons)-assigned, 'rosteredCount': sum(r['rostered'] for r in prows),
                       'teams': [{'id': r.team_id, 'name': r.team_name, 'currentShift': _attendance_roster_shift_v4(r.team_id, r, day)} for r in teams.values()]
                                + [{'id': 'GENERAL', 'name': 'General', 'currentShift': 'GENERAL'}]}}


def punch(db, user, p, day, sh, definition):
    require(user, 'PERSON_ATTENDANCE'); open_shift(db, day, sh)
    ids = p.get('ids', [])
    if not isinstance(ids, list) or not 1 <= len(ids) <= 500 or len(set(ids)) != len(ids):
        raise HTTPException(422, 'Choose 1â€“500 distinct employees.')
    action, status = p.get('action'), p.get('status')
    now = now_local(); live_day, live_shift, _ = operating_context(now)
    if action == 'STATUS':
        if (day, sh) != (live_day, live_shift) and sh != 'GENERAL':
            raise HTTPException(409, 'Live punches require the current operating shift.')
        start, end = shift_bounds(day, definition)
        if not start <= now <= end:
            raise HTTPException(409, 'Live punches require the selected shift time window.')
        if status not in {'PRESENT', 'ABSENT', 'LEAVE', 'WEEKLY OFF'}:
            raise HTTPException(422, 'Invalid attendance status.')
    elif action == 'OUT':
        start, end = shift_bounds(day, definition)
        if not start <= now <= end + timedelta(minutes=120):
            raise HTTPException(409, 'OUT is outside the shift/grace window.')
    else:
        raise HTTPException(422, 'Invalid attendance action.')
    existing = {r.employee_id: r for r in rows(db, PersonAttendance, day, sh)}
    result = []
    for key in ids:
        person = active_resource(db, Person, key)
        row = existing.get(key)
        if action == 'OUT' and (not row or row.status != 'PRESENT' or not row.in_at):
            raise HTTPException(409, f'{key} must be PRESENT before OUT.')
        assignments = [r for r in rows(db, ShiftCrew, day, sh) if r.status == 'ACTIVE' and r.employee_id == key]
        if (action == 'OUT' or status != 'PRESENT') and assignments:
            raise HTTPException(409, f'Release {key} from crew before changing attendance or recording OUT.')
        if not row:
            row = PersonAttendance(operating_date=day, shift=sh, employee_id=key, status=status)
            db.add(row)
        before = person_row(person, row)
        if action == 'OUT':
            if not row.out_at:  # Repeated OUT must preserve the original time.
                row.out_at = now
                row.worked_hours = Decimal(str(round((now-row.in_at).total_seconds()/3600, 2)))
                row.out_source = 'SERVER_PUNCH'; row.review_status = 'OK'
        elif status == 'PRESENT':
            if row.out_at:
                raise HTTPException(409, f'{key} has already clocked OUT. An audited correction is required.')
            row.in_at = row.in_at or now; row.status = status
        else:
            if row.in_at:
                raise HTTPException(409, f'{key} already has an IN time. Do not erase attendance history.')
            row.status = status
        row.entered_by = user.login_id; row.entered_at = now
        audit(db, user, 'PUNCH', 'person_attendance', key, {'before': before, 'after': person_row(person, row)})
        result.append(person_row(person, row))
    return {'ok': True, 'message': 'Attendance saved.', 'rows': result}


def decimal_meter(value):
    if value is None or value == '':
        return None
    try:
        result = Decimal(str(value))
        if not result.is_finite() or result < 0 or result >= Decimal('1000000000000') or result.as_tuple().exponent < -2:
            raise ValueError()
        return result
    except (InvalidOperation, ValueError):
        raise HTTPException(422, 'Meters must be non-negative numbers with at most two decimals.')


def equipment_save(db, user, p, day, sh):
    require(user, 'EQUIPMENT_ATTENDANCE'); open_shift(db, day, sh)
    current = {r.machine_id: r for r in rows(db, EquipmentAttendance, day, sh)}
    for item in batch_rows(p):
        e = active_resource(db, Equipment, item.get('id'))
        status, condition = item.get('attendance'), item.get('condition')
        if status not in {'PRESENT', 'ABSENT'} or (status == 'PRESENT' and condition not in {'WORKING', 'BREAKDOWN', 'MAINTENANCE', 'STANDBY'}):
            raise HTTPException(422, f'Invalid equipment status/condition: {e.machine_id}')
        opening, closing = decimal_meter(item.get('opening')), decimal_meter(item.get('closing'))
        if status == 'PRESENT' and opening is None:
            raise HTTPException(422, f'Opening meter required for {e.machine_id}.')
        if closing is not None and (opening is None or closing < opening):
            raise HTTPException(422, f'Closing meter must be at least the opening meter: {e.machine_id}')
        if busy(db, e.machine_id):
            raise HTTPException(409, f'{e.machine_id} has an active loading/hauling trip.')
        if status != 'PRESENT' or condition != 'WORKING':
            if any(r.status == 'ACTIVE' and r.machine_id == e.machine_id for r in rows(db, ShiftCrew, day, sh)+rows(db, ShiftDeployment, day, sh)):
                raise HTTPException(409, f'Release deployment and crew before changing {e.machine_id}.')
        row = current.get(e.machine_id)
        if not row:
            row = EquipmentAttendance(operating_date=day, shift=sh, machine_id=e.machine_id, status=status)
            db.add(row)
        before = {'status': row.status, 'condition': row.condition, 'opening': row.opening_meter, 'closing': row.closing_meter}
        row.status = status; row.condition = condition if status == 'PRESENT' else None
        row.meter_type = 'KMR' if e.group in {'TRANSPORT', 'HSD_TANKER'} else 'HMR'
        row.opening_meter = opening; row.closing_meter = closing
        row.run_meter = closing-opening if opening is not None and closing is not None else None
        row.remarks = short(item.get('remarks', '')); row.entered_by = user.login_id; row.entered_at = now_local()
        audit(db, user, 'EQUIPMENT_ATTENDANCE', 'equipment_attendance', e.machine_id, {'before': before, 'after': item})
    return {'ok': True, 'message': 'Equipment attendance saved.'}


def short(value):
    if not isinstance(value, str) or len(value) > 250:
        raise HTTPException(422, 'Text must be 250 characters or fewer.')
    return value


def batch_rows(p):
    result = p.get('rows', [])
    if not isinstance(result, list) or not 1 <= len(result) <= 500 or not all(isinstance(x, dict) for x in result):
        raise HTTPException(422, 'Choose 1â€“500 rows.')
    keys = [x.get('id', x.get('machineId')) for x in result]
    if len(keys) != len(set(keys)):
        raise HTTPException(422, 'Duplicate rows in request.')
    return result


def compatible(person, equipment):
    role = ' '.join(person.role.upper().split())
    allowed = ('DRIVER', 'TRUCK OPERATOR', 'TIPPER OPERATOR', 'DUMPER OPERATOR', 'TRANSPORT OPERATOR', 'HEMM OPERATOR') if equipment.group in {'TRANSPORT', 'HSD_TANKER'} else ('EXCAVATOR OPERATOR', 'LOADER OPERATOR', 'HEMM OPERATOR')
    return any(x in role for x in allowed)


def setup_row(db, user, p, day, sh):
    require(user, 'SHIFT_CONTROL'); open_shift(db, day, sh)
    machine = active_resource(db, Equipment, p.get('machineId'))
    ea = db.scalar(select(EquipmentAttendance).where(EquipmentAttendance.operating_date == day,
                   EquipmentAttendance.shift == sh, EquipmentAttendance.machine_id == machine.machine_id))
    if not ea or ea.status != 'PRESENT':
        raise HTTPException(409, 'Mark equipment PRESENT first.')
    if p.get('condition') != 'WORKING' or ea.condition != 'WORKING':
        raise HTTPException(409, 'Only WORKING equipment can be deployed; set condition in Attendance.')
    person = active_resource(db, Person, p.get('employeeId'))
    location = active_resource(db, Location, p.get('locationId'))
    pa = db.scalar(select(PersonAttendance).where(PersonAttendance.operating_date == day,
                   PersonAttendance.shift == sh, PersonAttendance.employee_id == person.employee_id))
    if not pa or pa.status != 'PRESENT' or not pa.in_at or pa.out_at:
        raise HTTPException(409, 'Crew must be PRESENT with IN recorded and no OUT.')
    if not compatible(person, machine):
        raise HTTPException(409, 'Crew role is incompatible with this equipment group.')
    if busy(db, machine.machine_id):
        raise HTTPException(409, 'Finish the active trip before changing deployment.')
    crews = [r for r in rows(db, ShiftCrew, day, sh) if r.status == 'ACTIVE']
    if any(r.employee_id == person.employee_id and r.machine_id != machine.machine_id for r in crews):
        raise HTTPException(409, 'This employee is already assigned elsewhere. Release that assignment first.')
    now = now_local(); reason = short(p.get('reason', 'Shift setup'))
    same_crew = False
    for r in crews:
        if r.machine_id == machine.machine_id:
            if r.employee_id == person.employee_id:
                same_crew = True
            else:
                r.status = 'CLOSED'; r.to_at = now; r.changed_by = user.login_id; r.changed_at = now; r.reason = reason
    if not same_crew:
        db.add(ShiftCrew(operating_date=day, shift=sh, machine_id=machine.machine_id,
                        employee_id=person.employee_id, from_at=now, status='ACTIVE',
                        reason=reason, changed_by=user.login_id, changed_at=now))
    same_dep = False
    for r in rows(db, ShiftDeployment, day, sh):
        if r.machine_id == machine.machine_id and r.status == 'ACTIVE':
            if r.location_id == location.location_id:
                same_dep = True
            else:
                r.status = 'CLOSED'; r.to_at = now; r.changed_by=user.login_id; r.changed_at=now; r.reason=reason
    if not same_dep:
        db.add(ShiftDeployment(operating_date=day, shift=sh, machine_id=machine.machine_id,
                              location_id=location.location_id, from_at=now, status='ACTIVE',
                              reason=reason, changed_by=user.login_id, changed_at=now))
    audit(db, user, 'DEPLOY', 'equipment', machine.machine_id, p)
    db.flush()
    return {'ok': True, 'message': 'Crew and deployment saved.'}


def release(db, user, p, day, sh):
    require(user, 'SHIFT_CONTROL'); open_shift(db, day, sh)
    key = p.get('machineId'); active_resource(db, Equipment, key)
    reason = short(p.get('reason', '')).strip()
    if not reason:
        raise HTTPException(422, 'A release reason is required.')
    if busy(db, key):
        raise HTTPException(409, 'Finish the active trip before releasing equipment.')
    for row in rows(db, ShiftCrew, day, sh)+rows(db, ShiftDeployment, day, sh):
        if row.machine_id == key and row.status == 'ACTIVE':
            row.status = 'CLOSED'; row.to_at = now_local(); row.changed_by = user.login_id
            row.changed_at = now_local(); row.reason = reason
    audit(db, user, 'RELEASE', 'equipment', key, {'reason': reason})
    return {'ok': True, 'message': 'Crew and deployment released.'}


def shift_desk(db, user, day, sh):
    require(user, 'SHIFT_CONTROL')
    persons = {p.employee_id: p for p in db.scalars(select(Person).where(Person.active))}
    equipment = {e.machine_id: e for e in db.scalars(select(Equipment).where(Equipment.active))}
    locations = {l.location_id: l for l in db.scalars(select(Location).where(Location.active))}
    crews = {r.machine_id: r for r in rows(db, ShiftCrew, day, sh) if r.status == 'ACTIVE'}
    deps = {r.machine_id: r for r in rows(db, ShiftDeployment, day, sh) if r.status == 'ACTIVE'}
    out = []
    for row in rows(db, EquipmentAttendance, day, sh):
        e = equipment.get(row.machine_id)
        if row.status != 'PRESENT' or not e:
            continue
        c, d = crews.get(e.machine_id), deps.get(e.machine_id)
        out.append({'id': e.machine_id, 'type': e.type, 'group': e.group, 'condition': row.condition,
                    'employeeId': c.employee_id if c else '', 'locationId': d.location_id if d else '',
                    'locationName': locations[d.location_id].location_name if d and d.location_id in locations else '',
                    'fromTime': time_text(d.from_at) if d else ''})
    return {'date': day, 'shift': sh, 'equipment': sorted(out, key=lambda r:r['id']),
            'persons': [person_row(persons[r.employee_id], r) for r in rows(db, PersonAttendance, day, sh)
                        if r.status == 'PRESENT' and r.in_at and not r.out_at and r.employee_id in persons],
            'locations': [{'id': l.location_id, 'name': l.location_name} for l in locations.values()]}


def close_status(db, user, day, sh, definition):
    require(user, 'SHIFT_CONTROL')
    pa = [r for r in rows(db, PersonAttendance, day, sh) if r.status == 'PRESENT']
    ea = [r for r in rows(db, EquipmentAttendance, day, sh) if r.status == 'PRESENT']
    trips = rows(db, LoadTrip, day, sh)
    open_loads = sum(t.status in {'LOADING', 'HAULING'} for t in trips)
    batches = rows(db, WbImportBatch, day, sh)
    confirmed = [b for b in batches if b.confirmed_at and b.status == 'CONFIRMED']
    review = sum(b.review_rows for b in confirmed)
    state = db.scalar(select(ShiftState).where(ShiftState.operating_date == day, ShiftState.shift == sh))
    blockers = []
    missing_out = sum(not r.out_at for r in pa)
    missing_open = sum(r.opening_meter is None for r in ea)
    missing_close = sum(r.closing_meter is None for r in ea)
    for count, label in [(missing_out, 'person OUT missing'), (missing_open, 'opening meters missing'),
                         (missing_close, 'closing meters missing'), (open_loads, 'trips still active'), (review, 'WB rows need review')]:
        if count: blockers.append(f'{count} {label}')
    wb_required = any(t.vehicle_id for t in trips)
    if wb_required and not confirmed: blockers.append('WB confirmation required')
    _, end = shift_bounds(day, definition)
    if now_local() < end: blockers.append('Shift has not ended')
    if state and state.status == 'CLOSED': blockers.append('Shift is already closed')
    return {'date': day, 'shift': sh, 'shiftStart': str(definition.start_time), 'shiftEnd': str(definition.end_time),
            'state': state.status if state else 'OPEN', 'personsPresent': len(pa), 'equipmentPresent': len(ea),
            'missingOut': missing_out, 'missingOpening': missing_open, 'missingClosing': missing_close,
            'openLoads': open_loads, 'wbRequired': wb_required, 'wbConfirmed': bool(confirmed),
            'wbReviewRows': review, 'blockers': blockers, 'canClose': not blockers,
            'warnings': ['Automatic OUT records require HR review.'] if any(r.review_status == 'REVIEW' for r in pa) else []}


def _date_param(value, label):
    return _parse_ui_date_v2(value, label)


def _money(value):
    return float(value or 0)


def dashboard_desk(db, user, p):
    dash_started = perf_counter()
    require(user, 'DASHBOARD')
    live_day, live_shift, _ = operating_context()
    mode = str(p.get('mode') or 'CURRENT_SHIFT').upper()
    if mode not in {'CURRENT_SHIFT', 'TODAY', 'YESTERDAY', '7D', 'MTD', 'CUSTOM'}:
        raise HTTPException(422, 'Choose a valid dashboard period.')
    selected_shift = str(p.get('shift') or ('ALL' if mode != 'CURRENT_SHIFT' else live_shift)).upper()
    if selected_shift != 'ALL':
        definition = db.get(ShiftMaster, selected_shift)
        if not definition or not definition.active:
            raise HTTPException(422, 'Choose an active shift or ALL.')
        require(user, shift=selected_shift)
    if mode == 'CURRENT_SHIFT':
        start_day = end_day = live_day; selected_shift = live_shift; require(user, shift=selected_shift)
    elif mode == 'TODAY':
        start_day = end_day = _date_param(p.get('toDate') or live_day, 'date')
    elif mode == 'YESTERDAY':
        start_day = end_day = live_day - timedelta(days=1)
    elif mode == '7D':
        end_day = _date_param(p.get('toDate') or live_day, 'end date'); start_day = end_day - timedelta(days=6)
    elif mode == 'MTD':
        end_day = _date_param(p.get('toDate') or live_day, 'end date'); start_day = end_day.replace(day=1)
    else:
        start_day = _date_param(p.get('fromDate') or live_day, 'from date')
        end_day = _date_param(p.get('toDate') or live_day, 'to date')
    if end_day < start_day or (end_day - start_day).days > 61:
        raise HTTPException(422, 'Dashboard range must be 1–62 days.')

    allowed_shifts = None
    if selected_shift == 'ALL' and not user.admin and 'ALL' not in user.shifts.split(','):
        allowed_shifts = [x for x in user.shifts.split(',') if x]

    def scoped_range(model, first, last, shift_override=None):
        stmt = select(model).where(model.operating_date >= first, model.operating_date <= last)
        sh = selected_shift if shift_override is None else shift_override
        if sh != 'ALL':
            stmt = stmt.where(model.shift == sh)
        elif allowed_shifts:
            stmt = stmt.where(model.shift.in_(allowed_shifts))
        return db.scalars(stmt).all()

    def authoritative_wb(first, last, shift_override=None):
        batch_rows = [b for b in scoped_range(WbImportBatch, first, last, shift_override)
                      if b.status == 'CONFIRMED' and b.confirmed_at]
        active = {}
        for b in batch_rows:
            key = (b.operating_date, b.shift)
            if key not in active or b.confirmed_at > active[key].confirmed_at:
                active[key] = b
        ids = [b.batch_id for b in active.values()]
        movements = db.scalars(select(WbMovement).where(
            WbMovement.batch_id.in_(ids), WbMovement.row_status == 'VALID'
        )).all() if ids else []
        return active, movements

    def material_label(w):
        return (w.material_name or w.material_code or 'Unmapped').strip()

    def tonnes(w):
        return float((w.net_kg or Decimal('0')) / Decimal('1000'))

    def duration_minutes(start, finish):
        if not start or not finish:
            return None
        a, b = aware(start), aware(finish)
        value = (b - a).total_seconds() / 60
        return value if value >= 0 else None

    def add_metric(store, key, t=0.0, trips=1):
        row = store.setdefault(key, {'label': key, 'trips': 0, 'tonnes': 0.0})
        row['trips'] += trips; row['tonnes'] += t
        return row

    def has_token(value, *tokens):
        text_value = str(value or '').upper()
        return any(tok in text_value for tok in tokens)

    pa = scoped_range(PersonAttendance, start_day, end_day)
    ea = scoped_range(EquipmentAttendance, start_day, end_day)
    trips = scoped_range(LoadTrip, start_day, end_day)
    hsd = scoped_range(HsdIssue, start_day, end_day)
    active_batch, wb_all = authoritative_wb(start_day, end_day)

    # Dashboard is a read path. Reuse persisted canonical WB mappings in one
    # query instead of re-canonicalizing every movement on every refresh.
    wb_keys_all = [w.movement_key for w in wb_all]
    canonical_rows = list(db.scalars(
        select(TiomWbCanonical).where(TiomWbCanonical.movement_key.in_(wb_keys_all))
    )) if wb_keys_all else []
    wb_canonical_all = {x.movement_key: x for x in canonical_rows}
    missing_canonical = [w for w in wb_all if w.movement_key not in wb_canonical_all]
    if missing_canonical:
        wb_canonical_all.update(canonicalize_wb_rows(db, missing_canonical))
    source_options = location_options(db, 'SOURCE')
    destination_options = location_options(db, 'DESTINATION')
    available_materials = sorted({material_label(w) for w in wb_all})
    available_sources = source_options
    available_destinations = destination_options
    available_vehicles = sorted({(w.vehicle_id or w.vehicle_raw or 'Unknown').strip() for w in wb_all})
    source_ids = {x['id'] for x in source_options}
    destination_ids = {x['id'] for x in destination_options}
    material_filter = str(p.get('materialFilter') or '').strip()
    source_filter = str(p.get('sourceFilter') or '').strip()
    destination_filter = str(p.get('destinationFilter') or '').strip()
    vehicle_filter = str(p.get('vehicleFilter') or '').strip()
    if material_filter and material_filter not in available_materials: material_filter = ''
    if source_filter and source_filter not in source_ids: source_filter = ''
    if destination_filter and destination_filter not in destination_ids: destination_filter = ''
    if vehicle_filter and vehicle_filter not in available_vehicles: vehicle_filter = ''
    wb = [w for w in wb_all if
          (not material_filter or material_label(w) == material_filter) and
          (not source_filter or (wb_canonical_all.get(w.movement_key) and wb_canonical_all[w.movement_key].source_location_id == source_filter)) and
          (not destination_filter or (wb_canonical_all.get(w.movement_key) and wb_canonical_all[w.movement_key].destination_location_id == destination_filter)) and
          (not vehicle_filter or (w.vehicle_id or w.vehicle_raw or 'Unknown').strip() == vehicle_filter)]
    wb_canonical = {w.movement_key: wb_canonical_all.get(w.movement_key) for w in wb}
    if vehicle_filter:
        trips = [t for t in trips if (t.vehicle_id or '').strip() == vehicle_filter]
        hsd = [x for x in hsd if (x.machine_id or '').strip() == vehicle_filter]

    trip_ids = [t.trip_id for t in trips]
    movement_keys = [w.movement_key for w in wb]
    matches = []
    if trip_ids or movement_keys:
        clauses = []
        if trip_ids: clauses.append(LoadWbMatch.trip_id.in_(trip_ids))
        if movement_keys: clauses.append(LoadWbMatch.movement_key.in_(movement_keys))
        from sqlalchemy import or_
        matches = db.scalars(select(LoadWbMatch).where(or_(*clauses))).all()
    trip_scope, movement_scope = set(trip_ids), set(movement_keys)
    matches = [m for m in matches if m.trip_id in trip_scope and m.movement_key in movement_scope]
    matched_trip = {m.trip_id for m in matches if m.status in {'MATCHED','MANUAL_MATCH','LIKELY_MATCH'}}
    matched_move = {m.movement_key for m in matches if m.status in {'MATCHED','MANUAL_MATCH','LIKELY_MATCH'}}
    exact_keys = {m.movement_key for m in matches if m.status in {'MATCHED','MANUAL_MATCH'}}
    likely_keys = {m.movement_key for m in matches if m.status == 'LIKELY_MATCH'} - exact_keys

    trip_by_id = {t.trip_id: t for t in trips}
    wb_by_key = {w.movement_key: w for w in wb}
    exact_pairs = []
    seen_movements = set()
    for m in matches:
        if m.status not in {'MATCHED','MANUAL_MATCH'} or m.movement_key in seen_movements:
            continue
        t, w = trip_by_id.get(m.trip_id), wb_by_key.get(m.movement_key)
        if t and w:
            seen_movements.add(m.movement_key); exact_pairs.append((t,w))

    wb_trips = len(wb)
    wb_tonnes = sum(tonnes(w) for w in wb)
    avg_payload = wb_tonnes / wb_trips if wb_trips else 0.0
    filtered_exact_trip_ids = {t.trip_id for t,w in exact_pairs}
    field_vehicle_trips = [t for t in trips if t.vehicle_id and (not material_filter or t.trip_id in filtered_exact_trip_ids)]
    wb_unmatched = sum(w.movement_key not in matched_move for w in wb)
    trip_no_wb = sum(t.status == 'CLOSED' and t.trip_id not in matched_trip for t in field_vehicle_trips)

    present_people = [r for r in pa if r.status == 'PRESENT']
    present_equipment = [r for r in ea if r.status == 'PRESENT']
    working_equipment = [r for r in present_equipment if (r.condition or '').upper() == 'WORKING']
    hsd_litres = sum((x.litres or Decimal('0')) for x in hsd)
    hsd_cost = sum((x.amount or Decimal('0')) for x in hsd)
    open_loading = sum(t.status == 'LOADING' for t in trips)
    in_transit = sum(t.status == 'HAULING' for t in trips)

    dep_stmt = select(ShiftDeployment).where(ShiftDeployment.operating_date >= start_day, ShiftDeployment.operating_date <= end_day)
    if selected_shift != 'ALL': dep_stmt = dep_stmt.where(ShiftDeployment.shift == selected_shift)
    elif allowed_shifts: dep_stmt = dep_stmt.where(ShiftDeployment.shift.in_(allowed_shifts))
    deps = db.scalars(dep_stmt).all()
    deployed = sum(r.status == 'ACTIVE' for r in deps) if start_day == end_day and selected_shift != 'ALL' else len(deps)

    equipment_map = {e.machine_id: e for e in db.scalars(select(Equipment)).all()}
    location_map = {l.location_id: l for l in db.scalars(select(Location)).all()}

    # TIOM production truth: fresh ROM movement is the 100% mine-production base.
    # Plant feed is a separate physical throughput measure. Rehandling / old stock
    # may feed a plant but must never increase fresh production or recovery.
    production_facts={}
    production_fact_trips={}
    production_fact_rows={}
    excluded_movement_keys=set()
    excluded_codes={'PROJECT_AREA_FINES_TO_STACK','PRODUCT_REHANDLED','ROM_REHANDLED','ROM_STOCK_TO_STOCK','ROM_STOCK_TO_PLANT_FEED','OLD_STOCK_BLEND_10_40'}
    old_stock_excluded=Decimal('0')
    old_stock_trips=0
    production_hourly={h:{'hour':f'{h:02d}:00','rom':0.0,'final':0.0} for h in range(24)}
    hourly_process_facts={h:{} for h in range(24)}
    final_output_codes={'SCREEN_FINES','SCREEN_5_18','CRUSHER_FINES','CRUSHER_5_18'}

    def _wb_side_text(row, side):
        canon=wb_canonical_all.get(row.movement_key)
        loc_id=(canon.source_location_id if side=='SOURCE' else canon.destination_location_id) if canon else None
        loc=location_map.get(loc_id) if loc_id else None
        raw=row.source_raw if side=='SOURCE' else row.destination_raw
        return ' '.join([
            str(raw or ''),str(loc_id or ''),
            str(loc.location_name if loc else ''),
            str(loc.location_type if loc else '')
        ]).upper()

    def _dashboard_process_facts(row):
        facts=list(tiom_wb_report_contributions(row))
        codes={code for code,_qty in facts}
        qty=(row.net_kg or Decimal('0'))/Decimal('1000')
        if qty<=0:
            return facts
        mat=(' '.join([str(row.material_code or ''),str(row.material_name or '')])).upper()
        src_txt=_wb_side_text(row,'SOURCE')
        dst_txt=_wb_side_text(row,'DESTINATION')
        source_is_stock=any(x in src_txt for x in ('STOCK','STACK'))
        source_is_project=('FINE' in mat and (('PROJECT' in src_txt and 'FINE' in src_txt) or (' PA ' in f' {src_txt} ' and 'FINE' in src_txt)))

        # Canonical master semantics override ambiguous raw labels: a movement
        # out of stock/project area is rehandling, never fresh production.
        if source_is_stock or source_is_project:
            facts=[(code,val) for code,val in facts if code not in final_output_codes and code!='ROM']
            codes={code for code,_qty in facts}
            rehandle_code='ROM_REHANDLED' if 'ROM' in mat else 'PRODUCT_REHANDLED'
            if rehandle_code not in codes:
                facts.append((rehandle_code,qty)); codes.add(rehandle_code)

        # Enrich physical plant-feed facts from canonical destination semantics.
        if 'ROM' in mat and any(x in dst_txt for x in ('MSP','SCREEN')) and 'MSP_FEED' not in codes:
            facts.append(('MSP_FEED',qty)); codes.add('MSP_FEED')
        intermediate={'LUMPS_FROM_SCREEN','LUMPS_FEED_TO_CRUSHER','LUMPS_TO_CRUSHER_FROM_STOCK','SPILLAGE_TO_CRUSHER','PRODUCT_REHANDLED','OLD_STOCK_BLEND_10_40'}
        if codes.intersection(intermediate) and any(x in dst_txt for x in ('CRUSH','OCP')) and 'CRUSHER_FEED' not in codes:
            facts.append(('CRUSHER_FEED',qty))
        return facts

    for row in wb_all:
        facts=_dashboard_process_facts(row)
        fact_codes={code for code,_qty in facts}
        if fact_codes.intersection(excluded_codes) and row.movement_key not in excluded_movement_keys:
            excluded_movement_keys.add(row.movement_key)
            old_stock_excluded+=(row.net_kg or Decimal('0'))/Decimal('1000')
        stamp=aware(row.weigh_at).astimezone(TZ) if row.weigh_at else None
        seen_codes=set()
        for code,qty in facts:
            q=Decimal(qty or 0)
            production_facts[code]=production_facts.get(code,Decimal('0'))+q
            production_fact_rows.setdefault(code,[]).append(row)
            if code not in seen_codes:
                production_fact_trips[code]=production_fact_trips.get(code,0)+1
                seen_codes.add(code)
            if stamp:
                hour_store=hourly_process_facts[stamp.hour]
                hour_store[code]=hour_store.get(code,Decimal('0'))+q
            if stamp and code=='ROM':
                production_hourly[stamp.hour]['rom']+=float(q)
    old_stock_trips=len(excluded_movement_keys)
    for hour,hfacts in hourly_process_facts.items():
        screen_hour=hfacts.get('SCREEN_FINES',Decimal('0'))+hfacts.get('SCREEN_5_18',Decimal('0'))
        crusher_gross_hour=hfacts.get('CRUSHER_FINES',Decimal('0'))+hfacts.get('CRUSHER_5_18',Decimal('0'))
        crusher_feed_hour=hfacts.get('CRUSHER_FEED',Decimal('0'))
        crusher_blend_hour=hfacts.get('CRUSHER_BLEND_FEED',Decimal('0'))
        crusher_fresh_hour,_blend_hour=allocate_blended_crusher_output(crusher_gross_hour,crusher_feed_hour,crusher_blend_hour)
        production_hourly[hour]['final']=float(screen_hour+crusher_fresh_hour)

    # Fresh mine ROM is the 100% production base. MSP feed is deliberately
    # separate because it can contain ROM drawn from prior stock.
    msp_feed=production_facts.get('MSP_FEED',Decimal('0'))
    fresh_rom=production_facts.get('ROM',Decimal('0'))
    rom_input=fresh_rom if fresh_rom>0 else msp_feed
    rom_basis='FRESH_ROM' if fresh_rom>0 else 'MSP_FEED_FALLBACK'

    screen_fines=production_facts.get('SCREEN_FINES',Decimal('0'))
    screen_518=production_facts.get('SCREEN_5_18',Decimal('0'))
    crusher_fines=production_facts.get('CRUSHER_FINES',Decimal('0'))
    crusher_clo=production_facts.get('CRUSHER_5_18',Decimal('0'))
    screen_lumps=production_facts.get('LUMPS_FROM_SCREEN',Decimal('0'))
    spillage_to_crusher=production_facts.get('SPILLAGE_TO_CRUSHER',Decimal('0'))
    spillage_recycle=production_facts.get('SPILLAGE_RECYCLE_MSP',Decimal('0'))
    spillage_wip=production_facts.get('SPILLAGE_WIP',Decimal('0'))
    screen_intermediate=screen_lumps+spillage_to_crusher
    crusher_feed_fact=production_facts.get('CRUSHER_FEED',Decimal('0'))
    crusher_blend_feed=production_facts.get('CRUSHER_BLEND_FEED',Decimal('0'))

    screen_direct=screen_fines+screen_518
    screen_accounted=screen_direct+screen_intermediate+spillage_wip
    crusher_final=crusher_fines+crusher_clo
    crusher_fresh_final,crusher_blend_output=allocate_blended_crusher_output(
        crusher_final,crusher_feed_fact,crusher_blend_feed
    )
    final_production=screen_direct+crusher_fresh_final
    # Keep physical gross-output evidence distinct from an unverified fresh share.
    # A missing crusher-feed quantity must never silently prove 100% fresh output.
    allocation_pending=(crusher_final>0 and (crusher_feed_fact<=0 or crusher_blend_feed>crusher_feed_fact))
    allocation_status=('CRUSHER_FEED_TONNAGE_MISSING' if crusher_final>0 and crusher_feed_fact<=0
                       else 'BLEND_FEED_EXCEEDS_TOTAL' if crusher_blend_feed>crusher_feed_fact
                       else 'PROPORTIONAL_ALLOCATION' if crusher_blend_feed>0
                       else 'NO_RECORDED_BLEND')

    def prod_pct(value,base=rom_input):
        return round(float(Decimal(value or 0)/base*Decimal('100')),2) if base and base>0 else None

    # Bench context is fallback only for legacy Source Masters that do not carry
    # RL in the canonical Source name/ID.
    ctx_stmt=select(TiomSourceContext).where(
        TiomSourceContext.operating_date>=start_day,
        TiomSourceContext.operating_date<=end_day,
        TiomSourceContext.active.is_(True)
    )
    if selected_shift!='ALL':
        ctx_stmt=ctx_stmt.where(TiomSourceContext.shift==selected_shift)
    elif allowed_shifts:
        ctx_stmt=ctx_stmt.where(TiomSourceContext.shift.in_(allowed_shifts))
    source_contexts={(x.operating_date,x.shift,x.source_location_id):x.bench_rl_m for x in db.scalars(ctx_stmt)}
    # One approved master lookup per distinct path/bench for dashboard WB + MIS.
    # This avoids re-querying Route and Lead Masters for every trip of a large shift.
    dashboard_lead_cache={}
    def dashboard_lead(source_id, bench_rl, dest_id, mode='WITH_WB', wb_linked=False):
        key=(source_id,bench_rl,dest_id,mode,bool(wb_linked))
        if key not in dashboard_lead_cache:
            dashboard_lead_cache[key]=_tiom_resolve_lead(
                db,source_id,bench_rl,dest_id,mode,wb_linked=wb_linked)
        return dashboard_lead_cache[key]

    def _unique_wb_rows(*groups):
        out=[]; seen=set()
        for group in groups:
            for row in group or []:
                key=row.movement_key
                if key in seen: continue
                seen.add(key); out.append(row)
        return out

    def _wb_active_hours(rows):
        bins=set()
        for row in rows:
            if row.weigh_at:
                bins.add(aware(row.weigh_at).astimezone(TZ).strftime('%Y-%m-%d %H'))
        return len(bins)

    def _plant_lead_stats(rows):
        active_hours=_wb_active_hours(rows)
        resolved=0; missing=0; qty=0.0; ton_km=0.0; trip_km=0.0
        for row in rows:
            canon=wb_canonical_all.get(row.movement_key)
            src_id=canon.source_location_id if canon else None
            dst_id=canon.destination_location_id if canon else None
            bench=_tiom_bench_rl_from_location(db,src_id) if src_id else None
            if bench is None and src_id:
                bench=source_contexts.get((row.operating_date,row.shift,src_id))
            lead=dashboard_lead(src_id,bench,dst_id,'WITH_WB',wb_linked=True)
            if lead.get('leadKm') is None:
                missing+=1; continue
            lk=float(lead['leadKm']); q=tonnes(row)
            resolved+=1; qty+=q; ton_km+=q*lk; trip_km+=lk
        avg_lead=(ton_km/qty) if qty>0 else None
        return {
            'activeHours':active_hours,
            'tripsPerHour':round(len(rows)/active_hours,2) if active_hours else None,
            'avgLeadKm':round(avg_lead,3) if avg_lead is not None else None,
            'tripKmPerHour':round(trip_km/active_hours,2) if active_hours and resolved else None,
            'tonKmPerHour':round(ton_km/active_hours,1) if active_hours and resolved else None,
            'leadResolvedTrips':resolved,'leadMissingTrips':missing,
        }

    msp_feed_rows=_unique_wb_rows(production_fact_rows.get('MSP_FEED',[]))
    screen_final_rows=_unique_wb_rows(production_fact_rows.get('SCREEN_FINES',[]),production_fact_rows.get('SCREEN_5_18',[]))
    msp_transfer_rows=_unique_wb_rows(production_fact_rows.get('LUMPS_FEED_TO_CRUSHER',[]),production_fact_rows.get('SPILLAGE_TO_CRUSHER',[]))
    msp_recycle_rows=_unique_wb_rows(production_fact_rows.get('SPILLAGE_RECYCLE_MSP',[]))
    msp_wip_rows=_unique_wb_rows(production_fact_rows.get('SPILLAGE_WIP',[]))
    screen_output_rows=_unique_wb_rows(screen_final_rows,msp_transfer_rows,msp_wip_rows)
    crusher_feed_rows=_unique_wb_rows(production_fact_rows.get('CRUSHER_FEED',[]))
    crusher_blend_rows=_unique_wb_rows(production_fact_rows.get('CRUSHER_BLEND_FEED',[]))
    crusher_output_rows=_unique_wb_rows(production_fact_rows.get('CRUSHER_FINES',[]),production_fact_rows.get('CRUSHER_5_18',[]))

    def _msp_plant(row,side):
        txt=_wb_side_text(row,side)
        match=re.search(r'\bMSP\s*(?:PLANT\s*)?-?\s*(\d+)\b',txt)
        return f'MSP-{match.group(1)}' if match else None

    # Process-plant hours come from the central SiteAssetMeter, not the
    # paused attendance workflow. Link plant/Crusher equipment by an explicit
    # shift deployment location, falling back to an unambiguous equipment label.
    plant_meter_query=select(SiteAssetMeter).where(
        SiteAssetMeter.site_id=='TIOM',SiteAssetMeter.meter_type=='HMR',
        SiteAssetMeter.operating_date>=start_day,SiteAssetMeter.operating_date<=end_day)
    plant_deploy_query=select(TiomSourceDeployment).where(
        TiomSourceDeployment.operating_date>=start_day,
        TiomSourceDeployment.operating_date<=end_day,TiomSourceDeployment.active.is_(True))
    if selected_shift!='ALL':
        plant_meter_query=plant_meter_query.where(SiteAssetMeter.shift==selected_shift)
        plant_deploy_query=plant_deploy_query.where(TiomSourceDeployment.shift==selected_shift)
    elif allowed_shifts:
        plant_meter_query=plant_meter_query.where(SiteAssetMeter.shift.in_(allowed_shifts))
        plant_deploy_query=plant_deploy_query.where(TiomSourceDeployment.shift.in_(allowed_shifts))
    plant_meters=list(db.scalars(plant_meter_query))
    plant_locations={x.location_id:x for x in db.scalars(select(Location))}
    plant_deployments={}
    for assigned in db.scalars(plant_deploy_query):
        plant_deployments.setdefault((assigned.operating_date,assigned.shift,assigned.machine_id),set()).add(
            assigned.source_location_id)
    def _plant_run_hours(label):
        by_shift={}
        for meter in plant_meters:
            equipment=equipment_map.get(meter.asset_id)
            if not equipment:continue
            locations=plant_deployments.get((meter.operating_date,meter.shift,meter.asset_id),set())
            labels=[(plant_locations[src].location_name+' '+src) if src in plant_locations else src for src in locations]
            identity=' '.join([str(equipment.machine_id),str(equipment.type or ''),str(equipment.make_model or '')]).upper()
            joined=' '.join(labels).upper()
            matched=False
            if label.startswith('MSP-'):
                no=re.escape(label.split('-',1)[1])
                expression=rf'\bMSP\s*(?:PLANT\s*)?-?\s*{no}\b'
                matched=bool(re.search(expression,joined))
                if not locations and equipment.group=='PROCESSING':
                    matched=bool(re.search(expression,identity))
            elif label=='Crusher':
                matched=('CRUSH' in joined or 'OCP' in joined)
                if not locations and equipment.group=='PROCESSING':
                    matched=('CRUSH' in identity or 'OCP' in identity)
            if not matched:continue
            usage=(float(meter.usage) if meter.usage is not None else
                   (float(meter.closing_reading-meter.opening_reading)
                    if meter.closing_reading is not None and meter.opening_reading is not None
                       and meter.closing_reading>=meter.opening_reading else None))
            if usage is None or usage<0:continue
            key=(meter.operating_date,meter.shift)
            by_shift[key]=max(by_shift.get(key,0.0),usage)
        # Parallel components do not double-count a shift's running hours.
        # Across shifts/days, hours are added rather than taking one global max.
        return sum(by_shift.values()) if by_shift else None

    def _group_by_msp(rows,side):
        grouped={}
        for row in rows:
            label=_msp_plant(row,side) or 'MSP / Unmapped'
            grouped.setdefault(label,[]).append(row)
        return grouped

    feed_by_msp=_group_by_msp(msp_feed_rows,'DESTINATION')
    final_by_msp=_group_by_msp(screen_final_rows,'SOURCE')
    transfer_by_msp=_group_by_msp(msp_transfer_rows,'SOURCE')
    recycle_by_msp=_group_by_msp(msp_recycle_rows,'DESTINATION')
    wip_by_msp=_group_by_msp(msp_wip_rows,'SOURCE')
    msp_labels=sorted(set(feed_by_msp)|set(final_by_msp)|set(transfer_by_msp)|set(recycle_by_msp)|set(wip_by_msp),
                      key=lambda x:(x=='MSP / Unmapped',x))
    plant_performance=[]
    for label in msp_labels:
        feed_rows=_unique_wb_rows(feed_by_msp.get(label,[]))
        final_rows=_unique_wb_rows(final_by_msp.get(label,[]))
        transfer_rows=_unique_wb_rows(transfer_by_msp.get(label,[]))
        recycle_rows=_unique_wb_rows(recycle_by_msp.get(label,[]))
        wip_rows=_unique_wb_rows(wip_by_msp.get(label,[]))
        output_rows=_unique_wb_rows(final_rows,transfer_rows,wip_rows)
        feed_mt=sum(tonnes(r) for r in feed_rows)
        recycle_mt=sum(tonnes(r) for r in recycle_rows)
        gross_throughput=feed_mt+recycle_mt
        final_mt=sum(tonnes(r) for r in final_rows)
        transfer_mt=sum(tonnes(r) for r in transfer_rows)
        wip_mt=sum(tonnes(r) for r in wip_rows)
        accounted_mt=final_mt+transfer_mt+wip_mt
        run_hours=_plant_run_hours(label)
        tph=round(gross_throughput/run_hours,2) if run_hours and run_hours>0 else None
        recovery=round(final_mt/feed_mt*100,2) if feed_mt>0 else None
        plant_performance.append({
            'plant':label,'plantType':'MSP',
            'feedMt':round(feed_mt,2),'freshFeedMt':round(feed_mt,2),'feedTrips':len(feed_rows),
            'recycleMt':round(recycle_mt,2),'blendFeedMt':0.0,
            'grossThroughputMt':round(gross_throughput,2),
            'directFinalMt':round(final_mt,2),'outputMt':round(final_mt,2),'grossOutputMt':round(accounted_mt,2),
            'transferToCrusherMt':round(transfer_mt,2),'wipMt':round(wip_mt,2),'outputTrips':len(output_rows),
            'varianceMt':round(feed_mt-accounted_mt,2),
            'recoveryPct':recovery,'outputVsFeedPct':round(accounted_mt/feed_mt*100,2) if feed_mt>0 else None,
            'finalYieldPct':recovery,
            'runningHours':round(run_hours,2) if run_hours is not None else None,
            'tph':tph,'feedTph':tph,'outputTph':None,
            'runningHoursSource':'EQUIPMENT_HMR' if run_hours is not None else None,
            'avgLeadKm':None,'tripKmPerHour':None,'tonKmPerHour':None,'leadMissingTrips':0,
        })

    crusher_feed=crusher_feed_fact
    crusher_fresh_feed=max(Decimal('0'),crusher_feed-crusher_blend_feed)
    crusher_recovery=(float(crusher_final/crusher_feed*Decimal('100'))) if crusher_feed>0 else None
    crusher_run_hours=_plant_run_hours('Crusher')
    crusher_feed_tph=round(float(crusher_feed)/crusher_run_hours,2) if crusher_run_hours and crusher_run_hours>0 else None
    crusher_recovery_value=round(crusher_recovery,2) if crusher_recovery is not None else None
    plant_performance.append({
        'plant':'Crusher','plantType':'CRUSHER',
        'feedMt':round(float(crusher_fresh_feed),2),'freshFeedMt':round(float(crusher_fresh_feed),2),'feedTrips':len(crusher_feed_rows),
        'recycleMt':0.0,'blendFeedMt':round(float(crusher_blend_feed),2),'blendTrips':len(crusher_blend_rows),
        'grossThroughputMt':round(float(crusher_feed),2),
        'directFinalMt':round(float(crusher_fresh_final),2),'outputMt':round(float(crusher_fresh_final),2),
        'grossOutputMt':round(float(crusher_final),2),'blendOutputMt':round(float(crusher_blend_output),2),
        'transferToCrusherMt':0.0,'wipMt':0.0,'outputTrips':len(crusher_output_rows),
        'varianceMt':round(float(crusher_feed-crusher_final),2),
        'recoveryPct':crusher_recovery_value,'outputVsFeedPct':crusher_recovery_value,
        'finalYieldPct':round(float(crusher_fresh_final/crusher_fresh_feed*Decimal('100')),2) if crusher_fresh_feed>0 else None,
        'runningHours':round(crusher_run_hours,2) if crusher_run_hours is not None else None,
        'tph':crusher_feed_tph,'feedTph':crusher_feed_tph,'outputTph':None,
        'runningHoursSource':'EQUIPMENT_HMR' if crusher_run_hours is not None else None,
        'avgLeadKm':None,'tripKmPerHour':None,'tonKmPerHour':None,'leadMissingTrips':0,
    })

    haulage_perf={}
    all_haul_hours=set()
    def add_haulage(source_id,dest_id,source_label,dest_label,route_mode,qty,stamp,
                    lead_km=None,lead_status='NOT_CONFIGURED',cycle_min=None):
        mode=str(route_mode or 'UNSPECIFIED').upper()
        key=(str(source_id or source_label or 'UNMAPPED'),str(dest_id or dest_label or 'UNMAPPED'),mode)
        row=haulage_perf.setdefault(key,{
            'source':source_label or source_id or 'UNMAPPED',
            'destination':dest_label or dest_id or 'UNMAPPED',
            'sourceId':source_id or '','destinationId':dest_id or '',
            'routeMode':mode,'trips':0,'tonnes':0.0,'hourBins':set(),
            'leadResolvedTrips':0,'leadMissingTrips':0,'leadQty':0.0,
            'leadTonKm':0.0,'leadTripKm':0.0,'cycleSamples':[],'leadStatuses':set()
        })
        row['trips']+=1; row['tonnes']+=float(qty or 0)
        if stamp:
            local_stamp=aware(stamp).astimezone(TZ)
            hour_key=local_stamp.strftime('%Y-%m-%d %H')
            row['hourBins'].add(hour_key); all_haul_hours.add(hour_key)
        if lead_km is not None:
            lk=float(lead_km); q=float(qty or 0)
            row['leadResolvedTrips']+=1; row['leadQty']+=q
            row['leadTonKm']+=q*lk; row['leadTripKm']+=lk
        else:
            row['leadMissingTrips']+=1
        row['leadStatuses'].add(str(lead_status or 'NOT_CONFIGURED'))
        if cycle_min is not None:
            row['cycleSamples'].append(float(cycle_min))

    hourly = {h: {'hour': f'{h:02d}:00', 'trips': 0, 'tonnes': 0.0} for h in range(24)}
    materials, sources, destinations, routes, vehicles = {}, {}, {}, {}, {}
    material_buckets = {'ROM':0.0,'FINES':0.0,'CLO':0.0,'REJECT':0.0,'WASTE':0.0,'OTHER':0.0}
    crusher_rows, screen_rows = {}, {}
    for w in wb:
        t = tonnes(w); label = material_label(w)
        stamp = aware(w.weigh_at).astimezone(TZ) if w.weigh_at else None
        if stamp:
            hourly[stamp.hour]['trips'] += 1; hourly[stamp.hour]['tonnes'] += t
        add_metric(materials, label, t)
        canon=wb_canonical.get(w.movement_key)
        src_id=canon.source_location_id if canon else None; dst_id=canon.destination_location_id if canon else None
        src_obj=location_map.get(src_id) if src_id else None; dst_obj=location_map.get(dst_id) if dst_id else None
        src=src_obj.location_name if src_obj else 'UNMAPPED'; dst=dst_obj.location_name if dst_obj else 'UNMAPPED'
        add_metric(sources, src, t); add_metric(destinations, dst, t)
        route_key = f'{src_id or "UNMAPPED"} → {dst_id or "UNMAPPED"}'; rr = add_metric(routes, route_key, t); rr['source']=src; rr['destination']=dst; rr.setdefault('cycleSamples',[])
        bench_rl=_tiom_bench_rl_from_location(db,src_id) if src_id else None
        if bench_rl is None and src_id:
            bench_rl=source_contexts.get((w.operating_date,w.shift,src_id))
        lead_result=dashboard_lead(src_id,bench_rl,dst_id,'WITH_WB',wb_linked=True)
        add_haulage(src_id,dst_id,src,dst,'WITH_WB',t,stamp,lead_result.get('leadKm'),lead_result.get('status'))
        vk = (w.vehicle_id or w.vehicle_raw or 'Unknown').strip(); vr = add_metric(vehicles, vk, t); vr['vehicle']=vk
        u = label.upper()
        if 'WASTE' in u or 'OVERBURDEN' in u or re.search(r'(^|[^A-Z])OB([^A-Z]|$)',u): material_buckets['WASTE'] += t
        elif 'REJECT' in u: material_buckets['REJECT'] += t
        elif 'CLO' in u: material_buckets['CLO'] += t
        elif 'FINE' in u: material_buckets['FINES'] += t
        elif 'ROM' in u: material_buckets['ROM'] += t
        else: material_buckets['OTHER'] += t
        if has_token(src,'CRUSH','OCP') or has_token(dst,'CRUSH','OCP'):
            machine = dst if has_token(dst,'CRUSH','OCP') else src
            key=(machine,label); row=crusher_rows.setdefault(key,{'machine':machine,'material':label,'trips':0,'tonnes':0.0})
            row['trips']+=1;row['tonnes']+=t
        if has_token(src,'SCREEN','MSP-','MSP ') or has_token(dst,'SCREEN','MSP-','MSP '):
            machine = dst if has_token(dst,'SCREEN','MSP-','MSP ') else src
            key=(machine,label); row=screen_rows.setdefault(key,{'machine':machine,'material':label,'trips':0,'tonnes':0.0})
            row['trips']+=1;row['tonnes']+=t

    # Month-to-date material cumulative values and monthly management summary.
    month_start = end_day.replace(day=1)
    _, month_wb_all = authoritative_wb(month_start, end_day)
    month_material, month_days, month_production_days = {}, {}, {}
    month_process_facts,month_process_trips={},{}
    month_day_facts={}
    month_excluded_keys=set();month_old_stock_excluded=Decimal('0')
    final_codes={'SCREEN_FINES','SCREEN_5_18','CRUSHER_FINES','CRUSHER_5_18'}
    for w in month_wb_all:
        label = material_label(w); month_material[label] = month_material.get(label,0.0) + tonnes(w)
        key=str(w.operating_date); month_days[key]=month_days.get(key,0.0)+tonnes(w)
        facts=tiom_wb_report_contributions(w)
        fact_codes={code for code,_qty in facts}
        if fact_codes.intersection(excluded_codes) and w.movement_key not in month_excluded_keys:
            month_excluded_keys.add(w.movement_key)
            month_old_stock_excluded+=(w.net_kg or Decimal('0'))/Decimal('1000')
        day_store=month_day_facts.setdefault(key,{})
        seen_month=set()
        for code,qty in facts:
            q=Decimal(qty or 0)
            month_process_facts[code]=month_process_facts.get(code,Decimal('0'))+q
            day_store[code]=day_store.get(code,Decimal('0'))+q
            if code not in seen_month:
                month_process_trips[code]=month_process_trips.get(code,0)+1;seen_month.add(code)
    for key,facts in month_day_facts.items():
        screen_day=facts.get('SCREEN_FINES',Decimal('0'))+facts.get('SCREEN_5_18',Decimal('0'))
        crusher_gross_day=facts.get('CRUSHER_FINES',Decimal('0'))+facts.get('CRUSHER_5_18',Decimal('0'))
        crusher_fresh_day,_blend_day=allocate_blended_crusher_output(
            crusher_gross_day,facts.get('CRUSHER_FEED',Decimal('0')),facts.get('CRUSHER_BLEND_FEED',Decimal('0'))
        )
        month_production_days[key]=float(screen_day+crusher_fresh_day)
    for row in materials.values():
        row['monthTonnes'] = month_material.get(row['label'],0.0)

    # Route cycle time from exact WB↔field matches.
    for t,w in exact_pairs:
        canon=wb_canonical.get(w.movement_key)
        rk=f'{canon.source_location_id if canon and canon.source_location_id else "UNMAPPED"} → {canon.destination_location_id if canon and canon.destination_location_id else "UNMAPPED"}'
        cycle=duration_minutes(t.loading_start_at,t.unload_at)
        if rk in routes and cycle is not None: routes[rk]['cycleSamples'].append(cycle)
        if canon and cycle is not None:
            hk=(str(canon.source_location_id or 'UNMAPPED'),str(canon.destination_location_id or 'UNMAPPED'),'WITH_WB')
            if hk in haulage_perf:
                haulage_perf[hk]['cycleSamples'].append(float(cycle))

    # Exact-match loader/excavator/material attribution.
    machine_prod, operator_prod, matrix = {}, {}, {}
    for t,w in exact_pairs:
        amount=tonnes(w); mid=t.machine_id; eq=equipment_map.get(mid); typ=(eq.type if eq else 'LOADING') or 'LOADING'
        row=machine_prod.setdefault(mid,{'machine':mid,'type':typ,'trips':0,'tonnes':0.0,'loadMinutes':[],'cycleMinutes':[]})
        row['trips']+=1;row['tonnes']+=amount
        lm=duration_minutes(t.loading_start_at,t.loading_end_at); cm=duration_minutes(t.loading_start_at,t.unload_at)
        if lm is not None: row['loadMinutes'].append(lm)
        if cm is not None: row['cycleMinutes'].append(cm)
        if t.machine_operator_id:
            op=operator_prod.setdefault(t.machine_operator_id,{'operator':t.machine_operator_id,'trips':0,'tonnes':0.0})
            op['trips']+=1;op['tonnes']+=amount
        label=material_label(w); cell=matrix.setdefault(mid,{}).setdefault(label,{'trips':0,'tonnes':0.0});cell['trips']+=1;cell['tonnes']+=amount

    # Vehicle field cycle/loading and fuel attribution.
    vehicle_cycles, vehicle_loads = {}, {}
    for t in trips:
        if t.vehicle_id:
            cm=duration_minutes(t.loading_start_at,t.unload_at); lm=duration_minutes(t.loading_start_at,t.loading_end_at)
            if cm is not None: vehicle_cycles.setdefault(t.vehicle_id,[]).append(cm)
            if lm is not None: vehicle_loads.setdefault(t.vehicle_id,[]).append(lm)

    fuel_by_machine = {}
    for x in hsd:
        row=fuel_by_machine.setdefault(x.machine_id,{'machine':x.machine_id,'litres':0.0,'cost':0.0})
        row['litres'] += float(x.litres or 0); row['cost'] += float(x.amount or 0)

    # Latest equipment state in selected period for status/productivity KPIs.
    latest_att = {}
    for r in sorted(ea, key=lambda x:(x.operating_date, x.id or 0)):
        latest_att[r.machine_id]=r
    last_day_trip_ids = set()
    for t in trips:
        if t.operating_date == end_day:
            if t.machine_id: last_day_trip_ids.add(t.machine_id)
            if t.vehicle_id: last_day_trip_ids.add(t.vehicle_id)
    status_counts={'Running':0,'Idle':0,'Maintenance':0,'Fueling':0,'Waiting':0,'Breakdown':0}
    running_loaders=running_excavators=running_tippers=active_crushers=active_screens=0
    idle_equipment=0
    for mid,r in latest_att.items():
        if r.status != 'PRESENT': continue
        eq=equipment_map.get(mid); typ=((eq.type if eq else '') or '').upper(); grp=((eq.group if eq else '') or '').upper(); cond=(r.condition or '').upper()
        if 'BREAK' in cond: state='Breakdown'
        elif 'MAINT' in cond: state='Maintenance'
        elif 'FUEL' in cond: state='Fueling'
        elif 'WAIT' in cond: state='Waiting'
        elif cond=='WORKING' and mid in last_day_trip_ids: state='Running'
        elif cond=='WORKING': state='Idle'
        else: state='Idle'
        status_counts[state]+=1
        if state=='Idle': idle_equipment+=1
        if cond=='WORKING' and 'CRUSH' in typ: active_crushers+=1
        if cond=='WORKING' and ('SCREEN' in typ or 'SCREEN' in grp): active_screens+=1
        if state=='Running':
            if 'EXCAVATOR' in typ: running_excavators+=1
            elif 'LOADER' in typ or grp=='LOADING': running_loaders+=1
            if grp=='TRANSPORT' or 'TIPPER' in typ or 'DUMPER' in typ: running_tippers+=1

    production_equipment_present=sum(status_counts.values())
    equipment_util=(status_counts['Running']/production_equipment_present*100) if production_equipment_present else 0.0

    machine_rows=[]
    for mid,row in machine_prod.items():
        load_avg=sum(row['loadMinutes'])/len(row['loadMinutes']) if row['loadMinutes'] else None
        cycle_avg=sum(row['cycleMinutes'])/len(row['cycleMinutes']) if row['cycleMinutes'] else None
        machine_rows.append({'machine':mid,'type':row['type'],'trips':row['trips'],'tonnes':round(row['tonnes'],2),
                             'avgLoadingMin':round(load_avg,1) if load_avg is not None else None,
                             'avgCycleMin':round(cycle_avg,1) if cycle_avg is not None else None,
                             'idleMin':None,'avgBucketCount':None})
    machine_rows.sort(key=lambda x:x['tonnes'],reverse=True)
    loader_rows=[x for x in machine_rows if 'EXCAVATOR' not in str(x['type']).upper()]
    excavator_rows=[x for x in machine_rows if 'EXCAVATOR' in str(x['type']).upper()]

    vehicle_rows=[]
    for vk,row in vehicles.items():
        cycles=vehicle_cycles.get(vk,[]); fuel=fuel_by_machine.get(vk,{}).get('litres',0.0)
        att=latest_att.get(vk); cond=(att.condition or '').upper() if att else ''
        if 'BREAK' in cond: state='Breakdown'
        elif 'MAINT' in cond: state='Maintenance'
        elif 'FUEL' in cond: state='Fueling'
        elif vk in last_day_trip_ids: state='Running'
        else: state='Idle'
        vehicle_rows.append({'vehicle':vk,'trips':row['trips'],'tonnes':round(row['tonnes'],2),
                             'avgPayload':round(row['tonnes']/row['trips'],2) if row['trips'] else 0,
                             'fuel':round(fuel,1),'cycleTime':round(sum(cycles)/len(cycles),1) if cycles else None,'status':state})
    vehicle_rows.sort(key=lambda x:x['tonnes'],reverse=True)

    # Fuel breakdown with matched-tonnage efficiency when possible.
    tonnes_by_machine={r['machine']:r['tonnes'] for r in machine_rows}
    tonnes_by_vehicle={r['vehicle']:r['tonnes'] for r in vehicle_rows}
    fuel_rows=[]
    for mid,row in fuel_by_machine.items():
        eq=equipment_map.get(mid); typ=(eq.type or eq.group or '') if eq else ''
        ton=tonnes_by_machine.get(mid,tonnes_by_vehicle.get(mid,0.0))
        fuel_rows.append({'machine':mid,'type':typ or 'Other','litres':round(row['litres'],1),'tonnes':round(ton,2),
                          'litresPerTonne':round(row['litres']/ton,2) if ton else None})
    fuel_rows.sort(key=lambda x:x['litres'],reverse=True)

    # Shift comparison for the selected period.
    shift_comp={}
    for w in wb:
        r=shift_comp.setdefault(w.shift,{'shift':w.shift,'trips':0,'tonnes':0.0,'fuel':0.0})
        r['trips']+=1;r['tonnes']+=tonnes(w)
    for x in hsd:
        r=shift_comp.setdefault(x.shift,{'shift':x.shift,'trips':0,'tonnes':0.0,'fuel':0.0});r['fuel']+=float(x.litres or 0)
    shift_comp_rows=[]
    for r in shift_comp.values():
        r['avgPayload']=round(r['tonnes']/r['trips'],2) if r['trips'] else 0;r['tonnes']=round(r['tonnes'],2);r['fuel']=round(r['fuel'],1);shift_comp_rows.append(r)
    shift_comp_rows.sort(key=lambda x:x['shift'])

    # Always provide a 7-day management trend ending on the selected end date.
    trend_start=end_day-timedelta(days=6)
    _, trend_wb=authoritative_wb(trend_start,end_day)
    trend_hsd=scoped_range(HsdIssue,trend_start,end_day)
    seven={str(trend_start+timedelta(days=i)):{'date':str(trend_start+timedelta(days=i)),'trips':0,'tonnes':0.0,'fuel':0.0,'romInput':0.0,'finalProduction':0.0,'processFacts':{}} for i in range(7)}
    for w in trend_wb:
        r=seven[str(w.operating_date)];r['trips']+=1;r['tonnes']+=tonnes(w)
        for code,qty in tiom_wb_report_contributions(w):
            q=Decimal(qty or 0)
            facts=r['processFacts'];facts[code]=facts.get(code,Decimal('0'))+q
            if code == 'ROM':
                r['romInput']+=float(q)
    for r in seven.values():
        facts=r['processFacts']
        screen_day=facts.get('SCREEN_FINES',Decimal('0'))+facts.get('SCREEN_5_18',Decimal('0'))
        crusher_gross_day=facts.get('CRUSHER_FINES',Decimal('0'))+facts.get('CRUSHER_5_18',Decimal('0'))
        crusher_fresh_day,_blend_day=allocate_blended_crusher_output(
            crusher_gross_day,facts.get('CRUSHER_FEED',Decimal('0')),facts.get('CRUSHER_BLEND_FEED',Decimal('0'))
        )
        r['finalProduction']=float(screen_day+crusher_fresh_day)
    for x in trend_hsd:
        if str(x.operating_date) in seven: seven[str(x.operating_date)]['fuel']+=float(x.litres or 0)
    seven_rows=[dict(r,tonnes=round(r['tonnes'],2),fuel=round(r['fuel'],1)) for r in seven.values()]

    # Daily and monthly summary.
    best_day=max(month_production_days.items(),key=lambda x:x[1]) if month_production_days else ('',0.0)
    worst_day=min(month_production_days.items(),key=lambda x:x[1]) if month_production_days else ('',0.0)
    month_actual=sum(month_production_days.values()); elapsed_days=(end_day-month_start).days+1
    monthly={'actual':round(month_actual,2),'bestDay':best_day[0],'bestTonnes':round(best_day[1],2),
             'worstDay':worst_day[0],'worstTonnes':round(worst_day[1],2),
             'avgDaily':round(month_actual/elapsed_days,2) if elapsed_days else 0.0,'elapsedDays':elapsed_days,
             'basis':'FINAL_PRODUCTION'}

    load_samples=[duration_minutes(t.loading_start_at,t.loading_end_at) for t in trips]
    load_samples=[x for x in load_samples if x is not None]
    cycle_samples=[duration_minutes(t.loading_start_at,t.unload_at) for t in trips]
    cycle_samples=[x for x in cycle_samples if x is not None]
    avg_loading=sum(load_samples)/len(load_samples) if load_samples else None
    avg_cycle=sum(cycle_samples)/len(cycle_samples) if cycle_samples else None

    ore_tonnes=wb_tonnes-material_buckets['WASTE']-material_buckets['REJECT']
    production_balance=rom_input-final_production
    production_summary={
        'romInputMt':round(float(rom_input),2),'romPct':100.0 if rom_input>0 else None,'romBasis':rom_basis,
        'plantFeedMt':round(float(msp_feed),2),'plantFeedPct':prod_pct(msp_feed),
        'screenFinesMt':round(float(screen_fines),2),'screenFinesPct':prod_pct(screen_fines),
        'screen518Mt':round(float(screen_518),2),'screen518Pct':prod_pct(screen_518),
        'screenLumpsMt':round(float(screen_lumps),2),'screenLumpsPct':prod_pct(screen_lumps),
        'spillageToCrusherMt':round(float(spillage_to_crusher),2),'spillageToCrusherPct':prod_pct(spillage_to_crusher),
        'spillageRecycleMt':round(float(spillage_recycle),2),'spillageRecyclePct':prod_pct(spillage_recycle),
        'spillageWipMt':round(float(spillage_wip),2),'spillageWipPct':prod_pct(spillage_wip),
        'screenIntermediateMt':round(float(screen_intermediate),2),'screenIntermediatePct':prod_pct(screen_intermediate),
        'screenDirectMt':round(float(screen_direct),2),'screenDirectPct':prod_pct(screen_direct),
        'screenAccountedMt':round(float(screen_accounted),2),'screenAccountedPct':prod_pct(screen_accounted),
        'crusherFinesMt':round(float(crusher_fines),2),'crusherFinesPct':prod_pct(crusher_fines),
        'crusherCloMt':round(float(crusher_clo),2),'crusherCloPct':prod_pct(crusher_clo),
        'crusherGrossOutputMt':round(float(crusher_final),2),'crusherGrossOutputPct':prod_pct(crusher_final),
        'crusherFreshOutputMt':None if allocation_pending else round(float(crusher_fresh_final),2),
        'crusherFreshOutputPct':None if allocation_pending else prod_pct(crusher_fresh_final),
        'allocationStatus':allocation_status,'allocationPending':allocation_pending,
        'crusherBlendOutputMt':round(float(crusher_blend_output),2),'crusherBlendOutputPct':prod_pct(crusher_blend_output),
        'finalProductionMt':None if allocation_pending else round(float(final_production),2),
        'finalRecoveryPct':None if allocation_pending else prod_pct(final_production),
        'provisionalOutputMt':round(float(final_production),2) if allocation_pending else None,
        'crusherFeedMt':round(float(crusher_feed),2),'crusherFeedPct':prod_pct(crusher_feed),
        'crusherFreshFeedMt':round(float(crusher_fresh_feed),2),'crusherFreshFeedPct':prod_pct(crusher_fresh_feed),
        'crusherBlendFeedMt':round(float(crusher_blend_feed),2),'crusherBlendFeedPct':prod_pct(crusher_blend_feed),
        'crusherRecoveryPct':round(crusher_recovery,2) if crusher_recovery is not None else None,
        'oldStockExcludedMt':round(float(old_stock_excluded),2),'oldStockTrips':old_stock_trips,
        'oldStockVsRomPct':prod_pct(old_stock_excluded),
        'balanceMt':None if allocation_pending else round(float(production_balance),2),
        'balancePct':None if allocation_pending else prod_pct(production_balance),
        'note':'Fresh ROM is the 100% mine-production base. Genuine MSP lumps and large spillage sent to Crusher are fresh intermediate feed. CLO 10-40 is company old-stock blend, not NMTPL production. Crusher gross output is allocated pro-rata between fresh feed and old-stock blend because the mixed final product cannot be traced by feed stream after crushing. Same-MSP spillage is recycle and increases throughput/TPH, not fresh production.'
    }
    fuel_per_tonne=float(hsd_litres)/wb_tonnes if wb_tonnes else 0.0
    fuel_per_trip=float(hsd_litres)/wb_trips if wb_trips else 0.0
    unique_vehicles=len(vehicles); unique_loaders=len(loader_rows); unique_excavators=len(excavator_rows)

    # Material flow uses recorded stages only; crusher/screen/stock are WB-location classifications.
    stock_trips=sum(1 for w in wb if has_token(w.destination_raw,'STOCK') or re.fullmatch(r'S\d+',str(w.destination_raw or '').strip().upper()))
    stock_tonnes=sum(tonnes(w) for w in wb if has_token(w.destination_raw,'STOCK') or re.fullmatch(r'S\d+',str(w.destination_raw or '').strip().upper()))
    crusher_trips=sum(1 for w in wb if has_token(w.source_raw,'CRUSH','OCP') or has_token(w.destination_raw,'CRUSH','OCP'))
    crusher_tonnes=sum(tonnes(w) for w in wb if has_token(w.source_raw,'CRUSH','OCP') or has_token(w.destination_raw,'CRUSH','OCP'))
    screen_trips=sum(1 for w in wb if has_token(w.source_raw,'SCREEN','MSP-','MSP ') or has_token(w.destination_raw,'SCREEN','MSP-','MSP '))
    screen_tonnes=sum(tonnes(w) for w in wb if has_token(w.source_raw,'SCREEN','MSP-','MSP ') or has_token(w.destination_raw,'SCREEN','MSP-','MSP '))
    excavator_trip_count=sum(r['trips'] for r in excavator_rows); excavator_tonnes=sum(r['tonnes'] for r in excavator_rows)
    loader_trip_count=sum(r['trips'] for r in loader_rows); loader_tonnes=sum(r['tonnes'] for r in loader_rows)
    material_flow=[
        {'stage':'Excavator','trips':excavator_trip_count,'tonnes':round(excavator_tonnes,2)},
        {'stage':'Loader','trips':loader_trip_count,'tonnes':round(loader_tonnes,2)},
        {'stage':'Tipper','trips':wb_trips,'tonnes':round(wb_tonnes,2)},
        {'stage':'WB','trips':wb_trips,'tonnes':round(wb_tonnes,2)},
        {'stage':'Crusher','trips':crusher_trips,'tonnes':round(crusher_tonnes,2)},
        {'stage':'Screen','trips':screen_trips,'tonnes':round(screen_tonnes,2)},
        {'stage':'Stock','trips':stock_trips,'tonnes':round(stock_tonnes,2)},
    ]

    # Heatmap shape: exact matched tonnes only.
    matrix_materials=sorted({mat for values in matrix.values() for mat in values})
    heatmap=[]
    for mid,values in matrix.items():
        heatmap.append({'machine':mid,'type':(equipment_map.get(mid).type if equipment_map.get(mid) else ''),
                        'values':{m:{'trips':v['trips'],'tonnes':round(v['tonnes'],2)} for m,v in values.items()}})
    heatmap.sort(key=lambda x:x['machine'])

    # Relative alerts: no site thresholds are fabricated.
    low_payload=sum(r['trips']>=2 and avg_payload>0 and r['avgPayload']<avg_payload*0.8 for r in vehicle_rows)
    crusher_down=sum(1 for mid,r in latest_att.items() if 'CRUSH' in (((equipment_map.get(mid).type if equipment_map.get(mid) else '') or '').upper()) and has_token(r.condition,'BREAK','MAINT'))
    fuel_eff_values=[x['litresPerTonne'] for x in fuel_rows if x['litresPerTonne'] is not None]
    fleet_fuel_avg=(sum(fuel_eff_values)/len(fuel_eff_values)) if fuel_eff_values else 0
    fuel_high=sum(x['litresPerTonne'] is not None and fleet_fuel_avg and x['litresPerTonne']>fleet_fuel_avg*1.3 for x in fuel_rows)
    loader_idle=sum(1 for mid,r in latest_att.items() if (r.condition or '').upper()=='WORKING' and mid not in last_day_trip_ids and
                    ('LOADER' in (((equipment_map.get(mid).type if equipment_map.get(mid) else '') or '').upper()) or 'EXCAVATOR' in (((equipment_map.get(mid).type if equipment_map.get(mid) else '') or '').upper())))
    no_activity_30=0
    if end_day==live_day:
        now=now_local(); shift_def=db.get(ShiftMaster, live_shift); shift_start=datetime.combine(live_day,shift_def.start_time,TZ) if shift_def else now
        last_activity={}
        for t in trips:
            if t.operating_date!=live_day: continue
            stamp=t.unload_at or t.loading_end_at or t.loading_start_at
            for mid in [t.machine_id,t.vehicle_id]:
                if mid and (mid not in last_activity or aware(stamp)>aware(last_activity[mid])): last_activity[mid]=stamp
        for mid,r in latest_att.items():
            if r.operating_date!=live_day or r.status!='PRESENT' or (r.condition or '').upper()!='WORKING': continue
            eq=equipment_map.get(mid)
            if not eq or eq.group not in {'LOADING','TRANSPORT'}: continue
            stamp=last_activity.get(mid,shift_start)
            if (now-aware(stamp)).total_seconds()>1800: no_activity_30+=1

    review_rows=sum(b.review_rows or 0 for b in active_batch.values())
    missing_out=sum(r.status=='PRESENT' and r.in_at is not None and r.out_at is None for r in pa)
    missing_closing=sum(r.status=='PRESENT' and r.closing_meter is None for r in ea)
    breakdown=sum(r.status=='PRESENT' and has_token(r.condition,'BREAK') for r in ea)
    exceptions=[
        {'label':'Low payload vehicles (relative)','value':low_payload,'severity':'warn'},
        {'label':'No recorded activity >30 min','value':no_activity_30,'severity':'warn'},
        {'label':'No WB match','value':trip_no_wb,'severity':'warn'},
        {'label':'Crusher down / maintenance','value':crusher_down,'severity':'bad'},
        {'label':'High fuel intensity (relative)','value':fuel_high,'severity':'warn'},
        {'label':'Loader / excavator no recorded activity','value':loader_idle,'severity':'warn'},
        {'label':'WB review rows','value':review_rows,'severity':'bad'},
        {'label':'Breakdown equipment-shifts','value':breakdown,'severity':'bad'},
        {'label':'Missing OUT','value':missing_out,'severity':'bad'},
        {'label':'Closing meter pending','value':missing_closing,'severity':'warn'},
        {'label':'Valid WB unmatched','value':wb_unmatched,'severity':'warn'},
    ]
    for x in exceptions:
        if not x['value']: x['severity']='ok'

    # Top/bottom performers use recorded trips/tonnes only.
    performer_pool=[]
    if loader_rows:
        top=max(loader_rows,key=lambda x:(x['tonnes'],x['trips']));low=min(loader_rows,key=lambda x:(x['trips'],x['tonnes']))
        performer_pool.append({'category':'Top Loader','name':top['machine'],'trips':top['trips'],'tonnes':top['tonnes']})
        bottom_loader={'category':'Lowest Loader','name':low['machine'],'trips':low['trips'],'tonnes':low['tonnes']}
    else: bottom_loader=None
    if excavator_rows:
        top=max(excavator_rows,key=lambda x:(x['tonnes'],x['trips']));low=min(excavator_rows,key=lambda x:(x['trips'],x['tonnes']))
        performer_pool.append({'category':'Top Excavator','name':top['machine'],'trips':top['trips'],'tonnes':top['tonnes']})
        bottom_exc={'category':'Lowest Excavator','name':low['machine'],'trips':low['trips'],'tonnes':low['tonnes']}
    else: bottom_exc=None
    if vehicle_rows:
        top=max(vehicle_rows,key=lambda x:(x['tonnes'],x['trips']));low=min(vehicle_rows,key=lambda x:(x['trips'],x['tonnes']))
        performer_pool.append({'category':'Top Vehicle','name':top['vehicle'],'trips':top['trips'],'tonnes':top['tonnes']})
        bottom_vehicle={'category':'Lowest Vehicle','name':low['vehicle'],'trips':low['trips'],'tonnes':low['tonnes']}
    else: bottom_vehicle=None
    op_rows=list(operator_prod.values())
    if op_rows:
        top=max(op_rows,key=lambda x:(x['tonnes'],x['trips']));performer_pool.append({'category':'Top Operator','name':top['operator'],'trips':top['trips'],'tonnes':round(top['tonnes'],2)})
    bottom_performers=[x for x in [bottom_loader,bottom_exc,bottom_vehicle] if x]

    match_rate=((len(exact_keys)+len(likely_keys))/wb_trips*100) if wb_trips else 0.0
    avg_trips_vehicle=wb_trips/unique_vehicles if unique_vehicles else 0.0
    avg_trips_loader=sum(x['trips'] for x in loader_rows)/unique_loaders if unique_loaders else 0.0
    avg_trips_exc=sum(x['trips'] for x in excavator_rows)/unique_excavators if unique_excavators else 0.0
    avg_tonnes_loader=sum(x['tonnes'] for x in loader_rows)/unique_loaders if unique_loaders else 0.0
    avg_tonnes_exc=sum(x['tonnes'] for x in excavator_rows)/unique_excavators if unique_excavators else 0.0

    mat_rows=[]
    for row in sorted(materials.values(),key=lambda x:x['tonnes'],reverse=True):
        mat_rows.append({'label':row['label'],'trips':row['trips'],'tonnes':round(row['tonnes'],2),'monthTonnes':round(row.get('monthTonnes',0),2),
                         'avgPayload':round(row['tonnes']/row['trips'],2) if row['trips'] else 0,
                         'pct':round(row['tonnes']/wb_tonnes*100,1) if wb_tonnes else 0})
    month_rom_input=month_process_facts.get('ROM',Decimal('0'))
    if month_rom_input<=0:
        month_rom_input=month_process_facts.get('MSP_FEED',Decimal('0'))
    month_screen_direct=month_process_facts.get('SCREEN_FINES',Decimal('0'))+month_process_facts.get('SCREEN_5_18',Decimal('0'))
    month_crusher_gross=month_process_facts.get('CRUSHER_FINES',Decimal('0'))+month_process_facts.get('CRUSHER_5_18',Decimal('0'))
    month_crusher_fresh,month_blend_output=allocate_blended_crusher_output(
        month_crusher_gross,
        month_process_facts.get('CRUSHER_FEED',Decimal('0')),
        month_process_facts.get('CRUSHER_BLEND_FEED',Decimal('0'))
    )
    month_final_production=month_screen_direct+month_crusher_fresh
    old_stock_blend_10_40=production_facts.get('OLD_STOCK_BLEND_10_40',Decimal('0'))
    month_old_stock_blend_10_40=month_process_facts.get('OLD_STOCK_BLEND_10_40',Decimal('0'))
    other_old_stock_excluded=max(Decimal('0'),old_stock_excluded-old_stock_blend_10_40)
    month_other_old_stock_excluded=max(Decimal('0'),month_old_stock_excluded-month_old_stock_blend_10_40)
    period_rom_trips=production_fact_trips.get('ROM',0)
    if not period_rom_trips and rom_basis=='MSP_FEED_FALLBACK':
        period_rom_trips=production_fact_trips.get('MSP_FEED',0)
    month_rom_trips=month_process_trips.get('ROM',0)
    if not month_rom_trips and month_rom_input>0:
        month_rom_trips=month_process_trips.get('MSP_FEED',0)

    def _prod_material_row(label,qty,trips,month_qty,month_trips,pct_value=None,kind='OUTPUT'):
        q=Decimal(qty or 0); mq=Decimal(month_qty or 0)
        return {
            'label':label,'kind':kind,'trips':int(trips or 0),'tonnes':round(float(q),2),
            'monthTonnes':round(float(mq),2),'avgPayload':round(float(q)/trips,2) if trips else 0,
            'monthTrips':int(month_trips or 0),
            'pct':pct_value if pct_value is not None else prod_pct(q),
            'excluded':kind=='EXCLUDED'
        }

    production_material_rows=[
        _prod_material_row('Fresh ROM',rom_input,period_rom_trips,month_rom_input,month_rom_trips,100.0 if rom_input>0 else None,'INPUT'),
        _prod_material_row('Screen Fines',screen_fines,production_fact_trips.get('SCREEN_FINES',0),month_process_facts.get('SCREEN_FINES',0),month_process_trips.get('SCREEN_FINES',0)),
        _prod_material_row('Screen 5-18',screen_518,production_fact_trips.get('SCREEN_5_18',0),month_process_facts.get('SCREEN_5_18',0),month_process_trips.get('SCREEN_5_18',0)),
        _prod_material_row('MSP Lumps → Crusher (Fresh Intermediate)',screen_lumps,production_fact_trips.get('LUMPS_FROM_SCREEN',0),month_process_facts.get('LUMPS_FROM_SCREEN',0),month_process_trips.get('LUMPS_FROM_SCREEN',0),prod_pct(screen_lumps),'INTERMEDIATE'),
        _prod_material_row('Spillage → Crusher (Fresh Intermediate)',spillage_to_crusher,production_fact_trips.get('SPILLAGE_TO_CRUSHER',0),month_process_facts.get('SPILLAGE_TO_CRUSHER',0),month_process_trips.get('SPILLAGE_TO_CRUSHER',0),prod_pct(spillage_to_crusher),'INTERMEDIATE'),
        _prod_material_row('Spillage Re-screen / Recycle',spillage_recycle,production_fact_trips.get('SPILLAGE_RECYCLE_MSP',0),month_process_facts.get('SPILLAGE_RECYCLE_MSP',0),month_process_trips.get('SPILLAGE_RECYCLE_MSP',0),prod_pct(spillage_recycle),'INTERMEDIATE'),
        _prod_material_row('Crusher Fines (Gross Output)',crusher_fines,production_fact_trips.get('CRUSHER_FINES',0),month_process_facts.get('CRUSHER_FINES',0),month_process_trips.get('CRUSHER_FINES',0)),
        _prod_material_row('Crusher CLO 5-18 (Gross Output)',crusher_clo,production_fact_trips.get('CRUSHER_5_18',0),month_process_facts.get('CRUSHER_5_18',0),month_process_trips.get('CRUSHER_5_18',0)),
        _prod_material_row('CLO 10-40 Old Stock Blend → Crusher',old_stock_blend_10_40,production_fact_trips.get('OLD_STOCK_BLEND_10_40',0),month_old_stock_blend_10_40,month_process_trips.get('OLD_STOCK_BLEND_10_40',0),prod_pct(old_stock_blend_10_40),'EXCLUDED'),
        _prod_material_row('Old-Stock Share Allocated in Crusher Output',crusher_blend_output,0,month_blend_output,0,prod_pct(crusher_blend_output),'EXCLUDED'),
        _prod_material_row('Fresh Crusher Output (Net of Blend)',crusher_fresh_final,
            production_fact_trips.get('CRUSHER_FINES',0)+production_fact_trips.get('CRUSHER_5_18',0),
            month_crusher_fresh,
            month_process_trips.get('CRUSHER_FINES',0)+month_process_trips.get('CRUSHER_5_18',0),
            prod_pct(crusher_fresh_final),'NET'),
        _prod_material_row('Final Fresh Production',final_production,
            production_fact_trips.get('SCREEN_FINES',0)+production_fact_trips.get('SCREEN_5_18',0)+production_fact_trips.get('CRUSHER_FINES',0)+production_fact_trips.get('CRUSHER_5_18',0),
            month_final_production,
            month_process_trips.get('SCREEN_FINES',0)+month_process_trips.get('SCREEN_5_18',0)+month_process_trips.get('CRUSHER_FINES',0)+month_process_trips.get('CRUSHER_5_18',0),
            prod_pct(final_production),'TOTAL'),
        _prod_material_row('Other RH / Old Stock Movement (Excluded)',other_old_stock_excluded,max(0,old_stock_trips-production_fact_trips.get('OLD_STOCK_BLEND_10_40',0)),
            month_other_old_stock_excluded,max(0,len(month_excluded_keys)-month_process_trips.get('OLD_STOCK_BLEND_10_40',0)),
            prod_pct(other_old_stock_excluded),'EXCLUDED')
    ]
    if allocation_pending:
        # Retain gross crusher product measurements but suppress unsupported
        # NET and TOTAL fresh production rows until feed evidence is complete.
        for item in production_material_rows:
            if item['kind'] in {'NET','TOTAL'}:
                item['tonnes']=None
                item['pct']=None
                item['status']='ALLOCATION_PENDING'
        for item in plant_performance:
            if item.get('plantType')=='CRUSHER':
                item['outputMt']=None
                item['finalYieldPct']=None
                item['allocationStatus']=allocation_status

    source_rows=[{'label':r['label'],'trips':r['trips'],'tonnes':round(r['tonnes'],2),'avgPayload':round(r['tonnes']/r['trips'],2) if r['trips'] else 0,
                  'pct':round(r['tonnes']/wb_tonnes*100,1) if wb_tonnes else 0} for r in sorted(sources.values(),key=lambda x:x['tonnes'],reverse=True)]
    dest_rows=[{'label':r['label'],'trips':r['trips'],'tonnes':round(r['tonnes'],2),'avgPayload':round(r['tonnes']/r['trips'],2) if r['trips'] else 0,
                'pct':round(r['tonnes']/wb_tonnes*100,1) if wb_tonnes else 0} for r in sorted(destinations.values(),key=lambda x:x['tonnes'],reverse=True)]
    route_rows=[]
    for r in sorted(routes.values(),key=lambda x:x['tonnes'],reverse=True):
        samples=r.get('cycleSamples',[])
        route_rows.append({'source':r['source'],'destination':r['destination'],'trips':r['trips'],'tonnes':round(r['tonnes'],2),
                           'avgPayload':round(r['tonnes']/r['trips'],2) if r['trips'] else 0,
                           'avgTime':round(sum(samples)/len(samples),1) if samples else None})

    # TIOM MIS operational quantities are reported separately from authoritative WB tonnes.
    # Submitted manual rows are exposed to the dashboard in their own operational
    # views. WB-linked rows remain visible here as MIS evidence but are NEVER added
    # again to authoritative WB production totals.
    mis_stmt=select(TiomMisReport).where(
        TiomMisReport.operating_date>=start_day,TiomMisReport.operating_date<=end_day,TiomMisReport.status=='SUBMITTED'
    )
    draft_stmt=select(TiomMisReport).where(
        TiomMisReport.operating_date>=start_day,TiomMisReport.operating_date<=end_day,TiomMisReport.status=='DRAFT'
    )
    if selected_shift!='ALL':
        mis_stmt=mis_stmt.where(TiomMisReport.shift==selected_shift)
        draft_stmt=draft_stmt.where(TiomMisReport.shift==selected_shift)
    elif allowed_shifts:
        mis_stmt=mis_stmt.where(TiomMisReport.shift.in_(allowed_shifts))
        draft_stmt=draft_stmt.where(TiomMisReport.shift.in_(allowed_shifts))
    mis_reports=list(db.scalars(mis_stmt))
    mis_drafts=list(db.scalars(draft_stmt))
    if vehicle_filter:
        mis_reports=[r for r in mis_reports if r.vehicle_id==vehicle_filter]
        mis_drafts=[r for r in mis_drafts if r.vehicle_id==vehicle_filter]
    mis_report_map={r.report_id:r for r in mis_reports}
    mis_ids=list(mis_report_map)
    mis_rows=list(db.scalars(select(TiomMisTripRow).where(TiomMisTripRow.report_id.in_(mis_ids)))) if mis_ids else []
    mis_row_ids=[r.row_id for r in mis_rows]
    mis_details={x.row_id:x for x in db.scalars(select(TiomMisTripDetail).where(TiomMisTripDetail.row_id.in_(mis_row_ids)))} if mis_row_ids else {}
    mis_recs={x.row_id:x for x in db.scalars(select(TiomMisReconciliation).where(TiomMisReconciliation.row_id.in_(mis_row_ids)))} if mis_row_ids else {}
    mis_leads={x.row_id:x for x in db.scalars(select(TiomMisTripLead).where(TiomMisTripLead.row_id.in_(mis_row_ids)))} if mis_row_ids else {}
    mis_products={x.product_id:x for x in db.scalars(select(Product))}
    mis_trip_count=len(mis_rows); mis_ob_trips=0; mis_ob_qty=Decimal('0'); mis_rom_trips=0; mis_rom_qty=Decimal('0')
    mis_total_qty=Decimal('0'); mis_wb_linked=0; mis_factor_trips=0
    lead_resolved_trips=0; lead_missing_trips=0; lead_with_wb=0; lead_without_wb=0
    lead_covered_qty=Decimal('0'); lead_ton_km=Decimal('0')
    lead_review=[]; qty_pending=[]
    mis_materials={}; mis_sources={}; mis_destinations={}; mis_vehicles={}; mis_machines={}; lead_routes={}
    for r in mis_rows:
        d=mis_details.get(r.row_id)
        if not d: continue
        report=mis_report_map.get(r.report_id)
        prod=mis_products.get(d.material_id) if d.material_id else None
        mat=(prod.name if prod else (d.material_id or r.material_raw or 'Unmapped')).strip()
        src_obj=location_map.get(d.source_location_id) if d.source_location_id else None
        dst_obj=location_map.get(d.destination_location_id) if d.destination_location_id else None
        src=(src_obj.location_name if src_obj else (d.source_location_id or r.source_raw or 'Unknown')).strip()
        dst=(dst_obj.location_name if dst_obj else (d.destination_location_id or r.destination_raw or 'Unknown')).strip()
        vehicle=(report.vehicle_id if report else 'Unknown') or 'Unknown'
        machine=d.machine_id or 'Unknown'
        qty=Decimal(d.calculated_qty_mt or 0)
        if d.calculated_qty_mt is None and not (mis_recs.get(r.row_id) and mis_recs[r.row_id].wb_movement_key):
            qty_pending.append({'rowId':r.row_id,'date':str(report.operating_date) if report else '',
                                'shift':report.shift if report else '','vehicle':vehicle,'material':mat,
                                'source':src,'destination':dst,'reason':'NO_WEIGHT_OR_APPROVED_FACTOR'})
        mis_total_qty+=qty
        add_metric(mis_materials,mat,float(qty))
        add_metric(mis_sources,src,float(qty))
        add_metric(mis_destinations,dst,float(qty))
        add_metric(mis_vehicles,vehicle,float(qty))
        add_metric(mis_machines,machine,float(qty))
        rec=mis_recs.get(r.row_id)
        lead=mis_leads.get(r.row_id)
        # Resolve from CURRENT approved masters without writing historical rows.
        # Old snapshots may be missing despite a valid area/bench lead master.
        bench_rl=(lead.bench_rl_m if lead else None)
        if bench_rl is None and d.source_location_id:
            bench_rl=_tiom_bench_rl_from_location(db,d.source_location_id)
        mode_for_lead='WITH_WB' if rec and rec.wb_movement_key else 'WITHOUT_WB'
        fresh_lead=dashboard_lead(
            d.source_location_id,bench_rl,d.destination_location_id,mode_for_lead,
            wb_linked=bool(rec and rec.wb_movement_key))
        resolved_km=(fresh_lead.get('leadKm') if fresh_lead.get('status')=='OK' else None)
        lead_status=fresh_lead.get('status') or 'ROUTE_NOT_CONFIGURED'
        if resolved_km is None:
            lead_review.append({'rowId':r.row_id,'date':str(report.operating_date) if report else '',
                                'shift':report.shift if report else '','vehicle':vehicle,
                                'source':src,'destination':dst,'benchRl':bench_rl,
                                'routeMode':mode_for_lead,'reason':lead_status})
        if rec and rec.wb_movement_key: mis_wb_linked+=1
        else: mis_factor_trips+=1
        route_mode=mode_for_lead
        # WB-linked MIS is evidence only; authoritative WB already contributed
        # that trip to haulage performance. Non-WB rows (OB/internal haulage)
        # join the same route KPI using their form timestamps and snapshotted lead.
        if not (rec and rec.wb_movement_key):
            cycle=duration_minutes(r.loading_at,r.unloading_at)
            add_haulage(
                d.source_location_id,d.destination_location_id,src,dst,route_mode,float(qty),
                r.loading_at or r.unloading_at,
                resolved_km,
                lead_status,
                cycle
            )
        if route_mode=='WITH_WB': lead_with_wb+=1
        elif route_mode=='WITHOUT_WB': lead_without_wb+=1
        if resolved_km is not None:
            lead_km=Decimal(resolved_km)
            ton_km=qty*lead_km
            lead_resolved_trips+=1; lead_covered_qty+=qty; lead_ton_km+=ton_km
            rk=(src,dst,route_mode or 'UNSPECIFIED')
            lr=lead_routes.setdefault(rk,{'source':src,'destination':dst,'routeMode':route_mode or 'UNSPECIFIED','trips':0,'tonnes':Decimal('0'),'leadKmSum':Decimal('0'),'tonKm':Decimal('0')})
            lr['trips']+=1; lr['tonnes']+=qty; lr['leadKmSum']+=lead_km; lr['tonKm']+=ton_km
        else:
            lead_missing_trips+=1
        txt=(' '.join([str(d.material_id or ''),str(prod.name if prod else r.material_raw or '')])).upper()
        if re.search(r'(^|[^A-Z0-9])OB([^A-Z0-9]|$)',txt): mis_ob_trips+=1; mis_ob_qty+=qty
        elif 'ROM' in txt: mis_rom_trips+=1; mis_rom_qty+=qty
    def _mis_rows(store,key_name='label'):
        out=[]
        for x in sorted(store.values(),key=lambda z:(z['tonnes'],z['trips']),reverse=True):
            out.append({key_name:x['label'],'trips':x['trips'],'tonnes':round(x['tonnes'],2),
                        'avgPayload':round(x['tonnes']/x['trips'],2) if x['trips'] else 0})
        return out

    lead_route_rows=[]
    for x in sorted(lead_routes.values(),key=lambda z:(z['tonKm'],z['trips']),reverse=True):
        avg_lead=(x['tonKm']/x['tonnes']) if x['tonnes']>0 else (x['leadKmSum']/x['trips'] if x['trips'] else Decimal('0'))
        lead_route_rows.append({
            'source':x['source'],'destination':x['destination'],'routeMode':x['routeMode'],
            'trips':x['trips'],'tonnes':round(float(x['tonnes']),2),
            'avgLeadKm':round(float(avg_lead),3),'tonKm':round(float(x['tonKm']),2)
        })
    weighted_avg_lead=(lead_ton_km/lead_covered_qty) if lead_covered_qty>0 else Decimal('0')

    # Combined haulage productivity: authoritative WB trips + non-WB form trips.
    # TPH here intentionally means Trips Per Hour, using active clock-hour buckets.
    haulage_rows=[]
    combined_lead_qty=0.0; combined_ton_km=0.0; combined_trip_km=0.0
    combined_trips=0; combined_lead_missing=0
    for row in haulage_perf.values():
        active_hours=len(row['hourBins'])
        avg_lead=(row['leadTonKm']/row['leadQty']) if row['leadQty']>0 else None
        tph=(row['trips']/active_hours) if active_hours>0 else None
        trip_kmh=(row['leadTripKm']/active_hours) if active_hours>0 and row['leadResolvedTrips']>0 else None
        ton_kmh=(row['leadTonKm']/active_hours) if active_hours>0 and row['leadResolvedTrips']>0 else None
        cycles=row['cycleSamples']
        if row['leadMissingTrips']==0 and row['leadResolvedTrips']>0:
            lead_status='OK'
        elif row['leadResolvedTrips']>0:
            lead_status='PARTIAL'
        else:
            lead_status=sorted(row['leadStatuses'])[0] if row['leadStatuses'] else 'NOT_CONFIGURED'
        haulage_rows.append({
            'source':row['source'],'destination':row['destination'],'routeMode':row['routeMode'],
            'trips':row['trips'],'tonnes':round(row['tonnes'],2),'activeHours':active_hours,
            'tripsPerHour':round(tph,2) if tph is not None else None,
            'avgLeadKm':round(avg_lead,3) if avg_lead is not None else None,
            'tripKmPerHour':round(trip_kmh,2) if trip_kmh is not None else None,
            'tonKmPerHour':round(ton_kmh,2) if ton_kmh is not None else None,
            'avgCycleMin':round(sum(cycles)/len(cycles),1) if cycles else None,
            'leadResolvedTrips':row['leadResolvedTrips'],'leadMissingTrips':row['leadMissingTrips'],
            'leadStatus':lead_status,
        })
        combined_trips+=row['trips']; combined_lead_qty+=row['leadQty']
        combined_ton_km+=row['leadTonKm']; combined_trip_km+=row['leadTripKm']
        combined_lead_missing+=row['leadMissingTrips']
    haulage_rows.sort(key=lambda x:(x['trips'],x['tonnes']),reverse=True)
    combined_active_hours=len(all_haul_hours)
    combined_avg_lead=(combined_ton_km/combined_lead_qty) if combined_lead_qty>0 else None
    combined_tph=(combined_trips/combined_active_hours) if combined_active_hours>0 else None
    combined_trip_kmh=(combined_trip_km/combined_active_hours) if combined_active_hours>0 else None
    combined_ton_kmh=(combined_ton_km/combined_active_hours) if combined_active_hours>0 else None

    # Build 7-day MIS/lead trend independently of the currently selected period.
    # This mirrors the existing 7-day WB trend and preserves the selected shift scope.
    trend_mis_stmt=select(TiomMisReport).where(
        TiomMisReport.operating_date>=trend_start,TiomMisReport.operating_date<=end_day,
        TiomMisReport.status=='SUBMITTED'
    )
    if selected_shift!='ALL': trend_mis_stmt=trend_mis_stmt.where(TiomMisReport.shift==selected_shift)
    elif allowed_shifts: trend_mis_stmt=trend_mis_stmt.where(TiomMisReport.shift.in_(allowed_shifts))
    trend_reports=list(db.scalars(trend_mis_stmt))
    trend_report_map={x.report_id:x for x in trend_reports}
    trend_ids=list(trend_report_map)
    trend_rows=list(db.scalars(select(TiomMisTripRow).where(TiomMisTripRow.report_id.in_(trend_ids)))) if trend_ids else []
    trend_row_ids=[x.row_id for x in trend_rows]
    trend_details={x.row_id:x for x in db.scalars(select(TiomMisTripDetail).where(TiomMisTripDetail.row_id.in_(trend_row_ids)))} if trend_row_ids else {}
    trend_leads={x.row_id:x for x in db.scalars(select(TiomMisTripLead).where(TiomMisTripLead.row_id.in_(trend_row_ids)))} if trend_row_ids else {}
    for row in trend_rows:
        report=trend_report_map.get(row.report_id); detail=trend_details.get(row.row_id)
        if not report or not detail: continue
        key=str(report.operating_date)
        if key not in seven: continue
        qty=Decimal(detail.calculated_qty_mt or 0)
        seven[key]['misQty']=seven[key].get('misQty',0.0)+float(qty)
        lead=trend_leads.get(row.row_id)
        if lead and lead.lead_status=='OK' and lead.lead_km is not None:
            seven[key]['leadTonKm']=seven[key].get('leadTonKm',0.0)+float(qty*Decimal(lead.lead_km))

    cumulative_trips=0; cumulative_tonnes=0.0; cumulative_fuel=0.0; cumulative_mis=0.0; cumulative_lead_ton_km=0.0
    seven_rows=[]
    for key in sorted(seven):
        row=seven[key]
        row.setdefault('misQty',0.0); row.setdefault('leadTonKm',0.0)
        cumulative_trips+=int(row.get('trips') or 0)
        cumulative_tonnes+=float(row.get('tonnes') or 0)
        cumulative_fuel+=float(row.get('fuel') or 0)
        cumulative_mis+=float(row.get('misQty') or 0)
        cumulative_lead_ton_km+=float(row.get('leadTonKm') or 0)
        seven_rows.append({
            'date':row['date'],'trips':row['trips'],'tonnes':round(row['tonnes'],2),'fuel':round(row['fuel'],1),
            'romInput':round(row.get('romInput',0.0),2),'finalProduction':round(row.get('finalProduction',0.0),2),
            'misQty':round(row['misQty'],2),'leadTonKm':round(row['leadTonKm'],2),
            'cumulativeTrips':cumulative_trips,'cumulativeTonnes':round(cumulative_tonnes,2),
            'cumulativeFuel':round(cumulative_fuel,1),'cumulativeMisQty':round(cumulative_mis,2),
            'cumulativeLeadTonKm':round(cumulative_lead_ton_km,2)
        })

    month_rows=[]; month_cumulative=0.0
    for offset in range((end_day-month_start).days+1):
        md=month_start+timedelta(days=offset); daily=float(month_production_days.get(str(md),0.0))
        month_cumulative+=daily
        month_rows.append({'date':str(md),'tonnes':round(daily,2),'cumulativeTonnes':round(month_cumulative,2)})

    # Shift Production report quantities: calculated from submitted MIS + submitted Plant/Shifting rows.
    prod_periods=set()
    sp_stmt=select(TiomShiftProductionReport).where(TiomShiftProductionReport.operating_date>=start_day,TiomShiftProductionReport.operating_date<=end_day,TiomShiftProductionReport.status=='SUBMITTED')
    if selected_shift!='ALL': sp_stmt=sp_stmt.where(TiomShiftProductionReport.shift==selected_shift)
    elif allowed_shifts: sp_stmt=sp_stmt.where(TiomShiftProductionReport.shift.in_(allowed_shifts))
    for r in db.scalars(sp_stmt): prod_periods.add((r.operating_date,r.shift))
    for r in mis_reports: prod_periods.add((r.operating_date,r.shift))
    wb_period_stmt=select(WbImportBatch).where(
        WbImportBatch.operating_date>=start_day,WbImportBatch.operating_date<=end_day,
        WbImportBatch.status=='CONFIRMED',WbImportBatch.confirmed_at.is_not(None)
    )
    if selected_shift!='ALL': wb_period_stmt=wb_period_stmt.where(WbImportBatch.shift==selected_shift)
    elif allowed_shifts: wb_period_stmt=wb_period_stmt.where(WbImportBatch.shift.in_(allowed_shifts))
    for r in db.scalars(wb_period_stmt): prod_periods.add((r.operating_date,r.shift))
    shift_excavation=Decimal('0'); shift_processed=Decimal('0')
    for pd,ps in prod_periods:
        f=_tiom_shift_ftd(db,pd,ps,False)
        shift_excavation+=f.get('TOTAL_EXCAVATION',Decimal('0')); shift_processed+=f.get('TOTAL_PRODUCTION',Decimal('0'))

    drilling=_tiom_drilling_summary(db,start_day,end_day,selected_shift)
    drill_totals=drilling.get('totals',{})

    wb_source_unmapped=sum(1 for x in wb_canonical.values() if x and not x.source_location_id)
    wb_destination_unmapped=sum(1 for x in wb_canonical.values() if x and not x.destination_location_id)
    exceptions.append({'label':'WB source unmapped to Location Master','value':wb_source_unmapped,'severity':'warn' if wb_source_unmapped else 'ok'})
    exceptions.append({'label':'WB destination unmapped to Location Master','value':wb_destination_unmapped,'severity':'warn' if wb_destination_unmapped else 'ok'})
    exceptions.append({'label':'MIS rows missing lead','value':lead_missing_trips,'severity':'warn' if lead_missing_trips else 'ok'})
    exceptions.append({'label':'MIS trips awaiting tonnage','value':len(qty_pending),'severity':'warn' if qty_pending else 'ok'})
    crusher_out=[]
    for r in sorted(crusher_rows.values(),key=lambda x:x['tonnes'],reverse=True):
        crusher_out.append(dict(r,tonnes=round(r['tonnes'],2),avgFeed=round(r['tonnes']/r['trips'],2) if r['trips'] else 0,utilization=None))
    screen_out=[]
    for r in sorted(screen_rows.values(),key=lambda x:x['tonnes'],reverse=True):
        screen_out.append(dict(r,tonnes=round(r['tonnes'],2),recovery=None))

    # Mechanical is the single breakdown ledger. Clip each event against the
    # requested, permission-scoped shift windows, including overnight shifts.
    active_shifts={x.shift:x for x in db.scalars(
        select(ShiftMaster).where(ShiftMaster.active.is_(True)))}
    scoped_shifts=([selected_shift] if selected_shift!='ALL'
                   else [name for name in active_shifts
                         if not allowed_shifts or name in allowed_shifts])
    shift_windows=[]
    day_cursor=start_day
    while day_cursor<=end_day:
        for name in scoped_shifts:
            definition=active_shifts.get(name)
            if not definition: continue
            begin=datetime.combine(day_cursor,definition.start_time,TZ)
            finish=datetime.combine(day_cursor,definition.end_time,TZ)
            if finish<=begin: finish+=timedelta(days=1)
            shift_windows.append((begin,finish))
        day_cursor+=timedelta(days=1)
    breakdown_rows=[]
    active_breakdowns=0
    breakdown_hours=0.0
    if shift_windows:
        min_window=min(x[0] for x in shift_windows)
        max_window=max(x[1] for x in shift_windows)
        bd_query=select(MaintenanceBreakdown).where(
            MaintenanceBreakdown.site_id=='TIOM',
            MaintenanceBreakdown.reported_at<max_window,
            (MaintenanceBreakdown.released_at.is_(None)) |
            (MaintenanceBreakdown.released_at>min_window)
        )
        reference_now=now_local()
        for bd in db.scalars(bd_query):
            event_start=aware(bd.reported_at)
            event_end=aware(bd.released_at) if bd.released_at else reference_now
            if event_end<=event_start: continue
            overlap=sum(max(0.0,(min(event_end,stop)-max(event_start,start)).total_seconds())
                        for start,stop in shift_windows)
            if overlap<=0: continue
            hours=round(overlap/3600,2)
            breakdown_hours+=hours
            open_event=bd.status in {'OPEN','IN_PROGRESS'} and not bd.released_at
            active_breakdowns+=int(open_event)
            breakdown_rows.append({
                'id':bd.breakdown_id,'machine':bd.asset_id,'fault':bd.problem,
                'category':bd.problem_category,'status':bd.status,
                'reportedAt':bd.reported_at.isoformat() if bd.reported_at else None,
                'releasedAt':bd.released_at.isoformat() if bd.released_at else None,
                'periodDowntimeHours':hours,'isOpen':open_event
            })
    breakdown_rows.sort(key=lambda x:(x['isOpen'],x['periodDowntimeHours']),reverse=True)

    log.info(
        'TIOM dashboard built in %.2fs mode=%s period=%s..%s shift=%s wb=%d canonical_missing=%d',
        perf_counter()-dash_started, mode, start_day, end_day, selected_shift, len(wb_all), len(missing_canonical)
    )
    return {
        'fromDate':start_day,'toDate':end_day,'shift':selected_shift,'mode':mode,'materialFilter':material_filter,'sourceFilter':source_filter,'destinationFilter':destination_filter,'vehicleFilter':vehicle_filter,'availableMaterials':available_materials,'availableSources':available_sources,'availableDestinations':available_destinations,'availableVehicles':available_vehicles,
        'kpis':{
            'wbTrips':wb_trips,'wbTonnes':round(wb_tonnes,2),'avgPayload':round(avg_payload,2),'fieldTrips':len(field_vehicle_trips),
            'matched':len(exact_keys),'likely':len(likely_keys),'wbUnmatched':wb_unmatched,'tripNoWb':trip_no_wb,'matchRate':round(match_rate,1),
            'presentPeople':len(present_people),'workingEquipment':len(working_equipment),'hsdLitres':round(float(hsd_litres),1),'hsdCost':round(float(hsd_cost),2),
            'hsdPerTonne':round(fuel_per_tonne,2),'fuelPerTrip':round(fuel_per_trip,2),'materials':len(materials),'sources':len(sources),'destinations':len(destinations),
            'oreTonnes':round(max(0,ore_tonnes),2),'wasteTonnes':round(material_buckets['WASTE'],2),'romTonnes':round(material_buckets['ROM'],2),
            'finesTonnes':round(material_buckets['FINES'],2),'cloTonnes':round(material_buckets['CLO'],2),'rejectTonnes':round(material_buckets['REJECT'],2),
            'crusherFeed':round(float(crusher_feed),2),'screenFeed':round(float(msp_feed),2),'avgCycleTime':round(avg_cycle,1) if avg_cycle is not None else None,
            'avgLoadingTime':round(avg_loading,1) if avg_loading is not None else None,'avgUnloadingTime':None,'avgQueueTime':None,
            'activeCrushers':active_crushers,'activeScreens':active_screens,'runningLoaders':running_loaders,'runningExcavators':running_excavators,
            'runningTippers':running_tippers,'idleEquipment':idle_equipment,'equipmentUtilization':round(equipment_util,1),'tonPerTrip':round(avg_payload,2),
            'avgTripsVehicle':round(avg_trips_vehicle,1),'avgTripsLoader':round(avg_trips_loader,1),'avgTripsExcavator':round(avg_trips_exc,1),
            'avgTonnesLoader':round(avg_tonnes_loader,1),'avgTonnesExcavator':round(avg_tonnes_exc,1),'totalRoutes':len(routes),
            'deployed':deployed,'loading':open_loading,'inTransit':in_transit,
            'misReports':len(mis_reports),'misDrafts':len(mis_drafts),'misTrips':mis_trip_count,
            'misOperationalQty':round(float(mis_total_qty),2),'misWbLinkedTrips':mis_wb_linked,'misFactorTrips':mis_factor_trips,
            'leadResolvedTrips':lead_resolved_trips,'leadMissingTrips':lead_missing_trips,
            'leadWithWbTrips':lead_with_wb,'leadWithoutWbTrips':lead_without_wb,
            'avgLeadKm':round(float(weighted_avg_lead),3),'leadTonKm':round(float(lead_ton_km),2),
            'misObTrips':mis_ob_trips,'misObQty':round(float(mis_ob_qty),2),'misRomTrips':mis_rom_trips,'misRomQty':round(float(mis_rom_qty),2),
            'shiftExcavationMt':round(float(shift_excavation),2),'shiftProcessedMt':round(float(shift_processed),2),
            'romInputMt':production_summary['romInputMt'],'finalProductionMt':production_summary['finalProductionMt'],
            'finalRecoveryPct':production_summary['finalRecoveryPct'],'oldStockExcludedMt':production_summary['oldStockExcludedMt'],
            'haulageTrips':combined_trips,'haulageActiveHours':combined_active_hours,
            'tripsPerHour':round(combined_tph,2) if combined_tph is not None else None,
            'haulAvgLeadKm':round(combined_avg_lead,3) if combined_avg_lead is not None else None,
            'tripKmPerHour':round(combined_trip_kmh,2) if combined_trip_kmh is not None else None,
            'tonKmPerHour':round(combined_ton_kmh,2) if combined_ton_kmh is not None else None,
            'haulLeadMissingTrips':combined_lead_missing,
            'drillHoles':int(drill_totals.get('holes') or 0),'drillMeterage':round(float(drill_totals.get('totalMeterage') or 0),2),
            'drillMeterPerHour':round(float(drill_totals.get('meterPerHour') or 0),2) if drill_totals.get('meterPerHour') is not None else None,
            'drillHsdLitres':round(float(drill_totals.get('hsdLitres') or 0),1),'drillBreakdownHours':round(float(drill_totals.get('breakdownHours') or 0),2),
        },
        'hourly':[dict(x,tonnes=round(x['tonnes'],2)) for x in hourly.values()],
        'productionHourly':[{'hour':x['hour'],'rom':round(x['rom'],2),'final':round(x['final'],2)} for x in production_hourly.values()],
        'materials':mat_rows,'productionMaterials':production_material_rows,'plantPerformance':plant_performance,'sources':source_rows[:20],'destinations':dest_rows[:20],'routes':route_rows[:20],
        'misMaterials':_mis_rows(mis_materials)[:30],'misSources':_mis_rows(mis_sources)[:30],
        'misDestinations':_mis_rows(mis_destinations)[:30],
        'misVehicles':_mis_rows(mis_vehicles,'vehicle')[:40],'misMachines':_mis_rows(mis_machines,'machine')[:40],
        'leadRoutes':lead_route_rows[:50],'haulagePerformance':haulage_rows[:100],
        'exceptionDetails':{'lead':lead_review[:500],'tonnage':qty_pending[:500],
                            'leadTotal':len(lead_review),'tonnageTotal':len(qty_pending)},
        'breakdowns':{'rows':breakdown_rows[:200],
                      'totalEvents':len(breakdown_rows),'activeEvents':active_breakdowns,
                      'periodDowntimeHours':round(breakdown_hours,2)},
        'production':production_summary,'drilling':drilling,
        'crusher':crusher_out[:30],'screens':screen_out[:30],'loaders':loader_rows[:20],'excavators':excavator_rows[:20],
        'vehicles':vehicle_rows[:30],'fuelByEquipment':fuel_rows[:30],'shiftComparison':shift_comp_rows,'sevenDay':seven_rows,'monthTrend':month_rows,'monthly':monthly,
        'equipmentStatus':[{'label':k,'value':v} for k,v in status_counts.items()],
        'heatmapMaterials':matrix_materials,'heatmap':heatmap[:30],'materialFlow':material_flow,
        'topPerformers':performer_pool,'bottomPerformers':bottom_performers,'exceptions':exceptions,
        'unsupported':{
            'avgUnloadingTime':'Unload-start time is not captured; only load-end and unload-complete exist.',
            'avgQueueTime':'Queue entry/exit timestamps are not captured.',
            'avgBucketCount':'Bucket count is not captured in load_trip.',
            'crusherUtilization':'Crusher run/downtime or throughput hours are not captured against WB movements.',
            'plantTelemetry':'Plant feed/output is derived from confirmed WB movements. TPH is shown only where actual equipment HMR/run-hours are recorded; otherwise it remains blank.',
            'weather':'Site weather source/coordinates are not configured.'
        },
        'notes':[
            'WB remains the authoritative movement source, but total WB tonnes is not the primary production KPI. Fresh ROM is the 100% mine-production base; plant throughput and rehandling are reported separately.',
            'CLO 10-40 is company old-stock blending material. It is included in physical Crusher Feed/TPH but excluded from NMTPL fresh production. Its share of mixed crusher output is allocated pro-rata by feed tonnes.',
            'Genuine MSP lumps and large spillage sent to Crusher are fresh intermediate feed. Spillage sent back to an MSP is internal recycle: it increases gross plant throughput/TPH but not fresh feed or production.',
            'RH / old-stock movements, including Project Area / PA SRF material and ROM drawn from stock, are excluded from new-production and recovery KPIs.',
            'Trips Per Hour means trips divided by active clock-hour buckets. Lead-adjusted haulage also reports Trip-km/hr and Ton-km/hr; unresolved lead is shown, never guessed.',
            'MIS Manual Entry panels use SUBMITTED driver reports only. Draft reports are shown as pending counts and are excluded from production until Submit Shift Report is used.',
            'Lead KM is automatic: Bench RL comes from Face / Bench Setup and route type comes from WB linkage. Source + Destination + Route Master + Lead Master determine Lead KM; unresolved rows are excluded from Ton-km.',
            'MIS operational MT is displayed separately from authoritative WB tonnes. WB-linked MIS rows are evidence only and are never added again to WB production totals; unlinked rows use the approved trip factor.',
            'Loader/excavator tonnes and machine-material heatmap use exact MATCHED / MANUAL_MATCH field↔WB records only.',
            'Running/idle and utilization are activity proxies from the latest attendance condition plus recorded production activity; auxiliary machine work is not treated as proven idle.',
            'Crusher/screen panels classify WB source/destination labels/codes containing CRUSH, SCREEN or MSP; dedicated crusher/screen process telemetry is not yet stored.'
        ]
    }

# ---------------- Migrated operations: Production, WB, HSD, Masters ----------------

def _active_shift_maps(db, day, sh):
    p_att = {r.employee_id: r for r in rows(db, PersonAttendance, day, sh)}
    e_att = {r.machine_id: r for r in rows(db, EquipmentAttendance, day, sh)}
    crew = {r.machine_id: r for r in rows(db, ShiftCrew, day, sh) if r.status == 'ACTIVE'}
    dep = {r.machine_id: r for r in rows(db, ShiftDeployment, day, sh) if r.status == 'ACTIVE'}
    return p_att, e_att, crew, dep


def _crew_is_present(p_att, crew_row):
    if not crew_row:
        return False
    row = p_att.get(crew_row.employee_id)
    return bool(row and row.status == 'PRESENT' and row.in_at and not row.out_at)


def production_desk(db, user):
    """TIOM Site Entry desk.

    Attendance / Shift Control are intentionally not prerequisites in the current
    TIOM phase.  The familiar Site Entry UI is retained, but it now draws from
    active equipment masters directly so field production remains usable while
    Attendance and Shift Control are hidden.
    """
    require(user, 'PRODUCTION')
    day, sh, now = operating_context(); require(user, shift=sh)
    equipment = {e.machine_id: e for e in db.scalars(select(Equipment).where(Equipment.active))}
    locations = {l.location_id: l for l in db.scalars(select(Location).where(Location.active))}
    products = list(db.scalars(select(Product).where(Product.active).order_by(Product.name)))
    activities = list(db.scalars(select(ActivityMaster).where(ActivityMaster.active).order_by(ActivityMaster.activity)))
    active = list(db.scalars(select(LoadTrip).where(
        LoadTrip.operating_date == day, LoadTrip.shift == sh,
        LoadTrip.status.in_(['LOADING','HAULING'])
    ).order_by(LoadTrip.loading_start_at)))
    machine_busy = {t.machine_id for t in active if t.status == 'LOADING'}
    vehicle_busy = {t.vehicle_id for t in active if t.vehicle_id and t.status in {'LOADING','HAULING'}}
    machines=[]; vehicles=[]
    for mid,e in equipment.items():
        row={'id':mid,'type':e.type,'group':e.group,'locationId':'',
             'locationName':'Ready','crewId':'','busy':False}
        if e.group == 'LOADING':
            row['busy']=mid in machine_busy; machines.append(row)
        elif e.group == 'TRANSPORT':
            row['busy']=mid in vehicle_busy; row['vehicleNo']=e.vehicle_no or ''; row['doorNo']=e.door_no or ''; vehicles.append(row)
    recent=list(db.scalars(select(LoadTrip).where(
        LoadTrip.operating_date==day, LoadTrip.shift==sh
    ).order_by(LoadTrip.loading_start_at.desc()).limit(40)))
    return {'date':day,'shift':sh,'time':now.strftime('%H:%M:%S'),
            'locations':[{'id':x.location_id,'name':x.location_name,'type':x.location_type or ''} for x in locations.values()],
            'products':[{'id':x.product_id,'name':x.name} for x in products],
            'activities':[{'id':x.activity,'vehicleRequired':x.vehicle_required} for x in activities],
            'machines':sorted(machines,key=lambda x:x['id']), 'vehicles':sorted(vehicles,key=lambda x:x['id']),
            'activeTrips':[{'tripId':t.trip_id,'status':t.status,'machine':t.machine_id,'vehicle':t.vehicle_id or '',
                            'source':t.source_location_id,'destination':t.destination_location_id or '',
                            'material':t.material_id or '', 'start':time_text(t.loading_start_at),
                            'loaded':time_text(t.loading_end_at), 'seq':t.trip_seq or ''} for t in active],
            'recent':[{'tripId':t.trip_id,'status':t.status,'machine':t.machine_id,'vehicle':t.vehicle_id or '',
                       'source':t.source_location_id,'destination':t.destination_location_id or '',
                       'material':t.material_id or '', 'start':time_text(t.loading_start_at),
                       'loaded':time_text(t.loading_end_at), 'unloaded':time_text(t.unload_at), 'seq':t.trip_seq or ''} for t in recent]}


def start_production_trip(db, user, p):
    require(user,'PRODUCTION')
    day, sh, now = operating_context(); require(user,shift=sh); open_shift(db,day,sh)
    request_id=str(p.get('requestId') or '').strip()
    if not re.fullmatch(r'[A-Za-z0-9_.:-]{12,100}', request_id): raise HTTPException(422,'Missing/invalid request ID. Reload Production.')
    old=db.scalar(select(LoadTrip).where(LoadTrip.request_id==request_id))
    if old: return {'ok':True,'message':'Trip already saved.','tripId':old.trip_id,'status':old.status}
    source=str(p.get('sourceLocationId') or '').strip(); dest=str(p.get('destinationLocationId') or '').strip()
    machine=str(p.get('machineId') or '').strip(); vehicle=str(p.get('vehicleId') or '').strip()
    material=str(p.get('materialId') or '').strip(); activity=str(p.get('activity') or 'LOADING').strip().upper()
    active_resource(db,Location,source)
    if dest: active_resource(db,Location,dest)
    act=active_resource(db,ActivityMaster,activity)
    if material: active_resource(db,Product,material)
    m=active_resource(db,Equipment,machine)
    if m.group!='LOADING': raise HTTPException(422,'Choose a loader/excavator from the LOADING group.')
    if busy(db,machine): raise HTTPException(409,f'{machine} already has an active loading event.')
    driver_id=None
    if act.vehicle_required:
        if not vehicle: raise HTTPException(422,'This activity requires a vehicle.')
        if not dest: raise HTTPException(422,'Destination is required for a vehicle trip.')
        if not material: raise HTTPException(422,'Material is required for a vehicle trip.')
        v=active_resource(db,Equipment,vehicle)
        if v.group!='TRANSPORT': raise HTTPException(422,'Choose a TRANSPORT vehicle.')
        if busy(db,vehicle): raise HTTPException(409,f'{vehicle} is already loading/hauling.')
    else:
        vehicle=''
    seq=None
    if vehicle:
        seq=(db.scalar(select(func.count()).select_from(LoadTrip).where(
            LoadTrip.operating_date==day,LoadTrip.shift==sh,LoadTrip.vehicle_id==vehicle)) or 0)+1
    trip=LoadTrip(trip_id=str(uuid4()),operating_date=day,shift=sh,source_location_id=source,
                  destination_location_id=dest or None,activity=activity,machine_id=machine,
                  machine_operator_id=None,vehicle_id=vehicle or None,vehicle_driver_id=driver_id,
                  trip_seq=seq,material_id=material or None,loading_start_at=now,status='LOADING',
                  request_id=request_id,created_by=user.login_id,created_at=now,updated_at=now)
    db.add(trip); audit(db,user,'START_TRIP','load_trip',trip.trip_id,{'machine':machine,'vehicle':vehicle,'source':source,'destination':dest,'material':material,'mode':'TIOM_SITE_ENTRY_NO_ATTENDANCE'})
    return {'ok':True,'message':'Loading started.','tripId':trip.trip_id,'status':'LOADING','tripSeq':seq}


def transition_trip(db,user,p,action):
    require(user,'PRODUCTION')
    trip=db.get(LoadTrip,str(p.get('tripId') or ''))
    if not trip: raise HTTPException(404,'Trip not found.')
    require(user,shift=trip.shift)
    now=now_local()
    if action=='LOADED':
        if trip.status!='LOADING': return {'ok':True,'message':f'Trip is already {trip.status}.','status':trip.status}
        trip.loading_end_at=now; trip.status='HAULING' if trip.vehicle_id else 'CLOSED'; trip.updated_at=now
    elif action=='UNLOADED':
        if trip.status!='HAULING': raise HTTPException(409,'Only a HAULING trip can be marked unloaded.')
        trip.unload_at=now; trip.status='CLOSED'; trip.updated_at=now
    audit(db,user,action,'load_trip',trip.trip_id,{'status':trip.status})
    return {'ok':True,'message':'Loading completed; loader released.' if action=='LOADED' else 'Vehicle unloaded / ready.','status':trip.status}


def hsd_desk(db,user):
    require(user,'HSD')
    day,sh,now=operating_context(); require(user,shift=sh)
    tankers=list(db.scalars(select(HsdTanker).where(HsdTanker.active).order_by(HsdTanker.tanker_id)))
    stock={t.tanker_id: db.scalar(select(func.coalesce(func.sum(HsdPurchaseLot.litres_remaining),0)).where(HsdPurchaseLot.tanker_id==t.tanker_id)) for t in tankers}
    p_att,e_att,crew,dep=_active_shift_maps(db,day,sh)
    eq={e.machine_id:e for e in db.scalars(select(Equipment).where(Equipment.active))}
    loc={l.location_id:l for l in db.scalars(select(Location).where(Location.active))}
    machines=[]
    for mid,e in eq.items():
        if e.group=='HSD_TANKER': continue
        a=e_att.get(mid)
        if not a or a.status!='PRESENT' or a.condition!='WORKING': continue
        c=crew.get(mid); d=dep.get(mid)
        machines.append({'id':mid,'type':e.type,'group':e.group,'locationId':d.location_id if d else '',
                         'locationName':loc[d.location_id].location_name if d and d.location_id in loc else '',
                         'crewId':c.employee_id if c and _crew_is_present(p_att,c) else '',
                         'meterType':'KMR' if e.group=='TRANSPORT' else 'HMR'})
    lots=list(db.scalars(select(HsdPurchaseLot).order_by(HsdPurchaseLot.received_at.desc()).limit(20)))
    issues=list(db.scalars(select(HsdIssue).order_by(HsdIssue.issued_at.desc()).limit(30)))
    return {'date':day,'shift':sh,'time':now.strftime('%H:%M:%S'),
            'tankers':[{'id':t.tanker_id,'vehicleNo':t.vehicle_no or '','capacity':float(t.capacity_l),'stock':float(stock[t.tanker_id] or 0)} for t in tankers],
            'machines':machines,
            'receipts':[{'lotId':x.lot_id,'tanker':x.tanker_id,'time':time_text(x.received_at),'supplier':x.supplier,'invoice':x.invoice_no or '',
                         'litres':float(x.litres_received),'remaining':float(x.litres_remaining),'rate':float(x.rate_per_l),'amount':float(x.amount)} for x in lots],
            'issues':[{'issueId':x.issue_id,'tanker':x.tanker_id,'machine':x.machine_id,'time':time_text(x.issued_at),'litres':float(x.litres),
                       'amount':float(x.amount),'location':x.location_id or '','meterType':x.meter_type or '','meter':float(x.meter_reading) if x.meter_reading is not None else None,
                       'reference':x.reference or ''} for x in issues]}


def save_hsd_receipt(db,user,p):
    require(user,'HSD')
    day,sh,now=operating_context(); require(user,shift=sh); open_shift(db,day,sh)
    request_id=str(p.get('requestId') or '').strip()
    old=db.scalar(select(HsdPurchaseLot).where(HsdPurchaseLot.request_id==request_id)) if request_id else None
    if old:return {'ok':True,'message':'Receipt already saved.','lotId':old.lot_id}
    tanker=active_resource(db,HsdTanker,str(p.get('tankerId') or '').strip())
    try: litres=Decimal(str(p.get('litres'))); rate=Decimal(str(p.get('ratePerL')))
    except Exception: raise HTTPException(422,'Enter valid litres and â‚¹/L.')
    if litres<=0 or rate<=0: raise HTTPException(422,'Litres and â‚¹/L must be positive.')
    supplier=short(p.get('supplier','')).strip()
    if not supplier: raise HTTPException(422,'Supplier / Petrol Pump is required.')
    amount=(litres*rate).quantize(Decimal('0.01'))
    lot=HsdPurchaseLot(lot_id=str(uuid4()),tanker_id=tanker.tanker_id,received_at=now,supplier=supplier,
                       invoice_no=short(p.get('invoiceNo','')) or None,pump_location=short(p.get('pumpLocation','')) or None,
                       litres_received=litres,litres_remaining=litres,rate_per_l=rate,amount=amount,request_id=request_id)
    db.add(lot); audit(db,user,'HSD_RECEIPT','hsd_purchase_lot',lot.lot_id,{'tanker':tanker.tanker_id,'litres':str(litres),'rate':str(rate),'amount':str(amount)})
    return {'ok':True,'message':'Fuel received into tanker as a FIFO lot.','lotId':lot.lot_id,'amount':float(amount)}


def save_hsd_issue(db,user,p):
    require(user,'HSD')
    day,sh,now=operating_context(); require(user,shift=sh); open_shift(db,day,sh)
    request_id=str(p.get('requestId') or '').strip()
    old=db.scalar(select(HsdIssue).where(HsdIssue.request_id==request_id)) if request_id else None
    if old:return {'ok':True,'message':'Issue already saved.','issueId':old.issue_id,'amount':float(old.amount)}
    tanker=active_resource(db,HsdTanker,str(p.get('tankerId') or '').strip())
    machine=active_resource(db,Equipment,str(p.get('machineId') or '').strip())
    try: litres=Decimal(str(p.get('litres')))
    except Exception: raise HTTPException(422,'Enter valid litres.')
    if litres<=0: raise HTTPException(422,'Litres must be positive.')
    p_att,e_att,crew,dep=_active_shift_maps(db,day,sh)
    a=e_att.get(machine.machine_id)
    if not a or a.status!='PRESENT' or a.condition!='WORKING': raise HTTPException(409,'Receiving machine must be PRESENT + WORKING.')
    d=dep.get(machine.machine_id); c=crew.get(machine.machine_id)
    location=str(p.get('locationId') or (d.location_id if d else '')).strip() or None
    if location: active_resource(db,Location,location)
    meter_type='KMR' if machine.group=='TRANSPORT' else 'HMR'
    meter_raw=p.get('meterReading')
    meter=None if meter_raw in (None,'') else Decimal(str(meter_raw))
    issue=HsdIssue(issue_id=str(uuid4()),operating_date=day,shift=sh,tanker_id=tanker.tanker_id,machine_id=machine.machine_id,
                   litres=litres,amount=Decimal('0'),location_id=location,meter_type=meter_type,meter_reading=meter,
                   recipient_employee_id=c.employee_id if c and _crew_is_present(p_att,c) else None,
                   reference=short(p.get('reference','')) or None,issued_at=now,request_id=request_id)
    try: allocations,total=apply_fifo_issue(db,issue)
    except ValueError as exc: raise HTTPException(409,str(exc))
    audit(db,user,'HSD_ISSUE','hsd_issue',issue.issue_id,{'tanker':tanker.tanker_id,'machine':machine.machine_id,'litres':str(litres),'amount':str(total),
          'lots':[{'lot':lot.lot_id,'litres':str(qty),'rate':str(lot.rate_per_l)} for lot,qty,_ in allocations]})
    return {'ok':True,'message':'HSD issue saved with exact FIFO cost.','issueId':issue.issue_id,'amount':float(total),
            'allocations':[{'lotId':lot.lot_id,'litres':float(qty),'rate':float(lot.rate_per_l),'amount':float(amount)} for lot,qty,amount in allocations]}


def masters_desk(db,user):
    require(user,'MASTERS')
    ensure_wb_header_mapping(db)
    ensure_location_master_roles(db)
    _ensure_tiom_routes_from_leads(db)
    return {
        'options':[{'category':r.category,'value':r.value} for r in db.scalars(select(MasterOption).order_by(MasterOption.value))],
        'rotationTeams':[{'id':r.team_id,'name':r.team_name} for r in db.scalars(select(ShiftRotation).where(ShiftRotation.active))],
        'persons':[{'id':r.employee_id,'name':r.name,'role':r.role,'department':r.department or '','rotationTeam':r.rotation_team or '','active':r.active} for r in db.scalars(select(Person).order_by(Person.employee_id))],
        'equipment':[{
            'id':r.machine_id,'doorNo':r.door_no or '','vehicleNo':r.vehicle_no or '',
            'type':r.type,'group':r.group,'makeModel':r.make_model or '',
            'bucketCum':r.bucket_cum,'ratedPayloadT':r.rated_payload_t,
            'ratedOutputTph':r.rated_output_tph,'standingTareKg':r.standing_tare_kg,
            'ownership':r.ownership or '','active':r.active
        } for r in db.scalars(select(Equipment).order_by(Equipment.machine_id))],
        'locations':[{
            'id':r.location_id,'name':r.location_name,'role':r.role or 'UNCLASSIFIED',
            'type':r.location_type or '','active':r.active
        } for r in db.scalars(select(Location).order_by(Location.location_id))],
        'locationAliases':[{'id':r.alias,'locationId':r.location_id,'direction':r.direction or 'ANY','active':r.active} for r in db.scalars(select(LocationAlias).order_by(LocationAlias.alias))],
        'wbHeaders':[{'id':str(r.id),'field':r.canonical_field,'header':r.header_alias,'occurrence':r.occurrence,'priority':r.priority,'required':r.required,'active':r.active} for r in db.scalars(select(WbHeaderAlias).order_by(WbHeaderAlias.canonical_field,WbHeaderAlias.priority,WbHeaderAlias.id))],
        'products':[{'id':r.product_id,'name':r.name,'active':r.active} for r in db.scalars(select(Product).order_by(Product.product_id))],
        'activities':[{'id':r.activity,'vehicleRequired':r.vehicle_required,'active':r.active} for r in db.scalars(select(ActivityMaster).order_by(ActivityMaster.activity))],
        'routes':[{'id':r.route_id,'name':r.route_name,'sourceLocationId':r.source_location_id,'destinationLocationId':r.destination_location_id,'routeMode':r.route_mode,'leadBasis':r.lead_basis or 'BENCH_RL','fixedLeadKm':float(r.fixed_lead_km) if r.fixed_lead_km is not None else None,'materialScope':r.material_scope or '','via':r.via_text or '','leadEntries':db.scalar(select(func.count()).select_from(TiomLeadDistance).where(TiomLeadDistance.source_location_id==r.source_location_id,TiomLeadDistance.destination_location_id==r.destination_location_id,TiomLeadDistance.route_mode==r.route_mode,TiomLeadDistance.active.is_(True))) or 0,'active':r.active} for r in db.scalars(select(TiomRouteMaster).order_by(TiomRouteMaster.route_name,TiomRouteMaster.route_id))],
        'leadDistances':[{'id':r.lead_id,'routeId':_tiom_route_key(r.source_location_id,r.destination_location_id,r.route_mode),'sourceLocationId':r.source_location_id,'benchRl':r.bench_rl_m,'destinationLocationId':r.destination_location_id,'routeMode':r.route_mode,'leadKm':float(r.lead_km),'materialScope':r.material_scope or '','active':r.active} for r in db.scalars(select(TiomLeadDistance).order_by(TiomLeadDistance.source_location_id,TiomLeadDistance.destination_location_id,TiomLeadDistance.bench_rl_m,TiomLeadDistance.route_mode))],
        'drillSets':[{'id':r.drill_set_id,'name':r.drill_set_name,'drillMachineId':r.drill_machine_id,'compressorMachineId':r.compressor_machine_id,'notes':r.notes or '','active':r.active} for r in db.scalars(select(TiomDrillSet).order_by(TiomDrillSet.drill_set_name,TiomDrillSet.drill_set_id))],
        'tankers':[{'id':r.tanker_id,'vehicleNo':r.vehicle_no or '','capacity':float(r.capacity_l),'active':r.active} for r in db.scalars(select(HsdTanker).order_by(HsdTanker.tanker_id))]
    }


def save_master_record(db,user,p):
    require(user,admin=True)
    table=str(p.get('table') or '').upper()
    r=p.get('row') if isinstance(p.get('row'),dict) else {}
    intent=str(p.get('intent') or '').lower()

    def flag(v):
        return v if isinstance(v,bool) else str(v).lower() in {'true','1','yes','y','active'}

    def number(v,label):
        if v in (None,''): return None
        try: return float(v)
        except Exception: raise HTTPException(422,f'Enter valid {label}.')

    if table == 'ROUTE':
        key=short(str(r.get('id') or '')).strip() or str(uuid4())
        source_id=short(str(r.get('sourceLocationId') or '')).strip(); dest_id=short(str(r.get('destinationLocationId') or '')).strip()
        if not source_id or not dest_id: raise HTTPException(422,'Route Source and Destination are required.')
        active_resource(db,Location,source_id); active_resource(db,Location,dest_id)
        mode=_tiom_route_mode(r.get('routeMode'))
        if not mode: raise HTTPException(422,'Route type must be WITH_WB or WITHOUT_WB.')
        duplicate=db.scalar(select(TiomRouteMaster).where(TiomRouteMaster.source_location_id==source_id,TiomRouteMaster.destination_location_id==dest_id,TiomRouteMaster.route_mode==mode))
        existing=db.get(TiomRouteMaster,key)
        if intent=='create' and duplicate: raise HTTPException(409,'This Source / Destination / route type already exists in Route Master.')
        if intent=='update' and not existing: raise HTTPException(409,'Route no longer exists. Refresh Masters.')
        obj=existing or TiomRouteMaster(route_id=key,route_name='',source_location_id=source_id,destination_location_id=dest_id,route_mode=mode)
        obj.route_name=short(str(r.get('name') or '')).strip() or f'{source_id} → {dest_id}'
        obj.source_location_id=source_id; obj.destination_location_id=dest_id; obj.route_mode=mode
        basis=str(r.get('leadBasis') or getattr(obj,'lead_basis',None) or 'BENCH_RL').strip().upper()
        if basis not in {'BENCH_RL','FIXED'}: raise HTTPException(422,'Lead Basis must be BENCH_RL or FIXED.')
        fixed_raw=r.get('fixedLeadKm')
        fixed=None if fixed_raw in (None,'') else Decimal(str(fixed_raw))
        if fixed is not None and (not fixed.is_finite() or fixed<=0): raise HTTPException(422,'Fixed Lead KM must be positive.')
        if basis=='FIXED' and fixed is None and intent!='create':
            fixed=obj.fixed_lead_km
        obj.lead_basis=basis; obj.fixed_lead_km=fixed if basis=='FIXED' else None
        obj.material_scope=short(str(r.get('materialScope') or '')) or None; obj.via_text=short(str(r.get('via') or '')) or None
        obj.active=flag(r.get('active',True)); obj.entered_by=user.login_id; obj.entered_at=now_local(); db.add(obj)
        audit(db,user,'SAVE_MASTER',table,key,{'row':r})
        return {'ok':True,'message':f'Route {key} saved.'}

    if table == 'LEAD_DISTANCE':
        route_id=short(str(r.get('routeId') or '')).strip(); route=db.get(TiomRouteMaster,route_id)
        if not route or not route.active: raise HTTPException(422,'Choose an active Route Master record.')
        if str(route.lead_basis or 'BENCH_RL').upper()=='FIXED':
            raise HTTPException(422,'This is a FIXED route. Enter Lead KM in Route Master; Bench RL is not required.')
        bench=_tiom_parse_bench_rl(r.get('benchRl'))
        if bench is None: raise HTTPException(422,'Bench RL is required.')
        try: lead=Decimal(str(r.get('leadKm')))
        except Exception: raise HTTPException(422,'Enter valid Lead KM.')
        if not lead.is_finite() or lead<=0: raise HTTPException(422,'Lead KM must be positive.')
        existing=None
        if r.get('id'): existing=db.get(TiomLeadDistance,str(r.get('id')))
        if existing is None:
            existing=db.scalar(select(TiomLeadDistance).where(
                TiomLeadDistance.source_location_id==route.source_location_id,
                TiomLeadDistance.destination_location_id==route.destination_location_id,
                TiomLeadDistance.route_mode==route.route_mode,
                TiomLeadDistance.bench_rl_m==bench
            ))
        key=existing.lead_id if existing else f'{route.source_location_id}|{bench}|{route.destination_location_id}|{route.route_mode}'[:180]
        obj=existing or TiomLeadDistance(lead_id=key,source_location_id=route.source_location_id,bench_rl_m=bench,destination_location_id=route.destination_location_id,route_mode=route.route_mode,lead_km=lead)
        obj.source_location_id=route.source_location_id; obj.destination_location_id=route.destination_location_id; obj.route_mode=route.route_mode; obj.bench_rl_m=bench; obj.lead_km=lead
        obj.material_scope=short(str(r.get('materialScope') or route.material_scope or '')) or None; obj.active=flag(r.get('active',True)); obj.entered_by=user.login_id; obj.entered_at=now_local()
        db.add(obj); audit(db,user,'SAVE_MASTER',table,key,{'row':r})
        return {'ok':True,'message':f'Lead {float(lead):.3f} KM saved for Bench RL {bench}.'}

    if table == 'DRILL_SET':
        key=short(str(r.get('id') or '')).strip(); name=short(str(r.get('name') or '')).strip()
        drill_id=short(str(r.get('drillMachineId') or '')).strip(); comp_id=short(str(r.get('compressorMachineId') or '')).strip()
        if not key or not name or not drill_id or not comp_id: raise HTTPException(422,'Drill Set ID, name, Drill machine and Compressor are required.')
        if drill_id==comp_id: raise HTTPException(422,'Drill and Compressor must be different equipment records.')
        drill=active_resource(db,Equipment,drill_id); comp=active_resource(db,Equipment,comp_id)
        if drill.group!='DRILLING' and 'DRILL' not in str(drill.type or '').upper(): raise HTTPException(422,'Selected Drill machine must belong to DRILLING equipment.')
        existing=db.get(TiomDrillSet,key)
        if intent=='create' and existing: raise HTTPException(409,'Drill Set ID already exists.')
        if intent=='update' and not existing: raise HTTPException(409,'Drill Set no longer exists. Refresh Masters.')
        obj=existing or TiomDrillSet(drill_set_id=key,drill_set_name=name,drill_machine_id=drill_id,compressor_machine_id=comp_id)
        obj.drill_set_name=name; obj.drill_machine_id=drill_id; obj.compressor_machine_id=comp_id; obj.notes=short(str(r.get('notes') or '')) or None; obj.active=flag(r.get('active',True)); obj.entered_by=user.login_id; obj.entered_at=now_local(); db.add(obj)
        audit(db,user,'SAVE_MASTER',table,key,{'row':r})
        return {'ok':True,'message':f'Drill Set {key} saved.'}

    # Alias masters are deliberately additive and use existing SQL tables.
    if table == 'LOCATION_ALIAS':
        key=short(str(r.get('id') or '')).strip()
        if not key: raise HTTPException(422,'Alias is required.')
        existing=db.get(LocationAlias,key)
        if intent=='create' and existing: raise HTTPException(409,'This alias already exists.')
        if intent=='update' and not existing: raise HTTPException(409,'This alias no longer exists. Refresh.')
        location_id=short(str(r.get('locationId') or '')).strip()
        location=db.get(Location,location_id)
        if not location: raise HTTPException(422,'Choose a valid location.')
        direction=short(str(r.get('direction') or 'ANY')).strip().upper()
        if direction not in {'ANY','SOURCE','DESTINATION'}: raise HTTPException(422,'Direction must be ANY, SOURCE or DESTINATION.')
        obj=existing or LocationAlias(alias=key,location_id=location_id)
        obj.location_id=location_id; obj.direction=direction; obj.active=flag(r.get('active',True))
        db.add(obj); audit(db,user,'SAVE_MASTER',table,key,{'row':r})
        return {'ok':True,'message':f'Location alias {key} saved.'}

    if table == 'WBHEADER':
        ensure_wb_header_mapping(db)
        allowed={'move','date','shift','vehicle','matcode','matname','source_code','source_name','dest_code','dest_name','tare','gross','net','time'}
        field=short(str(r.get('field') or '')).strip()
        header=short(str(r.get('header') or '')).strip()
        if field not in allowed: raise HTTPException(422,'Choose a valid canonical WB field.')
        if not header: raise HTTPException(422,'Excel header alias is required.')
        try:
            occurrence=int(r.get('occurrence') or 1); priority=int(r.get('priority') or 100)
        except Exception:
            raise HTTPException(422,'Occurrence and priority must be whole numbers.')
        if not 1 <= occurrence <= 10: raise HTTPException(422,'Occurrence must be 1-10.')
        if not 1 <= priority <= 999: raise HTTPException(422,'Priority must be 1-999.')
        obj=None
        if intent=='update':
            try: rid=int(r.get('id'))
            except Exception: raise HTTPException(422,'WB header mapping ID is invalid.')
            obj=db.get(WbHeaderAlias,rid)
            if not obj: raise HTTPException(409,'This WB header mapping no longer exists. Refresh.')
        else:
            duplicate=db.scalar(select(WbHeaderAlias).where(
                WbHeaderAlias.canonical_field==field,
                WbHeaderAlias.header_alias==header,
                WbHeaderAlias.occurrence==occurrence
            ))
            if duplicate: raise HTTPException(409,'This WB header mapping already exists.')
            obj=WbHeaderAlias()
        obj.canonical_field=field; obj.header_alias=header; obj.occurrence=occurrence; obj.priority=priority
        obj.required=flag(r.get('required',False)); obj.active=flag(r.get('active',True))
        db.add(obj); db.flush()
        audit(db,user,'SAVE_MASTER',table,obj.id,{'row':r})
        return {'ok':True,'message':f'WB header mapping {field} ← {header} saved.'}

    limits = {'PERSON': {'id':40,'name':120,'role':80,'department':80,'rotationTeam':30},
              'EQUIPMENT': {'id':50,'doorNo':50,'vehicleNo':50,'type':60,'group':30,'makeModel':120,'ownership':30},
              'LOCATION': {'id':80,'name':160,'role':20,'type':60}, 'PRODUCT': {'id':50,'name':120},
              'ACTIVITY': {'id':60}, 'TANKER': {'id':50,'vehicleNo':50}}
    for field, limit in limits.get(table, {}).items():
        value = r.get(field, '')
        if not isinstance(value, str) or len(value.strip()) > limit:
            raise HTTPException(422, f'{field} must be text with at most {limit} characters.')

    if intent in {'create', 'update'}:
        models = {'PERSON': Person, 'EQUIPMENT': Equipment, 'LOCATION': Location, 'PRODUCT': Product, 'ACTIVITY': ActivityMaster, 'TANKER': HsdTanker}
        model = models.get(table)
        key = str(r.get('id') or '').strip()
        if table == 'ACTIVITY': key = key.upper()
        existing = db.get(model, key) if model and key else None
        if intent == 'create' and existing:
            raise HTTPException(409, 'This ID already exists. Choose Edit on the existing record.')
        if intent == 'update' and not existing:
            raise HTTPException(409, 'This record no longer exists. Refresh the list.')

    if table=='PERSON':
        key=short(r.get('id','')).strip()
        if not key: raise HTTPException(422,'Employee ID required.')
        obj=db.get(Person,key) or Person(employee_id=key,name='',role='')
        obj.name=short(r.get('name','')).strip(); obj.role=short(r.get('role','')).strip(); obj.department=short(r.get('department','')) or None
        obj.rotation_team=short(r.get('rotationTeam','')) or None; obj.active=flag(r.get('active',True))
        if obj.rotation_team:
            team = db.get(ShiftRotation, obj.rotation_team)
            if not team or not team.active: raise HTTPException(422, 'Choose an active rotation team.')
        if not obj.name or not obj.role: raise HTTPException(422,'Name and role are required.')
    elif table=='EQUIPMENT':
        key=short(r.get('id','')).strip()
        if not key: raise HTTPException(422,'Machine ID required.')
        obj=db.get(Equipment,key) or Equipment(machine_id=key,type='',group='OTHER')
        obj.door_no=short(r.get('doorNo','')) or None
        obj.vehicle_no=short(r.get('vehicleNo','')) or None
        obj.type=short(r.get('type','')).strip()
        obj.group=short(r.get('group','')).upper()
        obj.make_model=short(r.get('makeModel','')) or None
        obj.bucket_cum=number(r.get('bucketCum'),'bucket capacity')
        obj.rated_payload_t=number(r.get('ratedPayloadT'),'rated payload')
        obj.rated_output_tph=number(r.get('ratedOutputTph'),'rated output')
        obj.standing_tare_kg=number(r.get('standingTareKg'),'standing tare')
        obj.ownership=short(r.get('ownership','')) or None
        obj.active=flag(r.get('active',True))
        valid_groups={'LOADING','TRANSPORT','PROCESSING','HSD_TANKER','EARTHMOVING','DRILLING','SUPPORT','OTHER'}
        if not obj.type or obj.group not in valid_groups:
            raise HTTPException(422,'Enter type and a valid centralized equipment group.')
    elif table=='LOCATION':
        key=short(r.get('id','')).strip()
        if not key: raise HTTPException(422,'Location ID required.')
        obj=db.get(Location,key) or Location(location_id=key,location_name='')
        obj.location_name=short(r.get('name','')).strip()
        role=short(r.get('role','UNCLASSIFIED')).strip().upper() or 'UNCLASSIFIED'
        if role not in {'SOURCE','DESTINATION','BOTH','UNCLASSIFIED'}:
            raise HTTPException(422,'Location role must be SOURCE, DESTINATION, BOTH or UNCLASSIFIED.')
        obj.role=role
        obj.location_type=short(r.get('type','')) or None
        obj.active=flag(r.get('active',True))
        if not obj.location_name: raise HTTPException(422,'Location name required.')
    elif table=='PRODUCT':
        key=short(r.get('id','')).strip()
        if not key: raise HTTPException(422,'Product ID required.')
        obj=db.get(Product,key) or Product(product_id=key,name='')
        obj.name=short(r.get('name','')).strip(); obj.active=flag(r.get('active',True))
        if not obj.name: raise HTTPException(422,'Product name required.')
    elif table=='ACTIVITY':
        key=short(r.get('id','')).strip().upper()
        if not key: raise HTTPException(422,'Activity required.')
        obj=db.get(ActivityMaster,key) or ActivityMaster(activity=key)
        obj.vehicle_required=flag(r.get('vehicleRequired',False)); obj.active=flag(r.get('active',True))
    elif table=='TANKER':
        key=short(r.get('id','')).strip()
        if not key: raise HTTPException(422,'Tanker ID required.')
        obj=db.get(HsdTanker,key) or HsdTanker(tanker_id=key,capacity_l=Decimal('0'))
        obj.vehicle_no=short(r.get('vehicleNo','')) or None
        try: obj.capacity_l=Decimal(str(r.get('capacity') or 0))
        except Exception: raise HTTPException(422,'Enter tanker capacity.')
        if not obj.capacity_l.is_finite() or obj.capacity_l<=0: raise HTTPException(422,'Tanker capacity must be positive.')
        obj.active=flag(r.get('active',True))
    else:
        raise HTTPException(422,'Unknown master table.')

    db.add(obj); audit(db,user,'SAVE_MASTER',table,key,{'row':r})
    return {'ok':True,'message':f'{table} {key} saved.'}


OPTION_CATEGORIES = {'PERSON.role':80, 'PERSON.department':80, 'EQUIPMENT.type':60, 'EQUIPMENT.ownership':30, 'LOCATION.type':60}


def save_master_option(db, user, p):
    require(user, admin=True)
    category = str(p.get('category') or '')
    value = short(p.get('value', '')).strip()
    if category not in OPTION_CATEGORIES or not value or len(value) > OPTION_CATEGORIES[category]:
        raise HTTPException(422, 'Enter a valid option within the field length limit.')
    if not db.get(MasterOption, (category, value)):
        db.add(MasterOption(category=category, value=value))
        audit(db, user, 'ADD_MASTER_OPTION', category, value, {})
    return {'ok': True, 'message': 'Dropdown option saved.', 'value': value}


def save_rotation_team(db, user, p):
    require(user, admin=True)
    key = short(p.get('id','')).strip().upper()
    name = short(p.get('name','')).strip()
    shift = str(p.get('shift') or '').upper()
    if not key or len(key)>30 or not name or len(name)>80 or shift not in {'A','B','C'}:
        raise HTTPException(422, 'Enter a team ID, name and starting shift A, B or C.')
    if db.get(ShiftRotation,key): raise HTTPException(409, 'Team ID already exists.')
    day = _date_param(p.get('monday'), 'anchor Monday')
    if day.weekday()!=0: raise HTTPException(422, 'The anchor date must be a Monday.')
    db.add(ShiftRotation(team_id=key,team_name=name,anchor_monday=day,anchor_shift=shift,rotation_pattern='A,C,B',active=True))
    audit(db,user,'ADD_ROTATION_TEAM','SHIFT_ROTATION',key,{'monday':str(day),'shift':shift})
    return {'ok':True,'message':'Rotation team saved.', 'value':key}


def _norm_header(v):
    return re.sub(r'[^A-Z0-9]+','',str(v or '').upper())


def _norm_vehicle(v):
    return re.sub(r'[^A-Z0-9]+','',str(v or '').upper())


def _cell_decimal(v):
    if v is None or str(v).strip()=='': return None
    try: return Decimal(str(v).replace(',','').strip())
    except Exception: return None


def _clock_value(v):
    if isinstance(v, datetime): return v.time().replace(tzinfo=None)
    if isinstance(v, dtime): return v.replace(tzinfo=None)
    if isinstance(v, (int,float)) and 0 <= float(v) < 1:
        sec=int(round(float(v)*86400))%86400; return dtime(sec//3600,(sec%3600)//60,sec%60)
    txt=str(v or '').strip()
    for fmt in ('%H:%M:%S','%H:%M','%I:%M:%S %p','%I:%M %p'):
        try:return datetime.strptime(txt,fmt).time()
        except ValueError: pass
    try:return datetime.fromisoformat(txt).time()
    except Exception:return None


def _date_value(v):
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    txt=str(v or '').strip()
    for fmt in ('%Y-%m-%d','%d/%m/%Y','%d-%m-%Y','%d.%m.%Y'):
        try:return datetime.strptime(txt,fmt).date()
        except ValueError: pass
    return None


def _wb_weigh_at(day, definition, clock):
    if not clock:return None
    start,end=shift_bounds(day,definition); stamp=datetime.combine(day,clock,TZ)
    if end.date()>day and stamp<start: stamp += timedelta(days=1)
    return stamp if start <= stamp < end else None


def wb_desk(db,user,day=None,sh=None):
    require(user,'WB')
    if day is None:
        day,sh,_=operating_context()
    else:
        definition=None
        day=_parse_ui_date_v2(day, 'operating date'); sh=str(sh).upper(); require(user,shift=sh)
    batches=list(db.scalars(select(WbImportBatch).where(WbImportBatch.operating_date==day,WbImportBatch.shift==sh).order_by(WbImportBatch.confirmed_at.desc().nullslast())))
    chosen=None
    for b in batches:
        if b.status=='PREVIEW': chosen=b; break
    if not chosen:
        for b in batches:
            if b.status=='CONFIRMED': chosen=b; break
    movements=[]
    if chosen:
        movements=list(db.scalars(select(WbMovement).where(WbMovement.batch_id==chosen.batch_id).order_by(WbMovement.weigh_at)))
    return {'date':day,'shift':sh,'batches':[{'batchId':b.batch_id,'fileName':b.file_name,'status':b.status,'valid':b.valid_rows,'review':b.review_rows,
             'confirmedAt':b.confirmed_at.isoformat() if b.confirmed_at else ''} for b in batches[:12]],
            'selected':({'batchId':chosen.batch_id,'fileName':chosen.file_name,'status':chosen.status,'valid':chosen.valid_rows,'review':chosen.review_rows} if chosen else None),
            'rows':[{'key':w.movement_key,'status':w.row_status,'issue':w.issue or '','vehicle':w.vehicle_raw,'vehicleId':w.vehicle_id or '',
                     'time':time_text(w.weigh_at),'material':w.material_name or w.material_code or '', 'source':w.source_raw or '', 'destination':w.destination_raw or '',
                     'tareKg':float(w.tare_kg),'grossKg':float(w.gross_kg),'netKg':float(w.net_kg),'tonnes':float(w.net_kg/Decimal('1000'))} for w in movements]}



# NMTPL_WB_MONITOR_DASHBOARD_V1
def wb_monitor(db, user, options=None):
    # Compact WB/Gmail coverage monitor for the WB tab. Read-only.
    require(user, 'WB')
    options = options if isinstance(options, dict) else {}
    try:
        days = int(options.get('days', 14))
    except Exception:
        days = 14
    days = max(7, min(days, 45))

    now = now_local()
    today = now.date()
    start_day = today - timedelta(days=days - 1)

    permitted = set(user.shifts.split(',')) if user.shifts else set()
    definitions = [
        d for d in db.scalars(
            select(ShiftMaster).where(ShiftMaster.active).order_by(ShiftMaster.start_time)
        )
        if user.admin or 'ALL' in permitted or d.shift in permitted
    ]

    batches = list(db.scalars(
        select(WbImportBatch).where(
            WbImportBatch.operating_date >= start_day,
            WbImportBatch.operating_date <= today,
            WbImportBatch.shift.in_([d.shift for d in definitions] or ['__NONE__']),
            WbImportBatch.status.in_(['PREVIEW', 'CONFIRMED']),
        )
    ))

    by_key = {}
    for b in batches:
        key = (b.operating_date, b.shift)
        slot = by_key.setdefault(key, {'confirmed': None, 'preview': None})
        if b.status == 'CONFIRMED':
            old = slot['confirmed']
            old_ts = old.confirmed_at if old and old.confirmed_at else datetime.min.replace(tzinfo=TZ)
            new_ts = b.confirmed_at if b.confirmed_at else datetime.min.replace(tzinfo=TZ)
            if old is None or new_ts >= old_ts:
                slot['confirmed'] = b
        elif b.status == 'PREVIEW':
            slot['preview'] = b

    counts = {'UPDATED': 0, 'PENDING': 0, 'REVIEW': 0, 'MISSING': 0, 'NOT DUE': 0}
    coverage = []
    for offset in range(days):
        day = today - timedelta(days=offset)
        row = {'date': day.isoformat(), 'shifts': {}}
        for definition in definitions:
            key = (day, definition.shift)
            slot = by_key.get(key, {})
            confirmed = slot.get('confirmed')
            preview = slot.get('preview')
            start_at, end_at = shift_bounds(day, definition)
            grace_end = end_at + timedelta(hours=1)

            if preview is not None:
                status = 'REVIEW' if int(preview.review_rows or 0) > 0 else 'PENDING'
                source = preview
            elif confirmed is not None:
                status = 'UPDATED'
                source = confirmed
            elif now < end_at:
                status = 'NOT DUE'
                source = None
            elif now < grace_end:
                status = 'PENDING'
                source = None
            else:
                status = 'MISSING'
                source = None

            counts[status] = counts.get(status, 0) + 1
            row['shifts'][definition.shift] = {
                'status': status,
                'fileName': source.file_name if source else '',
                'valid': int(source.valid_rows or 0) if source else 0,
                'review': int(source.review_rows or 0) if source else 0,
                'confirmedAt': (
                    confirmed.confirmed_at.isoformat()
                    if confirmed and confirmed.confirmed_at else ''
                ),
                'shiftEnd': end_at.isoformat(),
                'graceEnd': grace_end.isoformat(),
            }
        coverage.append(row)

    latest_confirmed = db.scalar(
        select(WbImportBatch).where(
            WbImportBatch.status == 'CONFIRMED'
        ).order_by(WbImportBatch.confirmed_at.desc().nullslast()).limit(1)
    )

    root = Path(__file__).resolve().parents[2]
    status_path = root / 'logs' / 'wb_gmail' / 'status.json'
    config_path = root / 'config' / 'wb_gmail_config.json'
    gmail = {
        'state': 'UNKNOWN',
        'lastCheck': '',
        'pollSeconds': 180,
        'messagesSeen': 0,
        'counts': {},
        'archiveSuccess': None,
    }

    try:
        cfg = json.loads(config_path.read_text(encoding='utf-8-sig')) if config_path.exists() else {}
        gmail['pollSeconds'] = int(cfg.get('poll_seconds', 180))
        gmail['archiveSuccess'] = bool(cfg.get('archive_success', False))
    except Exception:
        pass

    try:
        payload = json.loads(status_path.read_text(encoding='utf-8-sig')) if status_path.exists() else {}
        state = str(payload.get('state') or 'UNKNOWN').upper()
        raw_poll = str(payload.get('last_poll') or payload.get('updated_at') or '')
        last_dt = None
        if raw_poll:
            try:
                last_dt = datetime.fromisoformat(raw_poll.replace('Z', '+00:00'))
                if last_dt.tzinfo is None:
                    last_dt = last_dt.replace(tzinfo=TZ)
                local_dt = last_dt.astimezone(TZ)
                gmail['lastCheck'] = local_dt.strftime('%d-%b-%Y %H:%M:%S')
            except Exception:
                gmail['lastCheck'] = raw_poll

        if state == 'RUNNING' and last_dt is not None:
            stale_after = max(600, gmail['pollSeconds'] * 3 + 60)
            age = (now - last_dt.astimezone(TZ)).total_seconds()
            if age > stale_after:
                state = 'STALE'

        gmail['state'] = state
        gmail['messagesSeen'] = int(payload.get('messages_seen') or 0)
        gmail['counts'] = payload.get('counts') if isinstance(payload.get('counts'), dict) else {}
    except Exception as exc:
        gmail['state'] = 'ERROR'
        gmail['error'] = str(exc)

    last_wb = None
    if latest_confirmed:
        last_wb = {
            'date': latest_confirmed.operating_date.isoformat(),
            'shift': latest_confirmed.shift,
            'fileName': latest_confirmed.file_name,
            'valid': int(latest_confirmed.valid_rows or 0),
            'review': int(latest_confirmed.review_rows or 0),
            'confirmedAt': latest_confirmed.confirmed_at.isoformat() if latest_confirmed.confirmed_at else '',
        }

    return {
        'now': now.isoformat(),
        'days': days,
        'shifts': [d.shift for d in definitions],
        'gmail': gmail,
        'counts': counts,
        'lastWb': last_wb,
        'coverage': coverage,
    }


@router.post('/wb-upload')
async def wb_upload(request: Request, operatingDate: str = Form(...), shift: str = Form(...), file: UploadFile = File(...), db: Session = Depends(get_db)):
    csrf(request); user=get_user(db,request); require(user,'WB')
    day,sh,definition=context(db,user,operatingDate,shift); open_shift(db,day,sh); lock(db)
    content=await file.read()
    if len(content)>15*1024*1024: raise HTTPException(413,'WB file is larger than 15 MB.')
    digest=hashlib.sha256(content).hexdigest()
    same_batches=list(db.scalars(select(WbImportBatch).where(
        WbImportBatch.operating_date==day, WbImportBatch.shift==sh,
        WbImportBatch.file_hash==digest, WbImportBatch.status.in_(['PREVIEW','CONFIRMED'])
    )))
    confirmed_same=next((b for b in same_batches if b.status=='CONFIRMED'),None)
    if confirmed_same:
        return {'ok':True,'message':'This exact WB file is already confirmed.','batchId':confirmed_same.batch_id,**wb_desk(db,user,day,sh)}
    # Parser/header rules can improve over time. A PREVIEW is non-authoritative, so
    # supersede an older preview and parse the exact file again using current rules.
    for old_preview in same_batches:
        if old_preview.status=='PREVIEW': old_preview.status='REPLACED'
    try:
        wb=load_workbook(BytesIO(content),data_only=True,read_only=True)
        ws,header_row,idx,header_diag=find_wb_sheet(wb,db)
    except ValueError as exc:
        raise HTTPException(422,str(exc)) from exc
    except Exception as exc:
        raise HTTPException(422,'Unable to read XLSX file.') from exc

    eq=list(db.scalars(select(Equipment)))
    vehicle_map={norm_vehicle(e.machine_id):e.machine_id for e in eq}
    vehicle_map.update({norm_vehicle(e.vehicle_no):e.machine_id for e in eq if e.vehicle_no})
    for a in db.scalars(select(VehicleAlias).where(VehicleAlias.active)):
        vehicle_map[norm_vehicle(a.alias)]=a.machine_id
    location_resolver=build_location_resolver(db)

    batch=WbImportBatch(batch_id=str(uuid4()),operating_date=day,shift=sh,file_name=file.filename or 'WB.xlsx',file_hash=digest,status='PREVIEW',valid_rows=0,review_rows=0)
    db.add(batch); db.flush(); valid=review=0
    new_wb_rows=[]
    for rno,row in enumerate(ws.iter_rows(min_row=header_row+1,values_only=True),header_row+1):
        if not any(v not in (None,'') for v in row): continue
        issues=[]; move=str(row_value(row,idx,'move') or rno).strip(); vehicle_raw=str(row_value(row,idx,'vehicle') or '').strip()
        if not vehicle_raw: issues.append('Vehicle Number missing')
        row_day=_date_value(row_value(row,idx,'date')) if idx.get('date') is not None else day
        raw_shift=str(row_value(row,idx,'shift') or '').strip().upper()
        if raw_shift and raw_shift[0:1] in {'A','B','C'} and raw_shift[0]!=sh: issues.append(f'Shift {raw_shift} != selected {sh}')
        clock=_clock_value(row_value(row,idx,'time')); weigh_at=_wb_weigh_at(day,definition,clock)
        if not weigh_at:
            issues.append('Invalid or out-of-shift GW Created Time')
        else:
            expected_day=expected_movement_date(day,clock,definition)
            if row_day and expected_day and row_day!=expected_day:
                issues.append(f'Movement Date {row_day} != expected calendar date {expected_day} for operating date {day} / Shift {sh}')
        tare=_cell_decimal(row_value(row,idx,'tare'))
        gross=_cell_decimal(row_value(row,idx,'gross'))
        net=_cell_decimal(row_value(row,idx,'net'))
        tare = tare if tare is not None else Decimal('0')
        if gross is None and net is not None: gross=tare+net
        if net is None and gross is not None and gross>=tare: net=gross-tare
        gross = gross if gross is not None else Decimal('0')
        if gross<=0 or net is None or net<=0 or gross<tare: issues.append('Invalid tare/gross/net weight')

        source=location_resolver.resolve(row_value(row,idx,'source_code'),row_value(row,idx,'source_name'),'SOURCE')
        dest=location_resolver.resolve(row_value(row,idx,'dest_code'),row_value(row,idx,'dest_name'),'DESTINATION')
        if not source: issues.append('Source location missing')
        if not dest: issues.append('Destination location missing')

        key=f'{batch.batch_id}:{move}:{rno}'
        status='REVIEW' if issues else 'VALID'
        source_code_raw=str(row_value(row,idx,'source_code') or '').strip()
        source_name_raw=str(row_value(row,idx,'source_name') or '').strip()
        dest_code_raw=str(row_value(row,idx,'dest_code') or '').strip()
        dest_name_raw=str(row_value(row,idx,'dest_name') or '').strip()
        movement=WbMovement(
            movement_key=key,batch_id=batch.batch_id,operating_date=day,shift=sh,movement_no=move,vehicle_raw=vehicle_raw,
            vehicle_id=vehicle_map.get(norm_vehicle(vehicle_raw)),
            material_code=str(row_value(row,idx,'matcode') or '').strip() or None,
            material_name=str(row_value(row,idx,'matname') or '').strip() or None,
            source_raw=source_code_raw or source_name_raw or None,destination_raw=dest_code_raw or dest_name_raw or None,
            tare_kg=tare,gross_kg=gross,net_kg=net or Decimal('0'),
            weigh_at=weigh_at or datetime.combine(day,definition.start_time,TZ),row_status=status,
            issue='; '.join(issues) or None
        )
        db.add(movement); new_wb_rows.append(movement)
        if status=='VALID': valid+=1
        else: review+=1
    batch.valid_rows=valid;batch.review_rows=review
    db.flush()
    canonicalize_wb_rows(db,new_wb_rows)
    audit(db,user,'WB_PREVIEW','wb_import_batch',batch.batch_id,{
        'file':batch.file_name,'sheet':ws.title,'valid':valid,'review':review,
        'headerMappings':header_diag.get('matched',{}),
        'duplicateHeaders':header_diag.get('duplicates',{})
    })
    db.commit()
    return {'ok':True,'message':f'WB preview loaded from {ws.title}: {valid} VALID, {review} REVIEW.','batchId':batch.batch_id,**wb_desk(db,user,day,sh)}


def confirm_wb_batch(db,user,p):
    require(user,'WB'); batch=db.get(WbImportBatch,str(p.get('batchId') or ''))
    if not batch: raise HTTPException(404,'WB batch not found.')
    require(user,shift=batch.shift); open_shift(db,batch.operating_date,batch.shift)
    if batch.status=='CONFIRMED': return {'ok':True,'message':'WB batch already confirmed.',**wb_desk(db,user,batch.operating_date,batch.shift)}
    if batch.valid_rows<=0: raise HTTPException(409,'Cannot confirm a WB batch with zero VALID trips.')
    prior=list(db.scalars(select(WbImportBatch).where(WbImportBatch.operating_date==batch.operating_date,WbImportBatch.shift==batch.shift,WbImportBatch.status=='CONFIRMED')))
    for old in prior:
        keys=list(db.scalars(select(WbMovement.movement_key).where(WbMovement.batch_id==old.batch_id)))
        if keys: db.execute(delete(LoadWbMatch).where(LoadWbMatch.movement_key.in_(keys)))
        old.status='REPLACED'
    batch.status='CONFIRMED'; batch.confirmed_at=now_local(); db.flush()
    result=auto_reconcile(db,batch.operating_date,batch.shift)
    audit(db,user,'WB_CONFIRM','wb_import_batch',batch.batch_id,{'valid':batch.valid_rows,'review':batch.review_rows,'reconciliation':result})
    return {'ok':True,'message':f'WB confirmed. {result["matched"]} auto-matched; {result["wb_unmatched"]} WB unmatched; {result["field_without_wb"]} field trips without WB.',
            'reconciliation':result,**wb_desk(db,user,batch.operating_date,batch.shift)}


# NMTPL_WB_RPC_ROSTER_V3_START
def wb_history_v3(db, user, limit=60):
    """Read-only SQL WB batch history. Does not require an operating-date argument."""
    require(user, 'WB')
    try:
        limit = int(limit or 60)
    except Exception:
        limit = 60
    limit = max(10, min(limit, 200))
    batches = list(db.scalars(
        select(WbImportBatch)
        .where(WbImportBatch.status.in_(['PREVIEW', 'CONFIRMED', 'REPLACED']))
        .order_by(
            WbImportBatch.operating_date.desc(),
            WbImportBatch.confirmed_at.desc().nullslast(),
            WbImportBatch.shift.asc(),
        )
        .limit(limit)
    ))
    rows_out = [{
        'batchId': b.batch_id,
        'date': b.operating_date.isoformat(),
        'shift': b.shift,
        'fileName': b.file_name,
        'status': b.status,
        'valid': int(b.valid_rows or 0),
        'review': int(b.review_rows or 0),
        'confirmedAt': b.confirmed_at.isoformat() if b.confirmed_at else '',
    } for b in batches]
    return {'ok': True, 'rows': rows_out, 'history': rows_out, 'items': rows_out, 'batches': rows_out, 'queue': rows_out}


def wb_monitor_v3(db, user, options=None):
    """WB/Gmail coverage monitor with no selected-date dependency."""
    require(user, 'WB')
    options = options if isinstance(options, dict) else {}
    try:
        days = int(options.get('days', 14))
    except Exception:
        days = 14
    days = max(7, min(days, 45))

    now = now_local()
    today = now.date()
    start_day = today - timedelta(days=days - 1)
    permitted = set(user.shifts.split(',')) if user.shifts else set()
    definitions = [
        d for d in db.scalars(select(ShiftMaster).where(ShiftMaster.active).order_by(ShiftMaster.start_time))
        if user.admin or 'ALL' in permitted or d.shift in permitted
    ]
    shift_ids = [d.shift for d in definitions]
    batches = list(db.scalars(select(WbImportBatch).where(
        WbImportBatch.operating_date >= start_day,
        WbImportBatch.operating_date <= today,
        WbImportBatch.shift.in_(shift_ids or ['__NONE__']),
        WbImportBatch.status.in_(['PREVIEW', 'CONFIRMED']),
    )))

    by_key = {}
    for b in batches:
        slot = by_key.setdefault((b.operating_date, b.shift), {'confirmed': None, 'preview': None})
        if b.status == 'CONFIRMED':
            old = slot['confirmed']
            old_ts = old.confirmed_at if old and old.confirmed_at else datetime.min.replace(tzinfo=TZ)
            new_ts = b.confirmed_at if b.confirmed_at else datetime.min.replace(tzinfo=TZ)
            if old is None or new_ts >= old_ts:
                slot['confirmed'] = b
        elif b.status == 'PREVIEW':
            slot['preview'] = b

    counts = {'UPDATED': 0, 'PENDING': 0, 'REVIEW': 0, 'MISSING': 0, 'NOT DUE': 0}
    coverage = []
    for offset in range(days):
        op_day = today - timedelta(days=offset)
        row = {'date': op_day.isoformat(), 'shifts': {}}
        for definition in definitions:
            slot = by_key.get((op_day, definition.shift), {})
            confirmed = slot.get('confirmed')
            preview = slot.get('preview')
            _, end_at = shift_bounds(op_day, definition)
            grace_end = end_at + timedelta(hours=1)
            if preview is not None:
                status = 'REVIEW' if int(preview.review_rows or 0) > 0 else 'PENDING'
                source = preview
            elif confirmed is not None:
                status = 'UPDATED'
                source = confirmed
            elif now < end_at:
                status = 'NOT DUE'
                source = None
            elif now < grace_end:
                status = 'PENDING'
                source = None
            else:
                status = 'MISSING'
                source = None
            counts[status] = counts.get(status, 0) + 1
            row['shifts'][definition.shift] = {
                'status': status,
                'fileName': source.file_name if source else '',
                'valid': int(source.valid_rows or 0) if source else 0,
                'review': int(source.review_rows or 0) if source else 0,
                'confirmedAt': confirmed.confirmed_at.isoformat() if confirmed and confirmed.confirmed_at else '',
                'shiftEnd': end_at.isoformat(),
                'graceEnd': grace_end.isoformat(),
            }
        coverage.append(row)

    latest = db.scalar(select(WbImportBatch).where(WbImportBatch.status == 'CONFIRMED').order_by(WbImportBatch.confirmed_at.desc().nullslast()).limit(1))
    root = Path(__file__).resolve().parents[2]
    status_path = root / 'logs' / 'wb_gmail' / 'status.json'
    config_path = root / 'config' / 'wb_gmail_config.json'
    gmail = {'state': 'UNKNOWN', 'lastCheck': '', 'pollSeconds': 180, 'messagesSeen': 0, 'counts': {}, 'archiveSuccess': None}
    try:
        cfg = json.loads(config_path.read_text(encoding='utf-8-sig')) if config_path.exists() else {}
        gmail['pollSeconds'] = int(cfg.get('poll_seconds', 180))
        gmail['archiveSuccess'] = bool(cfg.get('archive_success', False))
    except Exception:
        pass
    try:
        payload = json.loads(status_path.read_text(encoding='utf-8-sig')) if status_path.exists() else {}
        state = str(payload.get('state') or 'UNKNOWN').upper()
        raw_poll = str(payload.get('last_poll') or payload.get('updated_at') or '')
        last_dt = None
        if raw_poll:
            try:
                last_dt = datetime.fromisoformat(raw_poll.replace('Z', '+00:00'))
                if last_dt.tzinfo is None:
                    last_dt = last_dt.replace(tzinfo=TZ)
                gmail['lastCheck'] = last_dt.astimezone(TZ).strftime('%d-%b-%Y %H:%M:%S')
            except Exception:
                gmail['lastCheck'] = raw_poll
        if state == 'RUNNING' and last_dt is not None:
            stale_after = max(600, gmail['pollSeconds'] * 3 + 60)
            if (now - last_dt.astimezone(TZ)).total_seconds() > stale_after:
                state = 'STALE'
        gmail['state'] = state
        gmail['messagesSeen'] = int(payload.get('messages_seen') or 0)
        gmail['counts'] = payload.get('counts') if isinstance(payload.get('counts'), dict) else {}
    except Exception as exc:
        gmail['state'] = 'ERROR'
        gmail['error'] = str(exc)

    last_wb = None
    if latest:
        last_wb = {
            'date': latest.operating_date.isoformat(), 'shift': latest.shift,
            'fileName': latest.file_name, 'valid': int(latest.valid_rows or 0),
            'review': int(latest.review_rows or 0),
            'confirmedAt': latest.confirmed_at.isoformat() if latest.confirmed_at else '',
        }
    history = wb_history_v3(db, user, int(options.get('historyLimit', 40) or 40))['rows']
    return {'ok': True, 'now': now.isoformat(), 'days': days, 'shifts': shift_ids,
            'gmail': gmail, 'counts': counts, 'lastWb': last_wb,
            'coverage': coverage, 'history': history}
# NMTPL_WB_RPC_ROSTER_V3_END

class RPC(BaseModel):
    method: str = Field(max_length=80)
    args: list = Field(default_factory=list, max_length=5)



# NMTPL_WB_CONTROL_CENTER_V1
def wb_desk_exact(db, user, operating_date_value=None, shift_value=None):
    require(user, 'WB')
    raw_day = str(operating_date_value or '').strip()
    shift = str(shift_value or '').strip().upper()
    if not raw_day:
        raise HTTPException(422, 'Choose a valid operating date.')
    try:
        day = datetime.strptime(raw_day[:10], '%Y-%m-%d').date()
    except Exception:
        raise HTTPException(422, 'Choose a valid operating date.')
    if shift not in {'A', 'B', 'C'}:
        raise HTTPException(422, 'Choose a valid shift.')

    batches = list(db.scalars(
        select(WbImportBatch).where(
            WbImportBatch.operating_date == day,
            WbImportBatch.shift == shift,
            WbImportBatch.status.in_(['PREVIEW', 'CONFIRMED'])
        )
    ))

    previews = [b for b in batches if b.status == 'PREVIEW']
    confirmed = [b for b in batches if b.status == 'CONFIRMED']

    batch = None
    if previews:
        batch = previews[-1]
    elif confirmed:
        confirmed.sort(
            key=lambda b: b.confirmed_at or datetime.min.replace(tzinfo=TZ),
            reverse=True
        )
        batch = confirmed[0]

    if batch is None:
        return {
            'requestedDate': day.isoformat(),
            'requestedShift': shift,
            'selected': None,
            'rows': [],
        }

    if batch.operating_date != day or str(batch.shift).upper() != shift:
        return {
            'requestedDate': day.isoformat(),
            'requestedShift': shift,
            'selected': None,
            'rows': [],
            'warning': 'A mismatched WB batch was blocked by the exact Date + Shift guard.',
        }

    movements = list(db.scalars(
        select(WbMovement).where(WbMovement.batch_id == batch.batch_id)
    ))

    rows = []
    for w in movements:
        wt = ''
        if w.weigh_at:
            try:
                wt = w.weigh_at.strftime('%H:%M:%S')
            except Exception:
                wt = str(w.weigh_at)
        rows.append({
            'status': w.row_status or '',
            'vehicle': w.vehicle_raw or w.vehicle_id or '',
            'time': wt,
            'material': w.material_name or w.material_code or '',
            'source': w.source_raw or '',
            'destination': w.destination_raw or '',
            'tonnes': float((w.net_kg or Decimal('0')) / Decimal('1000')),
            'issue': w.issue or '',
            'movementKey': w.movement_key,
        })

    return {
        'requestedDate': day.isoformat(),
        'requestedShift': shift,
        'selected': {
            'batchId': batch.batch_id,
            'fileName': batch.file_name,
            'status': batch.status,
            'valid': int(batch.valid_rows or 0),
            'review': int(batch.review_rows or 0),
            'operatingDate': batch.operating_date.isoformat(),
            'shift': batch.shift,
        },
        'rows': rows,
    }


def wb_historical_queue(db, user, options=None):
    require(user, 'WB')
    import re
    root = Path(__file__).resolve().parents[2]
    report_dir = root / 'data' / 'wb_historical_backfill' / 'reports'
    summaries = sorted(report_dir.glob('historical_summary_*.txt'), reverse=True)
    if not summaries:
        return {'source': '', 'rows': [], 'counts': {}}

    source = summaries[0]
    text = source.read_text(encoding='utf-8', errors='replace')
    items = {}
    status_map = {
        'MISSING_GMAIL_AND_SQL': 'MISSING',
        'GMAIL_FOUND_NEEDS_REVIEW': 'REVIEW',
        'GMAIL_FOUND_NOT_IN_SQL': 'AVAILABLE',
        'REVIEW': 'REVIEW',
        'SQL_NOT_CONFIRMED': 'PENDING SQL',
        'SQL_ONLY_NO_GMAIL_FILE': 'CONFIRMED SQL',
        'COVERED': 'CONFIRMED',
    }

    patt = re.compile(
        r'^\s*(2026-\d{2}-\d{2}) Shift ([ABC]) \| ([A-Z_]+) \|.*?(?:candidate=(.*))?$',
        re.M
    )
    for m in patt.finditer(text):
        day, shift, raw = m.group(1), m.group(2), m.group(3)
        candidate = (m.group(4) or '').strip()
        items[(day, shift)] = {
            'date': day,
            'shift': shift,
            'status': status_map.get(raw, raw),
            'rawStatus': raw,
            'file': '' if candidate == '-' else candidate,
        }

    timeout = re.compile(
        r'^\s*2026-\d{2}-\d{2} [^|]*\| [^|]*\| ([^|]+) \| hint=(2026-\d{2}-\d{2}) / ([ABC]) \| Workbook parser exceeded 25s',
        re.M
    )
    for m in timeout.finditer(text):
        filename, day, shift = m.group(1).strip(), m.group(2), m.group(3)
        items[(day, shift)] = {
            'date': day,
            'shift': shift,
            'status': 'TIMEOUT',
            'rawStatus': 'PARSE_TIMEOUT',
            'file': filename,
        }

    rows = sorted(items.values(), key=lambda x: (x['date'], x['shift']), reverse=True)
    counts = {}
    for r in rows:
        counts[r['status']] = counts.get(r['status'], 0) + 1
    return {'source': source.name, 'rows': rows, 'counts': counts}


# ---------------- TIOM Phase 1: MIS paper reports + HSD management ----------------

def _tiom_shift_date_time(db, day, shift, value):
    raw=str(value or '').strip()
    if not raw:
        return None
    try:
        hh,mm=[int(x) for x in raw[:5].split(':')]
        tm=dtime(hh,mm)
    except Exception:
        raise HTTPException(422,'Enter time as HH:MM.')
    definition=db.get(ShiftMaster,shift)
    dt=datetime.combine(day,tm,TZ)
    if definition and definition.end_time <= definition.start_time and tm < definition.end_time:
        dt += timedelta(days=1)
    return dt


def _tiom_context(db,user,p):
    day=_parse_ui_date_v2((p or {}).get('date') or now_local().date(),'operating date')
    sh=str((p or {}).get('shift') or operating_context()[1]).strip().upper()
    definition=db.get(ShiftMaster,sh)
    if not definition or not definition.active:
        raise HTTPException(422,'Choose an active shift.')
    require(user,shift=sh)
    return day,sh,definition


def _ensure_tiom_trip_factors(db,user=None):
    existing=list(db.scalars(select(TiomTripFactor).where(TiomTripFactor.active)))
    codes={x.material_code.upper():x for x in existing}
    seed=[('OB',Decimal('40.000'),'OB quantity = trips × 40 MT'),('ROM',Decimal('45.000'),'Current ROM excavator report basis = trips × 45 MT'),('ROM_LUMPS',Decimal('45.000'),'Audited 28-Sep shift workbook: ROM lumps movement basis = trips × 45 MT'),('SPILLAGE',Decimal('45.000'),'Audited 28-Sep shift workbook: spillage movement basis = trips × 45 MT')]
    for code,value,note in seed:
        if code not in codes:
            row=TiomTripFactor(factor_id=str(uuid4()),material_code=code,factor_mt_per_trip=value,
                effective_from=date(2026,1,1),effective_to=None,active=True,notes=note,
                entered_by=(user.login_id if user else 'SYSTEM'),entered_at=now_local())
            db.add(row); db.flush(); codes[code]=row
    return list(codes.values())


def _tiom_material_factor(db, product, day):
    if not product:
        return None,None
    key=(' '.join([str(product.product_id or ''),str(product.name or '')])).upper()
    factors=list(db.scalars(select(TiomTripFactor).where(
        TiomTripFactor.active,TiomTripFactor.effective_from<=day
    ).order_by(TiomTripFactor.effective_from.desc())))
    for f in factors:
        if f.effective_to and f.effective_to < day:
            continue
        code=f.material_code.upper().strip()
        if code and (code==str(product.product_id).upper().strip() or re.search(r'(^|[^A-Z0-9])'+re.escape(code)+r'([^A-Z0-9]|$)',key)):
            return f,Decimal(f.factor_mt_per_trip)
    return None,None


def _tiom_parse_bench_rl(value):
    if value in (None, ''):
        return None
    try:
        raw=Decimal(str(value))
    except Exception:
        raise HTTPException(422,'Bench RL must be a whole number in metres.')
    if not raw.is_finite() or raw != raw.to_integral_value():
        raise HTTPException(422,'Bench RL must be a whole number in metres.')
    return int(raw)


def _tiom_bench_rl_from_location(db, source_id):
    """Read Bench RL from the canonical source itself (e.g. 15HA/RL-800).

    Source/Face locations already carry the bench level, so field users should
    never re-enter the same RL in a second section.
    """
    if not source_id:
        return None
    loc=db.get(Location,str(source_id))
    values=[str(source_id)]
    if loc:
        values.extend([str(loc.location_id or ''),str(loc.location_name or '')])
    for value in values:
        m=re.search(r'(?:^|[/\\\s_.-])RL\s*[-:=]?\s*(-?\d{2,4})(?:\D|$)',value.upper())
        if m:
            try:
                return int(m.group(1))
            except Exception:
                pass
    return None


def _tiom_location_display(label, location_id):
    label=str(label or '').strip(); location_id=str(location_id or '').strip()
    if not label:
        return location_id
    if not location_id or label.upper()==location_id.upper():
        return label
    return f'{label} · {location_id}'


def _tiom_route_mode(value):
    text_value=str(value or '').strip().upper().replace(' ','_').replace('-','_')
    aliases={'WITHWB':'WITH_WB','WITH_WB':'WITH_WB','WB':'WITH_WB',
             'WITHOUTWB':'WITHOUT_WB','WITHOUT_WB':'WITHOUT_WB','NO_WB':'WITHOUT_WB'}
    return aliases.get(text_value)


def _tiom_route_key(source_id, destination_id, route_mode):
    return f"{source_id}|{destination_id}|{route_mode}"[:180]


def _ensure_tiom_routes_from_leads(db):
    """Backfill one Route Master row for each Source/Destination/WB path.

    Lead Master can contain many Bench RL rows for the same route. Keep an
    in-memory key set so repeated Bench RL entries never add duplicate pending
    TiomRouteMaster objects before SQLAlchemy flushes the session.
    """
    existing_routes=list(db.scalars(select(TiomRouteMaster)))
    known={
        (
            str(r.source_location_id or ''),
            str(r.destination_location_id or ''),
            str(r.route_mode or '').upper(),
        )
        for r in existing_routes
    }
    created=0
    for lead in db.scalars(
        select(TiomLeadDistance)
        .where(TiomLeadDistance.active.is_(True))
        .order_by(
            TiomLeadDistance.source_location_id,
            TiomLeadDistance.destination_location_id,
            TiomLeadDistance.route_mode,
            TiomLeadDistance.bench_rl_m,
        )
    ):
        mode=_tiom_route_mode(lead.route_mode)
        if not mode:
            continue
        key=(str(lead.source_location_id or ''),str(lead.destination_location_id or ''),mode)
        if key in known:
            continue
        src=db.get(Location,lead.source_location_id); dst=db.get(Location,lead.destination_location_id)
        route_id=_tiom_route_key(lead.source_location_id,lead.destination_location_id,mode)
        db.add(TiomRouteMaster(
            route_id=route_id,
            route_name=f"{src.location_name if src else lead.source_location_id} → {dst.location_name if dst else lead.destination_location_id}",
            source_location_id=lead.source_location_id,
            destination_location_id=lead.destination_location_id,
            route_mode=mode,
            lead_basis='BENCH_RL',
            fixed_lead_km=None,
            material_scope=lead.material_scope or None,
            via_text='Production Weigh Bridge' if mode=='WITH_WB' else 'Direct / Without Production Weigh Bridge',
            active=True,entered_by='SYSTEM',entered_at=now_local()
        ))
        known.add(key)
        created+=1
    if created:
        db.flush()
    return created


def _ensure_tiom_routes_from_wb(db, wb_rows, canonical_map):
    """Auto-register confirmed WB Source → Destination pairs in Route Master.

    WB can prove the route pair but never the distance. Bench-coded mine routes
    use the Lead Matrix; internal plant/stack routes are created as FIXED with
    Lead Pending until management enters an approved KM.
    """
    existing=list(db.scalars(select(TiomRouteMaster)))
    known={(str(r.source_location_id or ''),str(r.destination_location_id or ''),str(r.route_mode or '').upper()):r for r in existing}
    created=0
    for wb in wb_rows or []:
        canon=canonical_map.get(wb.movement_key) if canonical_map else None
        src_id=canon.source_location_id if canon else None
        dst_id=canon.destination_location_id if canon else None
        if not src_id or not dst_id or src_id==dst_id:
            continue
        key=(str(src_id),str(dst_id),'WITH_WB')
        if key in known:
            continue
        src=db.get(Location,src_id); dst=db.get(Location,dst_id)
        basis='BENCH_RL' if _tiom_bench_rl_from_location(db,src_id) is not None else 'FIXED'
        route=TiomRouteMaster(
            route_id=_tiom_route_key(src_id,dst_id,'WITH_WB'),
            route_name=f"{src.location_name if src else src_id} → {dst.location_name if dst else dst_id}",
            source_location_id=src_id,destination_location_id=dst_id,route_mode='WITH_WB',
            lead_basis=basis,fixed_lead_km=None,material_scope=None,
            via_text='Production Weigh Bridge',active=True,entered_by='WB_AUTO',entered_at=now_local()
        )
        db.add(route); known[key]=route; created+=1
    if created:
        db.flush()
    return created


def _tiom_source_context_map(db,day,sh):
    return {x.source_location_id:x for x in db.scalars(select(TiomSourceContext).where(
        TiomSourceContext.operating_date==day,TiomSourceContext.shift==sh,TiomSourceContext.active.is_(True)
    ))}


def _tiom_source_context_rows(db,day,sh):
    locations={x.location_id:x for x in db.scalars(select(Location))}
    return [{
        'contextId':x.context_id,'sourceLocationId':x.source_location_id,
        'sourceLabel':locations.get(x.source_location_id).location_name if locations.get(x.source_location_id) else x.source_location_id,
        'benchRl':x.bench_rl_m,'notes':x.notes or ''
    } for x in db.scalars(select(TiomSourceContext).where(
        TiomSourceContext.operating_date==day,TiomSourceContext.shift==sh,TiomSourceContext.active.is_(True)
    ).order_by(TiomSourceContext.source_location_id))]


def _tiom_upsert_hmr(db,user,day,sh,machine_id,opening,closing,source_type,marker=''):
    if opening is None and closing is None:
        return None
    if opening is not None and closing is not None and closing < opening:
        raise HTTPException(422,f'{machine_id}: closing HMR cannot be below opening HMR.')
    meter=db.scalar(select(SiteAssetMeter).where(
        SiteAssetMeter.site_id=='TIOM',SiteAssetMeter.operating_date==day,SiteAssetMeter.shift==sh,
        SiteAssetMeter.asset_id==machine_id,SiteAssetMeter.meter_type=='HMR'
    ))
    incoming=(opening,closing)
    if meter:
        existing=(Decimal(meter.opening_reading) if meter.opening_reading is not None else None,
                  Decimal(meter.closing_reading) if meter.closing_reading is not None else None)
        if any(x is not None for x in existing) and existing!=incoming and source_type not in str(meter.source_type or ''):
            raise HTTPException(409,f'{machine_id}: a different HMR reading already exists for {day} Shift {sh}. Correct the original source instead of duplicating it.')
    else:
        meter=SiteAssetMeter(reading_id=str(uuid4()),site_id='TIOM',operating_date=day,shift=sh,
            asset_id=machine_id,meter_type='HMR',entered_by=user.login_id,entered_at=now_local())
        db.add(meter)
    meter.opening_reading=opening; meter.closing_reading=closing
    meter.usage=(closing-opening) if opening is not None and closing is not None else None
    meter.source_type=source_type; meter.remarks=marker or None
    meter.entered_by=user.login_id; meter.entered_at=now_local()
    return meter


def _tiom_drilling_summary(db,from_day,to_day,shift='ALL'):
    q=select(TiomDrillingShift).where(TiomDrillingShift.operating_date>=from_day,TiomDrillingShift.operating_date<=to_day)
    if shift!='ALL': q=q.where(TiomDrillingShift.shift==shift)
    rows=list(db.scalars(q.order_by(TiomDrillingShift.operating_date,TiomDrillingShift.shift,TiomDrillingShift.drill_set_id)))
    sets={x.drill_set_id:x for x in db.scalars(select(TiomDrillSet))}
    equipment={x.machine_id:x for x in db.scalars(select(Equipment))}
    ids=set()
    for st in sets.values(): ids.update([st.drill_machine_id,st.compressor_machine_id])
    hsd={mid:Decimal('0') for mid in ids}
    if ids:
        hq=select(HsdIssue).where(HsdIssue.operating_date>=from_day,HsdIssue.operating_date<=to_day,HsdIssue.machine_id.in_(ids))
        if shift!='ALL': hq=hq.where(HsdIssue.shift==shift)
        for x in db.scalars(hq): hsd[x.machine_id]=hsd.get(x.machine_id,Decimal('0'))+Decimal(x.litres or 0)
    agg={}
    total={'holes':0,'romObHoles':0,'bhjBhqHoles':0,'unclassifiedHoles':0,'romObMeterage':Decimal('0'),'bhjBhqMeterage':Decimal('0'),'totalMeterage':Decimal('0'),'drillHours':Decimal('0'),'compressorHours':Decimal('0'),'breakdownHours':Decimal('0'),'hsdLitres':Decimal('0')}
    for r in rows:
        st=sets.get(r.drill_set_id)
        a=agg.setdefault(r.drill_set_id,{'id':r.drill_set_id,'label':st.drill_set_name if st else r.drill_set_id,'holes':0,'romObHoles':0,'bhjBhqHoles':0,'unclassifiedHoles':0,'romObMeterage':Decimal('0'),'bhjBhqMeterage':Decimal('0'),'totalMeterage':Decimal('0'),'drillHours':Decimal('0'),'compressorHours':Decimal('0'),'breakdownHours':Decimal('0')})
        rom=Decimal(r.rom_ob_meterage or 0); bhj=Decimal(r.bhj_bhq_meterage or 0); met=rom+bhj
        dh=(Decimal(r.drill_close_hmr)-Decimal(r.drill_open_hmr)) if r.drill_open_hmr is not None and r.drill_close_hmr is not None and r.drill_close_hmr>=r.drill_open_hmr else Decimal('0')
        ch=(Decimal(r.compressor_close_hmr)-Decimal(r.compressor_open_hmr)) if r.compressor_open_hmr is not None and r.compressor_close_hmr is not None and r.compressor_close_hmr>=r.compressor_open_hmr else Decimal('0')
        bd=Decimal(r.breakdown_hours or 0)
        a['holes']+=int(r.holes or 0)
        if r.rom_ob_holes is None or r.bhj_bhq_holes is None:
            a['unclassifiedHoles']+=int(r.holes or 0)
        else:
            a['romObHoles']+=int(r.rom_ob_holes);a['bhjBhqHoles']+=int(r.bhj_bhq_holes)
        a['romObMeterage']+=rom; a['bhjBhqMeterage']+=bhj; a['totalMeterage']+=met; a['drillHours']+=dh; a['compressorHours']+=ch; a['breakdownHours']+=bd
        total['holes']+=int(r.holes or 0)
        if r.rom_ob_holes is None or r.bhj_bhq_holes is None:
            total['unclassifiedHoles']+=int(r.holes or 0)
        else:
            total['romObHoles']+=int(r.rom_ob_holes);total['bhjBhqHoles']+=int(r.bhj_bhq_holes)
        total['romObMeterage']+=rom; total['bhjBhqMeterage']+=bhj; total['totalMeterage']+=met; total['drillHours']+=dh; total['compressorHours']+=ch; total['breakdownHours']+=bd
    out=[]
    for sid,a in agg.items():
        st=sets.get(sid); fuel=(hsd.get(st.drill_machine_id,Decimal('0'))+hsd.get(st.compressor_machine_id,Decimal('0'))) if st else Decimal('0')
        a['hsdLitres']=fuel; total['hsdLitres']+=fuel
        a['meterPerHour']=(a['totalMeterage']/a['drillHours']) if a['drillHours']>0 else None
        a['litrePerMeter']=(fuel/a['totalMeterage']) if a['totalMeterage']>0 else None
        out.append({k:(round(float(v),3) if isinstance(v,Decimal) else v) for k,v in a.items()})
    total['meterPerHour']=(total['totalMeterage']/total['drillHours']) if total['drillHours']>0 else None
    total['litrePerMeter']=(total['hsdLitres']/total['totalMeterage']) if total['totalMeterage']>0 else None
    totals={k:(round(float(v),3) if isinstance(v,Decimal) else v) for k,v in total.items()}
    return {'fromDate':str(from_day),'toDate':str(to_day),'shift':shift,'rows':out,'totals':totals}


def _tiom_area_key(location_text):
    """Normalize recognised mine-area names, e.g. 3HA/RL-840 and 3_HA.

    Do not merge arbitrary unrelated locations just because the text resembles
    one another. Only a recognised <number>HA mine area is eligible.
    """
    m = re.search(r'(?<![A-Z0-9])(\d+)\s*[_ -]*HA(?=[^A-Z0-9]|$)',
                  str(location_text or '').upper())
    return f'{int(m.group(1))}HA' if m else None


def _tiom_area_route_fallback(db, source_id, dest_id, mode, bench_rl):
    """Resolve one exact area+destination+mode+bench Lead Master candidate.

    Never substitute an unapproved distance, a different WB mode, or a
    different bench. Conflicts stay unresolved; this function is read-only.
    """
    location = db.get(Location, str(source_id))
    source_area = (_tiom_area_key(source_id) or
                   _tiom_area_key(location.location_name if location else None))
    if not source_area:
        return None
    candidates = []
    routes = list(db.scalars(select(TiomRouteMaster).where(
        TiomRouteMaster.destination_location_id == dest_id,
        TiomRouteMaster.route_mode == mode,
        TiomRouteMaster.active.is_(True)
    )))
    for route in routes:
        if route.source_location_id == source_id:
            continue
        master_source = db.get(Location, route.source_location_id)
        master_area = (_tiom_area_key(route.source_location_id) or
                       _tiom_area_key(master_source.location_name if master_source else None))
        if master_area != source_area:
            continue
        basis = str(route.lead_basis or 'BENCH_RL').upper()
        if basis == 'FIXED':
            if route.fixed_lead_km is not None and Decimal(route.fixed_lead_km) > 0:
                candidates.append((route, None, Decimal(route.fixed_lead_km)))
        elif bench_rl is not None:
            rules = list(db.scalars(select(TiomLeadDistance).where(
                TiomLeadDistance.source_location_id == route.source_location_id,
                TiomLeadDistance.destination_location_id == dest_id,
                TiomLeadDistance.route_mode == mode,
                TiomLeadDistance.bench_rl_m == bench_rl,
                TiomLeadDistance.active.is_(True)
            )))
            for rule in rules:
                if rule.lead_km is not None and Decimal(rule.lead_km) > 0:
                    candidates.append((route, rule, Decimal(rule.lead_km)))
    if len(candidates) == 1:
        route, rule, km = candidates[0]
        return {'rule': rule, 'route': route, 'routeMode': mode,
                'leadKm': km, 'status': 'OK', 'sourceBasis': 'AREA_BENCH_MASTER',
                'masterSourceLocationId': route.source_location_id}
    if len(candidates) > 1:
        return {'rule': None, 'routeMode': mode, 'leadKm': None,
                'status': 'AMBIGUOUS_AREA_ROUTE'}
    return None


def _tiom_resolve_lead(db, source_id, bench_rl, dest_id, route_mode=None, wb_linked=False):
    if not source_id or not dest_id:
        return {'rule':None,'routeMode':'WITH_WB' if wb_linked else _tiom_route_mode(route_mode),
                'leadKm':None,'status':'MISSING_ROUTE'}
    mode='WITH_WB' if wb_linked else _tiom_route_mode(route_mode)
    if not mode:
        return {'rule':None,'routeMode':None,'leadKm':None,'status':'ROUTE_MODE_REQUIRED'}
    route=db.scalar(select(TiomRouteMaster).where(
        TiomRouteMaster.source_location_id==source_id,
        TiomRouteMaster.destination_location_id==dest_id,
        TiomRouteMaster.route_mode==mode,
        TiomRouteMaster.active.is_(True)
    ))
    route_rules=list(db.scalars(select(TiomLeadDistance).where(
        TiomLeadDistance.source_location_id==source_id,
        TiomLeadDistance.destination_location_id==dest_id,
        TiomLeadDistance.route_mode==mode,
        TiomLeadDistance.active.is_(True)
    )))
    if not route and not route_rules:
        fallback=_tiom_area_route_fallback(db,source_id,dest_id,mode,bench_rl)
        return fallback or {'rule':None,'routeMode':mode,'leadKm':None,'status':'ROUTE_NOT_CONFIGURED'}
    if route and str(route.lead_basis or 'BENCH_RL').upper()=='FIXED':
        if route.fixed_lead_km is not None and Decimal(route.fixed_lead_km)>0:
            return {'rule':None,'route':route,'routeMode':mode,'leadKm':Decimal(route.fixed_lead_km),'status':'OK'}
        return {'rule':None,'route':route,'routeMode':mode,'leadKm':None,'status':'FIXED_LEAD_NOT_CONFIGURED'}
    if bench_rl is None:
        return {'rule':None,'route':route,'routeMode':mode,'leadKm':None,'status':'MISSING_BENCH_RL'}
    rules=[r for r in route_rules if r.bench_rl_m==bench_rl]
    if len(rules)==1:
        r=rules[0]
        return {'rule':r,'routeMode':mode,'leadKm':Decimal(r.lead_km),'status':'OK'}
    if len(rules)>1:
        return {'rule':None,'routeMode':mode,'leadKm':None,'status':'AMBIGUOUS_LEAD'}
    # A bench-specific location may have a placeholder route without a KM;
    # fall back to an approved area master only when exactly one rule matches.
    fallback=_tiom_area_route_fallback(db,source_id,dest_id,mode,bench_rl)
    return fallback or {'rule':None,'routeMode':mode,'leadKm':None,'status':'LEAD_NOT_CONFIGURED'}

def _tiom_asset_label(e):
    bits=[e.machine_id]
    if e.door_no: bits.append(e.door_no)
    if e.vehicle_no: bits.append(e.vehicle_no)
    if e.type: bits.append(e.type)
    return ' · '.join(dict.fromkeys(str(x) for x in bits if x))


def _tiom_report_rows(db, report_id):
    rows_=list(db.scalars(select(TiomMisTripRow).where(TiomMisTripRow.report_id==report_id).order_by(TiomMisTripRow.row_no)))
    details={x.row_id:x for x in db.scalars(select(TiomMisTripDetail).where(TiomMisTripDetail.row_id.in_([r.row_id for r in rows_]))) } if rows_ else {}
    return rows_,details


def _tiom_mis_productivity_summary(db,from_day,to_day,shift='ALL'):
    """ERP productivity view for a selected date range.

    Trips/tonnes come from submitted MIS rows. HMR/working hours come from the
    central SiteAssetMeter table, so deployment-only equipment is also visible.
    """
    if to_day < from_day:
        raise HTTPException(422,'To date cannot be before From date.')
    shift=str(shift or 'ALL').upper()
    report_q=select(TiomMisReport).where(
        TiomMisReport.operating_date>=from_day,
        TiomMisReport.operating_date<=to_day,
        TiomMisReport.status=='SUBMITTED'
    )
    if shift!='ALL':
        report_q=report_q.where(TiomMisReport.shift==shift)
    reports=list(db.scalars(report_q))
    report_ids=[r.report_id for r in reports]
    rows_=list(db.scalars(select(TiomMisTripRow).where(TiomMisTripRow.report_id.in_(report_ids)))) if report_ids else []
    details={x.row_id:x for x in db.scalars(select(TiomMisTripDetail).where(TiomMisTripDetail.row_id.in_([r.row_id for r in rows_]))) } if rows_ else {}
    eq={x.machine_id:x for x in db.scalars(select(Equipment).where(Equipment.active))}
    products={x.product_id:x for x in db.scalars(select(Product))}

    meter_q=select(SiteAssetMeter).where(
        SiteAssetMeter.site_id=='TIOM',
        SiteAssetMeter.operating_date>=from_day,
        SiteAssetMeter.operating_date<=to_day,
        SiteAssetMeter.meter_type=='HMR'
    )
    if shift!='ALL':
        meter_q=meter_q.where(SiteAssetMeter.shift==shift)
    meter_rows=list(db.scalars(meter_q.order_by(
        SiteAssetMeter.operating_date,SiteAssetMeter.shift,SiteAssetMeter.entered_at
    )))
    meters_by_machine={}
    for m in meter_rows:
        meters_by_machine.setdefault(m.asset_id,[]).append(m)

    agg={}
    for r in rows_:
        d=details.get(r.row_id)
        if not d or not d.machine_id: continue
        a=agg.setdefault(d.machine_id,{
            'machineId':d.machine_id,'romTrips':0,'romQty':Decimal('0'),
            'obTrips':0,'obQty':Decimal('0'),'otherTrips':0,'otherQty':Decimal('0')
        })
        prod=products.get(d.material_id) if d.material_id else None
        txt=(' '.join([str(d.material_id or ''),str(prod.name if prod else r.material_raw or '')])).upper()
        q=Decimal(d.calculated_qty_mt or 0)
        if re.search(r'(^|[^A-Z0-9])OB([^A-Z0-9]|$)',txt):
            a['obTrips']+=1; a['obQty']+=q
        elif 'ROM' in txt:
            a['romTrips']+=1; a['romQty']+=q
        else:
            a['otherTrips']+=1; a['otherQty']+=q

    # HMR-only loader/excavator appears, but unrelated trucks, compressors,
    # drilling and graders must not be displayed as zero-production loaders.
    for mid in meters_by_machine:
        machine=eq.get(mid)
        if not machine or not (machine.group=='LOADING' or 'LOADER' in str(machine.type or '').upper() or 'EXCAV' in str(machine.type or '').upper()):
            continue
        agg.setdefault(mid,{
            'machineId':mid,'romTrips':0,'romQty':Decimal('0'),
            'obTrips':0,'obQty':Decimal('0'),'otherTrips':0,'otherQty':Decimal('0')
        })

    out=[]
    for mid,a in sorted(agg.items()):
        machine=eq.get(mid)
        if not machine or not (machine.group=='LOADING' or 'LOADER' in str(machine.type or '').upper() or 'EXCAV' in str(machine.type or '').upper()):
            continue
        meter_list=meters_by_machine.get(mid,[])
        opening=None; closing=None; hrs=Decimal('0'); shifts_worked=set(); days_worked=set()
        for m in meter_list:
            if opening is None and m.opening_reading is not None:
                opening=Decimal(m.opening_reading)
            if m.closing_reading is not None:
                closing=Decimal(m.closing_reading)
            usage=Decimal(m.usage) if m.usage is not None else None
            if usage is None and m.opening_reading is not None and m.closing_reading is not None and m.closing_reading>=m.opening_reading:
                usage=Decimal(m.closing_reading)-Decimal(m.opening_reading)
            if usage is not None and usage>=0:
                hrs+=usage
            shifts_worked.add((m.operating_date,m.shift))
            days_worked.add(m.operating_date)

        total_qty=a['romQty']+a['obQty']+a['otherQty']
        total_trips=a['romTrips']+a['obTrips']+a['otherTrips']
        productivity=(total_qty/hrs) if hrs>0 and total_qty>0 else None
        e=eq.get(mid)
        out.append({
            **a,
            'label':_tiom_asset_label(e) if e else mid,
            'type':e.type if e else '',
            'group':e.group if e else '',
            'openingHmr':float(opening) if opening is not None else None,
            'closingHmr':float(closing) if closing is not None else None,
            'hrsRun':float(hrs),
            'shiftsWorked':len(shifts_worked),
            'daysWorked':len(days_worked),
            'totalTrips':total_trips,
            'totalQty':float(total_qty),
            'romQty':float(a['romQty']),
            'obQty':float(a['obQty']),
            'otherQty':float(a['otherQty']),
            'productivity':float(productivity) if productivity is not None else None
        })

    # Driver/tripper productivity belongs to a different equipment category.
    # Keep vehicle trips on their parent reports and never attribute them as
    # loading-machine trips simply because HMR was recorded for the truck.
    vehicle_agg={}
    for report in reports:
        meter_hrs=None
        if report.opening_kmr is not None and report.closing_kmr is not None and report.closing_kmr>=report.opening_kmr:
            meter_hrs=float(report.closing_kmr-report.opening_kmr)
        target=vehicle_agg.setdefault(report.vehicle_id,{'vehicleId':report.vehicle_id,'trips':0,'tonnes':Decimal('0'),'unweighedTrips':0,
                                                           'operatorIds':set(),'kmRun':0.0,'kmKnown':False})
        if report.operator_id:target['operatorIds'].add(report.operator_id)
        if meter_hrs is not None:target['kmRun']+=meter_hrs;target['kmKnown']=True
    report_for_id={r.report_id:r for r in reports}
    for row in rows_:
        parent=report_for_id.get(row.report_id)
        if not parent:continue
        v=vehicle_agg[parent.vehicle_id]
        detail=details.get(row.row_id)
        v['trips']+=1
        if detail and detail.calculated_qty_mt is not None:
            v['tonnes']+=Decimal(detail.calculated_qty_mt)
        else:v['unweighedTrips']+=1
    vehicle_rows=[{'vehicleId':x['vehicleId'],'label':_tiom_asset_label(eq[x['vehicleId']]) if x['vehicleId'] in eq else x['vehicleId'],
                   'trips':x['trips'],'tonnes':round(float(x['tonnes']),2),'unweighedTrips':x['unweighedTrips'],
                   'operatorIds':sorted(x['operatorIds']),'kmRun':round(x['kmRun'],2) if x['kmKnown'] else None,
                   'averagePayloadMt':round(float(x['tonnes'])/(x['trips']-x['unweighedTrips']),2) if x['trips']>x['unweighedTrips'] else None}
                  for x in vehicle_agg.values()]
    vehicle_rows.sort(key=lambda x:(-x['trips'],x['vehicleId']))
    totals={
        'hrsRun':sum((Decimal(str(x['hrsRun'])) for x in out),Decimal('0')),
        'romTrips':sum(x['romTrips'] for x in out),
        'romQty':sum((Decimal(str(x['romQty'])) for x in out),Decimal('0')),
        'obTrips':sum(x['obTrips'] for x in out),
        'obQty':sum((Decimal(str(x['obQty'])) for x in out),Decimal('0')),
        'totalTrips':sum(x['totalTrips'] for x in out),
        'totalQty':sum((Decimal(str(x['totalQty'])) for x in out),Decimal('0'))
    }
    totals['productivity']=(totals['totalQty']/totals['hrsRun']) if totals['hrsRun']>0 else None
    return {
        'fromDate':str(from_day),'toDate':str(to_day),'shift':shift,
        'rows':out,'vehicleRows':vehicle_rows,
        'totals':{k:(float(v) if isinstance(v,Decimal) else v) for k,v in totals.items()}
    }


def _tiom_mis_shift_summary(db,day,sh):
    return _tiom_mis_productivity_summary(db,day,day,sh)

def _tiom_source_deployments(db, day, sh):
    locations={x.location_id:x for x in db.scalars(select(Location))}
    equipment={x.machine_id:x for x in db.scalars(select(Equipment))}
    people={x.employee_id:x for x in db.scalars(select(Person))}
    source_contexts=_tiom_source_context_map(db,day,sh)
    meters={x.asset_id:x for x in db.scalars(select(SiteAssetMeter).where(
        SiteAssetMeter.site_id=='TIOM',SiteAssetMeter.operating_date==day,
        SiteAssetMeter.shift==sh,SiteAssetMeter.meter_type=='HMR'
    ))}
    rows=list(db.scalars(select(TiomSourceDeployment).where(
        TiomSourceDeployment.operating_date==day,
        TiomSourceDeployment.shift==sh,
        TiomSourceDeployment.active.is_(True)
    ).order_by(TiomSourceDeployment.from_at, TiomSourceDeployment.source_location_id, TiomSourceDeployment.machine_id)))
    out=[]
    for r in rows:
        loc=locations.get(r.source_location_id); eq=equipment.get(r.machine_id); meter=meters.get(r.machine_id)
        reset=bool(meter and str(meter.remarks or '').startswith('[METER_RESET]'))
        meter_note=(str(meter.remarks or '').replace('[METER_RESET]','',1).strip() if reset else (meter.remarks or '' if meter else ''))
        out.append({
            'deploymentId':r.deployment_id,
            'sourceLocationId':r.source_location_id,
            'sourceLabel':(f'{loc.location_name} · {loc.location_id}' if loc else r.source_location_id),
            'benchRl':_tiom_bench_rl_from_location(db,r.source_location_id) or (source_contexts.get(r.source_location_id).bench_rl_m if source_contexts.get(r.source_location_id) else None),
            'machineId':r.machine_id,
            'operatorId':r.operator_id or '',
            'operatorName':people[r.operator_id].name if r.operator_id in people else '',
            'machineLabel':(_tiom_asset_label(eq) if eq else r.machine_id),
            'activity':r.activity or 'EXCAVATION',
            'fromTime':r.from_at.astimezone(TZ).strftime('%H:%M') if r.from_at else '',
            'toTime':r.to_at.astimezone(TZ).strftime('%H:%M') if r.to_at else '',
            'openingHmr':float(meter.opening_reading) if meter and meter.opening_reading is not None else None,
            'closingHmr':float(meter.closing_reading) if meter and meter.closing_reading is not None else None,
            'hrsRun':float(meter.usage) if meter and meter.usage is not None else None,
            'meterReset':reset,
            'meterNote':meter_note,
            'notes':r.notes or ''
        })
    return out


def save_tiom_source_deployments(db,user,p):
    require(user,'PRODUCTION'); p=p or {}; day,sh,_=_tiom_context(db,user,p)
    incoming=p.get('rows') if isinstance(p.get('rows'),list) else []
    incoming_contexts=p.get('contexts') if isinstance(p.get('contexts'),list) else None
    parsed=[]; meter_by_machine={}; seen=set()
    if incoming_contexts is not None:
        contexts=[]; seen_sources=set()
        for idx,item in enumerate(incoming_contexts[:50],start=1):
            if not isinstance(item,dict): continue
            source_id=str(item.get('sourceLocationId') or '').strip()
            if not source_id and item.get('benchRl') in (None,''): continue
            if not source_id: raise HTTPException(422,f'Face row {idx}: select Source.')
            active_resource(db,Location,source_id)
            bench=_tiom_parse_bench_rl(item.get('benchRl'))
            if bench is None: raise HTTPException(422,f'Face row {idx}: Bench RL is required.')
            if source_id in seen_sources: raise HTTPException(409,f'Face row {idx}: Source {source_id} appears more than once. Enter Bench RL once per source / shift.')
            seen_sources.add(source_id); contexts.append((source_id,bench,short(str(item.get('notes') or ''))))
        db.execute(delete(TiomSourceContext).where(TiomSourceContext.operating_date==day,TiomSourceContext.shift==sh))
        for source_id,bench,notes in contexts:
            db.add(TiomSourceContext(context_id=str(uuid4()),operating_date=day,shift=sh,source_location_id=source_id,bench_rl_m=bench,active=True,notes=notes,entered_by=user.login_id,entered_at=now_local()))
        db.flush()
    def dec(value,label):
        if value in (None,''): return None
        try: return Decimal(str(value))
        except Exception: raise HTTPException(422,f'Enter valid {label}.')
    for idx,item in enumerate(incoming[:50],start=1):
        if not isinstance(item,dict): continue
        source_id=str(item.get('sourceLocationId') or '').strip()
        machine_id=str(item.get('machineId') or '').strip()
        operator_id=str(item.get('operatorId') or '').strip() or None
        if not source_id and not machine_id: continue
        if not source_id: raise HTTPException(422,f'Deployment row {idx}: select source.')
        if not machine_id: raise HTTPException(422,f'Deployment row {idx}: select equipment / machine.')
        source=active_resource(db,Location,source_id)
        machine=active_resource(db,Equipment,machine_id)
        if operator_id: active_resource(db,Person,operator_id)
        if machine.group in {'TRANSPORT','HSD_TANKER'}: raise HTTPException(422,f'Deployment row {idx}: choose working HMR equipment; transport/tanker equipment belongs in its own operational entry.')
        activity=short(str(item.get('activity') or '').strip().upper())[:40]
        if not activity: raise HTTPException(422,f'Deployment row {idx}: activity is required.')
        active_resource(db,ActivityMaster,activity)
        from_at=_tiom_shift_date_time(db,day,sh,item.get('fromTime'))
        to_at=_tiom_shift_date_time(db,day,sh,item.get('toTime'))
        if from_at and to_at and to_at < from_at:
            raise HTTPException(422,f'Deployment row {idx}: To time must be after From time within the operating shift.')
        opening=dec(item.get('openingHmr'),f'HMR opening on deployment row {idx}')
        closing=dec(item.get('closingHmr'),f'HMR closing on deployment row {idx}')
        reset=bool(item.get('meterReset'))
        meter_note=short(str(item.get('meterNote') or item.get('notes') or ''))
        if reset and not meter_note: raise HTTPException(422,f'Deployment row {idx}: enter a reason for HMR reset / rollover.')
        if opening is not None and closing is not None and closing < opening and not reset:
            raise HTTPException(422,f'Deployment row {idx}: closing HMR is below opening HMR. Tick Reset and enter the reset/rollover reason if this is genuine.')
        meter_value=(opening,closing,reset,meter_note)
        if machine.machine_id in meter_by_machine:
            old=meter_by_machine[machine.machine_id]
            if any(v is not None for v in (opening,closing)) or reset:
                if old!=meter_value and any(v is not None for v in old[:2]):
                    raise HTTPException(409,f'{machine.machine_id}: HMR is shift-level data. Enter one consistent opening/closing reading for this machine.')
        elif any(v is not None for v in (opening,closing)) or reset:
            meter_by_machine[machine.machine_id]=meter_value
        key=(source.location_id,machine.machine_id,from_at.isoformat() if from_at else '',to_at.isoformat() if to_at else '')
        if key in seen: raise HTTPException(409,f'Deployment row {idx}: duplicate source/machine period.')
        seen.add(key)
        parsed.append((source,machine,activity,from_at,to_at,short(str(item.get('notes') or '')),operator_id))
    if not parsed: raise HTTPException(422,'Add at least one source-machine deployment.')
    if incoming_contexts is not None:
        context_ids=set(_tiom_source_context_map(db,day,sh))
        missing=sorted({x[0].location_id for x in parsed if x[0].location_id not in context_ids})
        if missing: raise HTTPException(422,'Enter Bench RL once in Face / Bench Setup for: '+', '.join(missing))
    # Save submitted assignments individually; do not delete unrelated
    # machine allocations or meters when a field user updates one row.
    before=_tiom_source_deployments(db,day,sh)
    all_existing=list(db.scalars(select(TiomSourceDeployment).where(
        TiomSourceDeployment.operating_date==day,TiomSourceDeployment.shift==sh,
        TiomSourceDeployment.active.is_(True))))
    existing_by_key={}
    for record in all_existing:
        key=(record.source_location_id,record.machine_id,record.from_at,record.to_at)
        existing_by_key.setdefault(key,record)
    # Never allocate the same machine to two overlapping locations. A null
    # time is a whole-shift reservation and must be explicitly closed before a transfer.
    combined=list(all_existing)
    def conflicts(a_from,a_to,b_from,b_to):
        return not (a_from and a_to and b_from and b_to and
                    (a_to<=b_from or b_to<=a_from))
    for source,machine,activity,from_at,to_at,notes,operator_id in parsed:
        key=(source.location_id,machine.machine_id,from_at,to_at)
        existing=existing_by_key.get(key)
        for other in combined:
            if other is existing or other.machine_id!=machine.machine_id or other.source_location_id==source.location_id:
                continue
            if conflicts(from_at,to_at,other.from_at,other.to_at):
                raise HTTPException(409,f'{machine.machine_id} already allocated to {other.source_location_id} in this shift. Close the earlier time range before transferring.')
        entry=existing or TiomSourceDeployment(
            deployment_id=str(uuid4()),operating_date=day,shift=sh,
            source_location_id=source.location_id,machine_id=machine.machine_id,
            entered_by=user.login_id,entered_at=now_local())
        entry.operator_id=operator_id
        entry.activity=activity;entry.from_at=from_at;entry.to_at=to_at
        entry.notes=notes;entry.active=True;entry.entered_by=user.login_id;entry.entered_at=now_local()
        if not existing:
            db.add(entry);combined.append(entry);existing_by_key[key]=entry
    for machine_id,(opening,closing,reset,meter_note) in meter_by_machine.items():
        usage=(closing-opening) if opening is not None and closing is not None and closing>=opening and not reset else None
        meter=db.scalar(select(SiteAssetMeter).where(
            SiteAssetMeter.site_id=='TIOM',SiteAssetMeter.operating_date==day,SiteAssetMeter.shift==sh,
            SiteAssetMeter.asset_id==machine_id,SiteAssetMeter.meter_type=='HMR'
        ))
        if not meter:
            meter=SiteAssetMeter(reading_id=str(uuid4()),site_id='TIOM',operating_date=day,shift=sh,asset_id=machine_id,meter_type='HMR',entered_by=user.login_id,entered_at=now_local())
            db.add(meter)
        meter.opening_reading=opening; meter.closing_reading=closing; meter.usage=usage
        meter.source_type='TIOM_DEPLOYMENT'; meter.remarks=(('[METER_RESET] '+meter_note).strip() if reset else meter_note)
        meter.entered_by=user.login_id; meter.entered_at=now_local()
    db.flush()
    after=_tiom_source_deployments(db,day,sh)
    audit(db,user,'TIOM_SOURCE_DEPLOY','tiom_source_deployment',f'{day}:{sh}',{'before':before,'after':after,'hmrMachines':sorted(meter_by_machine)})
    return {'ok':True,'message':f'{len(after)} source-machine deployment row(s) saved. HMR stored at shift-machine level.','rows':after}


def _tiom_management_recipients(db):
    return [x.value for x in db.scalars(select(MasterOption).where(MasterOption.category=='TIOM_REPORT_EMAIL').order_by(MasterOption.value)) if x.value]


def save_tiom_management_recipients(db,user,p):
    require(user,admin=True); p=p or {}
    raw=p.get('recipients')
    values=raw if isinstance(raw,list) else re.split(r'[,;\s]+',str(raw or ''))
    cleaned=[]
    for value in values:
        email=str(value or '').strip().lower()
        if not email: continue
        if len(email)>80 or not re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+',email): raise HTTPException(422,f'Invalid email address: {email}')
        if email not in cleaned: cleaned.append(email)
    if len(cleaned)>20: raise HTTPException(422,'Maximum 20 management recipients.')
    db.execute(delete(MasterOption).where(MasterOption.category=='TIOM_REPORT_EMAIL'))
    for email in cleaned: db.add(MasterOption(category='TIOM_REPORT_EMAIL',value=email))
    audit(db,user,'TIOM_REPORT_RECIPIENTS','master_options','TIOM_REPORT_EMAIL',{'recipients':cleaned})
    return {'ok':True,'message':'Management recipients saved.','recipients':cleaned}


def send_tiom_management_image(db,user,p):
    require(user,'PRODUCTION'); p=p or {}
    recipients=_tiom_management_recipients(db)
    if not recipients: raise HTTPException(422,'Management recipients are not configured. Ask an administrator to add them in the report panel.')
    image_b64=str(p.get('imageBase64') or '').strip()
    if image_b64.startswith('data:image/png;base64,'): image_b64=image_b64.split(',',1)[1]
    if not image_b64 or len(image_b64)>10_000_000: raise HTTPException(422,'Report image is missing or too large.')
    try:
        image_bytes=base64.b64decode(image_b64,validate=True)
    except Exception as exc:
        raise HTTPException(422,'Invalid PNG image data.') from exc
    if not image_bytes.startswith(b'\x89PNG\r\n\x1a\n'): raise HTTPException(422,'Only PNG report images are accepted.')
    root=Path(__file__).resolve().parents[2]
    token=root/'config'/'gmail_token.json'
    if not token.exists(): raise HTTPException(503,'Gmail is not authorised on this server. Run authorize_wb_gmail.bat once.')
    try:
        from google.oauth2.credentials import Credentials
        from googleapiclient.discovery import build
        creds=Credentials.from_authorized_user_file(str(token))
        service=build('gmail','v1',credentials=creds,cache_discovery=False)
        subject=short(str(p.get('subject') or 'TIOM Management Report'))[:180]
        filename=re.sub(r'[^A-Za-z0-9_.-]+','_',str(p.get('fileName') or 'TIOM_Report.png'))[:120] or 'TIOM_Report.png'
        note=str(p.get('message') or 'TIOM management report attached as image.')[:1000]
        msg=EmailMessage(); msg['To']=', '.join(recipients); msg['Subject']=subject
        msg.set_content(note); msg.add_attachment(image_bytes,maintype='image',subtype='png',filename=filename)
        raw=base64.urlsafe_b64encode(msg.as_bytes()).decode('ascii')
        sent=service.users().messages().send(userId='me',body={'raw':raw}).execute()
    except HTTPException:
        raise
    except Exception as exc:
        log.exception('TIOM report email failed')
        raise HTTPException(502,'Gmail could not send the report image. Check Gmail authorisation and internet connection.') from exc
    audit(db,user,'TIOM_REPORT_EMAIL','management_report',sent.get('id',''),{'recipients':recipients,'subject':subject,'fileName':filename})
    return {'ok':True,'message':f'Report image sent to {len(recipients)} management recipient(s).','messageId':sent.get('id'),'recipients':recipients}


def remove_tiom_mis_report(db,user,p):
    require(user,'PRODUCTION'); p=p or {}
    report=db.get(TiomMisReport,str(p.get('reportId') or '').strip())
    if not report: raise HTTPException(404,'MIS report not found.')
    require(user,shift=report.shift)
    if report.status=='DRAFT':
        rows_,_= _tiom_report_rows(db,report.report_id)
        ids=[r.row_id for r in rows_]
        if ids:
            db.execute(delete(TiomMisReconciliation).where(TiomMisReconciliation.row_id.in_(ids)))
            db.execute(delete(TiomMisTripDetail).where(TiomMisTripDetail.row_id.in_(ids)))
        db.execute(delete(TiomMisTripRow).where(TiomMisTripRow.report_id==report.report_id))
        audit(db,user,'TIOM_MIS_DELETE_DRAFT','tiom_mis_report',report.report_id,{'vehicle':report.vehicle_id,'date':str(report.operating_date),'shift':report.shift})
        db.delete(report)
        return {'ok':True,'message':'Draft report deleted.'}
    require(user,admin=True)
    if report.status=='VOID': return {'ok':True,'message':'Report is already void.'}
    report.status='VOID'; report.version=(report.version or 1)+1
    reason=short(str(p.get('reason') or 'Voided by management'))
    report.notes=((report.notes or '')+'\nVOID: '+reason).strip()
    audit(db,user,'TIOM_MIS_VOID','tiom_mis_report',report.report_id,{'reason':reason})
    return {'ok':True,'message':'Submitted report marked VOID. Historical record retained.'}


def tiom_mis_desk(db,user,p):
    require(user,'PRODUCTION'); p=p or {}; day,sh,_=_tiom_context(db,user,p)
    factors=_ensure_tiom_trip_factors(db,user)
    _ensure_tiom_routes_from_leads(db)

    # Central Activity Master: make common HMR/deployment activities available
    # without creating a separate TIOM-only activity list.
    for activity in ('EXCAVATION','LOADING','REHANDLING','LEVELLING','DOZING','GRADING','DRILLING','CRUSHING','SCREENING','CLEANUP','SUPPORT','OTHER'):
        if not db.get(ActivityMaster,activity):
            db.add(ActivityMaster(activity=activity,vehicle_required=False,active=True))
    db.flush()

    equipment=list(db.scalars(select(Equipment).where(Equipment.active).order_by(Equipment.machine_id)))
    persons=list(db.scalars(select(Person).where(Person.active).order_by(Person.name)))
    locations=list(db.scalars(select(Location).where(Location.active).order_by(Location.location_name)))
    source_locations=location_options(db,'SOURCE')
    destination_locations=location_options(db,'DESTINATION')
    deployment_locations=list(source_locations)
    registered={x['id'] for x in deployment_locations}
    for location in locations:
        label=(str(location.location_name or '')+' '+str(location.location_id)).upper()
        if (location.location_type in {'SCREEN','CRUSHER','PLANT'} or re.search(r'\bMSP\s*-?\s*\d+\b',label)) and location.location_id not in registered:
            deployment_locations.append({'id':location.location_id,'label':location.location_name})
            registered.add(location.location_id)
    products=list(db.scalars(select(Product).where(Product.active).order_by(Product.name)))
    activities=list(db.scalars(select(ActivityMaster).where(ActivityMaster.active).order_by(ActivityMaster.activity)))
    lead_rules=list(db.scalars(select(TiomLeadDistance).where(TiomLeadDistance.active).order_by(
        TiomLeadDistance.source_location_id,TiomLeadDistance.destination_location_id,TiomLeadDistance.bench_rl_m,TiomLeadDistance.route_mode
    )))

    report_from=_parse_ui_date_v2(p.get('reportFrom'),'report from date') if p.get('reportFrom') else day.replace(day=1)
    report_to=_parse_ui_date_v2(p.get('reportTo'),'report to date') if p.get('reportTo') else day
    if report_to < report_from:
        raise HTTPException(422,'Report To date cannot be before From date.')
    report_shift=str(p.get('reportShift') or 'ALL').upper()
    valid_shifts={x.shift for x in db.scalars(select(ShiftMaster).where(ShiftMaster.active))}
    if report_shift!='ALL' and report_shift not in valid_shifts:
        raise HTTPException(422,'Choose a valid report shift.')

    report_q=select(TiomMisReport).where(
        TiomMisReport.operating_date>=report_from,
        TiomMisReport.operating_date<=report_to
    )
    if report_shift!='ALL':
        report_q=report_q.where(TiomMisReport.shift==report_shift)
    reports=list(db.scalars(report_q.order_by(
        TiomMisReport.operating_date.desc(),TiomMisReport.shift,TiomMisReport.entered_at.desc()
    ).limit(500)))

    # Carry forward ONLY readings from chronologically earlier shifts.
    # The previous implementation included current/future shifts on the same
    # date and ordered by edit timestamp, which could silently cross-fill HMR.
    shift_order={x.shift:i for i,x in enumerate(sorted(
        db.scalars(select(ShiftMaster).where(ShiftMaster.active.is_(True))),
        key=lambda x:x.start_time))}
    preceding={name for name,idx in shift_order.items() if idx<shift_order.get(sh,0)}
    def prior_meter_date(row):
        return row.operating_date<day or (row.operating_date==day and row.shift in preceding)
    def previous_key(row):
        return (row.operating_date,shift_order.get(row.shift,-1),
                row.entered_at.isoformat() if row.entered_at else '')
    historical_reports=[r for r in db.scalars(select(TiomMisReport).where(
        TiomMisReport.closing_kmr.is_not(None),TiomMisReport.operating_date<=day,
        TiomMisReport.status=='SUBMITTED')) if prior_meter_date(r)]
    prev_kmr={}
    for r in sorted(historical_reports,key=previous_key,reverse=True):
        prev_kmr.setdefault(r.vehicle_id,float(r.closing_kmr))
    historical_meters=[m for m in db.scalars(select(SiteAssetMeter).where(
        SiteAssetMeter.site_id=='TIOM',SiteAssetMeter.meter_type=='HMR',
        SiteAssetMeter.closing_reading.is_not(None),SiteAssetMeter.operating_date<=day
    )) if prior_meter_date(m)]
    prev_hmr={}
    for m in sorted(historical_meters,key=previous_key,reverse=True):
        prev_hmr.setdefault(m.asset_id,float(m.closing_reading))

    trip_machines=[e for e in equipment if e.group=='LOADING']
    deployment_machines=[e for e in equipment if e.group not in {'TRANSPORT','HSD_TANKER'}]
    summary_today=_tiom_mis_productivity_summary(db,day,day,sh)
    summary_range=_tiom_mis_productivity_summary(db,report_from,report_to,report_shift)

    return {'date':str(day),'shift':sh,
        'reportFrom':str(report_from),'reportTo':str(report_to),'reportShift':report_shift,
        'vehicles':[{
            'id':e.machine_id,'label':_tiom_asset_label(e),
            'previousKmr':prev_kmr.get(e.machine_id),'previousHmr':prev_hmr.get(e.machine_id)
        } for e in equipment if e.group=='TRANSPORT'],
        'machines':[{
            'id':e.machine_id,'label':_tiom_asset_label(e),'group':e.group,'type':e.type,
            'previousHmr':prev_hmr.get(e.machine_id)
        } for e in trip_machines],
        'deploymentMachines':[{
            'id':e.machine_id,'label':_tiom_asset_label(e),'group':e.group,'type':e.type,
            'previousHmr':prev_hmr.get(e.machine_id)
        } for e in deployment_machines],
        'activities':[{'id':x.activity,'label':x.activity,'vehicleRequired':x.vehicle_required} for x in activities],
        'routes':[{'id':x.route_id,'label':x.route_name,'sourceLocationId':x.source_location_id,'destinationLocationId':x.destination_location_id,'routeMode':x.route_mode,'materialScope':x.material_scope or ''} for x in db.scalars(select(TiomRouteMaster).where(TiomRouteMaster.active.is_(True)).order_by(TiomRouteMaster.route_name))],
        'leadRules':[{
            'id':x.lead_id,'sourceLocationId':x.source_location_id,'benchRl':x.bench_rl_m,
            'destinationLocationId':x.destination_location_id,'routeMode':x.route_mode,
            'leadKm':float(x.lead_km),'materialScope':x.material_scope or ''
        } for x in lead_rules],
        'deploymentLocations':[{'id':x['id'],'label':_tiom_location_display(x['label'],x['id'])} for x in deployment_locations],
        'sourceContexts':_tiom_source_context_rows(db,day,sh),  # legacy compatibility only
        'operators':[{'id':x.employee_id,'label':f'{x.name} · {x.employee_id} · {x.role}'} for x in persons],
        'locations':[{'id':x.location_id,'label':_tiom_location_display(x.location_name,x.location_id),'role':x.role or 'UNCLASSIFIED'} for x in locations],
        'sourceLocations':[{'id':x['id'],'label':_tiom_location_display(x['label'],x['id']),'benchRl':_tiom_bench_rl_from_location(db,x['id'])} for x in source_locations],
        'destinationLocations':[{'id':x['id'],'label':_tiom_location_display(x['label'],x['id'])} for x in destination_locations],
        'products':[{'id':x.product_id,'label':f'{x.name} · {x.product_id}'} for x in products],
        'factors':[{'id':x.factor_id,'materialCode':x.material_code,'factor':float(x.factor_mt_per_trip),'effectiveFrom':str(x.effective_from),'effectiveTo':str(x.effective_to) if x.effective_to else '','active':x.active,'notes':x.notes or ''} for x in factors],
        'deployments':_tiom_source_deployments(db,day,sh),
        'managementRecipients':_tiom_management_recipients(db),
        'reports':[{
            'reportId':x.report_id,'date':str(x.operating_date),'shift':x.shift,
            'vehicleId':x.vehicle_id,'operatorId':x.operator_id or '','status':x.status,
            'paperRef':x.paper_ref or '',
            'openingKmr':float(x.opening_kmr) if x.opening_kmr is not None else None,
            'closingKmr':float(x.closing_kmr) if x.closing_kmr is not None else None,
            'openingHmr':float(x.opening_hmr) if x.opening_hmr is not None else None,
            'closingHmr':float(x.closing_hmr) if x.closing_hmr is not None else None,
            'enteredAt':x.entered_at.strftime('%d-%m %H:%M') if x.entered_at else ''
        } for x in reports],
        'summary':summary_today,
        'summaryToday':summary_today,
        'summaryRange':summary_range}

def get_tiom_wb_suggestions(db,user,p):
    """Confirmed WB movements available to prefill one driver/tripper report."""
    require(user,'PRODUCTION'); p=p or {}; day,sh,_=_tiom_context(db,user,p)
    vehicle_id=str(p.get('vehicleId') or '').strip()
    if not vehicle_id: return {'rows':[],'batch':None,'message':'Select a tripper/dumper.'}
    vehicle=active_resource(db,Equipment,vehicle_id)
    if vehicle.group!='TRANSPORT': raise HTTPException(422,'Choose a tripper/dumper vehicle.')
    report_id=str(p.get('reportId') or '').strip()
    own_row_ids=set()
    if report_id:
        own_row_ids=set(db.scalars(select(TiomMisTripRow.row_id).where(TiomMisTripRow.report_id==report_id)))
    batch,wb_rows=tiom_authoritative_wb(db,day,sh)
    if not batch: return {'rows':[],'batch':None,'message':'No confirmed WB batch for this date / shift.'}
    reconciliations=list(db.scalars(select(TiomMisReconciliation).where(TiomMisReconciliation.wb_movement_key.is_not(None))))
    linked={r.wb_movement_key:r for r in reconciliations if r.wb_movement_key}
    deployments=list(db.scalars(select(TiomSourceDeployment).where(
        TiomSourceDeployment.operating_date==day,TiomSourceDeployment.shift==sh,TiomSourceDeployment.active.is_(True)
    )))
    aliases={norm_vehicle(x.alias) for x in db.scalars(select(VehicleAlias).where(
        VehicleAlias.machine_id==vehicle.machine_id,VehicleAlias.active.is_(True)))}
    out=[];other_linked=0;other_vehicle=0
    for wb in wb_rows:
        matched=tiom_wb_vehicle_matches(wb,vehicle) or bool(
            aliases.intersection({norm_vehicle(wb.vehicle_raw),norm_vehicle(wb.vehicle_id)}))
        if not matched:
            other_vehicle+=1;continue
        rec=linked.get(wb.movement_key)
        if rec and rec.row_id not in own_row_ids:
            other_linked+=1;continue
        item=tiom_wb_payload(db,wb)
        source_id=item.get('sourceLocationId') or ''
        machines=sorted({d.machine_id for d in deployments if d.source_location_id==source_id}) if source_id else []
        item['machineId']=machines[0] if len(machines)==1 else ''
        item['machineChoices']=machines
        item['linkedToThisReport']=bool(rec and rec.row_id in own_row_ids)
        out.append(item)
    return {'rows':out,'batch':{'batchId':batch.batch_id,'fileName':batch.file_name,'confirmedAt':batch.confirmed_at.isoformat() if batch.confirmed_at else ''},
            'diagnostics':{'batchValid':len(wb_rows),'otherVehicle':other_vehicle,'alreadyLinked':other_linked,
                           'available':len(out),'unmappedRows':sum(not (x.get('materialId') and x.get('sourceLocationId') and x.get('destinationLocationId')) for x in out)},
            'message':f'{len(out)} available; {other_linked} already linked to another MIS report; {other_vehicle} WB movements belong to other vehicles. Latest confirmed batch: {batch.file_name}.'}

def get_tiom_mis_report(db,user,p):
    require(user,'PRODUCTION'); report=db.get(TiomMisReport,str((p or {}).get('reportId') or ''))
    if not report: raise HTTPException(404,'MIS report not found.')
    require(user,shift=report.shift)
    rows_,details=_tiom_report_rows(db,report.report_id)
    row_ids=[r.row_id for r in rows_]
    recs={r.row_id:r for r in db.scalars(select(TiomMisReconciliation).where(TiomMisReconciliation.row_id.in_(row_ids)))} if row_ids else {}
    leads={r.row_id:r for r in db.scalars(select(TiomMisTripLead).where(TiomMisTripLead.row_id.in_(row_ids)))} if row_ids else {}
    wb_keys=[r.wb_movement_key for r in recs.values() if r.wb_movement_key]
    wbmap={w.movement_key:w for w in db.scalars(select(WbMovement).where(WbMovement.movement_key.in_(wb_keys)))} if wb_keys else {}
    meters=list(db.scalars(select(SiteAssetMeter).where(SiteAssetMeter.site_id=='TIOM',SiteAssetMeter.operating_date==report.operating_date,SiteAssetMeter.shift==report.shift,SiteAssetMeter.source_type=='TIOM_MIS')))
    trip_rows=[]
    for r in rows_:
        d=details.get(r.row_id); rec=recs.get(r.row_id); wb=wbmap.get(rec.wb_movement_key) if rec and rec.wb_movement_key else None; lead=leads.get(r.row_id)
        trip_rows.append({'rowNo':r.row_no,'loadingTime':r.loading_at.strftime('%H:%M') if r.loading_at else '','unloadingTime':r.unloading_at.strftime('%H:%M') if r.unloading_at else '',
            'materialId':d.material_id if d else '','sourceLocationId':d.source_location_id if d else '','destinationLocationId':d.destination_location_id if d else '','machineId':d.machine_id if d else '',
            'factor':float(d.factor_mt_per_trip) if d and d.factor_mt_per_trip is not None else None,'qtyMt':float(d.calculated_qty_mt) if d and d.calculated_qty_mt is not None else None,'remarks':r.remarks or '',
            'benchRl':lead.bench_rl_m if lead else None,'routeMode':lead.route_mode if lead else '','leadKm':float(lead.lead_km) if lead and lead.lead_km is not None else None,'leadStatus':lead.lead_status if lead else 'NOT_CAPTURED',
            'wbMovementKey':rec.wb_movement_key if rec and rec.wb_movement_key else '', 'wbMovementNo':wb.movement_no if wb else '', 'quantitySource':'WB' if wb else 'TRIP_FACTOR'})
    hmr_run=(report.closing_hmr-report.opening_hmr) if report.opening_hmr is not None and report.closing_hmr is not None and report.closing_hmr>=report.opening_hmr else None
    return {'reportId':report.report_id,'date':str(report.operating_date),'shift':report.shift,'vehicleId':report.vehicle_id,'operatorId':report.operator_id or '',
        'openingKmr':float(report.opening_kmr) if report.opening_kmr is not None else None,
        'closingKmr':float(report.closing_kmr) if report.closing_kmr is not None else None,
        'openingHmr':float(report.opening_hmr) if report.opening_hmr is not None else None,
        'closingHmr':float(report.closing_hmr) if report.closing_hmr is not None else None,
        'hmrRun':float(hmr_run) if hmr_run is not None else None,
        'paperRef':report.paper_ref or '','notes':report.notes or '','status':report.status,
        'rows':trip_rows,
        'meters':[], 'hmrSource':'CENTRAL_SITE_ASSET_METER'}

def save_tiom_mis_report(db,user,p,submit=False):
    require(user,'PRODUCTION'); p=p or {}; day,sh,_=_tiom_context(db,user,p); open_shift(db,day,sh); _ensure_tiom_trip_factors(db,user)
    vehicle=active_resource(db,Equipment,str(p.get('vehicleId') or '').strip())
    if vehicle.group!='TRANSPORT': raise HTTPException(422,'Choose a tripper/dumper vehicle.')
    operator_id=str(p.get('operatorId') or '').strip() or None
    if operator_id: active_resource(db,Person,operator_id)
    report_id=str(p.get('reportId') or '').strip(); report=db.get(TiomMisReport,report_id) if report_id else None
    if report and report.status in {'SUBMITTED','VOID'}: raise HTTPException(409,'Submitted/void report is immutable. Void the submitted report and create a new corrected report.')
    if not report:
        report=TiomMisReport(report_id=str(uuid4()),operating_date=day,shift=sh,vehicle_id=vehicle.machine_id,operator_id=operator_id,status='DRAFT',entered_by=user.login_id,entered_at=now_local(),version=1)
        db.add(report); db.flush()
    else:
        if report.operating_date!=day or report.shift!=sh: raise HTTPException(409,'Report date/shift cannot be changed. Create a new report.')
        report.version=(report.version or 1)+1; report.vehicle_id=vehicle.machine_id; report.operator_id=operator_id
    def dec(v,label):
        if v in (None,''): return None
        try:return Decimal(str(v))
        except:raise HTTPException(422,f'Enter valid {label}.')
    report.opening_kmr=dec(p.get('openingKmr'),'opening KMR'); report.closing_kmr=dec(p.get('closingKmr'),'closing KMR')
    if report.opening_kmr is not None and report.closing_kmr is not None and report.closing_kmr < report.opening_kmr:
        raise HTTPException(422,'Closing KMR cannot be below opening KMR.')
    report.opening_hmr=dec(p.get('openingHmr'),'opening HMR'); report.closing_hmr=dec(p.get('closingHmr'),'closing HMR')
    if report.opening_hmr is not None and report.closing_hmr is not None and report.closing_hmr < report.opening_hmr:
        raise HTTPException(422,'Closing HMR cannot be below opening HMR.')

    # Driver HMR uses the same central shift-level meter ledger as Shift Deployment.
    if report.opening_hmr is not None or report.closing_hmr is not None:
        marker=f'[DRIVER_REPORT:{report.report_id}]'
        meter=db.scalar(select(SiteAssetMeter).where(
            SiteAssetMeter.site_id=='TIOM',SiteAssetMeter.operating_date==day,SiteAssetMeter.shift==sh,
            SiteAssetMeter.asset_id==vehicle.machine_id,SiteAssetMeter.meter_type=='HMR'
        ))
        if meter:
            existing_marker=str(meter.remarks or '')
            same_report=marker in existing_marker
            incoming=(report.opening_hmr,report.closing_hmr)
            existing=(Decimal(meter.opening_reading) if meter.opening_reading is not None else None,
                      Decimal(meter.closing_reading) if meter.closing_reading is not None else None)
            if not same_report and any(x is not None for x in existing) and existing!=incoming:
                raise HTTPException(409,f'{vehicle.machine_id}: HMR already exists for {day} Shift {sh}. Use the existing central reading or correct the original source.')
        else:
            meter=SiteAssetMeter(
                reading_id=str(uuid4()),site_id='TIOM',operating_date=day,shift=sh,
                asset_id=vehicle.machine_id,meter_type='HMR',entered_by=user.login_id,entered_at=now_local()
            )
            db.add(meter)
        meter.opening_reading=report.opening_hmr
        meter.closing_reading=report.closing_hmr
        meter.usage=(report.closing_hmr-report.opening_hmr) if report.opening_hmr is not None and report.closing_hmr is not None else None
        meter.source_type='TIOM_DRIVER_REPORT'
        meter.remarks=marker
        meter.entered_by=user.login_id
        meter.entered_at=now_local()

    report.paper_ref=short(str(p.get('paperRef') or '')); report.notes=short(str(p.get('notes') or ''))
    old_rows=list(db.scalars(select(TiomMisTripRow).where(TiomMisTripRow.report_id==report.report_id)))
    if old_rows:
        ids=[r.row_id for r in old_rows]
        db.execute(delete(TiomMisReconciliation).where(TiomMisReconciliation.row_id.in_(ids)))
        db.execute(delete(TiomMisTripLead).where(TiomMisTripLead.row_id.in_(ids)))
        db.execute(delete(TiomMisTripDetail).where(TiomMisTripDetail.row_id.in_(ids)))
        db.execute(delete(TiomMisTripRow).where(TiomMisTripRow.report_id==report.report_id))
        db.flush()
    rows_payload=p.get('rows') if isinstance(p.get('rows'),list) else []
    deployments=list(db.scalars(select(TiomSourceDeployment).where(
        TiomSourceDeployment.operating_date==day,TiomSourceDeployment.shift==sh,TiomSourceDeployment.active.is_(True)
    )))
    batch,wb_rows=tiom_authoritative_wb(db,day,sh)
    wbmap={w.movement_key:w for w in wb_rows}
    already_linked={r.wb_movement_key:r.row_id for r in db.scalars(select(TiomMisReconciliation).where(TiomMisReconciliation.wb_movement_key.is_not(None)))}
    source_contexts=_tiom_source_context_map(db,day,sh)
    kept=0; linked_count=0; lead_ok_count=0; lead_missing_count=0
    for i,item in enumerate(rows_payload[:80],start=1):
        if not isinstance(item,dict): continue
        wb_key=str(item.get('wbMovementKey') or '').strip()
        material_id=str(item.get('materialId') or '').strip(); source_id=str(item.get('sourceLocationId') or '').strip(); dest_id=str(item.get('destinationLocationId') or '').strip(); machine_id=str(item.get('machineId') or '').strip(); lt=str(item.get('loadingTime') or '').strip(); ut=str(item.get('unloadingTime') or '').strip()
        if not any([wb_key,material_id,source_id,dest_id,machine_id,lt,ut]): continue
        wb=wbmap.get(wb_key) if wb_key else None
        if wb_key and not wb: raise HTTPException(409,f'Row {i}: selected WB movement is not in the active confirmed WB batch for {day} Shift {sh}.')
        if wb and not tiom_wb_vehicle_matches(wb,vehicle): raise HTTPException(422,f'Row {i}: WB movement {wb.movement_no} belongs to a different vehicle.')
        if wb_key and wb_key in already_linked: raise HTTPException(409,f'Row {i}: WB movement {wb.movement_no} is already linked to another MIS trip.')
        if wb:
            suggested=tiom_wb_payload(db,wb)
            material_id=material_id or str(suggested.get('materialId') or '')
            source_id=source_id or str(suggested.get('sourceLocationId') or '')
            dest_id=dest_id or str(suggested.get('destinationLocationId') or '')
        if not all([material_id,source_id,dest_id]): raise HTTPException(422,f'Row {i}: material, source and destination are required. Map the WB master values if they were not recognised automatically.')
        prod=active_resource(db,Product,material_id); active_resource(db,Location,source_id); active_resource(db,Location,dest_id)
        dep_machines=[d.machine_id for d in deployments if d.source_location_id==source_id]
        if not machine_id and len(set(dep_machines))==1:
            machine_id=dep_machines[0]
        if not machine_id:
            raise HTTPException(422,f'Row {i}: choose the loader / excavator used at {source_id}.')
        machine=active_resource(db,Equipment,machine_id)
        if machine.group in {'TRANSPORT','HSD_TANKER'}:
            raise HTTPException(422,f'Row {i}: choose working loading equipment, not transport/tanker equipment.')
        loading=_tiom_shift_date_time(db,day,sh,lt) if lt else None; unloading=_tiom_shift_date_time(db,day,sh,ut) if ut else None
        if loading and unloading and unloading<loading:
            unloading+=timedelta(days=1)
        dep_matches=[d for d in deployments if d.source_location_id==source_id and d.machine_id==machine_id]
        if not dep_matches:
            activity='EXCAVATION' if 'EXCAV' in str(machine.type or '').upper() else 'LOADING'
            if not db.get(ActivityMaster,activity):
                db.add(ActivityMaster(activity=activity,vehicle_required=False,active=True))
                db.flush()
            auto_dep=TiomSourceDeployment(
                deployment_id=str(uuid4()),operating_date=day,shift=sh,
                source_location_id=source_id,machine_id=machine_id,activity=activity,
                from_at=None,to_at=None,active=True,
                notes='Auto-created from MIS trip entry; HMR/auxiliary details can be added in Machine Setup.',
                entered_by=user.login_id,entered_at=now_local()
            )
            db.add(auto_dep); db.flush()
            deployments.append(auto_dep)
        if wb:
            factor=None; qty=Decimal(wb.net_kg or 0)/Decimal('1000')
            material_raw=wb.material_name or wb.material_code or prod.name
            source_raw=wb.source_raw or source_id; destination_raw=wb.destination_raw or dest_id
        else:
            factor_row,factor=_tiom_material_factor(db,prod,day); qty=factor if factor is not None else None
            material_raw=prod.name; source_raw=source_id; destination_raw=dest_id
        row_id=str(uuid4()); r=TiomMisTripRow(row_id=row_id,report_id=report.report_id,row_no=i,loading_at=loading,unloading_at=unloading,material_raw=material_raw,source_raw=source_raw,destination_raw=destination_raw,remarks=short(str(item.get('remarks') or '')),created_at=now_local())
        db.add(r); db.flush()
        db.add(TiomMisTripDetail(row_id=row_id,machine_id=machine_id,material_id=material_id,source_location_id=source_id,destination_location_id=dest_id,factor_mt_per_trip=factor,calculated_qty_mt=qty,entered_at=now_local()))
        source_ctx=source_contexts.get(source_id)
        bench_rl=_tiom_bench_rl_from_location(db,source_id)
        if bench_rl is None and source_ctx:
            bench_rl=source_ctx.bench_rl_m
        requested_route='WITH_WB' if wb else 'WITHOUT_WB'
        lead_result=_tiom_resolve_lead(db,source_id,bench_rl,dest_id,requested_route,wb_linked=bool(wb))
        lead_rule=lead_result.get('rule')
        db.add(TiomMisTripLead(
            row_id=row_id,bench_rl_m=bench_rl,route_mode=lead_result.get('routeMode'),
            lead_km=lead_result.get('leadKm'),lead_rule_id=lead_rule.lead_id if lead_rule else None,
            lead_status=lead_result.get('status') or 'NOT_CONFIGURED',entered_at=now_local()
        ))
        if lead_result.get('status')=='OK': lead_ok_count+=1
        else: lead_missing_count+=1
        if wb:
            db.flush()
            db.add(TiomMisReconciliation(reconciliation_id=str(uuid4()),row_id=row_id,field_trip_id=None,wb_movement_key=wb.movement_key,match_status='MATCHED',confidence=100,reason='Linked from confirmed WB movement in MIS driver entry.',reconciled_by=user.login_id,reconciled_at=now_local()))
            linked_count+=1
        kept+=1
    if kept==0: raise HTTPException(422,'Add at least one trip row.')
    # Compatibility only: old cached Phase 1.x clients may still send `meters`.
    # The current UI no longer has a separate HMR table; any legacy payload is
    # migrated into the same shift-level SiteAssetMeter used by Shift Deployment.
    meters=p.get('meters') if isinstance(p.get('meters'),list) else []
    deployed_machine_ids={d.machine_id for d in deployments}
    legacy_hmr=0
    for item in meters[:100]:
        if not isinstance(item,dict) or not str(item.get('assetId') or '').strip(): continue
        asset=active_resource(db,Equipment,str(item.get('assetId')).strip())
        if asset.machine_id not in deployed_machine_ids: raise HTTPException(422,f'{asset.machine_id}: save the machine in Shift Deployment before entering HMR.')
        opening=dec(item.get('opening'),'opening HMR'); closing=dec(item.get('closing'),'closing HMR')
        if opening is None and closing is None: continue
        if opening is not None and closing is not None and closing<opening: raise HTTPException(422,f'{asset.machine_id}: closing HMR cannot be below opening HMR in legacy entry. Use the Shift Deployment Reset option for a genuine reset/rollover.')
        usage=(closing-opening) if opening is not None and closing is not None else None
        meter=db.scalar(select(SiteAssetMeter).where(SiteAssetMeter.site_id=='TIOM',SiteAssetMeter.operating_date==day,SiteAssetMeter.shift==sh,SiteAssetMeter.asset_id==asset.machine_id,SiteAssetMeter.meter_type=='HMR'))
        if not meter:
            meter=SiteAssetMeter(reading_id=str(uuid4()),site_id='TIOM',operating_date=day,shift=sh,asset_id=asset.machine_id,meter_type='HMR',entered_by=user.login_id,entered_at=now_local()); db.add(meter)
        meter.opening_reading=opening; meter.closing_reading=closing; meter.usage=usage; meter.source_type='TIOM_DEPLOYMENT'; meter.remarks=short(str(item.get('remarks') or 'Legacy MIS HMR migrated to Shift Deployment')); meter.entered_by=user.login_id; meter.entered_at=now_local(); legacy_hmr+=1
    report.status='SUBMITTED' if submit else 'DRAFT'
    if submit: report.submitted_by=user.login_id; report.submitted_at=now_local()
    audit(db,user,'TIOM_MIS_SUBMIT' if submit else 'TIOM_MIS_DRAFT','tiom_mis_report',report.report_id,{'date':str(day),'shift':sh,'vehicle':vehicle.machine_id,'rows':kept,'wbLinkedRows':linked_count,'leadResolvedRows':lead_ok_count,'leadMissingRows':lead_missing_count,'legacyHmrMigrated':legacy_hmr})
    db.flush()
    return {'ok':True,'message':('MIS shift report submitted.' if submit else 'MIS draft saved.'),'reportId':report.report_id,'status':report.status,'wbLinkedRows':linked_count,'leadResolvedRows':lead_ok_count,'leadMissingRows':lead_missing_count,'summary':_tiom_mis_shift_summary(db,day,sh)}

def save_tiom_trip_factor(db,user,p):
    require(user,admin=True); p=p or {}; code=str(p.get('materialCode') or '').strip().upper()
    if not re.fullmatch(r'[A-Z0-9 _/.-]{1,60}',code): raise HTTPException(422,'Enter a valid material code.')
    try:value=Decimal(str(p.get('factorMtPerTrip')))
    except:raise HTTPException(422,'Enter a valid MT/trip factor.')
    if value<=0: raise HTTPException(422,'Factor must be greater than zero.')
    eff=_parse_ui_date_v2(p.get('effectiveFrom') or now_local().date(),'effective date')
    for r in db.scalars(select(TiomTripFactor).where(TiomTripFactor.material_code==code,TiomTripFactor.active)):
        if r.effective_to is None and r.effective_from<=eff: r.effective_to=eff-timedelta(days=1)
    row=TiomTripFactor(factor_id=str(uuid4()),material_code=code,factor_mt_per_trip=value,effective_from=eff,effective_to=None,active=True,notes=short(str(p.get('notes') or '')),entered_by=user.login_id,entered_at=now_local()); db.add(row)
    audit(db,user,'TIOM_TRIP_FACTOR','tiom_trip_factor',row.factor_id,{'material':code,'factor':str(value),'effectiveFrom':str(eff)})
    return {'ok':True,'message':f'{code} factor saved at {value} MT/trip.'}



# ---------------- TIOM Phase 1.4: plant/shifting entry + automatic shift production report ----------------
TIOM_SHIFT_REPORT_LINES = [
    ('EXCAVATION','TOTAL_EXCAVATION','Total Excavation',True),
    ('EXCAVATION','ROM','ROM',False),
    ('EXCAVATION','ROM_LUMPS','ROM LUMPS',False),
    ('EXCAVATION','SUBGRADE_DUMP','SUBGRADE DUMP',False),
    ('EXCAVATION','SUBGRADE_FEED_PLANT','SUBGRADE FEED PLANT',False),
    ('EXCAVATION','WASTE','WASTE',False),
    ('EXCAVATION','ROM_STOCK_YARD','ROM STOCK YARD',False),
    ('EXCAVATION','ROM_STOCK_TO_PLANT_FEED','ROM STOCK TO PLANT FEED',False),
    ('EXCAVATION','SPILLAGE','Spillage',False),
    ('PRODUCTION (PROCESSED ORE)','TOTAL_PRODUCTION','Total Production',True),
    ('PRODUCTION (PROCESSED ORE)','SCREEN_FINES','Screen Fines',False),
    ('PRODUCTION (PROCESSED ORE)','SCREEN_5_18','Screen 5-18MM',False),
    ('PRODUCTION (PROCESSED ORE)','LUMPS_FROM_SCREEN','Lumps from Screen',False),
    ('PRODUCTION (PROCESSED ORE)','LUMPS_SHIFTED_TO_STOCK','Lumps Shifted to Stock',True),
    ('CRUSHER PRODUCTION','LUMPS_TO_CRUSHER_FROM_STOCK','Lumps shifted to Crusher from Stock',False),
    ('CRUSHER PRODUCTION','LUMPS_FEED_TO_CRUSHER','Lumps feed to Crusher',False),
    ('CRUSHER PRODUCTION','CRUSHER_FINES','Crusher Fines',False),
    ('CRUSHER PRODUCTION','CRUSHER_5_18','Crusher 5-18mm',False),
    ('RE-SCREENING OUT PUT SHIFTING','FIVE_40_FEED_TO_PLANT','5-40mm feed to Plant',True),
    ('RE-SCREENING OUT PUT SHIFTING','RE_SCREEN_FINES','Re Screen Fines',False),
    ('RE-SCREENING OUT PUT SHIFTING','RE_SCREEN_5_18','Re Screen 5-18 mm',False),
    ('STACKING','SCREEN_FINES_SHIFTED','Screen Fines shifted',False),
    ('STACKING','CRUSHER_FINES_SHIFTED','Crusher Fines shifted',False),
    ('STACKING','SC_5_18_SHIFTED','SC 5-18 mm',False),
    ('RE-HANDLING','UNSCREENED_5_18','UNSCREENED 5-18(58-60)',False),
    ('RE-HANDLING','PROJECT_AREA_FINES_TO_STACK','Project Area Fines to Stack',False),
    ('RE-HANDLING','SHIFT_5_40_TANKURA','5-40mm shifted (TANKURA)',False),
]

TIOM_SHIFT_MOVEMENTS = [
    ('ROM_STOCK_YARD','ROM to Stock Yard',None),
    ('ROM_STOCK_TO_PLANT_FEED','ROM Stock to Plant Feed',None),
    ('SPILLAGE','Spillage','SPILLAGE'),
    ('SCREEN_FINES','Screen Fines / Plant Output',None),
    ('SCREEN_5_18','Screen 5-18 MM / Plant Output',None),
    ('LUMPS_FROM_SCREEN','Lumps from Screen','ROM_LUMPS'),
    ('LUMPS_TO_CRUSHER_FROM_STOCK','Lumps shifted to Crusher from Stock','ROM_LUMPS'),
    ('LUMPS_FEED_TO_CRUSHER','Lumps feed to Crusher',None),
    ('CRUSHER_FINES','Crusher Fines',None),
    ('CRUSHER_5_18','Crusher 5-18 mm',None),
    ('RE_SCREEN_FINES','Re Screen Fines',None),
    ('RE_SCREEN_5_18','Re Screen 5-18 mm',None),
    ('SCREEN_FINES_SHIFTED','Screen Fines shifted for Stacking',None),
    ('CRUSHER_FINES_SHIFTED','Crusher Fines shifted for Stacking',None),
    ('SC_5_18_SHIFTED','SC 5-18 shifted for Stacking',None),
    ('UNSCREENED_5_18','Unscreened 5-18 (58-60)',None),
    ('PROJECT_AREA_FINES_TO_STACK','Project Area Fines to Stack',None),
    ('SHIFT_5_40_TANKURA','5-40 mm shifted (TANKURA)',None),
]
TIOM_SHIFT_MOVEMENT_MAP = {x[0]:x for x in TIOM_SHIFT_MOVEMENTS}
TIOM_SHIFT_ORDER = {'GENERAL':0,'A':1,'B':2,'C':3}


def _tiom_shift_report_derive(values):
    out={k:Decimal(str(v or 0)) for k,v in (values or {}).items()}
    out['TOTAL_EXCAVATION']=sum(out.get(k,Decimal('0')) for k in ['ROM','ROM_LUMPS','SUBGRADE_DUMP','SUBGRADE_FEED_PLANT','WASTE'])
    # Final saleable/stack production only. Screen lumps are intermediate crusher feed.
    out['TOTAL_PRODUCTION']=sum(out.get(k,Decimal('0')) for k in ['SCREEN_FINES','SCREEN_5_18','CRUSHER_FINES','CRUSHER_5_18'])
    out['LUMPS_SHIFTED_TO_STOCK']=max(Decimal('0'),out.get('LUMPS_FROM_SCREEN',Decimal('0'))-out.get('LUMPS_FEED_TO_CRUSHER',Decimal('0')))
    out['FIVE_40_FEED_TO_PLANT']=out.get('RE_SCREEN_FINES',Decimal('0'))+out.get('RE_SCREEN_5_18',Decimal('0'))
    return out


def _tiom_factor_by_code(db, code, day):
    row=db.scalar(select(TiomTripFactor).where(
        TiomTripFactor.material_code==str(code).upper(),TiomTripFactor.active,
        TiomTripFactor.effective_from<=day,
        (TiomTripFactor.effective_to.is_(None)) | (TiomTripFactor.effective_to>=day)
    ).order_by(TiomTripFactor.effective_from.desc()).limit(1))
    return Decimal(row.factor_mt_per_trip) if row else None


def _tiom_classify_mis_line(product, source, destination):
    p=(' '.join([getattr(product,'product_id','') or '',getattr(product,'name','') or ''])).upper()
    src=(' '.join([getattr(source,'location_id','') or '',getattr(source,'location_name','') or ''])).upper()
    dst=(' '.join([getattr(destination,'location_id','') or '',getattr(destination,'location_name','') or ''])).upper()
    if re.search(r'(^|[^A-Z0-9])(OB|WASTE)([^A-Z0-9]|$)',p): return 'WASTE'
    if 'SUBGRADE' in p or re.search(r'(^|[^A-Z0-9])SG([^A-Z0-9]|$)',p):
        return 'SUBGRADE_FEED_PLANT' if any(x in dst for x in ['FEED','PLANT','MSP']) else 'SUBGRADE_DUMP'
    if 'LUMP' in p and 'ROM' in p:
        # Raw ROM lumps hauled from mine/excavation belong under Excavation.
        # Lumps coming from MSP/screen/plant are Processed Ore output.
        if any(x in src for x in ['MSP','SCREEN','PLANT']): return 'LUMPS_FROM_SCREEN'
        return 'ROM_LUMPS'
    if 'ROM' in p:
        if 'STOCK' in src and any(x in dst for x in ['PLANT','FEED','MSP']): return 'ROM_STOCK_TO_PLANT_FEED'
        if 'STOCK' in dst: return 'ROM_STOCK_YARD'
        return 'ROM'
    if 'SPILL' in p or 'SPILL' in dst: return 'SPILLAGE'
    return None


def _tiom_shift_ftd_bundle(db, day, sh, include_draft=False):
    vals={code:Decimal('0') for _,code,_,_ in TIOM_SHIFT_REPORT_LINES}
    sources={code:{} for _,code,_,_ in TIOM_SHIFT_REPORT_LINES}
    warnings=[]
    def add(code,qty,source):
        if not code or code not in vals: return
        q=Decimal(qty or 0)
        if q==0: return
        vals[code]=vals.get(code,Decimal('0'))+q
        sources.setdefault(code,{})[source]=sources.setdefault(code,{}).get(source,Decimal('0'))+q

    # 1) Confirmed WB is authoritative for weighed movements.
    batch,wb_rows=tiom_authoritative_wb(db,day,sh)
    for wb in wb_rows:
        for code,qty in tiom_wb_report_contributions(wb): add(code,qty,'WB')

    # 2) Submitted MIS supplies only non-WB trips (OB/internal haulage/etc.).
    reps=list(db.scalars(select(TiomMisReport).where(TiomMisReport.operating_date==day,TiomMisReport.shift==sh,TiomMisReport.status=='SUBMITTED')))
    ids=[r.report_id for r in reps]
    if ids:
        rows_=list(db.scalars(select(TiomMisTripRow).where(TiomMisTripRow.report_id.in_(ids))))
        row_ids=[r.row_id for r in rows_]
        details={x.row_id:x for x in db.scalars(select(TiomMisTripDetail).where(TiomMisTripDetail.row_id.in_(row_ids)))} if row_ids else {}
        recs={x.row_id:x for x in db.scalars(select(TiomMisReconciliation).where(TiomMisReconciliation.row_id.in_(row_ids)))} if row_ids else {}
        pmap={x.product_id:x for x in db.scalars(select(Product))}; lmap={x.location_id:x for x in db.scalars(select(Location))}
        for r in rows_:
            # WB-linked MIS rows are evidence/metadata only; WB already supplied quantity.
            if recs.get(r.row_id) and recs[r.row_id].wb_movement_key: continue
            d=details.get(r.row_id)
            if not d: continue
            code=_tiom_classify_mis_line(pmap.get(d.material_id),lmap.get(d.source_location_id),lmap.get(d.destination_location_id))
            if code: add(code,Decimal(d.calculated_qty_mt or 0),'MIS')

    # 3) Manual Plant Output / Exception entry is only a fallback. If WB or MIS already
    # supplies the same line, ignore the old/manual duplicate and surface a warning.
    report=db.scalar(select(TiomShiftProductionReport).where(TiomShiftProductionReport.operating_date==day,TiomShiftProductionReport.shift==sh))
    if report and (include_draft or report.status=='SUBMITTED'):
        for m in db.scalars(select(TiomShiftProductionMovement).where(TiomShiftProductionMovement.report_id==report.report_id)):
            q=Decimal(m.qty_mt or 0)
            if vals.get(m.movement_code,Decimal('0'))>0:
                warnings.append(f'{m.movement_code}: manual exception ignored because WB/MIS already supplies this report line.')
                continue
            add(m.movement_code,q,'MANUAL')

    vals=_tiom_shift_report_derive(vals)
    for code in ['TOTAL_EXCAVATION','TOTAL_PRODUCTION','LUMPS_SHIFTED_TO_STOCK','FIVE_40_FEED_TO_PLANT']:
        if vals.get(code,Decimal('0')):
            sources[code]={'CALC':vals[code]}
    return {'values':vals,'sources':sources,'warnings':warnings,'wbBatchId':batch.batch_id if batch else None,'wbRows':len(wb_rows)}


def _tiom_shift_ftd(db, day, sh, include_draft=False):
    return _tiom_shift_ftd_bundle(db,day,sh,include_draft)['values']

def _tiom_baseline_map(db, day):
    rows_=list(db.scalars(select(TiomShiftReportBaseline).where(TiomShiftReportBaseline.effective_date<=day).order_by(TiomShiftReportBaseline.effective_date)))
    out={}; dates={}
    for r in rows_:
        if r.line_code not in dates or r.effective_date>=dates[r.line_code]:
            out[r.line_code]=Decimal(r.opening_qty_mt or 0); dates[r.line_code]=r.effective_date
    base_date=max(dates.values()) if dates else day
    return out,base_date


def _tiom_shift_before(d1,s1,d2,s2):
    return d1<d2 or (d1==d2 and TIOM_SHIFT_ORDER.get(str(s1).upper(),99)<TIOM_SHIFT_ORDER.get(str(s2).upper(),99))


def _tiom_shift_report_data(db, day, sh, include_draft=True):
    baseline,base_date=_tiom_baseline_map(db,day)
    previous={code:Decimal(baseline.get(code,0)) for _,code,_,_ in TIOM_SHIFT_REPORT_LINES}
    periods=set()
    for r in db.scalars(select(TiomShiftProductionReport).where(TiomShiftProductionReport.operating_date>=base_date,TiomShiftProductionReport.operating_date<=day,TiomShiftProductionReport.status=='SUBMITTED')):
        if _tiom_shift_before(r.operating_date,r.shift,day,sh): periods.add((r.operating_date,r.shift))
    for r in db.scalars(select(TiomMisReport).where(TiomMisReport.operating_date>=base_date,TiomMisReport.operating_date<=day,TiomMisReport.status=='SUBMITTED')):
        if _tiom_shift_before(r.operating_date,r.shift,day,sh): periods.add((r.operating_date,r.shift))
    for r in db.scalars(select(WbImportBatch).where(WbImportBatch.operating_date>=base_date,WbImportBatch.operating_date<=day,WbImportBatch.status=='CONFIRMED',WbImportBatch.confirmed_at.is_not(None))):
        if _tiom_shift_before(r.operating_date,r.shift,day,sh): periods.add((r.operating_date,r.shift))
    for pd,ps in sorted(periods,key=lambda x:(x[0],TIOM_SHIFT_ORDER.get(str(x[1]).upper(),99))):
        f=_tiom_shift_ftd(db,pd,ps,False)
        for _,code,_,_ in TIOM_SHIFT_REPORT_LINES: previous[code]=previous.get(code,Decimal('0'))+f.get(code,Decimal('0'))
    bundle=_tiom_shift_ftd_bundle(db,day,sh,include_draft); ftd=bundle['values']; source_map=bundle['sources']
    lines=[]
    for sec,code,label,derived in TIOM_SHIFT_REPORT_LINES:
        prev=previous.get(code,Decimal('0')); cur=ftd.get(code,Decimal('0')); parts=source_map.get(code,{})
        source='+'.join(k for k in ['WB','MIS','MANUAL','CALC'] if parts.get(k)) or '—'
        lines.append({'section':sec,'code':code,'label':label,'derived':derived,'previous':float(prev),'ftd':float(cur),'cumulative':float(prev+cur),
                      'source':source,'sourceBreakdown':{k:float(v) for k,v in parts.items()}})
    report=db.scalar(select(TiomShiftProductionReport).where(TiomShiftProductionReport.operating_date==day,TiomShiftProductionReport.shift==sh))
    return {'date':str(day),'shift':sh,'status':report.status if report else 'NEW','reportId':report.report_id if report else '',
            'remarks':report.remarks or '' if report else '','baselineDate':str(base_date),'lines':lines,
            'warnings':bundle['warnings'],'wbBatchId':bundle['wbBatchId'],'wbRows':bundle['wbRows']}

def _tiom_shift_range_report(db, from_day, to_day, report_shift='ALL'):
    if to_day < from_day:
        raise HTTPException(422,'Report To date cannot be before From date.')
    report_shift=str(report_shift or 'ALL').upper()
    active_shifts=[x.shift for x in db.scalars(select(ShiftMaster).where(ShiftMaster.active).order_by(ShiftMaster.start_time))]
    shifts=active_shifts if report_shift=='ALL' else [report_shift]
    totals={code:Decimal('0') for _,code,_,_ in TIOM_SHIFT_REPORT_LINES}
    sources={code:{} for _,code,_,_ in TIOM_SHIFT_REPORT_LINES}
    warnings=[]; wb_rows=0; periods=0
    cur=from_day
    while cur<=to_day:
        for sh in shifts:
            bundle=_tiom_shift_ftd_bundle(db,cur,sh,False)
            periods+=1; wb_rows+=int(bundle.get('wbRows') or 0)
            for code,val in (bundle.get('values') or {}).items():
                if code in totals: totals[code]+=Decimal(str(val or 0))
            for code,parts in (bundle.get('sources') or {}).items():
                for src,val in (parts or {}).items():
                    sources.setdefault(code,{})[src]=sources.setdefault(code,{}).get(src,Decimal('0'))+Decimal(str(val or 0))
            warnings.extend(bundle.get('warnings') or [])
        cur+=timedelta(days=1)
    lines=[]
    for sec,code,label,derived in TIOM_SHIFT_REPORT_LINES:
        parts=sources.get(code,{})
        source='+'.join(k for k in ['WB','MIS','MANUAL','CALC'] if parts.get(k)) or '—'
        lines.append({'section':sec,'code':code,'label':label,'derived':derived,'qty':float(totals.get(code,0)),
                      'source':source,'sourceBreakdown':{k:float(v) for k,v in parts.items()}})
    return {'fromDate':str(from_day),'toDate':str(to_day),'shift':report_shift,'periods':periods,'wbRows':wb_rows,
            'lines':lines,'warnings':sorted(set(warnings))}


def get_tiom_shift_production_desk(db,user,p):
    require(user,'PRODUCTION'); p=p or {}; day,sh,_=_tiom_context(db,user,p); _ensure_tiom_trip_factors(db,user)
    report_from=_parse_ui_date_v2(p.get('reportFrom'),'report from date') if p.get('reportFrom') else day.replace(day=1)
    report_to=_parse_ui_date_v2(p.get('reportTo'),'report to date') if p.get('reportTo') else day
    report_shift=str(p.get('reportShift') or 'ALL').upper()
    active_shifts={x.shift for x in db.scalars(select(ShiftMaster).where(ShiftMaster.active))}
    if report_shift!='ALL' and report_shift not in active_shifts:
        raise HTTPException(422,'Choose a valid report shift.')
    if report_to < report_from:
        raise HTTPException(422,'Report To date cannot be before From date.')

    report=db.scalar(select(TiomShiftProductionReport).where(TiomShiftProductionReport.operating_date==day,TiomShiftProductionReport.shift==sh))
    movements=[]
    if report:
        for m in db.scalars(select(TiomShiftProductionMovement).where(TiomShiftProductionMovement.report_id==report.report_id).order_by(TiomShiftProductionMovement.row_no)):
            movements.append({'rowNo':m.row_no,'movementCode':m.movement_code,'sourceLocationId':m.source_location_id or '',
                              'destinationLocationId':m.destination_location_id or '','trips':m.trips,'qtyMt':float(m.qty_mt or 0),'remarks':m.remarks or ''})
    locations=[{'id':x.location_id,'label':f'{x.location_name} · {x.location_id}'} for x in db.scalars(select(Location).where(Location.active).order_by(Location.location_name))]
    bundle=_tiom_shift_ftd_bundle(db,day,sh,False); auto_sources=bundle['sources']; auto_values=bundle['values']
    options=[]
    for code,label,factor_code in TIOM_SHIFT_MOVEMENTS:
        factor=_tiom_factor_by_code(db,factor_code,day) if factor_code else None
        parts=auto_sources.get(code,{})
        auto=any(parts.get(k,Decimal('0')) for k in ('WB','MIS'))
        options.append({'code':code,'label':label,'factorCode':factor_code or '','factor':float(factor) if factor is not None else None,
                        'autoManaged':bool(auto),'autoQty':float(auto_values.get(code,0)) if auto else 0,
                        'autoSource':'+'.join(k for k in ('WB','MIS') if parts.get(k))})
    base,_=_tiom_baseline_map(db,day)
    baseline_rows=[{'code':code,'label':label,'value':float(base.get(code,0))} for _,code,label,_ in TIOM_SHIFT_REPORT_LINES]
    return {'date':str(day),'shift':sh,'reportFrom':str(report_from),'reportTo':str(report_to),'reportShift':report_shift,
            'reportId':report.report_id if report else '','status':report.status if report else 'NEW',
            'remarks':report.remarks or '' if report else '','movements':movements,'movementOptions':options,'locations':locations,
            'report':_tiom_shift_report_data(db,day,sh,True),
            'rangeReport':_tiom_shift_range_report(db,report_from,report_to,report_shift),
            'baselineRows':baseline_rows,'managementRecipients':_tiom_management_recipients(db)}


def save_tiom_shift_production(db,user,p,submit=False):
    require(user,'PRODUCTION'); p=p or {}; day,sh,_=_tiom_context(db,user,p); open_shift(db,day,sh); _ensure_tiom_trip_factors(db,user)
    report=db.scalar(select(TiomShiftProductionReport).where(TiomShiftProductionReport.operating_date==day,TiomShiftProductionReport.shift==sh))
    if report and report.status=='SUBMITTED' and not user.admin: raise HTTPException(409,'Submitted shift production report is locked. Ask Management to correct it.')
    if not report:
        report=TiomShiftProductionReport(report_id=str(uuid4()),operating_date=day,shift=sh,status='DRAFT',entered_by=user.login_id,entered_at=now_local(),version=1)
        db.add(report); db.flush()
    else: report.version=(report.version or 1)+1
    report.remarks=str(p.get('remarks') or '').strip()[:1000]
    db.execute(delete(TiomShiftProductionMovement).where(TiomShiftProductionMovement.report_id==report.report_id))
    rows_=p.get('rows') if isinstance(p.get('rows'),list) else []
    auto_bundle=_tiom_shift_ftd_bundle(db,day,sh,False); auto_sources=auto_bundle['sources']
    kept=0
    for idx,item in enumerate(rows_[:80],start=1):
        if not isinstance(item,dict): continue
        code=str(item.get('movementCode') or '').strip().upper(); src=str(item.get('sourceLocationId') or '').strip() or None; dst=str(item.get('destinationLocationId') or '').strip() or None
        trips_raw=item.get('trips'); qty_raw=item.get('qtyMt'); note=short(str(item.get('remarks') or ''))
        if not any([code,src,dst,trips_raw not in (None,''),qty_raw not in (None,''),note]): continue
        if code not in TIOM_SHIFT_MOVEMENT_MAP: raise HTTPException(422,f'Row {idx}: choose a valid movement.')
        parts=auto_sources.get(code,{})
        if parts.get('WB') or parts.get('MIS'):
            srcname='+'.join(k for k in ('WB','MIS') if parts.get(k))
            raise HTTPException(409,f'Row {idx}: {TIOM_SHIFT_MOVEMENT_MAP[code][1]} is already supplied automatically by {srcname}. Do not enter it again.')
        if src: active_resource(db,Location,src)
        if dst: active_resource(db,Location,dst)
        trips=None
        if trips_raw not in (None,''):
            try: trips=int(trips_raw)
            except: raise HTTPException(422,f'Row {idx}: trips must be a whole number.')
            if trips<0: raise HTTPException(422,f'Row {idx}: trips cannot be negative.')
        qty=None
        if qty_raw not in (None,''):
            try: qty=Decimal(str(qty_raw))
            except: raise HTTPException(422,f'Row {idx}: enter a valid quantity MT.')
        if qty is None and trips:
            factor_code=TIOM_SHIFT_MOVEMENT_MAP[code][2]; factor=_tiom_factor_by_code(db,factor_code,day) if factor_code else None
            if factor is None: raise HTTPException(422,f'Row {idx}: enter Qty MT; no automatic factor is configured for {TIOM_SHIFT_MOVEMENT_MAP[code][1]}.')
            qty=Decimal(trips)*factor
        if qty is None or qty<0: raise HTTPException(422,f'Row {idx}: Qty MT is required and cannot be negative.')
        db.add(TiomShiftProductionMovement(movement_id=str(uuid4()),report_id=report.report_id,row_no=idx,movement_code=code,
            source_location_id=src,destination_location_id=dst,trips=trips,qty_mt=qty,remarks=note,entered_at=now_local())); kept+=1
    if kept==0 and not any(auto_bundle['values'].values()) and not report.remarks: raise HTTPException(422,'No automatic WB/MIS data found. Add an exception / plant-output row or a shift remark.')
    report.status='SUBMITTED' if submit else 'DRAFT'
    if submit: report.submitted_by=user.login_id; report.submitted_at=now_local()
    audit(db,user,'TIOM_SHIFT_PROD_SUBMIT' if submit else 'TIOM_SHIFT_PROD_DRAFT','tiom_shift_production_report',report.report_id,{'date':str(day),'shift':sh,'rows':kept})
    db.flush()
    return {'ok':True,'message':'Management shift report submitted.' if submit else 'Management report draft saved.','reportId':report.report_id,'status':report.status,'report':_tiom_shift_report_data(db,day,sh,True)}


def save_tiom_shift_baselines(db,user,p):
    require(user,admin=True); p=p or {}; eff=_parse_ui_date_v2(p.get('effectiveDate') or '2026-09-28','baseline date')
    valid={code for _,code,_,_ in TIOM_SHIFT_REPORT_LINES}; rows_=p.get('rows') if isinstance(p.get('rows'),list) else []
    for item in rows_:
        if not isinstance(item,dict): continue
        code=str(item.get('code') or '').strip().upper()
        if code not in valid: continue
        try: value=Decimal(str(item.get('value') or 0))
        except: raise HTTPException(422,f'Invalid opening cumulative for {code}.')
        row=db.get(TiomShiftReportBaseline,{'line_code':code,'effective_date':eff})
        if not row:
            row=TiomShiftReportBaseline(line_code=code,effective_date=eff,opening_qty_mt=value,entered_by=user.login_id,entered_at=now_local());db.add(row)
        else: row.opening_qty_mt=value;row.entered_by=user.login_id;row.entered_at=now_local()
    audit(db,user,'TIOM_SHIFT_BASELINE','tiom_shift_report_baseline',str(eff),{'rows':len(rows_)})
    return {'ok':True,'message':'Opening cumulative values saved.','effectiveDate':str(eff)}

def _tiom_meter_type(e):
    txt=' '.join([e.group or '',e.type or '',e.make_model or '',e.machine_id or '']).upper()
    return 'KMR' if e.group in {'TRANSPORT','HSD_TANKER'} or any(k in txt for k in ['VEHICLE','TANKER','BOLERO','SCORPIO','FORTUNER','BUS','CAMPER','VAN']) else 'HMR'


def _tiom_hsd_category(e):
    txt=' '.join([e.group or '',e.type or '',e.make_model or '',e.machine_id or '']).upper()
    if 'TANKER' in txt or 'SPRINKLER' in txt: return 'DIESEL & WATER TANKER'
    if any(k in txt for k in ['BOLERO','SCORPIO','FORTUNER','BUS','CAMPER','SERVICE VAN','LIGHT VEHICLE']): return 'LIGHT VEHICLES'
    if 'EXCAVATOR' in txt or e.machine_id.upper().startswith('EX-'): return 'EXCAVATOR'
    if 'DOZER' in txt or 'DOZAR' in txt: return 'DOZER'
    if 'LOADER' in txt: return 'WHEEL LOADER'
    if 'DRILL' in txt or 'COMPRESSOR' in txt: return 'DRILL MACHINE'
    if 'GRADER' in txt: return 'GRADER'
    if 'PLANT' in txt or 'METSO' in txt: return 'PLANT'
    if e.group=='TRANSPORT': return 'TIPPER'
    return 'OTHER EQUIPMENT FOR MINES'


def _tiom_hsd_book_balance(db,tanker_id,before):
    receipts=db.scalar(select(func.coalesce(func.sum(HsdPurchaseLot.litres_received),0)).where(
        HsdPurchaseLot.tanker_id==tanker_id,HsdPurchaseLot.received_at<=before
    )) or Decimal('0')
    issues=db.scalar(select(func.coalesce(func.sum(HsdIssue.litres),0)).where(
        HsdIssue.tanker_id==tanker_id,HsdIssue.issued_at<=before
    )) or Decimal('0')
    return Decimal(receipts)-Decimal(issues)


def _tiom_hsd_previous_meter(db,machine_id,before):
    issue=db.scalar(select(HsdIssue).where(HsdIssue.machine_id==machine_id,HsdIssue.meter_reading.is_not(None),HsdIssue.issued_at<before).order_by(HsdIssue.issued_at.desc()).limit(1))
    return Decimal(issue.meter_reading) if issue and issue.meter_reading is not None else None


def _tiom_hsd_context_anchor(day, definition):
    now=now_local()
    start=datetime.combine(day,definition.start_time,TZ)
    end=datetime.combine(day,definition.end_time,TZ)
    if definition.end_time <= definition.start_time:
        end += timedelta(days=1)
    if start <= now <= end:
        return now
    return end


def _tiom_hsd_event_time(db,day,shift,value):
    exact=_tiom_shift_date_time(db,day,shift,value)
    if exact is not None:
        return exact
    live_day,live_shift,_now=operating_context()
    if day==live_day and shift==live_shift:
        return now_local()
    raise HTTPException(422,'Filling / receipt time is required for a previous date or previous shift entry.')


def tiom_hsd_previous_meter_info(db,user,p):
    require(user,'HSD'); p=p or {}; day,sh,definition=_tiom_context(db,user,p)
    machine=active_resource(db,Equipment,str(p.get('machineId') or '').strip())
    event=_tiom_shift_date_time(db,day,sh,p.get('time')) or _tiom_hsd_context_anchor(day,definition)
    issue=db.scalar(select(HsdIssue).where(
        HsdIssue.machine_id==machine.machine_id,
        HsdIssue.meter_reading.is_not(None),
        HsdIssue.issued_at<event
    ).order_by(HsdIssue.issued_at.desc()).limit(1))
    return {
        'machineId':machine.machine_id,
        'meterType':_tiom_meter_type(machine),
        'previousMeter':float(issue.meter_reading) if issue and issue.meter_reading is not None else None,
        'previousDate':str(issue.operating_date) if issue else None,
        'previousShift':issue.shift if issue else None,
        'previousTime':time_text(issue.issued_at) if issue else None,
        'previousLitres':float(issue.litres) if issue else None,
    }


def save_tiom_hsd_receipt(db,user,p):
    require(user,'HSD'); p=p or {}; day,sh,_=_tiom_context(db,user,p); request_id=str(p.get('requestId') or '').strip()
    if not request_id: request_id='tiom-hsd-r-'+uuid4().hex
    old=db.scalar(select(HsdPurchaseLot).where(HsdPurchaseLot.request_id==request_id))
    if old:return {'ok':True,'message':'Receipt already saved.','lotId':old.lot_id}
    tanker=active_resource(db,HsdTanker,str(p.get('tankerId') or '').strip())
    try: litres=Decimal(str(p.get('litres'))); rate=Decimal(str(p.get('ratePerL') or 0)); discount=Decimal(str(p.get('discountPerL') or 0))
    except: raise HTTPException(422,'Enter valid HSD quantity/rate/discount.')
    receipt_type=str(p.get('receiptType') or 'RECEIPT').upper()
    if receipt_type not in {'OPENING','RECEIPT'}: raise HTTPException(422,'Receipt type must be Opening or Receipt.')
    if litres<=0: raise HTTPException(422,'HSD quantity must be greater than zero.')
    if receipt_type=='RECEIPT' and rate<=0: raise HTTPException(422,'Rate/L must be greater than zero for a receipt.')
    if discount<0 or discount>rate: raise HTTPException(422,'Discount/L cannot be negative or above rate.')
    supplier=short(str(p.get('supplier') or ('OPENING STOCK' if receipt_type=='OPENING' else '')).strip())
    if not supplier: raise HTTPException(422,'Supplier is required.')
    gross=(litres*rate).quantize(Decimal('0.01')); net=(litres*(rate-discount)).quantize(Decimal('0.01'))
    event=_tiom_hsd_event_time(db,day,sh,p.get('time'))
    lot=HsdPurchaseLot(lot_id=str(uuid4()),tanker_id=tanker.tanker_id,received_at=event,supplier=supplier,invoice_no=short(str(p.get('invoiceNo') or '')) or None,pump_location=short(str(p.get('pumpLocation') or '')) or None,litres_received=litres,litres_remaining=litres,rate_per_l=rate,amount=gross,request_id=request_id)
    db.add(lot); db.flush(); db.add(TiomHsdReceiptDetail(lot_id=lot.lot_id,operating_date=day,shift=sh,receipt_type=receipt_type,discount_per_l=discount,net_amount=net,remarks=short(str(p.get('remarks') or '')),entered_by=user.login_id,entered_at=now_local()))
    audit(db,user,'TIOM_HSD_RECEIPT','hsd_purchase_lot',lot.lot_id,{'date':str(day),'type':receipt_type,'litres':str(litres),'rate':str(rate),'discount':str(discount),'net':str(net)})
    return {'ok':True,'message':'Opening HSD saved.' if receipt_type=='OPENING' else 'HSD receipt saved.','lotId':lot.lot_id,'grossAmount':float(gross),'netAmount':float(net)}


def save_tiom_hsd_issue(db,user,p):
    require(user,'HSD'); p=p or {}; day,sh,_=_tiom_context(db,user,p); request_id=str(p.get('requestId') or '').strip() or ('tiom-hsd-i-'+uuid4().hex)
    old=db.scalar(select(HsdIssue).where(HsdIssue.request_id==request_id))
    if old:return {'ok':True,'message':'HSD filling already saved.','issueId':old.issue_id}
    tanker=active_resource(db,HsdTanker,str(p.get('tankerId') or '').strip()); machine=active_resource(db,Equipment,str(p.get('machineId') or '').strip())
    try: litres=Decimal(str(p.get('litres'))); current=Decimal(str(p.get('currentMeter')))
    except: raise HTTPException(422,'Enter valid HSD litres and current meter.')
    if litres<=0: raise HTTPException(422,'HSD litres must be greater than zero.')
    event=_tiom_hsd_event_time(db,day,sh,p.get('time'))
    # Previous HMR/KMR is authoritative from the immediately preceding saved
    # HSD entry. Never trust a manually supplied previous reading.
    prev=_tiom_hsd_previous_meter(db,machine.machine_id,event)
    if prev is not None and current<prev: raise HTTPException(422,'Current meter cannot be below previous meter.')
    dup=db.scalar(select(HsdIssue).where(
        HsdIssue.machine_id==machine.machine_id,
        HsdIssue.tanker_id==tanker.tanker_id,
        HsdIssue.meter_reading==current,
        HsdIssue.litres==litres,
        HsdIssue.issued_at>=event-timedelta(minutes=5),
        HsdIssue.issued_at<=event+timedelta(minutes=5)
    ).limit(1))
    if dup:
        raise HTTPException(409,'Possible duplicate HSD filling: same machine, tanker, meter and litres already saved near this time.')
    usage=(current-prev) if prev is not None else None; meter_type=_tiom_meter_type(machine)
    efficiency=None; unit='KMPL' if meter_type=='KMR' else 'HSD/HR'
    if usage is not None and usage>0:
        efficiency=(usage/litres) if meter_type=='KMR' else (litres/usage)
    issue=HsdIssue(issue_id=str(uuid4()),operating_date=day,shift=sh,tanker_id=tanker.tanker_id,machine_id=machine.machine_id,litres=litres,amount=Decimal('0'),location_id=None,meter_type=meter_type,meter_reading=current,recipient_employee_id=None,reference=short(str(p.get('reference') or '')) or None,issued_at=event,request_id=request_id)
    try: allocations,total=apply_fifo_issue(db,issue)
    except ValueError as exc: raise HTTPException(409,str(exc))
    entry_source=str(p.get('entrySource') or 'OFFICE_MIS').strip().upper()
    if entry_source not in {'OFFICE_MIS','FIELD_SUPERVISOR'}:
        entry_source='OFFICE_MIS'
    db.add(TiomHsdIssueDetail(issue_id=issue.issue_id,previous_meter_reading=prev,usage=usage,efficiency=efficiency,efficiency_unit=unit,entered_by=user.login_id,entered_at=now_local()))
    audit(db,user,'TIOM_HSD_FILL','hsd_issue',issue.issue_id,{'date':str(day),'shift':sh,'machine':machine.machine_id,'litres':str(litres),'meterType':meter_type,'previous':str(prev) if prev is not None else None,'current':str(current),'usage':str(usage) if usage is not None else None,'efficiency':str(efficiency) if efficiency is not None else None,'entrySource':entry_source})
    return {'ok':True,'message':'HSD filling saved.','issueId':issue.issue_id,'previousMeter':float(prev) if prev is not None else None,'usage':float(usage) if usage is not None else None,'efficiency':float(efficiency) if efficiency is not None else None,'efficiencyUnit':unit,'amount':float(total),'entrySource':entry_source}


def save_tiom_hsd_issue_batch(db,user,p):
    require(user,'HSD'); p=p or {}
    rows=p.get('rows') if isinstance(p.get('rows'),list) else []
    base_request=str(p.get('requestId') or '').strip() or ('tiom-hsd-b-'+uuid4().hex)
    tanker_id=str(p.get('tankerId') or '').strip()
    if not tanker_id: raise HTTPException(422,'Select supplying tanker.')
    saved=[]
    for idx,row in enumerate(rows[:100],start=1):
        if not isinstance(row,dict): continue
        machine_id=str(row.get('machineId') or '').strip()
        current=row.get('currentMeter'); litres=row.get('litres')
        if not machine_id and current in (None,'') and litres in (None,''): continue
        if not machine_id: raise HTTPException(422,f'Row {idx}: select equipment.')
        item=dict(row)
        item.update({'date':p.get('date'),'shift':p.get('shift'),'tankerId':tanker_id,'requestId':f'{base_request}-{idx:03d}','entrySource':p.get('entrySource') or 'OFFICE_MIS'})
        try:
            result=save_tiom_hsd_issue(db,user,item)
        except HTTPException as exc:
            raise HTTPException(exc.status_code,f'Row {idx}: {exc.detail}') from exc
        saved.append({'row':idx,**result})
    if not saved: raise HTTPException(422,'Add at least one HSD filling row.')
    return {'ok':True,'message':f'{len(saved)} HSD filling row(s) saved.','rows':saved}


def save_tiom_hsd_stock_check(db,user,p):
    require(user,'HSD'); p=p or {}; day,sh,definition=_tiom_context(db,user,p)
    tanker=active_resource(db,HsdTanker,str(p.get('tankerId') or '').strip())
    try: physical=Decimal(str(p.get('physicalLitres')))
    except Exception: raise HTTPException(422,'Enter valid physical / dip stock litres.')
    if physical<0: raise HTTPException(422,'Physical / dip stock cannot be negative.')
    event=_tiom_hsd_event_time(db,day,sh,p.get('time'))
    book=_tiom_hsd_book_balance(db,tanker.tanker_id,event)
    variance=physical-book
    existing=db.scalar(select(TiomHsdStockCheck).where(
        TiomHsdStockCheck.operating_date==day,
        TiomHsdStockCheck.shift==sh,
        TiomHsdStockCheck.tanker_id==tanker.tanker_id
    ))
    row=existing or TiomHsdStockCheck(
        check_id=str(uuid4()),operating_date=day,shift=sh,tanker_id=tanker.tanker_id,
        checked_at=event,book_litres=book,physical_litres=physical,variance_litres=variance,
        entered_by=user.login_id,entered_at=now_local()
    )
    row.checked_at=event; row.book_litres=book; row.physical_litres=physical; row.variance_litres=variance
    row.remarks=short(str(p.get('remarks') or '')) or None; row.entered_by=user.login_id; row.entered_at=now_local()
    db.add(row)
    audit(db,user,'TIOM_HSD_STOCK_CHECK','tiom_hsd_stock_check',row.check_id,{
        'date':str(day),'shift':sh,'tanker':tanker.tanker_id,'book':str(book),
        'physical':str(physical),'variance':str(variance)
    })
    return {'ok':True,'message':'Physical / dip stock saved.','checkId':row.check_id,'bookLitres':float(book),'physicalLitres':float(physical),'varianceLitres':float(variance)}


def _tiom_hsd_report(db,day,from_day=None,to_day=None):
    from_day=from_day or day
    to_day=to_day or day
    if to_day < from_day:
        raise HTTPException(422,'Report To date cannot be before From date.')
    range_start=datetime.combine(from_day,dtime.min,TZ)
    range_end=datetime.combine(to_day,dtime.min,TZ)+timedelta(days=1)
    eq={e.machine_id:e for e in db.scalars(select(Equipment).where(Equipment.active))}
    issues=list(db.scalars(select(HsdIssue).where(
        HsdIssue.issued_at>=range_start,HsdIssue.issued_at<range_end
    ).order_by(HsdIssue.issued_at)))
    details={x.issue_id:x for x in db.scalars(select(TiomHsdIssueDetail).where(
        TiomHsdIssueDetail.issue_id.in_([i.issue_id for i in issues])
    ))} if issues else {}

    categories={}
    totals_by_machine={}
    for i in issues:
        e=eq.get(i.machine_id)
        cat=_tiom_hsd_category(e) if e else 'OTHER EQUIPMENT FOR MINES'
        key=(cat,i.machine_id)
        row=categories.setdefault(key,{
            'category':cat,'machineId':i.machine_id,
            'label':_tiom_asset_label(e) if e else i.machine_id,
            'meterType':i.meter_type or '','previous':None,'shifts':{}
        })
        d=details.get(i.issue_id)
        if d and row['previous'] is None and d.previous_meter_reading is not None:
            row['previous']=float(d.previous_meter_reading)
        shrow=row['shifts'].setdefault(i.shift,{
            'meter':None,'hsd':0.0,'usage':0.0,'efficiency':None,
            'unit':d.efficiency_unit if d else ('KMPL' if i.meter_type=='KMR' else 'HSD/HR')
        })
        if i.meter_reading is not None:
            shrow['meter']=float(i.meter_reading)
        shrow['hsd']+=float(i.litres)
        if d and d.usage is not None:
            shrow['usage']+=float(d.usage)
        mt=totals_by_machine.setdefault(i.machine_id,{'usage':0.0,'hsd':0.0,'meterType':i.meter_type or ''})
        mt['hsd']+=float(i.litres)
        if d and d.usage is not None:
            mt['usage']+=float(d.usage)

    rows_out=[]
    for key,row in sorted(categories.items()):
        for sh,v in row['shifts'].items():
            if v['usage']>0 and v['hsd']>0:
                v['efficiency']=(v['usage']/v['hsd']) if row['meterType']=='KMR' else (v['hsd']/v['usage'])
        m=totals_by_machine.get(row['machineId'],{'usage':0,'hsd':0})
        m_eff=(m['usage']/m['hsd']) if row['meterType']=='KMR' and m['hsd'] else ((m['hsd']/m['usage']) if row['meterType']!='KMR' and m['usage'] else None)
        row['mtd']={'usage':m['usage'],'hsd':m['hsd'],'efficiency':m_eff,'unit':'KMPL' if row['meterType']=='KMR' else 'HSD/HR'}
        rows_out.append(row)

    receipts=list(db.scalars(select(HsdPurchaseLot).where(
        HsdPurchaseLot.received_at<range_end
    ).order_by(HsdPurchaseLot.received_at)))
    receipt_details={x.lot_id:x for x in db.scalars(select(TiomHsdReceiptDetail).where(
        TiomHsdReceiptDetail.lot_id.in_([r.lot_id for r in receipts])
    ))} if receipts else {}
    opening_receipts=sum((Decimal(r.litres_received) for r in receipts if aware(r.received_at)<range_start),Decimal('0'))
    opening_issues=sum((Decimal(i.litres) for i in db.scalars(select(HsdIssue).where(HsdIssue.issued_at<range_start))),Decimal('0'))
    balance=opening_receipts-opening_issues
    procurement=[]
    cur=from_day
    while cur<=to_day:
        a=datetime.combine(cur,dtime.min,TZ); b=a+timedelta(days=1)
        day_receipts=[r for r in receipts if a<=aware(r.received_at)<b]
        rec_qty=sum((Decimal(r.litres_received) for r in day_receipts),Decimal('0'))
        gross=sum((Decimal(r.amount) for r in day_receipts),Decimal('0'))
        net=sum((Decimal(receipt_details[r.lot_id].net_amount) if r.lot_id in receipt_details else Decimal(r.amount) for r in day_receipts),Decimal('0'))
        discount=gross-net
        issued=sum((Decimal(i.litres) for i in issues if a<=aware(i.issued_at)<b),Decimal('0'))
        balance+=rec_qty-issued
        refs=', '.join([r.invoice_no or '' for r in day_receipts if r.invoice_no])
        procurement.append({'date':str(cur),'received':float(rec_qty),'grossAmount':float(gross),'discountAmount':float(discount),'netAmount':float(net),'issued':float(issued),'balance':float(balance),'billNo':refs})
        cur+=timedelta(days=1)

    total=sum((Decimal(i.litres) for i in issues),Decimal('0'))
    light=sum((Decimal(i.litres) for i in issues if eq.get(i.machine_id) and _tiom_hsd_category(eq[i.machine_id])=='LIGHT VEHICLES'),Decimal('0'))
    return {'date':str(day),'fromDate':str(from_day),'toDate':str(to_day),'rows':rows_out,'procurement':procurement,
            'totalHsd':float(total),'minesTotalHsd':float(total-light),'lightVehicleHsd':float(light)}

def tiom_drilling_desk(db,user,p):
    require(user,'PRODUCTION'); p=p or {}; day,sh,_=_tiom_context(db,user,p)
    report_from=_parse_ui_date_v2(p.get('reportFrom'),'report from date') if p.get('reportFrom') else day.replace(day=1)
    report_to=_parse_ui_date_v2(p.get('reportTo'),'report to date') if p.get('reportTo') else day
    if report_to < report_from: raise HTTPException(422,'Report To date cannot be before From date.')
    report_shift=str(p.get('reportShift') or 'ALL').upper()
    source_locations=location_options(db,'SOURCE')
    sets=list(db.scalars(select(TiomDrillSet).where(TiomDrillSet.active.is_(True)).order_by(TiomDrillSet.drill_set_name)))
    equipment={x.machine_id:x for x in db.scalars(select(Equipment))}
    current={x.drill_set_id:x for x in db.scalars(select(TiomDrillingShift).where(TiomDrillingShift.operating_date==day,TiomDrillingShift.shift==sh))}
    rows=[]
    for st in sets:
        r=current.get(st.drill_set_id)
        rows.append({'drillSetId':st.drill_set_id,'label':st.drill_set_name,
            'drillMachineId':st.drill_machine_id,'drillMachineLabel':_tiom_asset_label(equipment.get(st.drill_machine_id)) if equipment.get(st.drill_machine_id) else st.drill_machine_id,
            'compressorMachineId':st.compressor_machine_id,'compressorMachineLabel':_tiom_asset_label(equipment.get(st.compressor_machine_id)) if equipment.get(st.compressor_machine_id) else st.compressor_machine_id,
            'sourceLocationId':r.source_location_id if r else '', 'holes':r.holes if r else 0,
            'romObHoles':r.rom_ob_holes if r else None,'bhjBhqHoles':r.bhj_bhq_holes if r else None,
            'romObMeterage':float(r.rom_ob_meterage) if r else 0,'bhjBhqMeterage':float(r.bhj_bhq_meterage) if r else 0,
            'drillOpenHmr':float(r.drill_open_hmr) if r and r.drill_open_hmr is not None else None,'drillCloseHmr':float(r.drill_close_hmr) if r and r.drill_close_hmr is not None else None,
            'compressorOpenHmr':float(r.compressor_open_hmr) if r and r.compressor_open_hmr is not None else None,'compressorCloseHmr':float(r.compressor_close_hmr) if r and r.compressor_close_hmr is not None else None,
            'breakdownHours':float(r.breakdown_hours) if r and r.breakdown_hours is not None else None,'remarks':r.remarks or '' if r else ''})
    return {'date':str(day),'shift':sh,'reportFrom':str(report_from),'reportTo':str(report_to),'reportShift':report_shift,
        'sourceLocations':[{'id':x['id'],'label':f"{x['label']} · {x['id']}"} for x in source_locations],
        'rows':rows,'summaryToday':_tiom_drilling_summary(db,day,day,sh),'summaryRange':_tiom_drilling_summary(db,report_from,report_to,report_shift)}


def save_tiom_drilling_shift(db,user,p):
    require(user,'PRODUCTION'); p=p or {}; day,sh,_=_tiom_context(db,user,p)
    incoming=p.get('rows') if isinstance(p.get('rows'),list) else []
    def dec(v,label,zero=False):
        if v in (None,''): return Decimal('0') if zero else None
        try: x=Decimal(str(v))
        except Exception: raise HTTPException(422,f'Enter valid {label}.')
        if not x.is_finite() or x<0: raise HTTPException(422,f'{label} cannot be negative.')
        return x
    saved=0
    for idx,item in enumerate(incoming[:20],start=1):
        if not isinstance(item,dict): continue
        set_id=str(item.get('drillSetId') or '').strip(); source_id=str(item.get('sourceLocationId') or '').strip()
        holes_raw=item.get('holes'); rom_raw=item.get('romObMeterage'); bhj_raw=item.get('bhjBhqMeterage')
        if not set_id: continue
        st=db.get(TiomDrillSet,set_id)
        if not st or not st.active: raise HTTPException(422,f'Drilling row {idx}: choose an active Drill Set.')
        if source_id: active_resource(db,Location,source_id)
        try: holes=int(holes_raw or 0)
        except Exception: raise HTTPException(422,f'Drilling row {idx}: Holes must be a whole number.')
        if holes<0: raise HTTPException(422,f'Drilling row {idx}: Holes cannot be negative.')
        rom_holes_raw=item.get('romObHoles');bhj_holes_raw=item.get('bhjBhqHoles')
        if (rom_holes_raw in (None,'')) != (bhj_holes_raw in (None,'')):
            raise HTTPException(422,f'Drilling row {idx}: enter both ROM/OB and BHJ/BHQ hole counts (use zero where none).')
        if rom_holes_raw not in (None,'') and bhj_holes_raw not in (None,''):
            try:rom_holes=int(rom_holes_raw);bhj_holes=int(bhj_holes_raw)
            except Exception:raise HTTPException(422,f'Drilling row {idx}: Enter integer hole counts.')
            if rom_holes<0 or bhj_holes<0:raise HTTPException(422,f'Drilling row {idx}: Hole counts must be nonnegative.')
            holes=rom_holes+bhj_holes
        else:rom_holes=None;bhj_holes=None
        rom=dec(rom_raw,f'ROM/OB meterage row {idx}',True); bhj=dec(bhj_raw,f'BHJ/BHQ meterage row {idx}',True)
        do=dec(item.get('drillOpenHmr'),f'Drill HMR opening row {idx}'); dc=dec(item.get('drillCloseHmr'),f'Drill HMR closing row {idx}')
        co=dec(item.get('compressorOpenHmr'),f'Compressor HMR opening row {idx}'); cc=dec(item.get('compressorCloseHmr'),f'Compressor HMR closing row {idx}')
        bd=dec(item.get('breakdownHours'),f'Breakdown hours row {idx}')
        if do is not None and dc is not None and dc<do: raise HTTPException(422,f'Drilling row {idx}: Drill closing HMR cannot be below opening HMR.')
        if co is not None and cc is not None and cc<co: raise HTTPException(422,f'Drilling row {idx}: Compressor closing HMR cannot be below opening HMR.')
        existing=db.scalar(select(TiomDrillingShift).where(TiomDrillingShift.operating_date==day,TiomDrillingShift.shift==sh,TiomDrillingShift.drill_set_id==set_id))
        obj=existing or TiomDrillingShift(drilling_id=str(uuid4()),operating_date=day,shift=sh,drill_set_id=set_id,entered_by=user.login_id,entered_at=now_local())
        obj.source_location_id=source_id or None; obj.holes=holes
        obj.rom_ob_holes=rom_holes;obj.bhj_bhq_holes=bhj_holes
        obj.rom_ob_meterage=rom; obj.bhj_bhq_meterage=bhj
        obj.drill_open_hmr=do; obj.drill_close_hmr=dc; obj.compressor_open_hmr=co; obj.compressor_close_hmr=cc
        obj.breakdown_hours=bd; obj.remarks=short(str(item.get('remarks') or '')); obj.entered_by=user.login_id; obj.entered_at=now_local(); db.add(obj)
        _tiom_upsert_hmr(db,user,day,sh,st.drill_machine_id,do,dc,'TIOM_DRILLING',f'[DRILL_SET:{set_id}]')
        _tiom_upsert_hmr(db,user,day,sh,st.compressor_machine_id,co,cc,'TIOM_DRILLING',f'[DRILL_SET:{set_id}]')
        saved+=1
    if not saved: raise HTTPException(422,'Add at least one Drill Set row.')
    audit(db,user,'TIOM_DRILLING_SHIFT','tiom_drilling_shift',f'{day}:{sh}',{'rows':saved})
    return {'ok':True,'message':f'{saved} drilling row(s) saved. Drill and Compressor HMR were written to the central meter ledger.'}


def tiom_hsd_desk(db,user,p):
    require(user,'HSD'); p=p or {}; day,sh,definition=_tiom_context(db,user,p)
    report_from=_parse_ui_date_v2(p.get('reportFrom'),'report from date') if p.get('reportFrom') else day.replace(day=1)
    report_to=_parse_ui_date_v2(p.get('reportTo'),'report to date') if p.get('reportTo') else day
    if report_to < report_from: raise HTTPException(422,'Report To date cannot be before From date.')
    tankers=list(db.scalars(select(HsdTanker).where(HsdTanker.active).order_by(HsdTanker.tanker_id)))
    equipment=list(db.scalars(select(Equipment).where(Equipment.active).order_by(Equipment.machine_id)))
    stock={t.tanker_id:db.scalar(select(func.coalesce(func.sum(HsdPurchaseLot.litres_remaining),0)).where(HsdPurchaseLot.tanker_id==t.tanker_id)) for t in tankers}

    # Prefill Previous HMR/KMR from the last saved HSD issue before the selected
    # shift context. Exact row time is rechecked again on save.
    anchor=_tiom_hsd_context_anchor(day,definition)
    prev={}
    for e in equipment:
        issue=db.scalar(select(HsdIssue).where(
            HsdIssue.machine_id==e.machine_id,
            HsdIssue.meter_reading.is_not(None),
            HsdIssue.issued_at<anchor
        ).order_by(HsdIssue.issued_at.desc()).limit(1))
        prev[e.machine_id]=float(issue.meter_reading) if issue and issue.meter_reading is not None else None

    day_start=datetime.combine(day,dtime.min,TZ); day_end=day_start+timedelta(days=1)
    opening_receipts=db.scalar(select(func.coalesce(func.sum(HsdPurchaseLot.litres_received),0)).where(HsdPurchaseLot.received_at<day_start)) or Decimal('0')
    opening_issues=db.scalar(select(func.coalesce(func.sum(HsdIssue.litres),0)).where(HsdIssue.issued_at<day_start)) or Decimal('0')
    opening_book=Decimal(opening_receipts)-Decimal(opening_issues)

    day_lots=list(db.scalars(select(HsdPurchaseLot).where(HsdPurchaseLot.received_at>=day_start,HsdPurchaseLot.received_at<day_end)))
    day_lot_ids=[x.lot_id for x in day_lots]
    day_details={x.lot_id:x for x in db.scalars(select(TiomHsdReceiptDetail).where(TiomHsdReceiptDetail.lot_id.in_(day_lot_ids)))} if day_lot_ids else {}
    opening_set=sum((Decimal(x.litres_received) for x in day_lots if day_details.get(x.lot_id) and day_details[x.lot_id].receipt_type=='OPENING'),Decimal('0'))
    received=sum((Decimal(x.litres_received) for x in day_lots if not day_details.get(x.lot_id) or day_details[x.lot_id].receipt_type!='OPENING'),Decimal('0'))
    issued_day=db.scalar(select(func.coalesce(func.sum(HsdIssue.litres),0)).where(HsdIssue.operating_date==day)) or Decimal('0')
    issued_shift=db.scalar(select(func.coalesce(func.sum(HsdIssue.litres),0)).where(HsdIssue.operating_date==day,HsdIssue.shift==sh)) or Decimal('0')
    available_now=sum((Decimal(stock[t.tanker_id] or 0) for t in tankers),Decimal('0'))
    closing_book=opening_book+opening_set+received-Decimal(issued_day)

    recent=list(db.scalars(select(HsdIssue).where(
        HsdIssue.operating_date==day,HsdIssue.shift==sh
    ).order_by(HsdIssue.issued_at.desc()).limit(12)))
    stock_checks={x.tanker_id:x for x in db.scalars(select(TiomHsdStockCheck).where(
        TiomHsdStockCheck.operating_date==day,TiomHsdStockCheck.shift==sh
    ))}
    recent_details={x.issue_id:x for x in db.scalars(select(TiomHsdIssueDetail).where(
        TiomHsdIssueDetail.issue_id.in_([r.issue_id for r in recent])
    ))} if recent else {}

    return {'date':str(day),'shift':sh,'reportFrom':str(report_from),'reportTo':str(report_to),
            'tankers':[{'id':t.tanker_id,'label':f'{t.tanker_id}'+(f' · {t.vehicle_no}' if t.vehicle_no else ''),'stock':float(stock[t.tanker_id] or 0)} for t in tankers],
            'equipment':[{'id':e.machine_id,'label':_tiom_asset_label(e),'category':_tiom_hsd_category(e),'meterType':_tiom_meter_type(e),'previousMeter':prev.get(e.machine_id)} for e in equipment],
            'entrySummary':{'openingBook':float(opening_book),'openingSet':float(opening_set),'received':float(received),'issuedDay':float(issued_day),'issuedShift':float(issued_shift),'closingBook':float(closing_book),'availableNow':float(available_now)},
            'stockChecks':[{'checkId':x.check_id,'tankerId':x.tanker_id,'time':time_text(x.checked_at),'bookLitres':float(x.book_litres),'physicalLitres':float(x.physical_litres),'varianceLitres':float(x.variance_litres),'remarks':x.remarks or '','enteredBy':x.entered_by} for x in stock_checks.values()],
            'recentIssues':[{'issueId':x.issue_id,'time':time_text(x.issued_at),'machineId':x.machine_id,'tankerId':x.tanker_id,'litres':float(x.litres),'meterType':x.meter_type or '','meter':float(x.meter_reading) if x.meter_reading is not None else None,'previousMeter':float(recent_details[x.issue_id].previous_meter_reading) if x.issue_id in recent_details and recent_details[x.issue_id].previous_meter_reading is not None else None,'enteredBy':recent_details[x.issue_id].entered_by if x.issue_id in recent_details else ''} for x in recent],
            'managementRecipients':_tiom_management_recipients(db),
            'report':_tiom_hsd_report(db,day,report_from,report_to)}




# WB reporting is read-only and uses the latest confirmed batch per date/shift.
# Previous uploaded versions remain in history but must not double-count movement.
def _tiom_wb_history_rows(db,user,p):
    require(user,'WB')
    p=p or {}
    today=operating_context()[0]
    from_day=_parse_ui_date_v2(p.get('fromDate'),'WB report From') if p.get('fromDate') else today-timedelta(days=6)
    to_day=_parse_ui_date_v2(p.get('toDate'),'WB report To') if p.get('toDate') else today
    if to_day<from_day or (to_day-from_day).days>366:
        raise HTTPException(422,'WB report range must be chronological and at most 367 days.')
    sh=str(p.get('shift') or 'ALL').upper()
    all_shifts={x.shift for x in db.scalars(select(ShiftMaster).where(ShiftMaster.active.is_(True)))}
    if sh!='ALL' and sh not in all_shifts:raise HTTPException(422,'Invalid shift filter.')
    allowed=all_shifts if user.admin or 'ALL' in user.shifts.split(',') else all_shifts.intersection(user.shifts.split(','))
    if sh!='ALL':
        require(user,shift=sh)
        allowed={sh}
    if not allowed: return [],{'fromDate':str(from_day),'toDate':str(to_day),'shift':sh}
    batches=list(db.scalars(select(WbImportBatch).where(
        WbImportBatch.operating_date>=from_day,WbImportBatch.operating_date<=to_day,
        WbImportBatch.status=='CONFIRMED',WbImportBatch.confirmed_at.is_not(None),
        WbImportBatch.shift.in_(allowed)).order_by(
        WbImportBatch.operating_date,WbImportBatch.shift,WbImportBatch.confirmed_at.desc()
    )))
    latest={}
    for b in batches:latest.setdefault((b.operating_date,b.shift),b)
    ids=[x.batch_id for x in latest.values()]
    raw=list(db.scalars(select(WbMovement).where(WbMovement.batch_id.in_(ids)).order_by(
        WbMovement.operating_date,WbMovement.shift,WbMovement.weigh_at,WbMovement.movement_no
    ))) if ids else []
    filters={key:str(p.get(key) or '').strip().upper() for key in
             ('vehicleFilter','materialFilter','sourceFilter','destinationFilter','status')}
    def eligible(row):
        return (filters['status'] in ('','ALL') or str(row.row_status or '').upper()==filters['status']) and (
            not filters['vehicleFilter'] or filters['vehicleFilter'] in (str(row.vehicle_raw or '')+' '+str(row.vehicle_id or '')).upper()) and (
            not filters['materialFilter'] or filters['materialFilter'] in (str(row.material_code or '')+' '+str(row.material_name or '')).upper()) and (
            not filters['sourceFilter'] or filters['sourceFilter'] in str(row.source_raw or '').upper()) and (
            not filters['destinationFilter'] or filters['destinationFilter'] in str(row.destination_raw or '').upper())
    rows=[r for r in raw if eligible(r)]
    meta={'fromDate':str(from_day),'toDate':str(to_day),'shift':sh,
          'batchCount':len(latest),'filteredRecords':len(rows),'latestConfirmedOnly':True}
    return rows,meta


def get_tiom_wb_history(db,user,p):
    rows,meta=_tiom_wb_history_rows(db,user,p)
    valid=[r for r in rows if r.row_status=='VALID']
    mt=sum((Decimal(r.net_kg or 0) for r in valid),Decimal('0'))/Decimal('1000')
    distinct={str(r.vehicle_id or r.vehicle_raw or '') for r in valid if r.vehicle_id or r.vehicle_raw}
    groups={}
    for r in valid:
        m=str(r.material_name or r.material_code or 'Unclassified')
        group=groups.setdefault(m,{'material':m,'trips':0,'tonnes':Decimal('0')})
        group['trips']+=1;group['tonnes']+=Decimal(r.net_kg or 0)/Decimal('1000')
    return {**meta,'kpis':{
        'validMovements':len(valid),'reviewMovements':sum(r.row_status=='REVIEW' for r in rows),
        'validTonnes':round(float(mt),2),'uniqueVehicles':len(distinct),
        'avgPayloadMt':round(float(mt/len(valid)),2) if valid else None,
        'batches':meta['batchCount']},
        'materials':[{**g,'tonnes':round(float(g['tonnes']),2)} for g in sorted(groups.values(),key=lambda g:g['tonnes'],reverse=True)[:12]],
        'rows':[{'date':str(x.operating_date),'shift':x.shift,'movementNo':x.movement_no,
                 'vehicle':x.vehicle_raw,'material':x.material_name or x.material_code,
                 'source':x.source_raw,'destination':x.destination_raw,
                 'netMt':round(float(Decimal(x.net_kg or 0)/Decimal('1000')),3),
                 'status':x.row_status,'issue':x.issue or ''}
                for x in rows[:120]],
        'previewCount':min(120,len(rows))}


@router.get('/wb/full-export')
def tiom_full_wb_export(request:Request,fromDate:str='',toDate:str='',shift:str='ALL',
                       vehicleFilter:str='',materialFilter:str='',sourceFilter:str='',
                       destinationFilter:str='',status:str='ALL',db:Session=Depends(get_db)):
    user=get_user(db,request)
    rows,meta=_tiom_wb_history_rows(db,user,{
        'fromDate':fromDate,'toDate':toDate,'shift':shift,'vehicleFilter':vehicleFilter,
        'materialFilter':materialFilter,'sourceFilter':sourceFilter,
        'destinationFilter':destinationFilter,'status':status})
    if len(rows)>50000:
        raise HTTPException(422,'More than 50,000 WB rows matched. Narrow the date range before exporting.')
    book=Workbook();sheet=book.active;sheet.title='WB Full Data'
    sheet.append(['Operating Date','Shift','WB Movement No','Vehicle','Vehicle Master ID',
                  'Material Code','Material Name','Source','Destination','Tare KG','Gross KG',
                  'Net KG','Net MT','Weigh Time','Status','Exception','WB Batch'])
    def safe_cell(value):
        v=str(value or '')
        return "'"+v if v.lstrip().startswith(('=','+','-','@')) else v
    for x in rows:
        sheet.append([str(x.operating_date),safe_cell(x.shift),safe_cell(x.movement_no),
            safe_cell(x.vehicle_raw),safe_cell(x.vehicle_id),safe_cell(x.material_code),
            safe_cell(x.material_name),safe_cell(x.source_raw),safe_cell(x.destination_raw),
            float(x.tare_kg or 0),float(x.gross_kg or 0),float(x.net_kg or 0),
            round(float(Decimal(x.net_kg or 0)/Decimal('1000')),3),
            x.weigh_at.astimezone(TZ).strftime('%Y-%m-%d %H:%M:%S') if x.weigh_at else '',
            safe_cell(x.row_status),safe_cell(x.issue),safe_cell(x.batch_id)])
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for col in 'ABCDEFGHIJKLMNOPQ':sheet.column_dimensions[col].width=17
    for col in ['G','H','I','P']:sheet.column_dimensions[col].width=25
    buf=BytesIO();book.save(buf);buf.seek(0)
    name=f"TIOM_WB_Full_{meta['fromDate']}_to_{meta['toDate']}.xlsx"
    return StreamingResponse(buf,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
      headers={'Content-Disposition':f'attachment; filename="{name}"',
               'Cache-Control':'no-store'})

@router.post('/rpc')
def rpc(p: RPC, request: Request, db: Session = Depends(get_db)):
    csrf(request)
    user = get_user(db, request)
    read_methods = {'getBootstrap', 'getLiveContext', 'getDashboard', 'getAttendanceDesk', 'getShiftControl', 'getShiftCloseStatus', 'getUserAdminData', 'getProductionDesk', 'getHsdDesk', 'getMastersDesk', 'getWbDesk', 'getWbMonitor', 'getWbHistoryQueue', 'getImportHistory', 'getTiomMisDesk', 'getTiomMisReport', 'getTiomWbSuggestions', 'getTiomWbHistory', 'getTiomShiftProductionDesk', 'getTiomDrillingDesk', 'getTiomHsdDesk', 'getTiomHsdPreviousMeter'}
    if p.method not in read_methods:
        lock(db)
    try:
        result = dispatch(db, user, p.method, p.args)
        db.commit()
        return result
    except HTTPException:
        db.rollback()
        raise
    except (ValueError, TypeError, KeyError, AttributeError) as exc:
        db.rollback()
        raise HTTPException(422, 'Invalid request fields.') from exc
    except SQLAlchemyError as exc:
        db.rollback()
        ref=uuid4().hex[:8].upper()
        log.exception('Database error during RPC %s [%s]', p.method, ref)
        raise HTTPException(500, f'Database save failed. Nothing was committed. Error reference {ref}. Run validate_tiom_phase1_fix1.ps1 and check the server log.') from exc
    except Exception as exc:
        db.rollback()
        ref=uuid4().hex[:8].upper()
        log.exception('Unexpected error during RPC %s [%s]', p.method, ref)
        raise HTTPException(500, f'Operation failed. Nothing was committed. Error reference {ref}. Check the server log.') from exc


def dispatch(db, user, method, args):
    if method == 'getLiveContext':
        day, sh, now = operating_context()
        return {'date': day, 'shift': sh, 'time': now.strftime('%H:%M:%S')}
    if method == 'getBootstrap':
        day, current, now = operating_context()
        shifts = [r.shift for r in db.scalars(select(ShiftMaster).where(ShiftMaster.active))
                  if user.admin or 'ALL' in user.shifts.split(',') or r.shift in user.shifts.split(',')]
        assigned=set(user.modules.split(',')) if user.modules else set()
        modules = []
        if user.admin or assigned & {'PERSON_ATTENDANCE', 'EQUIPMENT_ATTENDANCE'}: modules.append('ATTENDANCE')
        for code in ['DASHBOARD','SHIFT_CONTROL','PRODUCTION','WB','HSD','MASTERS']:
            if user.admin or code in assigned: modules.append(code)
        # Mechanical uses the newer site-permission model. Surface it in the
        # dedicated TIOM launcher so field users can reach /tiom/mechanical
        # after normal login without duplicating legacy module assignments.
        if user.admin or any(has_permission(db, user, 'TIOM', 'MECHANICAL', action) for action in ('VIEW','CREATE','EDIT','EXPORT','ALL')):
            modules.append('MECHANICAL')
        return {'ok': True, 'today': day, 'shift': current if current in shifts else (shifts[0] if shifts else ''),
                'time': now.strftime('%H:%M:%S'), 'user': {'loginId': user.login_id, 'name': user.name,
                'isManagement': user.admin, 'modules': modules}, 'masters': {'shifts': shifts}}
    if method == 'saveMasterOption': return save_master_option(db,user,args[0] if args else {})
    if method == 'saveRotationTeam': return save_rotation_team(db,user,args[0] if args else {})
    if method == 'getDashboard':
        return dashboard_desk(db, user, args[0] if args and isinstance(args[0], dict) else {})
    if method == 'getProductionDesk': return production_desk(db,user)
    if method == 'getTiomMisDesk': return tiom_mis_desk(db,user,args[0] if args and isinstance(args[0],dict) else {})
    if method == 'getTiomMisReport': return get_tiom_mis_report(db,user,args[0] if args and isinstance(args[0],dict) else {})
    if method == 'getTiomWbSuggestions': return get_tiom_wb_suggestions(db,user,args[0] if args and isinstance(args[0],dict) else {})
    if method == 'getTiomWbHistory': return get_tiom_wb_history(db,user,args[0] if args and isinstance(args[0],dict) else {})
    if method == 'getTiomShiftProductionDesk': return get_tiom_shift_production_desk(db,user,args[0] if args and isinstance(args[0],dict) else {})
    if method == 'getTiomDrillingDesk': return tiom_drilling_desk(db,user,args[0] if args and isinstance(args[0],dict) else {})
    if method == 'getTiomHsdDesk': return tiom_hsd_desk(db,user,args[0] if args and isinstance(args[0],dict) else {})
    if method == 'getTiomHsdPreviousMeter': return tiom_hsd_previous_meter_info(db,user,args[0] if args and isinstance(args[0],dict) else {})
    if method == 'getHsdDesk': return hsd_desk(db,user)
    if method == 'getMastersDesk': return masters_desk(db,user)
    if method == 'getWbMonitor': return wb_monitor_v3(db, user, args[0] if args and isinstance(args[0], dict) else {})
    if method == 'getImportHistory': return wb_history_v3(db, user, (args[0].get('limit', 60) if args and isinstance(args[0], dict) else 60))
    if method.startswith('get') and 'wb' in method.lower() and 'histor' in method.lower():
        return wb_history_v3(db, user, (args[0].get('limit', 60) if args and isinstance(args[0], dict) else 60))
    if method == 'getWbDesk':
        if args and len(args)>=2 and args[0] and args[1]:
            return wb_desk_exact(db, user, args[0], args[1])
        return wb_desk(db,user)
    if method == 'getWbHistoryQueue':
        return wb_historical_queue(db, user, args[0] if args and isinstance(args[0], dict) else {})
    if method == 'startProductionTrip': return start_production_trip(db,user,args[0] if args else {})
    if method == 'markTripLoaded': return transition_trip(db,user,args[0] if args else {},'LOADED')
    if method == 'markTripUnloaded': return transition_trip(db,user,args[0] if args else {},'UNLOADED')
    if method == 'saveTiomMisDraft': return save_tiom_mis_report(db,user,args[0] if args else {},False)
    if method == 'submitTiomMisReport': return save_tiom_mis_report(db,user,args[0] if args else {},True)
    if method == 'saveTiomSourceDeployments': return save_tiom_source_deployments(db,user,args[0] if args else {})
    if method == 'saveTiomDrillingShift': return save_tiom_drilling_shift(db,user,args[0] if args else {})
    if method == 'removeTiomMisReport': return remove_tiom_mis_report(db,user,args[0] if args else {})
    if method == 'saveTiomManagementRecipients': return save_tiom_management_recipients(db,user,args[0] if args else {})
    if method == 'sendTiomManagementImage': return send_tiom_management_image(db,user,args[0] if args else {})
    if method == 'saveTiomTripFactor': return save_tiom_trip_factor(db,user,args[0] if args else {})
    if method == 'saveTiomShiftProductionDraft': return save_tiom_shift_production(db,user,args[0] if args else {},False)
    if method == 'submitTiomShiftProduction': return save_tiom_shift_production(db,user,args[0] if args else {},True)
    if method == 'saveTiomShiftBaselines': return save_tiom_shift_baselines(db,user,args[0] if args else {})
    if method == 'saveTiomHsdReceipt': return save_tiom_hsd_receipt(db,user,args[0] if args else {})
    if method == 'saveTiomHsdIssue': return save_tiom_hsd_issue(db,user,args[0] if args else {})
    if method == 'saveTiomHsdIssueBatch': return save_tiom_hsd_issue_batch(db,user,args[0] if args else {})
    if method == 'saveTiomHsdStockCheck': return save_tiom_hsd_stock_check(db,user,args[0] if args else {})
    if method == 'saveHsdReceipt': return save_hsd_receipt(db,user,args[0] if args else {})
    if method == 'saveHsdIssue': return save_hsd_issue(db,user,args[0] if args else {})
    if method == 'saveMasterRecord': return save_master_record(db,user,args[0] if args else {})
    if method == 'confirmWbBatch': return confirm_wb_batch(db,user,args[0] if args else {})
    if method in {'getUserAdminData', 'createLoginAccount', 'saveUserAccess', 'resetAccountPassword'}:
        return user_admin(db, user, method, args)
    if method == 'saveRotationAssignments':
        require(user, admin=True)
        for r in batch_rows(args[0]):
            person = active_resource(db, Person, r['id'])
            team = r.get('teamId', '')
            if team and team != 'GENERAL': active_resource(db, ShiftRotation, team)
            audit(db, user, 'ROSTER_TEAM', 'person', person.employee_id, {'before': person.rotation_team, 'after': team})
            person.rotation_team = team or None
        return {'ok': True, 'message': 'Rotation assignments saved.'}
    if method in {'getAttendanceDesk', 'getShiftControl', 'getShiftCloseStatus'}:
        day, sh, definition = context(db, user, args[0], args[1])
        if method == 'getAttendanceDesk': return attendance(db, user, day, sh, definition)
        if method == 'getShiftControl': return shift_desk(db, user, day, sh)
        return close_status(db, user, day, sh, definition)
    item = args[0] if args and isinstance(args[0], dict) else {}
    day, sh, definition = context(db, user, item.get('date'), item.get('shift'))
    if method == 'punchPersonAttendance': return punch(db, user, item, day, sh, definition)
    if method == 'saveEquipmentAttendance': return equipment_save(db, user, item, day, sh)
    if method == 'savePersonAttendance':
        require(user, 'PERSON_ATTENDANCE'); open_shift(db, day, sh)
        existing = {r.employee_id: r for r in rows(db, PersonAttendance, day, sh)}
        for r in batch_rows(item):
            row = existing.get(r['id'])
            if not row: raise HTTPException(409, 'Record attendance before saving remarks.')
            audit(db, user, 'ATTENDANCE_REMARK', 'person_attendance', row.id, {'before': row.remarks, 'after': r.get('remarks')})
            row.remarks = short(r.get('remarks', ''))
        return {'ok': True, 'message': 'Remarks saved.'}
    if method == 'applyShiftSetup': return setup_row(db, user, item, day, sh)
    if method == 'releaseEquipment': return release(db, user, item, day, sh)
    if method == 'applyShiftSetupBatch':
        skipped = []; done = 0
        for r in batch_rows(item):
            if not r.get('employeeId') or not r.get('locationId') or r.get('condition') != 'WORKING':
                skipped.append(r.get('machineId')); continue
            setup_row(db, user, r, day, sh); done += 1
        return {'ok': True, 'message': f'{done} deployed; {len(skipped)} incomplete rows skipped.', 'skipped': skipped}
    if method == 'autoCloseAttendance':
        require(user, 'PERSON_ATTENDANCE'); open_shift(db, day, sh)
        auto_close(db, user, day, sh, definition)
        return {'ok': True, 'message': 'Overdue OUT records checked; automatic OUT entries are marked REVIEW.'}
    if method == 'setShiftState':
        require(user, admin=True)
        if item.get('status') != 'CLOSED': raise HTTPException(409, 'Reopening is not available in this pilot.')
        status = close_status(db, user, day, sh, definition)
        if not status['canClose']: raise HTTPException(409, '; '.join(status['blockers']))
        reason = short(item.get('reason', '')).strip()
        if not reason: raise HTTPException(422, 'A closing reason is required.')
        state = db.scalar(select(ShiftState).where(ShiftState.operating_date == day, ShiftState.shift == sh))
        if not state:
            state = ShiftState(operating_date=day, shift=sh); db.add(state)
        state.status='CLOSED'; state.reason=reason; state.changed_by=user.login_id; state.changed_at=now_local()
        for r in rows(db, ShiftCrew, day, sh)+rows(db, ShiftDeployment, day, sh):
            if r.status == 'ACTIVE':
                r.status='CLOSED'; r.to_at=now_local(); r.reason=reason; r.changed_by=user.login_id; r.changed_at=now_local()
        audit(db, user, 'CLOSE_SHIFT', 'shift_state', f'{day}/{sh}', {'reason': reason})
        return {'ok': True, 'message': 'Shift closed.'}
    raise HTTPException(404, 'This operation has not been migrated yet.')


def user_admin(db, user, method, args):
    require(user, admin=True)
    if method == 'getUserAdminData':
        return [{'loginId': u.login_id, 'name': u.name, 'admin': u.admin, 'active': u.active,
                 'modules': u.modules.split(',') if u.modules else [], 'shifts': u.shifts.split(',')}
                for u in db.scalars(select(WebUser).order_by(WebUser.login_id))]
    item = args[0]; login_id = str(item.get('loginId', '')).lower()
    if not re.fullmatch(r'[a-z0-9_.-]{1,60}', login_id): raise HTTPException(422, 'Invalid user ID.')
    target = db.get(WebUser, login_id)
    if method == 'createLoginAccount':
        if target: raise HTTPException(409, 'User ID already exists.')
        name = item.get('name', '').strip()
        if not name or len(name) > 120: raise HTTPException(422, 'Name is required (120 characters maximum).')
        target = WebUser(login_id=login_id, name=name, password_hash=hash_password(item.get('password', '')),
                         modules='', shifts='', admin=False, active=True)
        db.add(target)
    elif not target:
        raise HTTPException(404, 'Account not found.')
    elif method == 'resetAccountPassword':
        target.password_hash = hash_password(item.get('password', ''))
        target.failed_attempts = 0; target.locked_until = None
    elif method == 'saveUserAccess':
        modules, shifts = item.get('modules', []), item.get('shifts', [])
        if not isinstance(modules, list) or not isinstance(shifts, list): raise HTTPException(422, 'Invalid access list.')
        if not set(modules) <= MODULES: raise HTTPException(422, 'Unknown module.')
        allowed = set(db.scalars(select(ShiftMaster.shift))) | {'ALL'}
        if not set(shifts) <= allowed: raise HTTPException(422, 'Unknown shift.')
        if type(item.get('admin')) is not bool or type(item.get('active')) is not bool: raise HTTPException(422, 'Invalid account flags.')
        if target.login_id == user.login_id and (not item['admin'] or not item['active']):
            raise HTTPException(409, 'You cannot disable or remove your own management access.')
        target.modules=','.join(sorted(set(modules))); target.shifts=','.join(sorted(set(shifts)))
        target.admin=item['admin']; target.active=item['active']
    if method != 'createLoginAccount':
        db.execute(delete(WebSession).where(WebSession.login_id == target.login_id))
    audit(db, user, method, 'web_user', login_id, {'account': login_id})
    return {'ok': True, 'message': 'Account saved. Changed accounts must sign in again.'}