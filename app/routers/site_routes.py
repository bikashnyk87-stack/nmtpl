"""KOCP Coal and OB movement masters and derived, shared physical locations.

Coal workbook's first column (Reject / Coal Stock / Siding) is a MOVEMENT GROUP,
not a physical source. OB workbook has real From and To locations. Never turn
distance ranges into numeric SiteLocation.distance_km or combine separate routes.
"""
from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from uuid import uuid4
import json
import re

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.auth import csrf, get_user
from app.db import get_db
from app.site_auth import has_permission, require_site
from app.site_models import MasterImportBatch, MasterImportRow, SiteAuditLog, SiteLocation, SiteMovementRoute

router = APIRouter(prefix="/api/sites", tags=["KOCP routes"])
_RANGE = re.compile(r"^\d+(?:\.\d+)?-\d+(?:\.\d+)?(?:/\d+(?:\.\d+)?-\d+(?:\.\d+)?)?$")
_KIND = {"COAL", "OB"}


def _access(db, request, site_id, action="VIEW"):
    user = get_user(db, request)
    site_id = require_site(db, user, site_id)
    if site_id != "KOCP":
        raise HTTPException(404, "Coal and OB routes are only configured for KOCP.")
    if not user.admin and not has_permission(db, user, site_id, "MASTERS", action):
        raise HTTPException(403, f"Masters {action.lower()} permission is required.")
    return user


def _kind(kind):
    key = str(kind).strip().upper()
    if key not in _KIND:
        raise HTTPException(404, "Unknown movement master.")
    return key


def _str(value):
    return " ".join(str(value or "").strip().split())


def _norm(value):
    return _str(value).casefold()


def _flag(value):
    return str(value if value is not None else "TRUE").strip().upper() in {"1", "TRUE", "YES", "Y", "ACTIVE"}


def _distance(value):
    raw = _str(value).replace(" ", "")
    if not _RANGE.fullmatch(raw):
        raise ValueError("Distance must be a range such as 0-1, 2-3 or 1-2/3-4.")
    for part in raw.split("/"):
        a, b = (float(x) for x in part.split("-"))
        if b <= a:
            raise ValueError("Distance range upper bound must exceed lower bound.")
    return raw


def _extract(kind, headers, row):
    raw = {str(headers[i]).strip().casefold(): row[i] for i in range(min(len(headers),len(row)))}
    def get(*names):
        for name in names:
            if name.casefold() in raw:
                return raw[name.casefold()]
        return None
    legacy = "distance km" in raw and "role" in raw and ("location name" in raw or "to location name" in raw)
    if legacy:
        group_value = row[0] if row else None
        destination = get("Location Name", "To Location Name")
        category = get("Role")
        distance = get("Distance KM")
        location_type = get("Location Type")
        active = get("Active")
        source = group_value if kind == "OB" else None
        group = group_value if kind == "COAL" else category
    else:
        group = get("Movement Group", "Route Group", "Group")
        source = get("Source", "From")
        destination = get("Destination", "To Location Name")
        category = get("Category", "Movement Category", "OB Category")
        distance = get("Distance Range", "Distance KM")
        location_type = get("Location Type")
        active = get("Active")
        if kind == "OB" and not group:
            group = category
    if kind == "COAL" and not group:
        group = category
    return {
        "kind": kind,
        "group": _str(group),
        "source": _str(source) if kind == "OB" else "",
        "destination": _str(destination),
        "category": _str(category),
        "distance": _str(distance),
        "locationType": _str(location_type),
        "active": _flag(active),
    }


def _issues(item):
    missing = [n for n in ("group", "destination", "category", "distance") if not item.get(n)]
    if item.get("kind") == "OB" and not item.get("source"):
        missing.append("source")
    issues = [f"{x} is required" for x in missing]
    for field, limit in (("group",160),("source",160),("destination",160),("category",160),("distance",50),("locationType",60)):
        if len(item.get(field) or "") > limit:
            issues.append(f"{field} exceeds {limit} characters")
    if item.get("distance"):
        try:
            _distance(item["distance"])
        except ValueError as exc:
            issues.append(str(exc))
    return issues


def _identity(item):
    parts = [item["kind"], _norm(item["group"]), _norm(item["source"]), _norm(item["destination"]),
             _norm(item["category"]), _distance(item["distance"])]
    return "KOCP-RTE-" + sha256(json.dumps(parts, ensure_ascii=False).encode()).hexdigest()[:24].upper()


def _route_dict(row):
    return {"routeId":row.route_id,"kind":row.movement_kind,"group":row.route_group,
            "category":row.category,"source":row.source_label,"destination":row.destination_label,
            "sourceLocationId":row.source_location_id,"destinationLocationId":row.destination_location_id,
            "distanceRange":row.distance_range,"locationType":row.location_type,"active":row.active}


def _physical_location(db, name, by_name, batch_id):
    key = _norm(name)
    loc = by_name.get(key)
    if loc:
        return loc
    code = "AUTO-" + sha256(key.encode()).hexdigest()[:20].upper()
    loc = SiteLocation(site_location_id=f"KOCP:{code}", site_id="KOCP", code=code,
                       name=name, active=True, import_batch_id=batch_id)
    by_name[key] = loc
    db.add(loc)
    return loc


def _audit(db,user,action,entity,entity_id,after):
    db.add(SiteAuditLog(site_id="KOCP",actor=user.login_id,action=action,
                        entity=entity,entity_id=entity_id,after_json=json.dumps(after)))


@router.get("/{site_id}/routes/{kind}/template.xlsx")
def route_template(site_id:str, kind:str,request:Request,db:Session=Depends(get_db)):
    _access(db,request,site_id,"VIEW"); kind=_kind(kind)
    wb=Workbook();ws=wb.active;ws.title=f"{kind} Routes"
    if kind=="COAL":
        ws.append(["Movement Group","Destination","Category","Location Type","Distance Range","Active"])
        ws.append(["Coal Stock","Coal Stock KN-5/KN-6","Stock","UNLOADING","1-2/3-4","TRUE"])
    else:
        ws.append(["Source","Destination","Category","Location Type","Distance Range","Active"])
        ws.append(["Working Face","Dump Yard","Soft OB (Without Drilling)","UNLOADING","0-1","TRUE"])
    ws.freeze_panes="A2"
    for col in "ABCDEF":ws.column_dimensions[col].width=28
    data=BytesIO();wb.save(data);data.seek(0)
    return StreamingResponse(data,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
         headers={"Content-Disposition":f'attachment; filename="NMTPL_KOCP_{kind}_Route_Master.xlsx"'})


@router.post("/{site_id}/routes/{kind}/preview")
async def route_preview(site_id:str,kind:str,request:Request,file:UploadFile=File(...),db:Session=Depends(get_db)):
    csrf(request);user=_access(db,request,site_id,"CREATE");kind=_kind(kind)
    payload=await file.read()
    if len(payload)>15*1024*1024:raise HTTPException(413,"Upload exceeds 15 MB.")
    try:
        wb=load_workbook(BytesIO(payload),read_only=True,data_only=True)
        ws=wb.worksheets[0]
        headers=[_str(c.value) for c in next(ws.iter_rows(min_row=1,max_row=1))]
    except Exception as exc:
        raise HTTPException(422,"Upload a valid XLSX workbook.") from exc
    norm_headers={h.casefold() for h in headers}
    legacy="role" in norm_headers and "distance km" in norm_headers and (
        "location name" in norm_headers or "to location name" in norm_headers)
    canonical={"destination","category","distance range"}<=norm_headers
    has_source=kind=="COAL" or ("source" in norm_headers or "from" in norm_headers) or legacy
    if not (legacy or canonical) or not has_source:
        raise HTTPException(422,"Wrong layout. Select the matching Coal/OB master and use its sample Excel.")
    staged=[];seen=set();errors=[]
    for no,cells in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
        if not any(v not in (None,"") for v in cells):continue
        item=_extract(kind,headers,cells)
        issues=_issues(item)
        if not issues:
            identity=_identity(item)
            if identity in seen:issues.append("Duplicate route in this workbook.")
            seen.add(identity)
        staged.append((no,item,issues))
        if issues:errors.append({"row":no,"errors":issues})
    if not staged:raise HTTPException(422,"No route rows in the workbook.")
    batch_id=f"MIMP-KOCP-{uuid4().hex[:12].upper()}"
    db.add(MasterImportBatch(batch_id=batch_id,site_id="KOCP",master_type=f"{kind}_ROUTE",
           file_name=file.filename or "routes.xlsx",file_hash=sha256(payload).hexdigest(),
           status="PREVIEW",total_rows=len(staged),valid_rows=len(staged)-len(errors),
           update_rows=0,error_rows=len(errors),uploaded_by=user.login_id))
    for no,item,issues in staged:
        db.add(MasterImportRow(batch_id=batch_id,row_no=no,row_json=json.dumps(item),
               status="ERROR" if issues else "VALID",error_text="; ".join(issues) if issues else None))
    _audit(db,user,"ROUTE_IMPORT_PREVIEW","master_import_batch",batch_id,
           {"kind":kind,"rows":len(staged),"errors":len(errors)})
    db.commit()
    return {"ok":not errors,"batchId":batch_id,"kind":kind,"totalRows":len(staged),
            "errorRows":len(errors),"errors":errors[:100],"sample":[{"row":n,**v} for n,v,_ in staged[:25]]}


@router.post("/{site_id}/routes/imports/{batch_id}/confirm")
def route_confirm(site_id:str,batch_id:str,request:Request,db:Session=Depends(get_db)):
    csrf(request);user=_access(db,request,site_id,"APPROVE")
    batch=db.get(MasterImportBatch,batch_id)
    if not batch or batch.site_id!="KOCP" or batch.master_type not in {"COAL_ROUTE","OB_ROUTE"}:
        raise HTTPException(404,"Route import batch not found.")
    if batch.status!="PREVIEW" or batch.error_rows:
        raise HTTPException(409,"Only a clean PREVIEW can be imported.")
    staged=list(db.scalars(select(MasterImportRow).where(MasterImportRow.batch_id==batch_id).order_by(MasterImportRow.row_no)))
    parsed=[];seen=set()
    for row in staged:
        item=json.loads(row.row_json)
        issues=_issues(item)
        route_id=_identity(item) if not issues else None
        if route_id in seen:issues.append("Duplicate route")
        if issues:raise HTTPException(422,f"Excel row {row.row_no}: {'; '.join(issues)}. Nothing imported.")
        seen.add(route_id);parsed.append((row,item,route_id))
    by_name={_norm(x.name):x for x in db.scalars(select(SiteLocation).where(SiteLocation.site_id=="KOCP").order_by(SiteLocation.site_location_id))}
    created=updated=0
    for row,item,route_id in parsed:
        source=_physical_location(db,item["source"],by_name,batch_id) if item["kind"]=="OB" else None
        destination=_physical_location(db,item["destination"],by_name,batch_id)
        route=db.get(SiteMovementRoute,route_id)
        if route is None:
            route=SiteMovementRoute(route_id=route_id,site_id="KOCP",movement_kind=item["kind"])
            db.add(route);created+=1
        else:
            updated+=1
        route.route_group=item["group"]
        route.category=item["category"]
        route.source_label=item["source"] or None
        route.destination_label=item["destination"]
        route.source_location_id=source.site_location_id if source else None
        route.destination_location_id=destination.site_location_id
        route.distance_range=_distance(item["distance"])
        route.location_type=item["locationType"] or None
        route.active=item["active"]
        route.import_batch_id=batch_id
        row.status="IMPORTED"
    batch.status="IMPORTED";batch.update_rows=updated
    batch.confirmed_by=user.login_id
    from datetime import datetime,timezone
    batch.confirmed_at=datetime.now(timezone.utc)
    _audit(db,user,"ROUTE_IMPORT_CONFIRM","master_import_batch",batch_id,
           {"created":created,"updated":updated,"rows":len(parsed)})
    db.commit()
    return {"ok":True,"created":created,"updated":updated,"rows":len(parsed)}


@router.get("/{site_id}/routes")
def list_routes(site_id:str,request:Request,kind:str|None=None,db:Session=Depends(get_db)):
    user=_access(db,request,site_id,"VIEW")
    stmt=select(SiteMovementRoute).where(SiteMovementRoute.site_id=="KOCP")
    if kind:stmt=stmt.where(SiteMovementRoute.movement_kind==_kind(kind))
    rows=db.scalars(stmt.order_by(SiteMovementRoute.movement_kind,SiteMovementRoute.route_group,
                                  SiteMovementRoute.category,SiteMovementRoute.distance_range)).all()
    return [_route_dict(r) for r in rows]


@router.get("/{site_id}/routes/locations")
def route_locations(site_id:str,request:Request,db:Session=Depends(get_db)):
    _access(db,request,site_id,"VIEW")
    rows=list(db.scalars(select(SiteLocation).where(SiteLocation.site_id=="KOCP").order_by(SiteLocation.name)))
    routes=list(db.scalars(select(SiteMovementRoute).where(SiteMovementRoute.site_id=="KOCP")))
    used={}
    for r in routes:
        for x in (r.source_location_id,r.destination_location_id):
            if x:used.setdefault(x,set()).add(r.movement_kind)
    return [{"locationId":r.site_location_id,"name":r.name,"code":r.code,
             "usedIn":sorted(used.get(r.site_location_id,set())),"active":r.active} for r in rows]


@router.post("/{site_id}/routes/{route_id}")
def edit_route(site_id:str,route_id:str,payload:dict,request:Request,db:Session=Depends(get_db)):
    csrf(request);user=_access(db,request,site_id,"EDIT")
    route=db.get(SiteMovementRoute,route_id)
    if not route or route.site_id!="KOCP":raise HTTPException(404,"Route not found.")
    before=_route_dict(route)
    if "active" in payload:route.active=bool(payload["active"])
    if "category" in payload:
        value=_str(payload["category"])
        if not value or len(value)>160:raise HTTPException(422,"Category must be 1 to 160 characters.")
        route.category=value
    if "distanceRange" in payload:
        try:route.distance_range=_distance(payload["distanceRange"])
        except ValueError as exc:raise HTTPException(422,str(exc)) from exc
    _audit(db,user,"ROUTE_EDIT","site_movement_route",route_id,{"before":before,"after":_route_dict(route)})
    db.commit()
    return {"ok":True,"record":_route_dict(route)}
