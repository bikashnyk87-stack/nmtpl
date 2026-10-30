"""KOCP Coal and OB material catalog import.

The field Excel workbooks contain 21 Coal and 11 OB rows, with material
names repeated for each distance range. These are not 32 distinct material
types: normalize to exactly the five canonical classifications while
preserving every original Excel row in audit staging.
"""
from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.auth import csrf, get_user
from app.db import get_db
from app.site_auth import has_permission, require_site
from app.site_models import (
    MasterImportBatch, MasterImportRow, SiteAuditLog, SiteMaterial, SiteMovementRoute
)
from app.routers.site_routes import _distance

router = APIRouter(prefix="/api/sites", tags=["KOCP material catalog"])

CATEGORIES = {
    "COAL": {
        "COAL": ("Coal", "COAL", "MT"),
        "REJECT": ("Reject", "REJECT", "MT"),
    },
    "OB": {
        "SOFT_OB": ("Soft OB (Without Drilling)", "OB", "CUM"),
        "HARD_OB": ("Hard OB (With Drilling)", "OB", "CUM"),
        "REHANDLING_OB": ("Re-handling OB", "OB", "CUM"),
    },
}


def _access(db, request, site_id, action):
    user=get_user(db,request)
    site_id=require_site(db,user,site_id)
    if site_id!="KOCP":
        raise HTTPException(404,"KOCP material catalog only.")
    if not user.admin and not has_permission(db,user,site_id,"MASTERS",action):
        raise HTTPException(403,f"MASTERS {action} permission required.")
    return user


def _kind(value):
    kind=str(value or "").upper().strip()
    if kind not in CATEGORIES:
        raise HTTPException(404,"Select Coal or OB materials.")
    return kind


def _clean(value):
    return " ".join(str(value or "").strip().split())


def _norm(value):
    return _clean(value).casefold()


def _classify(kind,name):
    text=_norm(name).replace("–","-").replace("—","-")
    if kind=="COAL":
        if "reject" in text:return "REJECT"
        if "coal" in text:return "COAL"
    else:
        if "re-hand" in text or "re hand" in text:return "REHANDLING_OB"
        if "soft" in text and "ob" in text:return "SOFT_OB"
        if "hard" in text and "ob" in text:return "HARD_OB"
    return None


def _as_row(kind,headers,cells):
    raw={_norm(headers[i]):cells[i] for i in range(min(len(headers),len(cells)))}
    def val(key):return raw.get(key.casefold())
    name=_clean(val("Material Name"))
    code=_classify(kind,name)
    distance=_clean(val("Distance KM"))
    unit=_clean(val("Default Unit")).upper()
    billable=_clean(val("Billable Unit")).upper() or unit
    return {
        "kind":kind,"materialType":code,"originalCode":_clean(val("Code")),
        "originalName":name,"originalGroup":_clean(val("Material Group")),
        "distanceRange":distance,
        "defaultUnit":unit,"billableUnit":billable,
        "active":_clean(val("Active")).upper() not in {"FALSE","NO","N","0","INACTIVE"},
    }


def _issues(item):
    issues=[]
    kind=item["kind"]
    code=item["materialType"]
    if not code:issues.append("Unrecognized material name. Expected Coal/Reject or Soft/Hard/Re-handling OB.")
    if code:
        expected=CATEGORIES[kind][code][2]
        if item["defaultUnit"]!=expected:
            issues.append(f"{kind} must use {expected} as Default Unit.")
        if item["billableUnit"]!=expected:
            issues.append(f"{kind} must use {expected} as Billable Unit.")
    distance=item["distanceRange"]
    if distance:
        try:_distance(distance)
        except ValueError as e:issues.append(str(e))
    if len(item["originalCode"])>50:issues.append("Original Code exceeds 50 characters.")
    return issues


def _material_existing(db,kind,code):
    name,group,unit=CATEGORIES[kind][code]
    records=list(db.scalars(select(SiteMaterial).where(SiteMaterial.site_id=="KOCP")))
    compatible=[]
    for row in records:
        if (_norm(row.name)==_norm(name) and str(row.default_unit).upper()==unit):
            compatible.append(row)
    if compatible:
        return sorted(compatible,key=lambda x:x.code)[0],None
    matched=next((x for x in records if str(x.code).upper()==code),None)
    if matched and (_norm(matched.name)!=_norm(name) or str(matched.default_unit).upper()!=unit):
        return None,f"Material code {code} already belongs to {matched.name}. Existing data will not be overwritten."
    return matched,None


def _catalog(db,kind=None):
    allrows=list(db.scalars(select(SiteMaterial).where(SiteMaterial.site_id=="KOCP").order_by(SiteMaterial.code)))
    result={}
    for k in ("COAL","OB"):
        if kind and k!=kind:continue
        result[k]=[]
        for code,(name,group,unit) in CATEGORIES[k].items():
            matched=next((x for x in allrows if _norm(x.name)==_norm(name) and str(x.default_unit).upper()==unit),None)
            if not matched: matched=next((x for x in allrows if str(x.code).upper()==code
                                             and str(x.default_unit).upper()==unit),None)
            result[k].append({"typeCode":code,"name":name,"unit":unit,
                "materialId":matched.site_material_id if matched else None,
                "storedCode":matched.code if matched else None,
                "active":matched.active if matched else False,
                "status":"EXISTS" if matched else "MISSING"})
    return result


@router.get("/{site_id}/material-types")
def material_types(site_id:str,request:Request,db:Session=Depends(get_db)):
    _access(db,request,site_id,"VIEW")
    return _catalog(db)


@router.get("/{site_id}/material-types/{kind}/template.xlsx")
def material_template(site_id:str,kind:str,request:Request,db:Session=Depends(get_db)):
    _access(db,request,site_id,"VIEW");kind=_kind(kind)
    wb=Workbook();ws=wb.active;ws.title=kind+" Material"
    ws.append(["Code","Material Name","Material Group","Default Unit","Billable Unit","Active"])
    for code,(name,group,unit) in CATEGORIES[kind].items():
        ws.append([code,name,group,unit,unit,"TRUE"])
    ws.freeze_panes="A2"
    for col in "ABCDEF":ws.column_dimensions[col].width=25
    note=wb.create_sheet("Instructions")
    note.append(["One row per material TYPE; route/distance variations are maintained in Coal/OB Route Master."])
    note.append(["Legacy 21-row Coal and 11-row OB spreadsheets are also accepted and automatically consolidated."])
    data=BytesIO();wb.save(data);data.seek(0)
    return StreamingResponse(data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":f'attachment; filename="NMTPL_KOCP_{kind}_Material_Master.xlsx"'})


@router.post("/{site_id}/material-types/{kind}/preview")
async def material_preview(site_id:str,kind:str,request:Request,file:UploadFile=File(...),db:Session=Depends(get_db)):
    csrf(request);user=_access(db,request,site_id,"CREATE");kind=_kind(kind)
    payload=await file.read()
    if len(payload)>15*1024*1024:raise HTTPException(413,"Excel exceeds 15 MB.")
    try:
        wb=load_workbook(BytesIO(payload),read_only=True,data_only=True)
        ws=wb.worksheets[0]
        headers=[_clean(c.value) for c in next(ws.iter_rows(min_row=1,max_row=1))]
    except Exception as e:raise HTTPException(422,"Upload a valid XLSX workbook.") from e
    needed={"material name","default unit"}
    if not needed <= {_norm(x) for x in headers}:
        raise HTTPException(422,"Expected Material Name and Default Unit columns. Use the matching Coal/OB material sample Excel.")
    staged=[]
    errors=[]
    grouped={}
    for no,cells in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
        if not any(x not in (None,"") for x in cells):continue
        item=_as_row(kind,headers,cells)
        issues=_issues(item)
        if issues:errors.append({"row":no,"errors":issues})
        elif item["materialType"]:
            grouped.setdefault(item["materialType"],[]).append(no)
        staged.append((no,item,issues))
    if not staged:raise HTTPException(422,"No material rows found.")
    conflicts=[]
    for code in grouped:
        _,conflict=_material_existing(db,kind,code)
        if conflict:conflicts.append(conflict)
    if conflicts:
        errors.extend({"row":0,"errors":[msg]} for msg in conflicts)
    # Route matching is informative, not a destructive source of truth. All
    # original codes, groups and ranges remain staged for audit.
    routes=list(db.scalars(select(SiteMovementRoute).where(
        SiteMovementRoute.site_id=="KOCP",SiteMovementRoute.movement_kind==kind)))
    match_rows=0
    missed=[]
    for no,item,issues in staged:
        if issues or not item["distanceRange"]:continue
        matching=[r for r in routes if r.distance_range==item["distanceRange"] and (
            (kind=="OB" and _classify(kind,r.category)==item["materialType"]) or
            (kind=="COAL" and (
                (item["materialType"]=="REJECT" and "reject" in _norm(r.route_group+" "+r.category))
                or (item["materialType"]=="COAL" and "reject" not in _norm(r.route_group+" "+r.category))
            ))
        )]
        if matching:match_rows+=1
        else:missed.append(no)
    batch_id=f"MIMP-KOCP-{uuid4().hex[:12].upper()}"
    db.add(MasterImportBatch(batch_id=batch_id,site_id="KOCP",master_type=f"{kind}_MATERIAL",
        file_name=file.filename or "materials.xlsx",file_hash=sha256(payload).hexdigest(),
        status="PREVIEW",total_rows=len(staged),
        valid_rows=len(staged)-len([x for x in staged if x[2]]),
        update_rows=0,error_rows=len(errors),uploaded_by=user.login_id,
        notes=json.dumps({"materialTypes":sorted(grouped),"routeMatchedRows":match_rows,
                          "routeUnmatchedRows":missed})))
    for no,item,issues in staged:
        db.add(MasterImportRow(batch_id=batch_id,row_no=no,row_json=json.dumps(item),
             status="ERROR" if issues else "VALID",error_text="; ".join(issues) if issues else None))
    db.add(SiteAuditLog(site_id="KOCP",actor=user.login_id,action="MATERIAL_TYPE_PREVIEW",
        entity="master_import_batch",entity_id=batch_id,
        after_json=json.dumps({"kind":kind,"rows":len(staged),"types":sorted(grouped),"errors":len(errors)})))
    db.commit()
    return {"ok":not errors,"batchId":batch_id,"kind":kind,"totalRows":len(staged),
         "materialTypes":[{"code":code,"name":CATEGORIES[kind][code][0],"rows":len(nos)} for code,nos in grouped.items()],
         "materialTypeCount":len(grouped),"routeMatchedRows":match_rows,
         "routeUnmatchedRows":missed[:30],"errorRows":len(errors),"errors":errors[:60]}


@router.post("/{site_id}/material-types/imports/{batch_id}/confirm")
def material_confirm(site_id:str,batch_id:str,request:Request,db:Session=Depends(get_db)):
    csrf(request);user=_access(db,request,site_id,"APPROVE")
    batch=db.get(MasterImportBatch,batch_id)
    if not batch or batch.site_id!="KOCP" or batch.master_type not in {"COAL_MATERIAL","OB_MATERIAL"}:
        raise HTTPException(404,"Material import batch not found.")
    if batch.status!="PREVIEW" or batch.error_rows:
        raise HTTPException(409,"Only an error-free material preview can be confirmed.")
    kind=batch.master_type.removesuffix("_MATERIAL")
    rows=list(db.scalars(select(MasterImportRow).where(MasterImportRow.batch_id==batch_id).order_by(MasterImportRow.row_no)))
    groups={}
    for row in rows:
        item=json.loads(row.row_json)
        issues=_issues(item)
        if item["kind"]!=kind:issues.append("Wrong material kind.")
        if issues:raise HTTPException(422,f"Row {row.row_no}: {'; '.join(issues)}. No material records imported.")
        groups.setdefault(item["materialType"],[]).append(row)
    plans=[]
    for code in groups:
        existing,conflict=_material_existing(db,kind,code)
        if conflict:raise HTTPException(409,conflict)
        plans.append((code,existing))
    created=reused=0
    for code,existing in plans:
        if existing:
            reused+=1
        else:
            name,group,unit=CATEGORIES[kind][code]
            record=SiteMaterial(site_material_id=f"KOCP:{code}",site_id="KOCP",code=code,
                name=name,material_group=group,default_unit=unit,billable_unit=unit,active=True)
            db.add(record);created+=1
    for row in rows:row.status="IMPORTED"
    batch.status="IMPORTED";batch.confirmed_by=user.login_id
    batch.confirmed_at=datetime.now(timezone.utc);batch.update_rows=reused
    db.add(SiteAuditLog(site_id="KOCP",actor=user.login_id,action="MATERIAL_TYPE_IMPORT",
        entity="master_import_batch",entity_id=batch_id,
        after_json=json.dumps({"sourceRows":len(rows),"distinctTypes":len(groups),"created":created,
                          "reusedExisting":reused,"kind":kind})))
    db.commit()
    return {"ok":True,"sourceRows":len(rows),"materialTypes":len(groups),
            "created":created,"reusedExisting":reused}
