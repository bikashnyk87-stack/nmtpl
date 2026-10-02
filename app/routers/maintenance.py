from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import uuid4
import json

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session

from app.auth import csrf, get_user
from app.db import get_db
from app.models import Equipment, EquipmentAttendance, HsdIssue, Location, Person, ShiftCrew, ShiftDeployment
from app.site_auth import require_permission, require_site
from app.site_models import Site, EquipmentSiteAssignment, SiteAssetAttendance, SiteAssetMeter, SiteAuditLog, SiteDeployment, TiomSourceDeployment
from app.services.site_context import resolve_site_context
from app.maintenance_models import (
    ComponentChangeHistory,
    EquipmentComponentSchedule,
    EquipmentDocument,
    MaintenanceBreakdown,
    MaintenanceComponentMaster,
    MaintenanceDefProfile,
    MaintenanceDefTransaction,
    MaintenanceJobCard,
    MaintenanceLubricantUsage,
    MaintenancePmExecution,
    MaintenancePmPlan,
    MaintenanceServiceHistory,
    MaintenanceServicePlan,
    MaintenanceSpareUsage,
    MaintenanceTyre,
    MaintenanceTyreFitment,
)
from app.maintenance_schemas import (
    BreakdownCloseIn,
    BreakdownIn,
    ComponentChangeIn,
    ComponentMasterIn,
    ComponentScheduleIn,
    DefProfileIn,
    DefTransactionIn,
    EquipmentDocumentIn,
    JobCardIn,
    JobCardUpdateIn,
    LubricantUsageIn,
    PmExecutionIn,
    PmPlanIn,
    ServiceCompleteIn,
    ServicePlanIn,
    TyreFitmentIn,
    TyreIn,
    TyreRemoveIn,
    UsageIn,
)

router = APIRouter(prefix="/api/maintenance", tags=["mechanical-maintenance"])


def _id(prefix: str, site_id: str) -> str:
    return f"{prefix}-{site_id}-{uuid4().hex[:14].upper()}"


def _audit(db: Session, user, site_id: str, action: str, entity: str, entity_id: str, before=None, after=None, reason=None):
    db.add(SiteAuditLog(
        site_id=site_id,
        actor=user.login_id,
        action=action,
        entity=entity,
        entity_id=entity_id,
        before_json=json.dumps(before, default=str) if before is not None else None,
        after_json=json.dumps(after, default=str) if after is not None else None,
        reason=reason,
    ))


def _access(db: Session, request: Request, site_id: str, action: str):
    user = get_user(db, request)
    site_id = require_site(db, user, site_id)
    require_permission(db, user, site_id, "MECHANICAL", action)
    return user, site_id


def _asset(db: Session, site_id: str, asset_id: str) -> Equipment:
    row = db.get(Equipment, asset_id)
    if not row or not row.active:
        raise HTTPException(422, "Choose an active equipment record.")
    if site_id != "TIOM":
        assigned = db.scalar(select(EquipmentSiteAssignment.id).where(
            EquipmentSiteAssignment.site_id == site_id,
            EquipmentSiteAssignment.machine_id == asset_id,
            EquipmentSiteAssignment.active.is_(True),
        ).limit(1))
        if not assigned:
            raise HTTPException(422, "Equipment is not assigned to this site.")
    return row


def _site_assets(db: Session, site_id: str):
    if site_id == "TIOM":
        return list(db.scalars(select(Equipment).where(Equipment.active.is_(True)).order_by(Equipment.machine_id)))
    return [x for (x,) in db.execute(
        select(Equipment).join(EquipmentSiteAssignment, EquipmentSiteAssignment.machine_id == Equipment.machine_id).where(
            EquipmentSiteAssignment.site_id == site_id,
            EquipmentSiteAssignment.active.is_(True),
            Equipment.active.is_(True),
        ).order_by(Equipment.machine_id)
    ).all()]


def _latest_meter(db: Session, site_id: str, asset_id: str, meter_type: str | None = None):
    stmt = select(SiteAssetMeter).where(SiteAssetMeter.site_id == site_id, SiteAssetMeter.asset_id == asset_id)
    if meter_type:
        stmt = stmt.where(SiteAssetMeter.meter_type == meter_type.upper())
    rows = list(db.scalars(stmt.order_by(SiteAssetMeter.operating_date.desc(), SiteAssetMeter.entered_at.desc()).limit(20)))
    for row in rows:
        reading = row.closing_reading if row.closing_reading is not None else row.opening_reading
        if reading is not None:
            return row.meter_type, Decimal(reading), row.operating_date, row.shift
    return None, None, None, None


def _current_operator_location(db: Session, site_id: str, asset_id: str, operating_date: date, shift: str):
    employee_id = None
    location_id = None
    if site_id == "TIOM":
        crew = db.scalar(select(ShiftCrew).where(
            ShiftCrew.operating_date == operating_date,
            ShiftCrew.shift == shift,
            ShiftCrew.machine_id == asset_id,
            ShiftCrew.status == "ACTIVE",
        ).order_by(ShiftCrew.from_at.desc()).limit(1))
        if crew:
            employee_id = crew.employee_id
        dep = db.scalar(select(ShiftDeployment).where(
            ShiftDeployment.operating_date == operating_date,
            ShiftDeployment.shift == shift,
            ShiftDeployment.machine_id == asset_id,
            ShiftDeployment.status == "ACTIVE",
        ).order_by(ShiftDeployment.from_at.desc()).limit(1))
        if dep:
            location_id = dep.location_id
        if not location_id:
            source_dep = db.scalar(select(TiomSourceDeployment).where(
                TiomSourceDeployment.operating_date == operating_date,
                TiomSourceDeployment.shift == shift,
                TiomSourceDeployment.machine_id == asset_id,
                TiomSourceDeployment.active.is_(True),
            ).order_by(TiomSourceDeployment.from_at.desc()).limit(1))
            if source_dep:
                location_id = source_dep.source_location_id
    else:
        dep = db.scalar(select(SiteDeployment).where(
            SiteDeployment.site_id == site_id,
            SiteDeployment.operating_date == operating_date,
            SiteDeployment.shift == shift,
            SiteDeployment.asset_id == asset_id,
            SiteDeployment.status == "ACTIVE",
        ).order_by(SiteDeployment.from_at.desc()).limit(1))
        if dep:
            employee_id = dep.employee_id
            location_id = dep.location_id
    return employee_id, location_id


def _running_assets(db: Session, site_id: str, operating_date: date, shift: str) -> set[str]:
    out = set()
    if site_id == "TIOM":
        out.update(x for x in db.scalars(select(TiomSourceDeployment.machine_id).where(
            TiomSourceDeployment.operating_date == operating_date,
            TiomSourceDeployment.shift == shift,
            TiomSourceDeployment.active.is_(True),
        )) if x)
        out.update(x for x in db.scalars(select(ShiftDeployment.machine_id).where(
            ShiftDeployment.operating_date == operating_date,
            ShiftDeployment.shift == shift,
            ShiftDeployment.status == "ACTIVE",
        )) if x)
    else:
        out.update(x for x in db.scalars(select(SiteDeployment.asset_id).where(
            SiteDeployment.site_id == site_id,
            SiteDeployment.operating_date == operating_date,
            SiteDeployment.shift == shift,
            SiteDeployment.status == "ACTIVE",
        )) if x)
    return out


def _idle_assets(db: Session, site_id: str, operating_date: date, shift: str) -> set[str]:
    idle = set()
    terms = {"IDLE", "STANDBY", "STAND BY", "NOT RUNNING"}
    if site_id == "TIOM":
        for row in db.scalars(select(EquipmentAttendance).where(
            EquipmentAttendance.operating_date == operating_date,
            EquipmentAttendance.shift == shift,
        )):
            if str(row.status or "").upper() in terms or str(row.condition or "").upper() in terms:
                idle.add(row.machine_id)
    else:
        for row in db.scalars(select(SiteAssetAttendance).where(
            SiteAssetAttendance.site_id == site_id,
            SiteAssetAttendance.operating_date == operating_date,
            SiteAssetAttendance.shift == shift,
        )):
            if str(row.status or "").upper() in terms or str(row.condition or "").upper() in terms:
                idle.add(row.asset_id)
    return idle


def _active_breakdowns(db: Session, site_id: str):
    return {x.asset_id: x for x in db.scalars(select(MaintenanceBreakdown).where(
        MaintenanceBreakdown.site_id == site_id,
        MaintenanceBreakdown.status.in_(["OPEN", "IN_PROGRESS"]),
    ).order_by(MaintenanceBreakdown.reported_at.desc()))}


def _active_jobs(db: Session, site_id: str):
    out = {}
    for x in db.scalars(select(MaintenanceJobCard).where(
        MaintenanceJobCard.site_id == site_id,
        MaintenanceJobCard.status.in_(["OPEN", "IN_PROGRESS"]),
    ).order_by(MaintenanceJobCard.opened_at.desc())):
        out.setdefault(x.asset_id, x)
    return out


def _display_asset(x: Equipment):
    return x.vehicle_no or x.door_no or x.machine_id


def _derive_statuses(db: Session, site_id: str, operating_date: date, shift: str):
    assets = _site_assets(db, site_id)
    bds = _active_breakdowns(db, site_id)
    jobs = _active_jobs(db, site_id)
    running = _running_assets(db, site_id, operating_date, shift)
    idle = _idle_assets(db, site_id, operating_date, shift)
    result = []
    for asset in assets:
        status = "AVAILABLE"
        since = None
        detail = None
        if asset.machine_id in bds:
            status = "BREAKDOWN"
            since = bds[asset.machine_id].reported_at
            detail = bds[asset.machine_id].problem
        elif asset.machine_id in jobs:
            job = jobs[asset.machine_id]
            jt = str(job.job_type or "").upper()
            status = "UNDER_SERVICE" if "SERVICE" in jt else ("PM" if jt in {"PM", "WASHING", "GREASING", "INSPECTION"} else "UNDER_MAINTENANCE")
            since = job.started_at or job.opened_at
            detail = job.job_type
        elif asset.machine_id in running:
            status = "RUNNING"
        elif asset.machine_id in idle:
            status = "IDLE"
        operator_id, location_id = _current_operator_location(db, site_id, asset.machine_id, operating_date, shift)
        mt, mr, md, ms = _latest_meter(db, site_id, asset.machine_id)
        result.append({
            "assetId": asset.machine_id,
            "display": _display_asset(asset),
            "vehicleNo": asset.vehicle_no,
            "doorNo": asset.door_no,
            "type": asset.type,
            "group": asset.group,
            "makeModel": asset.make_model,
            "status": status,
            "statusSince": since.isoformat() if since else None,
            "statusDetail": detail,
            "operatorId": operator_id,
            "locationId": location_id,
            "meterType": mt,
            "meterReading": float(mr) if mr is not None else None,
            "meterDate": md.isoformat() if md else None,
            "meterShift": ms,
        })
    return result


def _def_alerts(db: Session, site_id: str):
    if site_id != "TIOM":
        return []
    out = []
    for profile in db.scalars(select(MaintenanceDefProfile).where(MaintenanceDefProfile.site_id == site_id, MaintenanceDefProfile.active.is_(True))):
        last = db.scalar(select(MaintenanceDefTransaction).where(
            MaintenanceDefTransaction.site_id == site_id,
            MaintenanceDefTransaction.asset_id == profile.asset_id,
        ).order_by(MaintenanceDefTransaction.event_at.desc()).limit(1))
        if not last:
            out.append({"assetId": profile.asset_id, "state": "NO_ISSUE", "remainingL": None, "nextDueMeter": None, "basis": profile.rate_basis})
            continue
        base_qty = Decimal(last.quantity_l or profile.normal_issue_qty_l or 0)
        rate = Decimal(profile.expected_rate or 0)
        alert = Decimal(profile.alert_level_l or profile.minimum_level_l or 0)
        remaining = None
        next_due_meter = None
        consumed = None
        if rate > 0 and profile.rate_basis in {"PER_HOUR", "PER_KM"}:
            meter_type = "HMR" if profile.rate_basis == "PER_HOUR" else "KMR"
            _, current, _, _ = _latest_meter(db, site_id, profile.asset_id, meter_type)
            if last.meter_reading is not None and current is not None:
                delta = max(Decimal("0"), Decimal(current) - Decimal(last.meter_reading))
                consumed = delta * rate
                remaining = base_qty - consumed
                next_due_meter = Decimal(last.meter_reading) + (base_qty / rate)
        elif rate > 0 and profile.rate_basis == "PCT_HSD":
            hsd = db.scalar(select(func.coalesce(func.sum(HsdIssue.litres), 0)).where(
                HsdIssue.machine_id == profile.asset_id,
                HsdIssue.issued_at >= last.event_at,
            ))
            consumed = Decimal(hsd or 0) * rate / Decimal("100")
            remaining = base_qty - consumed
        state = "OK"
        if remaining is not None:
            if remaining <= 0:
                state = "DUE"
            elif remaining <= alert:
                state = "DUE_SOON"
        out.append({
            "assetId": profile.asset_id,
            "state": state,
            "remainingL": float(remaining) if remaining is not None else None,
            "estimatedConsumedL": float(consumed) if consumed is not None else None,
            "nextDueMeter": float(next_due_meter) if next_due_meter is not None else None,
            "basis": profile.rate_basis,
            "normalIssueQtyL": float(profile.normal_issue_qty_l) if profile.normal_issue_qty_l is not None else None,
            "lastIssueAt": last.event_at.isoformat(),
        })
    return out


def _calc_next(schedule_basis: str, interval_value, calendar_days, base_date: date | None, base_meter):
    next_date = None
    next_meter = None
    basis = (schedule_basis or "").upper()
    if basis in {"HMR", "KMR", "WHICHEVER_FIRST"} and interval_value is not None and base_meter is not None:
        next_meter = Decimal(base_meter) + Decimal(interval_value)
    if basis in {"DAYS", "MONTHS", "WHICHEVER_FIRST"} and calendar_days and base_date:
        next_date = base_date + timedelta(days=int(calendar_days))
    return next_date, next_meter


def _due_state(next_date, next_meter, warning_value, current_meter, today: date):
    due = False
    soon = False
    remaining = None
    if next_date:
        days = (next_date - today).days
        due = due or days <= 0
        soon = soon or (0 < days <= 30)
    if next_meter is not None and current_meter is not None:
        remaining = Decimal(next_meter) - Decimal(current_meter)
        due = due or remaining <= 0
        warning = Decimal(warning_value or 0)
        soon = soon or (remaining > 0 and warning > 0 and remaining <= warning)
    return ("OVERDUE" if due else "DUE_SOON" if soon else "OK"), remaining


@router.get("/{site_id}/dashboard")
def dashboard(site_id: str, request: Request, from_date: date | None = None, to_date: date | None = None, asset_id: str | None = None, status: str | None = None, equipment_type: str | None = None, db: Session = Depends(get_db)):
    user, site_id = _access(db, request, site_id, "VIEW")
    ctx = resolve_site_context(db, site_id)
    start = from_date or (ctx.operating_date - timedelta(days=29))
    end = to_date or ctx.operating_date
    if end < start or (end - start).days > 365:
        raise HTTPException(422, "Dashboard range must be 1–366 days.")
    all_assets = _derive_statuses(db, site_id, ctx.operating_date, ctx.shift)
    assets = list(all_assets)
    all_types = sorted({str(x.get("type") or "OTHER") for x in all_assets})
    all_statuses = sorted({str(x.get("status") or "AVAILABLE") for x in all_assets})
    if asset_id: assets = [x for x in assets if x["assetId"] == asset_id]
    if status: assets = [x for x in assets if x["status"] == status.upper()]
    if equipment_type: assets = [x for x in assets if str(x.get("type") or "").upper() == equipment_type.upper()]
    valid_assets = {x["assetId"] for x in assets}
    counts = {k: 0 for k in ["TOTAL", "RUNNING", "AVAILABLE", "BREAKDOWN", "UNDER_MAINTENANCE", "UNDER_SERVICE", "PM", "IDLE"]}
    counts["TOTAL"] = len(assets)
    for row in assets: counts[row["status"]] = counts.get(row["status"], 0) + 1
    current = {x["assetId"]: Decimal(str(x["meterReading"])) if x["meterReading"] is not None else None for x in assets}
    today = ctx.operating_date
    service_due = []
    for p in db.scalars(select(MaintenanceServicePlan).where(MaintenanceServicePlan.site_id == site_id, MaintenanceServicePlan.active.is_(True))):
        if valid_assets and p.asset_id and p.asset_id not in valid_assets: continue
        meter = current.get(p.asset_id) if p.asset_id else None
        state, remaining = _due_state(p.next_due_at, p.next_due_meter, p.warning_value, meter, today)
        if state != "OK": service_due.append({"id": p.plan_id, "assetId": p.asset_id, "name": p.service_name, "state": state, "nextDueAt": p.next_due_at.isoformat() if p.next_due_at else None, "nextDueMeter": float(p.next_due_meter) if p.next_due_meter is not None else None, "remaining": float(remaining) if remaining is not None else None})
    component_due = []; comps = {x.component_id: x for x in db.scalars(select(MaintenanceComponentMaster))}
    for p in db.scalars(select(EquipmentComponentSchedule).where(EquipmentComponentSchedule.site_id == site_id, EquipmentComponentSchedule.active.is_(True))):
        if valid_assets and p.asset_id and p.asset_id not in valid_assets: continue
        meter = current.get(p.asset_id) if p.asset_id else None
        state, remaining = _due_state(p.next_due_at, p.next_due_meter, p.warning_value, meter, today)
        if state != "OK": component_due.append({"id": p.schedule_id, "assetId": p.asset_id, "name": comps.get(p.component_id).component_name if comps.get(p.component_id) else p.component_id, "state": state, "nextDueAt": p.next_due_at.isoformat() if p.next_due_at else None, "nextDueMeter": float(p.next_due_meter) if p.next_due_meter is not None else None, "remaining": float(remaining) if remaining is not None else None})
    pm_due = []
    for p in db.scalars(select(MaintenancePmPlan).where(MaintenancePmPlan.site_id == site_id, MaintenancePmPlan.active.is_(True))):
        if valid_assets and p.asset_id and p.asset_id not in valid_assets: continue
        meter = current.get(p.asset_id) if p.asset_id else None
        state, remaining = _due_state(p.next_due_at, p.next_due_meter, p.warning_value, meter, today)
        if state != "OK": pm_due.append({"id": p.pm_plan_id, "assetId": p.asset_id, "name": p.plan_name, "activity": p.activity, "state": state, "nextDueAt": p.next_due_at.isoformat() if p.next_due_at else None, "nextDueMeter": float(p.next_due_meter) if p.next_due_meter is not None else None, "remaining": float(remaining) if remaining is not None else None})
    def_alerts = [x for x in _def_alerts(db, site_id) if not valid_assets or x["assetId"] in valid_assets]
    counts["DEF_DUE"] = sum(1 for x in def_alerts if x["state"] in {"DUE", "DUE_SOON"})

    bd_stmt = select(MaintenanceBreakdown).where(MaintenanceBreakdown.site_id == site_id, MaintenanceBreakdown.operating_date >= start, MaintenanceBreakdown.operating_date <= end)
    svc_stmt = select(MaintenanceServiceHistory).where(MaintenanceServiceHistory.site_id == site_id, func.date(MaintenanceServiceHistory.serviced_at) >= start, func.date(MaintenanceServiceHistory.serviced_at) <= end)
    pm_stmt = select(MaintenancePmExecution).where(MaintenancePmExecution.site_id == site_id, func.date(MaintenancePmExecution.completed_at) >= start, func.date(MaintenancePmExecution.completed_at) <= end)
    def_stmt = select(MaintenanceDefTransaction).where(MaintenanceDefTransaction.site_id == site_id, MaintenanceDefTransaction.operating_date >= start, MaintenanceDefTransaction.operating_date <= end)
    spare_stmt = select(MaintenanceSpareUsage).where(MaintenanceSpareUsage.site_id == site_id, func.date(MaintenanceSpareUsage.used_at) >= start, func.date(MaintenanceSpareUsage.used_at) <= end)
    lube_stmt = select(MaintenanceLubricantUsage).where(MaintenanceLubricantUsage.site_id == site_id, func.date(MaintenanceLubricantUsage.used_at) >= start, func.date(MaintenanceLubricantUsage.used_at) <= end)
    if asset_id:
        bd_stmt=bd_stmt.where(MaintenanceBreakdown.asset_id==asset_id); svc_stmt=svc_stmt.where(MaintenanceServiceHistory.asset_id==asset_id); pm_stmt=pm_stmt.where(MaintenancePmExecution.asset_id==asset_id); def_stmt=def_stmt.where(MaintenanceDefTransaction.asset_id==asset_id); spare_stmt=spare_stmt.where(MaintenanceSpareUsage.asset_id==asset_id); lube_stmt=lube_stmt.where(MaintenanceLubricantUsage.asset_id==asset_id)
    elif equipment_type or status:
        if valid_assets:
            bd_stmt=bd_stmt.where(MaintenanceBreakdown.asset_id.in_(valid_assets)); svc_stmt=svc_stmt.where(MaintenanceServiceHistory.asset_id.in_(valid_assets)); pm_stmt=pm_stmt.where(MaintenancePmExecution.asset_id.in_(valid_assets)); def_stmt=def_stmt.where(MaintenanceDefTransaction.asset_id.in_(valid_assets)); spare_stmt=spare_stmt.where(MaintenanceSpareUsage.asset_id.in_(valid_assets)); lube_stmt=lube_stmt.where(MaintenanceLubricantUsage.asset_id.in_(valid_assets))
    bds=list(db.scalars(bd_stmt)); services=list(db.scalars(svc_stmt)); pms=list(db.scalars(pm_stmt)); defs=list(db.scalars(def_stmt)); spares=list(db.scalars(spare_stmt)); lubes=list(db.scalars(lube_stmt))
    downtime=sum(float(x.downtime_hours or 0) for x in bds); closed=[x for x in bds if x.downtime_hours is not None]
    trend={}
    for x in bds:
        k=x.operating_date.isoformat(); r=trend.setdefault(k,{"date":k,"breakdowns":0,"downtimeHours":0.0}); r["breakdowns"]+=1; r["downtimeHours"]+=float(x.downtime_hours or 0)
    breakdown_trend=[dict(x,downtimeHours=round(x["downtimeHours"],2)) for x in sorted(trend.values(),key=lambda r:r["date"])]
    cost_spares=sum(float((x.quantity or 0)*(x.unit_cost or 0)) for x in spares); cost_lubes=sum(float((x.quantity or 0)*(x.unit_cost or 0)) for x in lubes)
    return {
        "siteId": site_id, "operatingDate": ctx.operating_date.isoformat(), "shift": ctx.shift, "fromDate": start.isoformat(), "toDate": end.isoformat(),
        "filters":{"assetId":asset_id or "","status":(status or "").upper(),"equipmentType":equipment_type or "","availableStatuses":all_statuses,"availableTypes":all_types},
        "counts": counts, "serviceDue": service_due[:50], "componentDue": component_due[:50], "pmDue": pm_due[:50], "defAlerts": def_alerts[:50], "assets": assets, "allAssets": all_assets,
        "historyKpis":{"breakdowns":len(bds),"openBreakdowns":sum(1 for x in bds if x.status in {"OPEN","IN_PROGRESS"}),"downtimeHours":round(downtime,2),"mttrHours":round(downtime/len(closed),2) if closed else 0,"services":len(services),"pmCompleted":len(pms),"defLitres":round(sum(float(x.quantity_l or 0) for x in defs),2),"spareCost":round(cost_spares,2),"lubricantCost":round(cost_lubes,2)},
        "breakdownTrend": breakdown_trend,
    }


@router.get("/{site_id}/assets")
def assets(site_id: str, request: Request, status: str | None = None, db: Session = Depends(get_db)):
    user, site_id = _access(db, request, site_id, "VIEW")
    ctx = resolve_site_context(db, site_id)
    rows = _derive_statuses(db, site_id, ctx.operating_date, ctx.shift)
    if status:
        rows = [x for x in rows if x["status"] == status.upper()]
    return rows


@router.get("/{site_id}/breakdowns")
def breakdowns(site_id: str, request: Request, from_date: date | None = None, to_date: date | None = None, status: str | None = None, asset_id: str | None = None, db: Session = Depends(get_db)):
    user, site_id = _access(db, request, site_id, "VIEW")
    stmt = select(MaintenanceBreakdown).where(MaintenanceBreakdown.site_id == site_id)
    if from_date: stmt = stmt.where(MaintenanceBreakdown.operating_date >= from_date)
    if to_date: stmt = stmt.where(MaintenanceBreakdown.operating_date <= to_date)
    if status: stmt = stmt.where(MaintenanceBreakdown.status == status.upper())
    if asset_id: stmt = stmt.where(MaintenanceBreakdown.asset_id == asset_id)
    rows = list(db.scalars(stmt.order_by(MaintenanceBreakdown.reported_at.desc()).limit(2000)))
    return [{
        "breakdownId": x.breakdown_id, "assetId": x.asset_id, "operatingDate": x.operating_date.isoformat(), "shift": x.shift,
        "reportedAt": x.reported_at.isoformat(), "problemCategory": x.problem_category, "problem": x.problem, "locationId": x.location_id,
        "meterType": x.meter_type, "meterReading": float(x.meter_reading) if x.meter_reading is not None else None,
        "status": x.status, "jobCardId": x.job_card_id, "releasedAt": x.released_at.isoformat() if x.released_at else None,
        "downtimeHours": float(x.downtime_hours) if x.downtime_hours is not None else None, "remarks": x.remarks,
    } for x in rows]


@router.post("/{site_id}/breakdowns")
def create_breakdown(site_id: str, p: BreakdownIn, request: Request, db: Session = Depends(get_db)):
    csrf(request); user, site_id = _access(db, request, site_id, "CREATE")
    asset = _asset(db, site_id, p.assetId)
    existing = db.scalar(select(MaintenanceBreakdown).where(MaintenanceBreakdown.site_id == site_id, MaintenanceBreakdown.asset_id == asset.machine_id, MaintenanceBreakdown.status.in_(["OPEN", "IN_PROGRESS"])).limit(1))
    if existing:
        raise HTTPException(409, f"{asset.machine_id} already has open breakdown {existing.breakdown_id}.")
    at = p.reportedAt or resolve_site_context(db, site_id).local_at
    ctx = resolve_site_context(db, site_id, at)
    mt, mr, _, _ = _latest_meter(db, site_id, asset.machine_id)
    operator_id, location_id = _current_operator_location(db, site_id, asset.machine_id, ctx.operating_date, ctx.shift)
    bid = _id("BD", site_id)
    jid = _id("JC", site_id)
    bd = MaintenanceBreakdown(
        breakdown_id=bid, site_id=site_id, asset_id=asset.machine_id, operating_date=ctx.operating_date, shift=ctx.shift,
        reported_at=at, problem_category=p.problemCategory.strip().upper(), problem=p.problem.strip(), location_id=p.locationId or location_id,
        meter_type=mt, meter_reading=mr, status="OPEN", job_card_id=jid, reported_by=user.login_id, remarks=p.remarks,
    )
    job = MaintenanceJobCard(
        job_card_id=jid, site_id=site_id, asset_id=asset.machine_id, source_type="BREAKDOWN", source_id=bid, job_type="BREAKDOWN",
        status="OPEN", opened_at=at, meter_type=mt, meter_reading=mr, complaint=p.problem.strip(), entered_by=user.login_id,
    )
    db.add_all([bd, job])
    _audit(db, user, site_id, "CREATE", "maintenance_breakdown", bid, after={"assetId": asset.machine_id, "jobCardId": jid, "problem": p.problem})
    db.commit()
    return {"ok": True, "breakdownId": bid, "jobCardId": jid, "status": "BREAKDOWN", "operatingDate": ctx.operating_date.isoformat(), "shift": ctx.shift, "meterType": mt, "meterReading": float(mr) if mr is not None else None, "operatorId": operator_id, "locationId": p.locationId or location_id}


@router.post("/{site_id}/breakdowns/{breakdown_id}/close")
def close_breakdown(site_id: str, breakdown_id: str, p: BreakdownCloseIn, request: Request, db: Session = Depends(get_db)):
    csrf(request); user, site_id = _access(db, request, site_id, "EDIT")
    bd = db.get(MaintenanceBreakdown, breakdown_id)
    if not bd or bd.site_id != site_id: raise HTTPException(404, "Breakdown not found.")
    if bd.status not in {"OPEN", "IN_PROGRESS"}: raise HTTPException(409, "Breakdown is already closed.")
    released = p.releasedAt or resolve_site_context(db, site_id).local_at
    reported = bd.reported_at
    # SQLite (installer smoke DB) may drop tzinfo from timezone columns; normalize to
    # the release timezone. PostgreSQL already returns timezone-aware values.
    if reported.tzinfo is None and released.tzinfo is not None:
        reported = reported.replace(tzinfo=released.tzinfo)
    elif reported.tzinfo is not None and released.tzinfo is None:
        released = released.replace(tzinfo=reported.tzinfo)
    if released < reported: raise HTTPException(422, "Release time cannot be before breakdown time.")
    hours = Decimal(str((released - reported).total_seconds() / 3600)).quantize(Decimal("0.001"))
    before = {"status": bd.status, "releasedAt": bd.released_at}
    bd.status = "RELEASED"; bd.released_at = released; bd.downtime_hours = hours; bd.released_by = user.login_id; bd.updated_at = datetime.utcnow()
    if p.remarks: bd.remarks = p.remarks
    job = db.get(MaintenanceJobCard, bd.job_card_id) if bd.job_card_id else None
    if job:
        job.status = "RELEASED"; job.completed_at = released; job.released_at = released; job.diagnosis = p.diagnosis; job.root_cause = p.rootCause; job.action_taken = p.actionTaken; job.mechanic_id = p.mechanicId; job.labour_hours = p.labourHours; job.external_cost = p.externalCost; job.remarks = p.remarks; job.updated_at = datetime.utcnow()
    _audit(db, user, site_id, "RELEASE", "maintenance_breakdown", breakdown_id, before=before, after={"status": bd.status, "releasedAt": released, "downtimeHours": hours})
    db.commit()
    return {"ok": True, "breakdownId": breakdown_id, "status": bd.status, "downtimeHours": float(hours), "jobCardId": bd.job_card_id}


@router.get("/{site_id}/job-cards")
def job_cards(site_id: str, request: Request, status: str | None = None, asset_id: str | None = None, from_date: date | None = None, to_date: date | None = None, db: Session = Depends(get_db)):
    user, site_id = _access(db, request, site_id, "VIEW")
    stmt = select(MaintenanceJobCard).where(MaintenanceJobCard.site_id == site_id)
    if status: stmt = stmt.where(MaintenanceJobCard.status == status.upper())
    if asset_id: stmt = stmt.where(MaintenanceJobCard.asset_id == asset_id)
    if from_date: stmt = stmt.where(func.date(MaintenanceJobCard.opened_at) >= from_date)
    if to_date: stmt = stmt.where(func.date(MaintenanceJobCard.opened_at) <= to_date)
    rows = list(db.scalars(stmt.order_by(MaintenanceJobCard.opened_at.desc()).limit(2000)))
    return [{"jobCardId": x.job_card_id, "assetId": x.asset_id, "sourceType": x.source_type, "sourceId": x.source_id, "jobType": x.job_type, "status": x.status, "openedAt": x.opened_at.isoformat(), "startedAt": x.started_at.isoformat() if x.started_at else None, "completedAt": x.completed_at.isoformat() if x.completed_at else None, "meterType": x.meter_type, "meterReading": float(x.meter_reading) if x.meter_reading is not None else None, "complaint": x.complaint, "diagnosis": x.diagnosis, "rootCause": x.root_cause, "actionTaken": x.action_taken, "mechanicId": x.mechanic_id, "labourHours": float(x.labour_hours) if x.labour_hours is not None else None, "externalCost": float(x.external_cost or 0), "remarks": x.remarks} for x in rows]


@router.post("/{site_id}/job-cards")
def create_job_card(site_id: str, p: JobCardIn, request: Request, db: Session = Depends(get_db)):
    csrf(request); user, site_id = _access(db, request, site_id, "CREATE")
    asset = _asset(db, site_id, p.assetId); at = p.openedAt or resolve_site_context(db, site_id).local_at
    mt, mr, _, _ = _latest_meter(db, site_id, asset.machine_id)
    jid = _id("JC", site_id)
    row = MaintenanceJobCard(job_card_id=jid, site_id=site_id, asset_id=asset.machine_id, source_type=p.sourceType.upper(), source_id=p.sourceId, job_type=p.jobType.upper(), status="OPEN", opened_at=at, meter_type=mt, meter_reading=mr, complaint=p.complaint, mechanic_id=p.mechanicId, remarks=p.remarks, entered_by=user.login_id)
    db.add(row); _audit(db, user, site_id, "CREATE", "maintenance_job_card", jid, after={"assetId": p.assetId, "jobType": p.jobType}); db.commit()
    return {"ok": True, "jobCardId": jid, "meterType": mt, "meterReading": float(mr) if mr is not None else None}


@router.patch("/{site_id}/job-cards/{job_card_id}")
def update_job_card(site_id: str, job_card_id: str, p: JobCardUpdateIn, request: Request, db: Session = Depends(get_db)):
    csrf(request); user, site_id = _access(db, request, site_id, "EDIT")
    row = db.get(MaintenanceJobCard, job_card_id)
    if not row or row.site_id != site_id: raise HTTPException(404, "Job card not found.")
    before = {"status": row.status}
    now = resolve_site_context(db, site_id).local_at
    row.status = p.status; row.diagnosis = p.diagnosis; row.root_cause = p.rootCause; row.action_taken = p.actionTaken; row.mechanic_id = p.mechanicId; row.labour_hours = p.labourHours
    if p.externalCost is not None: row.external_cost = p.externalCost
    row.remarks = p.remarks; row.updated_at = datetime.utcnow()
    if p.status == "IN_PROGRESS" and not row.started_at: row.started_at = now
    if p.status in {"COMPLETED", "RELEASED"}: row.completed_at = row.completed_at or now
    if p.status == "RELEASED": row.released_at = now
    _audit(db, user, site_id, "UPDATE", "maintenance_job_card", job_card_id, before=before, after={"status": row.status}); db.commit()
    return {"ok": True, "jobCardId": row.job_card_id, "status": row.status}


@router.get("/{site_id}/service-plans")
def service_plans(site_id: str, request: Request, db: Session = Depends(get_db)):
    user, site_id = _access(db, request, site_id, "VIEW")
    rows = list(db.scalars(select(MaintenanceServicePlan).where(MaintenanceServicePlan.site_id == site_id, MaintenanceServicePlan.active.is_(True)).order_by(MaintenanceServicePlan.service_name)))
    out=[]
    for x in rows:
        mt = "HMR" if x.schedule_basis in {"HMR","WHICHEVER_FIRST"} else "KMR" if x.schedule_basis=="KMR" else None
        _, current, _, _ = _latest_meter(db, site_id, x.asset_id, mt) if x.asset_id and mt else (None,None,None,None)
        state, rem = _due_state(x.next_due_at, x.next_due_meter, x.warning_value, current, date.today())
        out.append({"planId":x.plan_id,"assetId":x.asset_id,"equipmentType":x.equipment_type,"makeModel":x.make_model,"serviceName":x.service_name,"scheduleBasis":x.schedule_basis,"intervalValue":float(x.interval_value) if x.interval_value is not None else None,"calendarIntervalDays":x.calendar_interval_days,"warningValue":float(x.warning_value) if x.warning_value is not None else None,"lastServiceAt":x.last_service_at.isoformat() if x.last_service_at else None,"lastMeter":float(x.last_meter) if x.last_meter is not None else None,"nextDueAt":x.next_due_at.isoformat() if x.next_due_at else None,"nextDueMeter":float(x.next_due_meter) if x.next_due_meter is not None else None,"currentMeter":float(current) if current is not None else None,"dueState":state,"remaining":float(rem) if rem is not None else None})
    return out


@router.post("/{site_id}/service-plans")
def create_service_plan(site_id: str, p: ServicePlanIn, request: Request, db: Session = Depends(get_db)):
    csrf(request); user, site_id = _access(db, request, site_id, "EDIT")
    if p.assetId: _asset(db, site_id, p.assetId)
    next_date, next_meter = _calc_next(p.scheduleBasis, p.intervalValue, p.calendarIntervalDays, p.lastServiceAt, p.lastMeter)
    pid=_id("SVP",site_id); row=MaintenanceServicePlan(plan_id=pid,site_id=site_id,asset_id=p.assetId,equipment_type=p.equipmentType,make_model=p.makeModel,service_name=p.serviceName,schedule_basis=p.scheduleBasis,interval_value=p.intervalValue,calendar_interval_days=p.calendarIntervalDays,warning_value=p.warningValue,last_service_at=p.lastServiceAt,last_meter=p.lastMeter,next_due_at=next_date,next_due_meter=next_meter,entered_by=user.login_id)
    db.add(row); _audit(db,user,site_id,"CREATE","maintenance_service_plan",pid,after=p.model_dump()); db.commit(); return {"ok":True,"planId":pid}


@router.post("/{site_id}/service-plans/{plan_id}/complete")
def complete_service(site_id: str, plan_id: str, p: ServiceCompleteIn, request: Request, db: Session = Depends(get_db)):
    csrf(request); user, site_id = _access(db, request, site_id, "CREATE")
    plan=db.get(MaintenanceServicePlan,plan_id)
    if not plan or plan.site_id!=site_id: raise HTTPException(404,"Service plan not found.")
    if not plan.asset_id: raise HTTPException(422,"Complete service requires an equipment-specific plan.")
    at=p.servicedAt or resolve_site_context(db,site_id).local_at
    mt=p.meterType or ("KMR" if plan.schedule_basis=="KMR" else "HMR")
    mr=p.meterReading
    if mr is None:
        _,mr,_,_=_latest_meter(db,site_id,plan.asset_id,mt)
    next_date,next_meter=_calc_next(plan.schedule_basis,plan.interval_value,plan.calendar_interval_days,at.date(),mr)
    sid=_id("SVH",site_id); db.add(MaintenanceServiceHistory(service_id=sid,plan_id=plan_id,site_id=site_id,asset_id=plan.asset_id,job_card_id=p.jobCardId,serviced_at=at,meter_type=mt,meter_reading=mr,next_due_at=next_date,next_due_meter=next_meter,remarks=p.remarks,entered_by=user.login_id))
    plan.last_service_at=at.date(); plan.last_meter=mr; plan.next_due_at=next_date; plan.next_due_meter=next_meter; plan.updated_at=datetime.utcnow()
    _audit(db,user,site_id,"COMPLETE","maintenance_service_plan",plan_id,after={"serviceId":sid,"meterReading":mr,"nextDueMeter":next_meter,"nextDueAt":next_date}); db.commit(); return {"ok":True,"serviceId":sid,"nextDueAt":next_date.isoformat() if next_date else None,"nextDueMeter":float(next_meter) if next_meter is not None else None}


@router.get("/{site_id}/components")
def components(site_id: str, request: Request, db: Session = Depends(get_db)):
    user, site_id = _access(db, request, site_id, "VIEW")
    return [{"componentId":x.component_id,"componentName":x.component_name,"category":x.category,"partNumber":x.part_number,"unit":x.unit,"storeItemId":x.store_item_id} for x in db.scalars(select(MaintenanceComponentMaster).where(MaintenanceComponentMaster.active.is_(True)).order_by(MaintenanceComponentMaster.component_name))]


@router.post("/{site_id}/components")
def create_component(site_id: str, p: ComponentMasterIn, request: Request, db: Session = Depends(get_db)):
    csrf(request); user, site_id = _access(db, request, site_id, "EDIT")
    cid=_id("CMP",site_id); db.add(MaintenanceComponentMaster(component_id=cid,component_name=p.componentName.strip(),category=p.category,part_number=p.partNumber,unit=p.unit.upper(),store_item_id=p.storeItemId,entered_by=user.login_id)); _audit(db,user,site_id,"CREATE","maintenance_component_master",cid,after=p.model_dump()); db.commit(); return {"ok":True,"componentId":cid}


@router.get("/{site_id}/component-schedules")
def component_schedules(site_id: str, request: Request, db: Session = Depends(get_db)):
    user, site_id = _access(db, request, site_id, "VIEW")
    comps={x.component_id:x for x in db.scalars(select(MaintenanceComponentMaster))}; rows=[]
    for x in db.scalars(select(EquipmentComponentSchedule).where(EquipmentComponentSchedule.site_id==site_id,EquipmentComponentSchedule.active.is_(True)).order_by(EquipmentComponentSchedule.created_at.desc())):
        mt="KMR" if x.schedule_basis=="KMR" else "HMR" if x.schedule_basis in {"HMR","WHICHEVER_FIRST"} else None; _,current,_,_=_latest_meter(db,site_id,x.asset_id,mt) if x.asset_id and mt else (None,None,None,None); state,rem=_due_state(x.next_due_at,x.next_due_meter,x.warning_value,current,date.today()); c=comps.get(x.component_id)
        rows.append({"scheduleId":x.schedule_id,"componentId":x.component_id,"componentName":c.component_name if c else x.component_id,"assetId":x.asset_id,"equipmentType":x.equipment_type,"makeModel":x.make_model,"scheduleBasis":x.schedule_basis,"intervalValue":float(x.interval_value) if x.interval_value is not None else None,"calendarIntervalDays":x.calendar_interval_days,"warningValue":float(x.warning_value) if x.warning_value is not None else None,"requiredQty":float(x.required_qty),"lastChangedAt":x.last_changed_at.isoformat() if x.last_changed_at else None,"lastMeter":float(x.last_meter) if x.last_meter is not None else None,"nextDueAt":x.next_due_at.isoformat() if x.next_due_at else None,"nextDueMeter":float(x.next_due_meter) if x.next_due_meter is not None else None,"currentMeter":float(current) if current is not None else None,"dueState":state,"remaining":float(rem) if rem is not None else None})
    return rows


@router.post("/{site_id}/component-schedules")
def create_component_schedule(site_id: str,p:ComponentScheduleIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user,site_id=_access(db,request,site_id,"EDIT")
    if not db.get(MaintenanceComponentMaster,p.componentId): raise HTTPException(422,"Component not found.")
    if p.assetId: _asset(db,site_id,p.assetId)
    next_date,next_meter=_calc_next(p.scheduleBasis,p.intervalValue,p.calendarIntervalDays,p.lastChangedAt,p.lastMeter); sid=_id("CS",site_id); db.add(EquipmentComponentSchedule(schedule_id=sid,site_id=site_id,asset_id=p.assetId,equipment_type=p.equipmentType,make_model=p.makeModel,component_id=p.componentId,schedule_basis=p.scheduleBasis,interval_value=p.intervalValue,calendar_interval_days=p.calendarIntervalDays,warning_value=p.warningValue,required_qty=p.requiredQty,last_changed_at=p.lastChangedAt,last_meter=p.lastMeter,next_due_at=next_date,next_due_meter=next_meter,entered_by=user.login_id)); _audit(db,user,site_id,"CREATE","equipment_component_schedule",sid,after=p.model_dump()); db.commit(); return {"ok":True,"scheduleId":sid}


@router.post("/{site_id}/component-schedules/{schedule_id}/change")
def complete_component_change(site_id:str,schedule_id:str,p:ComponentChangeIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user,site_id=_access(db,request,site_id,"CREATE"); s=db.get(EquipmentComponentSchedule,schedule_id)
    if not s or s.site_id!=site_id: raise HTTPException(404,"Component schedule not found.")
    if not s.asset_id: raise HTTPException(422,"Component change requires an equipment-specific schedule.")
    at=p.changedAt or resolve_site_context(db,site_id).local_at; mt=p.meterType or ("KMR" if s.schedule_basis=="KMR" else "HMR"); mr=p.meterReading
    if mr is None: _,mr,_,_=_latest_meter(db,site_id,s.asset_id,mt)
    next_date,next_meter=_calc_next(s.schedule_basis,s.interval_value,s.calendar_interval_days,at.date(),mr); cid=_id("CCH",site_id); qty=p.quantity or s.required_qty
    db.add(ComponentChangeHistory(change_id=cid,schedule_id=schedule_id,site_id=site_id,asset_id=s.asset_id,job_card_id=p.jobCardId,changed_at=at,meter_type=mt,meter_reading=mr,quantity=qty,remarks=p.remarks,entered_by=user.login_id)); s.last_changed_at=at.date(); s.last_meter=mr; s.next_due_at=next_date; s.next_due_meter=next_meter; s.updated_at=datetime.utcnow()
    comp=db.get(MaintenanceComponentMaster,s.component_id); db.add(MaintenanceSpareUsage(usage_id=_id("SPU",site_id),site_id=site_id,asset_id=s.asset_id,job_card_id=p.jobCardId,component_id=s.component_id,store_item_id=comp.store_item_id if comp else None,used_at=at,description=comp.component_name if comp else s.component_id,quantity=qty,unit=comp.unit if comp else "NOS",remarks="Auto-posted from scheduled component change",entered_by=user.login_id))
    _audit(db,user,site_id,"COMPLETE","equipment_component_schedule",schedule_id,after={"changeId":cid,"nextDueAt":next_date,"nextDueMeter":next_meter}); db.commit(); return {"ok":True,"changeId":cid,"nextDueAt":next_date.isoformat() if next_date else None,"nextDueMeter":float(next_meter) if next_meter is not None else None}


@router.get("/{site_id}/pm-plans")
def pm_plans(site_id:str,request:Request,db:Session=Depends(get_db)):
    user,site_id=_access(db,request,site_id,"VIEW"); rows=[]
    for x in db.scalars(select(MaintenancePmPlan).where(MaintenancePmPlan.site_id==site_id,MaintenancePmPlan.active.is_(True)).order_by(MaintenancePmPlan.activity,MaintenancePmPlan.plan_name)):
        mt="KMR" if x.schedule_basis=="KMR" else "HMR" if x.schedule_basis in {"HMR","WHICHEVER_FIRST"} else None; _,current,_,_=_latest_meter(db,site_id,x.asset_id,mt) if x.asset_id and mt else (None,None,None,None); state,rem=_due_state(x.next_due_at,x.next_due_meter,x.warning_value,current,date.today())
        rows.append({"pmPlanId":x.pm_plan_id,"assetId":x.asset_id,"equipmentType":x.equipment_type,"activity":x.activity,"planName":x.plan_name,"scheduleBasis":x.schedule_basis,"intervalValue":float(x.interval_value) if x.interval_value is not None else None,"calendarIntervalDays":x.calendar_interval_days,"warningValue":float(x.warning_value) if x.warning_value is not None else None,"lastDoneAt":x.last_done_at.isoformat() if x.last_done_at else None,"lastMeter":float(x.last_meter) if x.last_meter is not None else None,"nextDueAt":x.next_due_at.isoformat() if x.next_due_at else None,"nextDueMeter":float(x.next_due_meter) if x.next_due_meter is not None else None,"currentMeter":float(current) if current is not None else None,"dueState":state,"remaining":float(rem) if rem is not None else None,"checklist":json.loads(x.checklist_json) if x.checklist_json else []})
    return rows


@router.post("/{site_id}/pm-plans")
def create_pm_plan(site_id:str,p:PmPlanIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user,site_id=_access(db,request,site_id,"EDIT");
    if p.assetId: _asset(db,site_id,p.assetId)
    next_date,next_meter=_calc_next(p.scheduleBasis,p.intervalValue,p.calendarIntervalDays,p.lastDoneAt,p.lastMeter); pid=_id("PMP",site_id); db.add(MaintenancePmPlan(pm_plan_id=pid,site_id=site_id,asset_id=p.assetId,equipment_type=p.equipmentType,activity=p.activity,plan_name=p.planName,schedule_basis=p.scheduleBasis,interval_value=p.intervalValue,calendar_interval_days=p.calendarIntervalDays,warning_value=p.warningValue,last_done_at=p.lastDoneAt,last_meter=p.lastMeter,next_due_at=next_date,next_due_meter=next_meter,checklist_json=json.dumps(p.checklist),entered_by=user.login_id)); _audit(db,user,site_id,"CREATE","maintenance_pm_plan",pid,after=p.model_dump()); db.commit(); return {"ok":True,"pmPlanId":pid}


@router.get("/{site_id}/pm-executions")
def pm_executions(site_id:str,request:Request,from_date:date|None=None,to_date:date|None=None,activity:str|None=None,db:Session=Depends(get_db)):
    user,site_id=_access(db,request,site_id,"VIEW"); stmt=select(MaintenancePmExecution).where(MaintenancePmExecution.site_id==site_id)
    if from_date: stmt=stmt.where(func.date(MaintenancePmExecution.completed_at)>=from_date)
    if to_date: stmt=stmt.where(func.date(MaintenancePmExecution.completed_at)<=to_date)
    if activity: stmt=stmt.where(MaintenancePmExecution.activity==activity.upper())
    return [{"executionId":x.execution_id,"pmPlanId":x.pm_plan_id,"assetId":x.asset_id,"activity":x.activity,"completedAt":x.completed_at.isoformat(),"meterType":x.meter_type,"meterReading":float(x.meter_reading) if x.meter_reading is not None else None,"result":x.result,"remarks":x.remarks} for x in db.scalars(stmt.order_by(MaintenancePmExecution.completed_at.desc()).limit(2000))]


@router.post("/{site_id}/pm-executions")
def create_pm_execution(site_id:str,p:PmExecutionIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user,site_id=_access(db,request,site_id,"CREATE"); _asset(db,site_id,p.assetId); at=p.completedAt or resolve_site_context(db,site_id).local_at; plan=db.get(MaintenancePmPlan,p.pmPlanId) if p.pmPlanId else None
    if plan and (plan.site_id!=site_id or plan.asset_id not in {None,p.assetId}): raise HTTPException(422,"PM plan does not match equipment.")
    mt=p.meterType or ("KMR" if plan and plan.schedule_basis=="KMR" else "HMR"); mr=p.meterReading
    if mr is None: _,mr,_,_=_latest_meter(db,site_id,p.assetId,mt)
    eid=_id("PME",site_id); db.add(MaintenancePmExecution(execution_id=eid,pm_plan_id=p.pmPlanId,site_id=site_id,asset_id=p.assetId,activity=p.activity,job_card_id=p.jobCardId,completed_at=at,meter_type=mt,meter_reading=mr,result=p.result,checklist_json=json.dumps(p.checklist) if p.checklist is not None else None,remarks=p.remarks,entered_by=user.login_id))
    if plan:
        nd,nm=_calc_next(plan.schedule_basis,plan.interval_value,plan.calendar_interval_days,at.date(),mr); plan.last_done_at=at.date();plan.last_meter=mr;plan.next_due_at=nd;plan.next_due_meter=nm;plan.updated_at=datetime.utcnow()
    _audit(db,user,site_id,"COMPLETE","maintenance_pm_execution",eid,after={"assetId":p.assetId,"activity":p.activity}); db.commit(); return {"ok":True,"executionId":eid}


@router.get("/{site_id}/def-summary")
def def_summary(site_id:str,request:Request,from_date:date|None=None,to_date:date|None=None,db:Session=Depends(get_db)):
    user,site_id=_access(db,request,site_id,"VIEW"); start=from_date or date.today().replace(day=1); end=to_date or date.today()
    if end<start: raise HTTPException(422,"to_date cannot be before from_date")
    def_totals={asset:Decimal(qty or 0) for asset,qty in db.execute(select(MaintenanceDefTransaction.asset_id,func.sum(MaintenanceDefTransaction.quantity_l)).where(MaintenanceDefTransaction.site_id==site_id,MaintenanceDefTransaction.operating_date>=start,MaintenanceDefTransaction.operating_date<=end).group_by(MaintenanceDefTransaction.asset_id)).all()}
    hsd_totals={}
    if site_id=="TIOM":
        hsd_totals={asset:Decimal(qty or 0) for asset,qty in db.execute(select(HsdIssue.machine_id,func.sum(HsdIssue.litres)).where(HsdIssue.operating_date>=start,HsdIssue.operating_date<=end).group_by(HsdIssue.machine_id)).all()}
    ids=sorted(set(def_totals)|set(hsd_totals)); rows=[]
    for asset in ids:
        d=def_totals.get(asset,Decimal("0")); h=hsd_totals.get(asset,Decimal("0")); pct=(d/h*Decimal("100")) if h>0 else None
        rows.append({"assetId":asset,"defL":float(d),"hsdL":float(h),"defPctHsd":float(pct) if pct is not None else None})
    return {"fromDate":start.isoformat(),"toDate":end.isoformat(),"rows":rows,"alerts":_def_alerts(db,site_id)}


@router.put("/{site_id}/def-profile")
def save_def_profile(site_id:str,p:DefProfileIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user,site_id=_access(db,request,site_id,"EDIT"); _asset(db,site_id,p.assetId); row=db.get(MaintenanceDefProfile,{"site_id":site_id,"asset_id":p.assetId})
    if not row: row=MaintenanceDefProfile(site_id=site_id,asset_id=p.assetId,updated_by=user.login_id);db.add(row)
    row.tank_capacity_l=p.tankCapacityL;row.normal_issue_qty_l=p.normalIssueQtyL;row.minimum_level_l=p.minimumLevelL;row.expected_rate=p.expectedRate;row.rate_basis=p.rateBasis;row.alert_level_l=p.alertLevelL;row.active=True;row.updated_by=user.login_id;row.updated_at=datetime.utcnow(); db.commit(); return {"ok":True}


@router.get("/{site_id}/def")
def def_rows(site_id:str,request:Request,from_date:date|None=None,to_date:date|None=None,asset_id:str|None=None,db:Session=Depends(get_db)):
    user,site_id=_access(db,request,site_id,"VIEW"); stmt=select(MaintenanceDefTransaction).where(MaintenanceDefTransaction.site_id==site_id)
    if from_date:stmt=stmt.where(MaintenanceDefTransaction.operating_date>=from_date)
    if to_date:stmt=stmt.where(MaintenanceDefTransaction.operating_date<=to_date)
    if asset_id:stmt=stmt.where(MaintenanceDefTransaction.asset_id==asset_id)
    return [{"defId":x.def_id,"assetId":x.asset_id,"eventAt":x.event_at.isoformat(),"operatingDate":x.operating_date.isoformat(),"shift":x.shift,"quantityL":float(x.quantity_l),"meterType":x.meter_type,"meterReading":float(x.meter_reading) if x.meter_reading is not None else None,"operatorId":x.operator_id,"issueReference":x.issue_reference,"remarks":x.remarks} for x in db.scalars(stmt.order_by(MaintenanceDefTransaction.event_at.desc()).limit(5000))]


@router.post("/{site_id}/def")
def create_def(site_id:str,p:DefTransactionIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user,site_id=_access(db,request,site_id,"CREATE"); _asset(db,site_id,p.assetId); at=p.eventAt or resolve_site_context(db,site_id).local_at; ctx=resolve_site_context(db,site_id,at); mt,mr,_,_=_latest_meter(db,site_id,p.assetId); operator,location=_current_operator_location(db,site_id,p.assetId,ctx.operating_date,ctx.shift); did=_id("DEF",site_id); db.add(MaintenanceDefTransaction(def_id=did,site_id=site_id,asset_id=p.assetId,event_at=at,operating_date=ctx.operating_date,shift=ctx.shift,quantity_l=p.quantityL,meter_type=mt,meter_reading=mr,operator_id=p.operatorId or operator,issue_reference=p.issueReference,remarks=p.remarks,entered_by=user.login_id)); _audit(db,user,site_id,"CREATE","maintenance_def_transaction",did,after={"assetId":p.assetId,"quantityL":p.quantityL}); db.commit(); return {"ok":True,"defId":did,"operatingDate":ctx.operating_date.isoformat(),"shift":ctx.shift,"meterType":mt,"meterReading":float(mr) if mr is not None else None,"operatorId":p.operatorId or operator}


@router.get("/{site_id}/spares")
def spare_rows(site_id:str,request:Request,from_date:date|None=None,to_date:date|None=None,db:Session=Depends(get_db)):
    user,site_id=_access(db,request,site_id,"VIEW"); stmt=select(MaintenanceSpareUsage).where(MaintenanceSpareUsage.site_id==site_id)
    if from_date:stmt=stmt.where(func.date(MaintenanceSpareUsage.used_at)>=from_date)
    if to_date:stmt=stmt.where(func.date(MaintenanceSpareUsage.used_at)<=to_date)
    return [{"usageId":x.usage_id,"assetId":x.asset_id,"jobCardId":x.job_card_id,"description":x.description,"quantity":float(x.quantity),"unit":x.unit,"unitCost":float(x.unit_cost) if x.unit_cost is not None else None,"storeIssueReference":x.store_issue_reference,"usedAt":x.used_at.isoformat()} for x in db.scalars(stmt.order_by(MaintenanceSpareUsage.used_at.desc()).limit(5000))]


@router.post("/{site_id}/spares")
def create_spare_usage(site_id:str,p:UsageIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user,site_id=_access(db,request,site_id,"CREATE"); _asset(db,site_id,p.assetId); uid=_id("SPU",site_id); now=resolve_site_context(db,site_id).local_at; db.add(MaintenanceSpareUsage(usage_id=uid,site_id=site_id,asset_id=p.assetId,job_card_id=p.jobCardId,component_id=p.componentId,store_item_id=p.storeItemId,used_at=now,description=p.description,quantity=p.quantity,unit=p.unit,unit_cost=p.unitCost,store_issue_reference=p.storeIssueReference,remarks=p.remarks,entered_by=user.login_id)); db.commit(); return {"ok":True,"usageId":uid}


@router.get("/{site_id}/lubricants")
def lubricant_rows(site_id:str,request:Request,from_date:date|None=None,to_date:date|None=None,db:Session=Depends(get_db)):
    user,site_id=_access(db,request,site_id,"VIEW"); stmt=select(MaintenanceLubricantUsage).where(MaintenanceLubricantUsage.site_id==site_id)
    if from_date:stmt=stmt.where(func.date(MaintenanceLubricantUsage.used_at)>=from_date)
    if to_date:stmt=stmt.where(func.date(MaintenanceLubricantUsage.used_at)<=to_date)
    return [{"usageId":x.usage_id,"assetId":x.asset_id,"jobCardId":x.job_card_id,"lubricantType":x.lubricant_type,"quantity":float(x.quantity),"unit":x.unit,"unitCost":float(x.unit_cost) if x.unit_cost is not None else None,"usedAt":x.used_at.isoformat(),"remarks":x.remarks} for x in db.scalars(stmt.order_by(MaintenanceLubricantUsage.used_at.desc()).limit(5000))]


@router.post("/{site_id}/lubricants")
def create_lubricant_usage(site_id:str,p:LubricantUsageIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user,site_id=_access(db,request,site_id,"CREATE"); _asset(db,site_id,p.assetId); uid=_id("LUB",site_id); now=resolve_site_context(db,site_id).local_at; db.add(MaintenanceLubricantUsage(usage_id=uid,site_id=site_id,asset_id=p.assetId,job_card_id=p.jobCardId,used_at=now,lubricant_type=p.lubricantType,quantity=p.quantity,unit=p.unit,unit_cost=p.unitCost,remarks=p.remarks,entered_by=user.login_id)); db.commit(); return {"ok":True,"usageId":uid}


@router.get("/{site_id}/tyres")
def tyres(site_id:str,request:Request,db:Session=Depends(get_db)):
    user,site_id=_access(db,request,site_id,"VIEW"); fits={x.tyre_id:x for x in db.scalars(select(MaintenanceTyreFitment).where(MaintenanceTyreFitment.site_id==site_id,MaintenanceTyreFitment.status=="FITTED"))}; return [{"tyreId":x.tyre_id,"serialNo":x.serial_no,"brand":x.brand,"tyreSize":x.tyre_size,"purchaseDate":x.purchase_date.isoformat() if x.purchase_date else None,"purchaseCost":float(x.purchase_cost) if x.purchase_cost is not None else None,"status":x.status,"assetId":fits.get(x.tyre_id).asset_id if fits.get(x.tyre_id) else None,"position":fits.get(x.tyre_id).position if fits.get(x.tyre_id) else None,"fitKmr":float(fits.get(x.tyre_id).fit_kmr) if fits.get(x.tyre_id) and fits.get(x.tyre_id).fit_kmr is not None else None} for x in db.scalars(select(MaintenanceTyre).where(MaintenanceTyre.active.is_(True)).order_by(MaintenanceTyre.serial_no))]


@router.post("/{site_id}/tyres")
def create_tyre(site_id:str,p:TyreIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user,site_id=_access(db,request,site_id,"EDIT"); tid=_id("TYR",site_id); db.add(MaintenanceTyre(tyre_id=tid,serial_no=p.serialNo,brand=p.brand,tyre_size=p.tyreSize,purchase_date=p.purchaseDate,purchase_cost=p.purchaseCost,entered_by=user.login_id)); db.commit(); return {"ok":True,"tyreId":tid}


@router.post("/{site_id}/tyres/fit")
def fit_tyre(site_id:str,p:TyreFitmentIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user,site_id=_access(db,request,site_id,"CREATE"); tyre=db.get(MaintenanceTyre,p.tyreId); _asset(db,site_id,p.assetId)
    if not tyre or not tyre.active: raise HTTPException(404,"Tyre not found.")
    active=db.scalar(select(MaintenanceTyreFitment).where(MaintenanceTyreFitment.tyre_id==p.tyreId,MaintenanceTyreFitment.status=="FITTED").limit(1))
    if active: raise HTTPException(409,"Tyre is already fitted.")
    at=p.fittedAt or resolve_site_context(db,site_id).local_at; kmr=p.fitKmr
    if kmr is None: _,kmr,_,_=_latest_meter(db,site_id,p.assetId,"KMR")
    fid=_id("TYF",site_id); db.add(MaintenanceTyreFitment(fitment_id=fid,site_id=site_id,tyre_id=p.tyreId,asset_id=p.assetId,axle=p.axle,position=p.position,fitted_at=at,fit_kmr=kmr,entered_by=user.login_id)); tyre.status="FITTED"; db.commit(); return {"ok":True,"fitmentId":fid,"fitKmr":float(kmr) if kmr is not None else None}


@router.post("/{site_id}/tyres/{fitment_id}/remove")
def remove_tyre(site_id:str,fitment_id:str,p:TyreRemoveIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user,site_id=_access(db,request,site_id,"CREATE"); fit=db.get(MaintenanceTyreFitment,fitment_id)
    if not fit or fit.site_id!=site_id: raise HTTPException(404,"Tyre fitment not found.")
    if fit.status!="FITTED": raise HTTPException(409,"Tyre is already removed.")
    at=p.removedAt or resolve_site_context(db,site_id).local_at; kmr=p.removalKmr
    if kmr is None: _,kmr,_,_=_latest_meter(db,site_id,fit.asset_id,"KMR")
    if fit.fit_kmr is not None and kmr is not None and kmr<fit.fit_kmr: raise HTTPException(422,"Removal KMR cannot be lower than fit KMR.")
    fit.removed_at=at;fit.removal_kmr=kmr;fit.running_km=(Decimal(kmr)-Decimal(fit.fit_kmr)) if fit.fit_kmr is not None and kmr is not None else None;fit.removal_reason=p.reason;fit.status="REMOVED";fit.updated_at=datetime.utcnow(); tyre=db.get(MaintenanceTyre,fit.tyre_id);
    if tyre: tyre.status="REMOVED"
    db.commit(); return {"ok":True,"runningKm":float(fit.running_km) if fit.running_km is not None else None}


@router.get("/{site_id}/documents")
def documents(site_id:str,request:Request,asset_id:str|None=None,db:Session=Depends(get_db)):
    user,site_id=_access(db,request,site_id,"VIEW"); stmt=select(EquipmentDocument).where(EquipmentDocument.site_id==site_id,EquipmentDocument.active.is_(True));
    if asset_id:stmt=stmt.where(EquipmentDocument.asset_id==asset_id)
    today=date.today(); return [{"documentId":x.document_id,"assetId":x.asset_id,"documentType":x.document_type,"referenceNo":x.reference_no,"validFrom":x.valid_from.isoformat() if x.valid_from else None,"validTo":x.valid_to.isoformat() if x.valid_to else None,"status":"EXPIRED" if x.valid_to and x.valid_to<today else "DUE_SOON" if x.valid_to and x.valid_to<=today+timedelta(days=30) else "VALID","remarks":x.remarks} for x in db.scalars(stmt.order_by(EquipmentDocument.valid_to))]


@router.post("/{site_id}/documents")
def create_document(site_id:str,p:EquipmentDocumentIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user,site_id=_access(db,request,site_id,"EDIT"); _asset(db,site_id,p.assetId); did=_id("DOC",site_id); db.add(EquipmentDocument(document_id=did,site_id=site_id,asset_id=p.assetId,document_type=p.documentType.upper(),reference_no=p.referenceNo,valid_from=p.validFrom,valid_to=p.validTo,remarks=p.remarks,entered_by=user.login_id)); db.commit(); return {"ok":True,"documentId":did}


@router.get("/{site_id}/history/{asset_id}")
def equipment_history(site_id:str,asset_id:str,request:Request,from_date:date|None=None,to_date:date|None=None,db:Session=Depends(get_db)):
    user,site_id=_access(db,request,site_id,"VIEW"); asset=_asset(db,site_id,asset_id); start=from_date or date.today()-timedelta(days=90); end=to_date or date.today()
    meters=list(db.scalars(select(SiteAssetMeter).where(SiteAssetMeter.site_id==site_id,SiteAssetMeter.asset_id==asset_id,SiteAssetMeter.operating_date>=start,SiteAssetMeter.operating_date<=end).order_by(SiteAssetMeter.operating_date.desc()).limit(1000)))
    bds=list(db.scalars(select(MaintenanceBreakdown).where(MaintenanceBreakdown.site_id==site_id,MaintenanceBreakdown.asset_id==asset_id,MaintenanceBreakdown.operating_date>=start,MaintenanceBreakdown.operating_date<=end).order_by(MaintenanceBreakdown.reported_at.desc())))
    jobs=list(db.scalars(select(MaintenanceJobCard).where(MaintenanceJobCard.site_id==site_id,MaintenanceJobCard.asset_id==asset_id,func.date(MaintenanceJobCard.opened_at)>=start,func.date(MaintenanceJobCard.opened_at)<=end).order_by(MaintenanceJobCard.opened_at.desc())))
    deff=list(db.scalars(select(MaintenanceDefTransaction).where(MaintenanceDefTransaction.site_id==site_id,MaintenanceDefTransaction.asset_id==asset_id,MaintenanceDefTransaction.operating_date>=start,MaintenanceDefTransaction.operating_date<=end).order_by(MaintenanceDefTransaction.event_at.desc())))
    return {"asset":{"assetId":asset.machine_id,"display":_display_asset(asset),"type":asset.type,"group":asset.group,"makeModel":asset.make_model},"meters":[{"date":x.operating_date.isoformat(),"shift":x.shift,"type":x.meter_type,"opening":float(x.opening_reading) if x.opening_reading is not None else None,"closing":float(x.closing_reading) if x.closing_reading is not None else None,"usage":float(x.usage) if x.usage is not None else None,"source":x.source_type} for x in meters],"breakdowns":[{"breakdownId":x.breakdown_id,"reportedAt":x.reported_at.isoformat(),"releasedAt":x.released_at.isoformat() if x.released_at else None,"problem":x.problem,"status":x.status,"downtimeHours":float(x.downtime_hours) if x.downtime_hours is not None else None} for x in bds],"jobCards":[{"jobCardId":x.job_card_id,"jobType":x.job_type,"status":x.status,"openedAt":x.opened_at.isoformat(),"actionTaken":x.action_taken} for x in jobs],"def":[{"eventAt":x.event_at.isoformat(),"quantityL":float(x.quantity_l),"meterReading":float(x.meter_reading) if x.meter_reading is not None else None} for x in deff]}


def _style_sheet(ws, title):
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    fill = PatternFill("solid", fgColor="0F766E")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True); cell.fill = fill; cell.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        letter = col[0].column_letter; width=min(42,max(11,max(len(str(c.value or "")) for c in col)+2)); ws.column_dimensions[letter].width=width


def _make_excel_datetimes_safe(wb: Workbook, timezone_name: str = "Asia/Kolkata") -> None:
    """Convert timezone-aware datetimes to site-local naive datetimes for Excel.

    PostgreSQL returns TIMESTAMPTZ values with tzinfo. openpyxl/Excel rejects timezone-aware
    datetime values at save time, so report exports must normalize them first.
    """
    from zoneinfo import ZoneInfo
    zone = ZoneInfo(timezone_name or "Asia/Kolkata")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, datetime) and value.tzinfo is not None:
                    cell.value = value.astimezone(zone).replace(tzinfo=None)
                    if cell.number_format == "General":
                        cell.number_format = "yyyy-mm-dd hh:mm:ss"



@router.get("/{site_id}/reports/preview")
def report_preview(site_id:str,request:Request,report:str="COMPLETE",from_date:date|None=None,to_date:date|None=None,asset_id:str|None=None,limit:int=250,db:Session=Depends(get_db)):
    user,site_id=_access(db,request,site_id,"VIEW"); start=from_date or date.today().replace(day=1); end=to_date or date.today(); report=report.upper(); limit=max(1,min(int(limit or 250),1000))
    if end<start: raise HTTPException(422,"to_date cannot be before from_date")
    def take(stmt): return list(db.scalars(stmt.limit(limit)))
    if report=="COMPLETE":
        def cnt(model,date_col):
            stmt=select(func.count()).select_from(model).where(model.site_id==site_id,date_col>=start,date_col<=end)
            if asset_id and hasattr(model,'asset_id'): stmt=stmt.where(model.asset_id==asset_id)
            return int(db.scalar(stmt) or 0)
        summary=[
            ["Breakdowns",cnt(MaintenanceBreakdown,MaintenanceBreakdown.operating_date)],
            ["Job Cards",cnt(MaintenanceJobCard,func.date(MaintenanceJobCard.opened_at))],
            ["DEF / Urea",cnt(MaintenanceDefTransaction,MaintenanceDefTransaction.operating_date)],
            ["PM / WGP",cnt(MaintenancePmExecution,func.date(MaintenancePmExecution.completed_at))],
            ["Spares",cnt(MaintenanceSpareUsage,func.date(MaintenanceSpareUsage.used_at))],
            ["Lubricants",cnt(MaintenanceLubricantUsage,func.date(MaintenanceLubricantUsage.used_at))],
            ["Services",cnt(MaintenanceServiceHistory,func.date(MaintenanceServiceHistory.serviced_at))],
            ["Component Changes",cnt(ComponentChangeHistory,func.date(ComponentChangeHistory.changed_at))],
        ]
        return {"report":report,"fromDate":start.isoformat(),"toDate":end.isoformat(),"columns":["Section","Records"],"rows":summary,"truncated":False}
    if report=="BREAKDOWN":
        stmt=select(MaintenanceBreakdown).where(MaintenanceBreakdown.site_id==site_id,MaintenanceBreakdown.operating_date>=start,MaintenanceBreakdown.operating_date<=end).order_by(MaintenanceBreakdown.reported_at.desc())
        if asset_id: stmt=stmt.where(MaintenanceBreakdown.asset_id==asset_id)
        rows=take(stmt); return {"report":report,"columns":["Equipment","Date","Shift","Status","Problem","Downtime Hrs"],"rows":[[x.asset_id,x.operating_date.isoformat(),x.shift,x.status,x.problem,float(x.downtime_hours) if x.downtime_hours is not None else None] for x in rows],"truncated":len(rows)>=limit}
    if report=="JOB_CARD":
        stmt=select(MaintenanceJobCard).where(MaintenanceJobCard.site_id==site_id,func.date(MaintenanceJobCard.opened_at)>=start,func.date(MaintenanceJobCard.opened_at)<=end).order_by(MaintenanceJobCard.opened_at.desc())
        if asset_id: stmt=stmt.where(MaintenanceJobCard.asset_id==asset_id)
        rows=take(stmt); return {"report":report,"columns":["Equipment","Job Card","Type","Opened","Status","Complaint","Action"],"rows":[[x.asset_id,x.job_card_id,x.job_type,x.opened_at.isoformat() if x.opened_at else None,x.status,x.complaint,x.action_taken] for x in rows],"truncated":len(rows)>=limit}
    if report=="DEF":
        stmt=select(MaintenanceDefTransaction).where(MaintenanceDefTransaction.site_id==site_id,MaintenanceDefTransaction.operating_date>=start,MaintenanceDefTransaction.operating_date<=end).order_by(MaintenanceDefTransaction.event_at.desc())
        if asset_id: stmt=stmt.where(MaintenanceDefTransaction.asset_id==asset_id)
        rows=take(stmt); return {"report":report,"columns":["Equipment","Date","Shift","Quantity L","Meter","Reading","Reference"],"rows":[[x.asset_id,x.operating_date.isoformat(),x.shift,float(x.quantity_l),x.meter_type,float(x.meter_reading) if x.meter_reading is not None else None,x.issue_reference] for x in rows],"truncated":len(rows)>=limit}
    if report=="PM":
        stmt=select(MaintenancePmExecution).where(MaintenancePmExecution.site_id==site_id,func.date(MaintenancePmExecution.completed_at)>=start,func.date(MaintenancePmExecution.completed_at)<=end).order_by(MaintenancePmExecution.completed_at.desc())
        if asset_id: stmt=stmt.where(MaintenancePmExecution.asset_id==asset_id)
        rows=take(stmt); return {"report":report,"columns":["Equipment","Activity","Completed","Result","Meter","Reading","Remarks"],"rows":[[x.asset_id,x.activity,x.completed_at.isoformat() if x.completed_at else None,x.result,x.meter_type,float(x.meter_reading) if x.meter_reading is not None else None,x.remarks] for x in rows],"truncated":len(rows)>=limit}
    if report=="SPARES":
        stmt=select(MaintenanceSpareUsage).where(MaintenanceSpareUsage.site_id==site_id,func.date(MaintenanceSpareUsage.used_at)>=start,func.date(MaintenanceSpareUsage.used_at)<=end).order_by(MaintenanceSpareUsage.used_at.desc())
        if asset_id: stmt=stmt.where(MaintenanceSpareUsage.asset_id==asset_id)
        rows=take(stmt); return {"report":report,"columns":["Equipment","Used At","Part / Spare","Qty","Unit","Unit Cost","Store Ref"],"rows":[[x.asset_id,x.used_at.isoformat() if x.used_at else None,x.description,float(x.quantity),x.unit,float(x.unit_cost) if x.unit_cost is not None else None,x.store_issue_reference] for x in rows],"truncated":len(rows)>=limit}
    if report=="LUBRICANTS":
        stmt=select(MaintenanceLubricantUsage).where(MaintenanceLubricantUsage.site_id==site_id,func.date(MaintenanceLubricantUsage.used_at)>=start,func.date(MaintenanceLubricantUsage.used_at)<=end).order_by(MaintenanceLubricantUsage.used_at.desc())
        if asset_id: stmt=stmt.where(MaintenanceLubricantUsage.asset_id==asset_id)
        rows=take(stmt); return {"report":report,"columns":["Equipment","Used At","Lubricant","Qty","Unit","Unit Cost","Remarks"],"rows":[[x.asset_id,x.used_at.isoformat() if x.used_at else None,x.lubricant_type,float(x.quantity),x.unit,float(x.unit_cost) if x.unit_cost is not None else None,x.remarks] for x in rows],"truncated":len(rows)>=limit}
    if report=="SERVICE":
        stmt=select(MaintenanceServiceHistory).where(MaintenanceServiceHistory.site_id==site_id,func.date(MaintenanceServiceHistory.serviced_at)>=start,func.date(MaintenanceServiceHistory.serviced_at)<=end).order_by(MaintenanceServiceHistory.serviced_at.desc())
        if asset_id: stmt=stmt.where(MaintenanceServiceHistory.asset_id==asset_id)
        rows=take(stmt); return {"report":report,"columns":["Equipment","Serviced At","Meter","Reading","Next Due Date","Next Due Meter","Job Card"],"rows":[[x.asset_id,x.serviced_at.isoformat() if x.serviced_at else None,x.meter_type,float(x.meter_reading) if x.meter_reading is not None else None,x.next_due_at.isoformat() if x.next_due_at else None,float(x.next_due_meter) if x.next_due_meter is not None else None,x.job_card_id] for x in rows],"truncated":len(rows)>=limit}
    if report=="COMPONENTS":
        stmt=select(ComponentChangeHistory).where(ComponentChangeHistory.site_id==site_id,func.date(ComponentChangeHistory.changed_at)>=start,func.date(ComponentChangeHistory.changed_at)<=end).order_by(ComponentChangeHistory.changed_at.desc())
        if asset_id: stmt=stmt.where(ComponentChangeHistory.asset_id==asset_id)
        rows=take(stmt); return {"report":report,"columns":["Equipment","Changed At","Meter","Reading","Qty","Job Card","Remarks"],"rows":[[x.asset_id,x.changed_at.isoformat() if x.changed_at else None,x.meter_type,float(x.meter_reading) if x.meter_reading is not None else None,float(x.quantity),x.job_card_id,x.remarks] for x in rows],"truncated":len(rows)>=limit}
    if report=="TYRES":
        stmt=select(MaintenanceTyreFitment).where(MaintenanceTyreFitment.site_id==site_id,func.date(MaintenanceTyreFitment.fitted_at)<=end,or_(MaintenanceTyreFitment.removed_at.is_(None),func.date(MaintenanceTyreFitment.removed_at)>=start)).order_by(MaintenanceTyreFitment.fitted_at.desc())
        if asset_id: stmt=stmt.where(MaintenanceTyreFitment.asset_id==asset_id)
        rows=take(stmt); return {"report":report,"columns":["Equipment","Tyre","Position","Fitted","Fit KMR","Removed","Running KM","Status"],"rows":[[x.asset_id,x.tyre_id,x.position,x.fitted_at.isoformat() if x.fitted_at else None,float(x.fit_kmr) if x.fit_kmr is not None else None,x.removed_at.isoformat() if x.removed_at else None,float(x.running_km) if x.running_km is not None else None,x.status] for x in rows],"truncated":len(rows)>=limit}
    if report=="COMPLIANCE":
        stmt=select(EquipmentDocument).where(EquipmentDocument.site_id==site_id,EquipmentDocument.active.is_(True)).order_by(EquipmentDocument.valid_to)
        if asset_id: stmt=stmt.where(EquipmentDocument.asset_id==asset_id)
        today=date.today(); rows=take(stmt); return {"report":report,"columns":["Equipment","Type","Reference","Valid To","Status","Remarks"],"rows":[[x.asset_id,x.document_type,x.reference_no,x.valid_to.isoformat() if x.valid_to else None,"EXPIRED" if x.valid_to and x.valid_to<today else "DUE_SOON" if x.valid_to and x.valid_to<=today+timedelta(days=30) else "VALID",x.remarks] for x in rows],"truncated":len(rows)>=limit}
    if report=="STATUS":
        ctx=resolve_site_context(db,site_id); rows=_derive_statuses(db,site_id,ctx.operating_date,ctx.shift); rows=[x for x in rows if not asset_id or x["assetId"]==asset_id][:limit]
        return {"report":report,"columns":["Equipment","Type","Status","Operator","Location","Meter","Reading"],"rows":[[x["assetId"],x["type"],x["status"],x["operatorId"],x["locationId"],x["meterType"],x["meterReading"]] for x in rows],"truncated":len(rows)>=limit}
    raise HTTPException(422,"Unknown report type.")

@router.get("/{site_id}/reports/export.xlsx")
def export_report(site_id:str,request:Request,report:str="COMPLETE",from_date:date|None=None,to_date:date|None=None,asset_id:str|None=None,db:Session=Depends(get_db)):
    user,site_id=_access(db,request,site_id,"EXPORT"); start=from_date or date.today().replace(day=1); end=to_date or date.today()
    if end<start: raise HTTPException(422,"to_date cannot be before from_date")
    report=report.upper(); wb=Workbook(); wb.remove(wb.active)
    if report in {"COMPLETE","BREAKDOWN"}:
        ws=wb.create_sheet("Breakdown"); ws.append(["Breakdown ID","Equipment","Date","Shift","Reported","Released","Status","Problem Category","Problem","Downtime Hrs","HMR/KMR","Reading"]); stmt=select(MaintenanceBreakdown).where(MaintenanceBreakdown.site_id==site_id,MaintenanceBreakdown.operating_date>=start,MaintenanceBreakdown.operating_date<=end)
        if asset_id:stmt=stmt.where(MaintenanceBreakdown.asset_id==asset_id)
        for x in db.scalars(stmt.order_by(MaintenanceBreakdown.reported_at)):
            ws.append([x.breakdown_id,x.asset_id,x.operating_date,x.shift,x.reported_at,x.released_at,x.status,x.problem_category,x.problem,float(x.downtime_hours) if x.downtime_hours is not None else None,x.meter_type,float(x.meter_reading) if x.meter_reading is not None else None]); _style_sheet(ws,"Breakdown")
    if report in {"COMPLETE","JOB_CARD"}:
        ws=wb.create_sheet("Job Cards"); ws.append(["Job Card","Equipment","Type","Source","Opened","Completed","Status","Complaint","Diagnosis","Root Cause","Action","Mechanic","Labour Hrs","External Cost"]); stmt=select(MaintenanceJobCard).where(MaintenanceJobCard.site_id==site_id,func.date(MaintenanceJobCard.opened_at)>=start,func.date(MaintenanceJobCard.opened_at)<=end)
        if asset_id:stmt=stmt.where(MaintenanceJobCard.asset_id==asset_id)
        for x in db.scalars(stmt.order_by(MaintenanceJobCard.opened_at)):
            ws.append([x.job_card_id,x.asset_id,x.job_type,x.source_type,x.opened_at,x.completed_at,x.status,x.complaint,x.diagnosis,x.root_cause,x.action_taken,x.mechanic_id,float(x.labour_hours) if x.labour_hours is not None else None,float(x.external_cost or 0)]); _style_sheet(ws,"Job Cards")
    if report in {"COMPLETE","DEF"}:
        ws=wb.create_sheet("DEF Urea"); ws.append(["DEF ID","Equipment","Date","Shift","Time","Quantity L","HMR/KMR","Reading","Operator","Reference","Remarks"]); stmt=select(MaintenanceDefTransaction).where(MaintenanceDefTransaction.site_id==site_id,MaintenanceDefTransaction.operating_date>=start,MaintenanceDefTransaction.operating_date<=end)
        if asset_id:stmt=stmt.where(MaintenanceDefTransaction.asset_id==asset_id)
        for x in db.scalars(stmt.order_by(MaintenanceDefTransaction.event_at)):
            ws.append([x.def_id,x.asset_id,x.operating_date,x.shift,x.event_at,float(x.quantity_l),x.meter_type,float(x.meter_reading) if x.meter_reading is not None else None,x.operator_id,x.issue_reference,x.remarks]); _style_sheet(ws,"DEF")
    if report in {"COMPLETE","DEF"} and site_id == "TIOM":
        ws=wb.create_sheet("DEF vs HSD Summary"); ws.append(["Equipment","DEF L","HSD L","DEF % of HSD"]); def_totals={asset:Decimal(qty or 0) for asset,qty in db.execute(select(MaintenanceDefTransaction.asset_id,func.sum(MaintenanceDefTransaction.quantity_l)).where(MaintenanceDefTransaction.site_id==site_id,MaintenanceDefTransaction.operating_date>=start,MaintenanceDefTransaction.operating_date<=end).group_by(MaintenanceDefTransaction.asset_id)).all()}; hsd_totals={asset:Decimal(qty or 0) for asset,qty in db.execute(select(HsdIssue.machine_id,func.sum(HsdIssue.litres)).where(HsdIssue.operating_date>=start,HsdIssue.operating_date<=end).group_by(HsdIssue.machine_id)).all()}
        ids=sorted(set(def_totals)|set(hsd_totals)); ids=[x for x in ids if not asset_id or x==asset_id]
        for aid in ids:
            d=def_totals.get(aid,Decimal("0")); h=hsd_totals.get(aid,Decimal("0")); pct=(d/h*Decimal("100")) if h>0 else None; ws.append([aid,float(d),float(h),float(pct) if pct is not None else None])
        _style_sheet(ws,"DEF vs HSD")
    if report in {"COMPLETE","PM"}:
        ws=wb.create_sheet("PM WGP"); ws.append(["Execution ID","Equipment","Activity","Completed","Result","HMR/KMR","Reading","Remarks"]); stmt=select(MaintenancePmExecution).where(MaintenancePmExecution.site_id==site_id,func.date(MaintenancePmExecution.completed_at)>=start,func.date(MaintenancePmExecution.completed_at)<=end)
        if asset_id:stmt=stmt.where(MaintenancePmExecution.asset_id==asset_id)
        for x in db.scalars(stmt.order_by(MaintenancePmExecution.completed_at)):
            ws.append([x.execution_id,x.asset_id,x.activity,x.completed_at,x.result,x.meter_type,float(x.meter_reading) if x.meter_reading is not None else None,x.remarks]); _style_sheet(ws,"PM")
    if report in {"COMPLETE","SPARES"}:
        ws=wb.create_sheet("Spares"); ws.append(["Usage ID","Equipment","Job Card","Used At","Part / Spare","Qty","Unit","Unit Cost","Store Issue Ref"]); stmt=select(MaintenanceSpareUsage).where(MaintenanceSpareUsage.site_id==site_id,func.date(MaintenanceSpareUsage.used_at)>=start,func.date(MaintenanceSpareUsage.used_at)<=end)
        if asset_id:stmt=stmt.where(MaintenanceSpareUsage.asset_id==asset_id)
        for x in db.scalars(stmt.order_by(MaintenanceSpareUsage.used_at)):
            ws.append([x.usage_id,x.asset_id,x.job_card_id,x.used_at,x.description,float(x.quantity),x.unit,float(x.unit_cost) if x.unit_cost is not None else None,x.store_issue_reference]); _style_sheet(ws,"Spares")
    if report in {"COMPLETE","LUBRICANTS"}:
        ws=wb.create_sheet("Lubricants"); ws.append(["Usage ID","Equipment","Job Card","Used At","Lubricant","Qty","Unit","Unit Cost","Remarks"]); stmt=select(MaintenanceLubricantUsage).where(MaintenanceLubricantUsage.site_id==site_id,func.date(MaintenanceLubricantUsage.used_at)>=start,func.date(MaintenanceLubricantUsage.used_at)<=end)
        if asset_id:stmt=stmt.where(MaintenanceLubricantUsage.asset_id==asset_id)
        for x in db.scalars(stmt.order_by(MaintenanceLubricantUsage.used_at)):
            ws.append([x.usage_id,x.asset_id,x.job_card_id,x.used_at,x.lubricant_type,float(x.quantity),x.unit,float(x.unit_cost) if x.unit_cost is not None else None,x.remarks]); _style_sheet(ws,"Lubricants")
    if report in {"COMPLETE","SERVICE"}:
        ws=wb.create_sheet("Service History"); ws.append(["Service ID","Plan ID","Equipment","Serviced At","HMR/KMR","Reading","Next Due Date","Next Due Meter","Job Card","Remarks"]); stmt=select(MaintenanceServiceHistory).where(MaintenanceServiceHistory.site_id==site_id,func.date(MaintenanceServiceHistory.serviced_at)>=start,func.date(MaintenanceServiceHistory.serviced_at)<=end)
        if asset_id:stmt=stmt.where(MaintenanceServiceHistory.asset_id==asset_id)
        for x in db.scalars(stmt.order_by(MaintenanceServiceHistory.serviced_at)):
            ws.append([x.service_id,x.plan_id,x.asset_id,x.serviced_at,x.meter_type,float(x.meter_reading) if x.meter_reading is not None else None,x.next_due_at,float(x.next_due_meter) if x.next_due_meter is not None else None,x.job_card_id,x.remarks]); _style_sheet(ws,"Service History")
    if report in {"COMPLETE","COMPONENTS"}:
        ws=wb.create_sheet("Component Changes"); ws.append(["Change ID","Schedule ID","Equipment","Changed At","HMR/KMR","Reading","Qty","Job Card","Remarks"]); stmt=select(ComponentChangeHistory).where(ComponentChangeHistory.site_id==site_id,func.date(ComponentChangeHistory.changed_at)>=start,func.date(ComponentChangeHistory.changed_at)<=end)
        if asset_id:stmt=stmt.where(ComponentChangeHistory.asset_id==asset_id)
        for x in db.scalars(stmt.order_by(ComponentChangeHistory.changed_at)):
            ws.append([x.change_id,x.schedule_id,x.asset_id,x.changed_at,x.meter_type,float(x.meter_reading) if x.meter_reading is not None else None,float(x.quantity),x.job_card_id,x.remarks]); _style_sheet(ws,"Component Changes")
    if report in {"COMPLETE","TYRES"}:
        ws=wb.create_sheet("Tyre History"); ws.append(["Fitment ID","Tyre ID","Equipment","Axle","Position","Fitted At","Fit KMR","Removed At","Removal KMR","Running KM","Reason","Status"]); stmt=select(MaintenanceTyreFitment).where(MaintenanceTyreFitment.site_id==site_id,func.date(MaintenanceTyreFitment.fitted_at)<=end,or_(MaintenanceTyreFitment.removed_at.is_(None),func.date(MaintenanceTyreFitment.removed_at)>=start))
        if asset_id:stmt=stmt.where(MaintenanceTyreFitment.asset_id==asset_id)
        for x in db.scalars(stmt.order_by(MaintenanceTyreFitment.fitted_at)):
            ws.append([x.fitment_id,x.tyre_id,x.asset_id,x.axle,x.position,x.fitted_at,float(x.fit_kmr) if x.fit_kmr is not None else None,x.removed_at,float(x.removal_kmr) if x.removal_kmr is not None else None,float(x.running_km) if x.running_km is not None else None,x.removal_reason,x.status]); _style_sheet(ws,"Tyre History")
    if report in {"COMPLETE","COMPLIANCE"}:
        ws=wb.create_sheet("Compliance"); ws.append(["Document ID","Equipment","Type","Reference","Valid From","Valid To","Current Status","Remarks"]); stmt=select(EquipmentDocument).where(EquipmentDocument.site_id==site_id,EquipmentDocument.active.is_(True))
        if asset_id:stmt=stmt.where(EquipmentDocument.asset_id==asset_id)
        today=date.today()
        for x in db.scalars(stmt.order_by(EquipmentDocument.asset_id,EquipmentDocument.document_type)):
            st="EXPIRED" if x.valid_to and x.valid_to<today else "DUE_SOON" if x.valid_to and x.valid_to<=today+timedelta(days=30) else "VALID"
            ws.append([x.document_id,x.asset_id,x.document_type,x.reference_no,x.valid_from,x.valid_to,st,x.remarks]); _style_sheet(ws,"Compliance")
    if report in {"COMPLETE","STATUS"}:
        ws=wb.create_sheet("Equipment Status"); ws.append(["Equipment","Display","Type","Group","Make Model","Status","Operator","Location","HMR/KMR","Reading","Meter Date","Shift"]); ctx=resolve_site_context(db,site_id); snapshot=_derive_statuses(db,site_id,ctx.operating_date,ctx.shift)
        if asset_id:snapshot=[x for x in snapshot if x["assetId"]==asset_id]
        for x in snapshot:
            ws.append([x["assetId"],x["display"],x["type"],x["group"],x["makeModel"],x["status"],x["operatorId"],x["locationId"],x["meterType"],x["meterReading"],x["meterDate"],x["meterShift"]]); _style_sheet(ws,"Equipment Status")
    if not wb.sheetnames:
        raise HTTPException(422,"Unknown report type.")
    site = db.get(Site, site_id)
    _make_excel_datetimes_safe(wb, site.timezone if site else "Asia/Kolkata")
    stream=BytesIO(); wb.save(stream); stream.seek(0); filename=f"TIOM_Mechanical_{report}_{start}_{end}.xlsx"
    return StreamingResponse(stream,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="{filename}"'})
