"""TIOM despatch ledger. Separate from internal WB and production movements."""
from __future__ import annotations

import csv
import hashlib
import json
import re
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from email.utils import parseaddr
from io import BytesIO, StringIO
from uuid import uuid4

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy import Boolean, Date, DateTime, Integer, LargeBinary, Numeric, String, Text, func, select
from sqlalchemy.orm import Mapped, Session, mapped_column

from app.auth import csrf, get_user, require
from app.db import Base, get_db

router = APIRouter(prefix="/api/tiom/despatch", tags=["TIOM Despatch"])

DEFAULT_SENDERS = (
    ("sanjaymahakud50@gmail.com", "Sanjay Mahakud"),
    ("dillip.sahu80@gmail.com", "Dillip Kumar Sahu"),
    ("chinmaykumarpanda@gmail.com", "Chinmay Panda"),
    ("binodsingh8018@gmail.com", "Binod Singh"),
    ("tarunnayak63@gmail.com", "Tarun Kumar Nayak"),
    ("sanjaykumarbarik1982@gmail.com", "Sanjay Kumar Barik"),
    ("niranjanbabu95@gmail.com", "Niranjan Mahanta"),
    ("rajugupta97093@gmail.com", "Bansilal Gupta"),
    ("sumantakumarmahanta@gmail.com", "Sumanta Kumar Mahanta"),
)


class DespatchSender(Base):
    __tablename__ = "tiom_despatch_sender"
    email: Mapped[str] = mapped_column(String(200), primary_key=True)
    name: Mapped[str] = mapped_column(String(160), nullable=False)
    active: Mapped[bool] = mapped_column(Boolean, default=True, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class DespatchTrip(Base):
    __tablename__ = "tiom_despatch_trip"
    epass_no: Mapped[str] = mapped_column(String(120), primary_key=True)
    report_date: Mapped[date] = mapped_column(Date, nullable=False, index=True)
    permit_no: Mapped[str] = mapped_column(String(80), nullable=False, index=True)
    permit_qty: Mapped[Decimal] = mapped_column(Numeric(14, 3), default=0)
    buyer: Mapped[str] = mapped_column(String(240), default="")
    vehicle_no: Mapped[str] = mapped_column(String(80), default="")
    transporter: Mapped[str] = mapped_column(String(240), default="")
    destination: Mapped[str] = mapped_column(String(240), default="")
    material_grade: Mapped[str] = mapped_column(String(160), default="")
    stack_no: Mapped[str] = mapped_column(String(100), default="")
    consignee: Mapped[str] = mapped_column(String(300), default="")
    tare_mt: Mapped[Decimal] = mapped_column(Numeric(12, 3), default=0)
    gross_mt: Mapped[Decimal] = mapped_column(Numeric(12, 3), default=0)
    net_mt: Mapped[Decimal] = mapped_column(Numeric(12, 3), default=0)
    dispatch_time: Mapped[str] = mapped_column(String(80), default="")
    status: Mapped[str] = mapped_column(String(24), default="ACTIVE", nullable=False, index=True)
    source: Mapped[str] = mapped_column(String(200), default="")
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class DespatchDocument(Base):
    __tablename__ = "tiom_despatch_document"
    id: Mapped[str] = mapped_column(String(40), primary_key=True)
    report_date: Mapped[date | None] = mapped_column(Date, index=True)
    filename: Mapped[str] = mapped_column(String(255), nullable=False)
    media_type: Mapped[str] = mapped_column(String(120), nullable=False)
    sha256: Mapped[str] = mapped_column(String(64), nullable=False, unique=True)
    image_data: Mapped[bytes] = mapped_column(LargeBinary, nullable=False)
    review_status: Mapped[str] = mapped_column(String(30), default="NEEDS_REVIEW")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class DespatchPermitDay(Base):
    __tablename__ = "tiom_despatch_permit_day"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    report_date: Mapped[date] = mapped_column(Date, nullable=False, index=True)
    permit_no: Mapped[str] = mapped_column(String(80), nullable=False, index=True)
    buyer: Mapped[str] = mapped_column(String(200), default="")
    destination: Mapped[str] = mapped_column(String(200), default="")
    material: Mapped[str] = mapped_column(String(120), default="")
    grade_percent: Mapped[str] = mapped_column(String(40), default="")
    permit_qty: Mapped[Decimal | None] = mapped_column(Numeric(14, 3))
    reported_trips: Mapped[int | None] = mapped_column(Integer)
    reported_mt: Mapped[Decimal | None] = mapped_column(Numeric(14, 3))
    balance_mt: Mapped[Decimal | None] = mapped_column(Numeric(14, 3))
    document_id: Mapped[str | None] = mapped_column(String(40))
    review_status: Mapped[str] = mapped_column(String(30), default="NEEDS_REVIEW")


class DespatchAudit(Base):
    __tablename__ = "tiom_despatch_audit"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    changed_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    actor: Mapped[str] = mapped_column(String(80))
    action: Mapped[str] = mapped_column(String(40))
    entity_key: Mapped[str] = mapped_column(String(160))
    reason: Mapped[str] = mapped_column(Text, default="")
    before_json: Mapped[str | None] = mapped_column(Text)
    after_json: Mapped[str | None] = mapped_column(Text)


def _user(request, db, edit=False):
    user = get_user(db, request)
    require(user, "PRODUCTION")
    if edit and not user.admin:
        raise HTTPException(403, "Despatch corrections and sender changes require management.")
    return user


def _seed_senders(db):
    existing = set(db.scalars(select(DespatchSender.email)).all())
    if not existing:
        db.add_all(DespatchSender(email=e, name=n, active=True) for e, n in DEFAULT_SENDERS)
        db.commit()


def _amount(value):
    try:
        out = Decimal(str(value if value is not None and value != "" else 0))
        if not out.is_finite():
            raise InvalidOperation()
        return out.quantize(Decimal("0.001"))
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError("Invalid dispatch quantity or weight")


def _date(value, fallback=None):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for pattern in ("%Y-%m-%d", "%d.%m.%y", "%d/%m/%Y", "%b %d %Y", "%b  %d %Y"):
        try:
            return datetime.strptime(str(value or "").strip(), pattern).date()
        except ValueError:
            pass
    if fallback:
        return fallback
    raise ValueError("Invalid dispatch date")


def _trip(d):
    def item(name):
        return str(d.get(name) or "").strip()
    epass = item("epass_no")
    if not epass or len(epass) > 120:
        raise ValueError("Missing or invalid E-Pass No.")
    day = _date(item("report_date"))
    if day < date(2026, 10, 1):
        raise ValueError("Despatch backfill starts from 2026-10-01")
    tare, gross, net = (_amount(item(k)) for k in ("tare_mt", "gross_mt", "net_mt"))
    if min(tare, gross, net) < 0 or abs((gross - tare) - net) > Decimal(".025"):
        raise ValueError("Gross - tare does not match net weight")
    return dict(
        epass_no=epass, report_date=day, permit_no=item("permit_no"),
        permit_qty=_amount(item("permit_qty")), buyer=item("buyer"),
        vehicle_no=item("vehicle_no"), transporter=item("transporter"),
        destination=item("destination"), material_grade=item("material_grade"),
        stack_no=item("stack_no"), consignee=item("consignee"),
        tare_mt=tare, gross_mt=gross, net_mt=net, dispatch_time=item("dispatch_time"),
        status="ACTIVE", source=item("source")[:200],
    )


def _extract(data, filename):
    from app.services.tiom_despatch_source_adapters import parse_despatch_sources
    return parse_despatch_sources(data, filename)


def _import_data(db, data, filename, actor):
    """Import verified weighments and enrich other source data by E-Pass ID.

    Detailed sources add E-Pass rows; transporter manifests may fill only
    missing transporter metadata, never fabricate weights or override a match.
    Both source types are atomic: a mismatch blocks the entire upload.
    """
    if len(data) > 12_000_000:
        raise ValueError("File exceeds 12 MB")
    raw = _extract(data, filename)
    if len(raw) > 100_000:
        raise ValueError("Too many rows")
    detailed = [row for row in raw if row.get("_role", "detailed") == "detailed"]
    enrichment = [row for row in raw if row.get("_role") == "enrichment"]
    seen, parsed, errors = set(), [], []
    for i, row in enumerate(detailed, 1):
        try:
            trip = _trip(row)
            key = trip["epass_no"].upper()
            if key in seen:
                raise ValueError("Repeated E-Pass within source")
            seen.add(key)
            parsed.append(trip)
        except ValueError as exc:
            errors.append({"row": i, "error": str(exc)})
    seen_aux = set()
    for i, row in enumerate(enrichment, 1):
        epass = str(row.get("epass_no") or "").strip()
        try:
            if epass.upper() in seen_aux:
                raise ValueError("Repeated E-Pass in transporter source")
            seen_aux.add(epass.upper())
            _date(row.get("report_date"))
            _amount(row.get("net_mt"))
            if not row.get("transporter") or not row.get("vehicle_no"):
                raise ValueError("Vehicle and transporter are mandatory for enrichment")
        except ValueError as exc:
            errors.append({"row": i + len(detailed), "error": str(exc)})
    if errors:
        return {"ok": False, "inserted": 0, "enriched": 0,
                "errors": errors[:30], "errorCount": len(errors), "reviewRequired": True}
    lookup_keys = list({x["epass_no"] for x in parsed} |
                       {str(x["epass_no"]) for x in enrichment})
    existing = {}
    for offset in range(0, len(lookup_keys), 500):
        for trip in db.scalars(select(DespatchTrip).where(
            DespatchTrip.epass_no.in_(lookup_keys[offset:offset + 500])
        )):
            existing[trip.epass_no.upper()] = trip
    new_trips, enriched, mismatches = [], 0, []
    for entry in parsed:
        key = entry["epass_no"].upper()
        old = existing.get(key)
        if old:
            if (abs(old.net_mt - entry["net_mt"]) > Decimal(".025") or
                    old.report_date != entry["report_date"] or
                    (old.vehicle_no and entry["vehicle_no"] and
                     old.vehicle_no.strip().upper() != entry["vehicle_no"].strip().upper())):
                mismatches.append(key + ": existing E-Pass weight/date/vehicle differs")
            continue
        trip = DespatchTrip(**entry)
        existing[key] = trip
        new_trips.append(trip)
    for entry in enrichment:
        epass = str(entry["epass_no"]).strip()
        old = existing.get(epass.upper())
        if not old:
            mismatches.append(epass + ": transporter manifest has no matching weighment E-Pass")
            continue
        if (old.report_date != _date(entry["report_date"]) or
                abs(old.net_mt - _amount(entry["net_mt"])) > Decimal(".025") or
                (old.vehicle_no.strip().upper() != str(entry["vehicle_no"]).strip().upper())):
            mismatches.append(epass + ": transporter manifest date/MT/vehicle mismatch")
            continue
        transporter = str(entry["transporter"]).strip()
        if old.transporter and old.transporter.strip().casefold() != transporter.casefold():
            mismatches.append(epass + ": conflicting transporter; original retained")
        elif not old.transporter:
            old.transporter = transporter
            enriched += 1
    if mismatches:
        db.rollback()
        return {"ok": False, "inserted": 0, "enriched": 0, "errorCount": len(mismatches),
                "errors": [{"row": 0, "error": x} for x in mismatches[:30]],
                "reviewRequired": True, "conflicts": mismatches[:30],
                "conflictCount": len(mismatches)}
    for trip in new_trips:
        db.add(trip)
    db.add(DespatchAudit(
        actor=actor, action="IMPORT", entity_key=filename[:160],
        reason="Schema-detected E-Pass import / cross-source enrichment",
        before_json=None, after_json=json.dumps({
            "sourceRows": len(raw), "detailedRows": len(detailed),
            "enrichmentRows": len(enrichment), "inserted": len(new_trips),
            "transporterEnriched": enriched,
            "duplicateEPasses": len(detailed) - len(new_trips),
        })
    ))
    db.commit()
    return {"ok": True, "inserted": len(new_trips), "enriched": enriched,
            "skippedExisting": len(detailed) - len(new_trips),
            "conflicts": [], "conflictCount": 0, "sourceRows": len(raw),
            "sourceType": "weighment" if detailed else "transporter_enrichment"}



def _as_dict(x):
    return {c.name: (str(getattr(x, c.name)) if isinstance(getattr(x, c.name), (Decimal, date, datetime)) else getattr(x, c.name))
            for c in x.__table__.columns if c.name not in {"image_data"}}


@router.get("/senders")
def senders(request: Request, db: Session = Depends(get_db)):
    _user(request, db)
    _seed_senders(db)
    return [_as_dict(x) for x in db.scalars(select(DespatchSender).order_by(DespatchSender.email))]


class SenderIn(BaseModel):
    email: str = Field(min_length=6, max_length=200)
    name: str = Field(default="", max_length=160)
    active: bool = True


@router.post("/senders")
def set_sender(p: SenderIn, request: Request, db: Session = Depends(get_db)):
    csrf(request)
    user = _user(request, db, edit=True)
    _seed_senders(db)
    email = p.email.strip().lower()
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        raise HTTPException(422, "Invalid sender email address")
    old = db.get(DespatchSender, email)
    before = _as_dict(old) if old else None
    if old:
        old.name = p.name.strip() or old.name
        old.active = p.active
    else:
        old = DespatchSender(email=email, name=p.name.strip() or email, active=p.active)
        db.add(old)
    db.add(DespatchAudit(actor=user.login_id, action="SENDER_CONFIG", entity_key=email,
                         reason="Sender allowlist administration", before_json=json.dumps(before),
                         after_json=json.dumps({"active": p.active, "name": old.name})))
    db.commit()
    return _as_dict(old)


@router.post("/upload")
async def upload_trips(request: Request, file: UploadFile = File(...),
                       db: Session = Depends(get_db)):
    csrf(request)
    user = _user(request, db, edit=True)
    data = await file.read(12_000_001)
    try:
        return _import_data(db, data, file.filename or "upload.xlsx", user.login_id)
    except ValueError as exc:
        raise HTTPException(422, str(exc))


@router.post("/upload-batch")
async def upload_multi_sources(request: Request, files: list[UploadFile] = File(...),
                               db: Session = Depends(get_db)):
    """Handle multiple source files with schema detection, never by filename."""
    csrf(request)
    user = _user(request, db, edit=True)
    if not files or len(files) > 8:
        raise HTTPException(422, "Choose 1-8 XLS, XLSX or CSV sources")
    from app.services.tiom_despatch_source_adapters import parse_despatch_sources
    staged = []
    seen_detailed, seen_aux = {}, {}
    for upload in files:
        filename = upload.filename or "source.xlsx"
        binary = await upload.read(12_000_001)
        if len(binary) > 12_000_000:
            raise HTTPException(422, filename + ": file exceeds 12 MB")
        try:
            records = parse_despatch_sources(binary, filename)
        except ValueError as exc:
            raise HTTPException(422, filename + ": " + str(exc)) from exc
        role = records[0]["_role"]
        for row in records:
            epass = row["epass_no"].upper()
            target = seen_detailed if role == "detailed" else seen_aux
            if epass in target:
                raise HTTPException(422, filename + ": duplicate E-Pass " + epass +
                                    " across uploaded " + role + " sources")
            target[epass] = row
        staged.append((0 if role == "detailed" else 1, filename, binary, role))
    # Validate complementing source values before the first database write.
    for epass, aux in seen_aux.items():
        primary = seen_detailed.get(epass)
        if primary and (
            primary["report_date"] != aux["report_date"] or
            _amount(primary["net_mt"]) != _amount(aux["net_mt"]) or
            primary["vehicle_no"].strip().upper() != aux["vehicle_no"].strip().upper()
        ):
            raise HTTPException(422, "Source reconciliation failed for E-Pass " + epass +
                                ": report date, MT or vehicle differ")
    results = []
    for _, filename, binary, role in sorted(staged, key=lambda item: item[0]):
        info = _import_data(db, binary, filename, user.login_id)
        results.append({"file": filename, "type": role, **info})
        if not info.get("ok"):
            # Existing validated imports are retained; report partial completion
            # explicitly instead of silently claiming all-or-nothing success.
            return {"ok": False, "partial": True, "results": results,
                    "inserted": sum(x.get("inserted", 0) for x in results),
                    "enriched": sum(x.get("enriched", 0) for x in results),
                    "errors": info.get("errors", []), "errorCount": info.get("errorCount", 0)}
    return {"ok": True, "files": len(results), "results": results,
            "inserted": sum(x.get("inserted", 0) for x in results),
            "enriched": sum(x.get("enriched", 0) for x in results),
            "skippedExisting": sum(x.get("skippedExisting", 0) for x in results),
            "conflictCount": 0}


class TripEdit(BaseModel):
    status: str | None = None
    net_mt: Decimal | None = None
    tare_mt: Decimal | None = None
    gross_mt: Decimal | None = None
    report_date: date | None = None
    permit_no: str | None = None
    buyer: str | None = None
    transporter: str | None = None
    destination: str | None = None
    material_grade: str | None = None
    vehicle_no: str | None = None
    reason: str = Field(min_length=5, max_length=500)


@router.post("/trip/{epass_no:path}/edit")
def edit_trip(epass_no: str, p: TripEdit, request: Request, db: Session = Depends(get_db)):
    csrf(request)
    user = _user(request, db, edit=True)
    trip = db.get(DespatchTrip, epass_no)
    if not trip:
        raise HTTPException(404, "E-pass not found")
    before = _as_dict(trip)
    changes = p.model_dump(exclude_none=True, exclude={"reason"})
    if not changes:
        raise HTTPException(422, "No correction supplied")
    for key, value in changes.items():
        if key == "status" and value not in {"ACTIVE", "CANCELLED"}:
            raise HTTPException(422, "Status must be ACTIVE or CANCELLED")
        setattr(trip, key, value)
    if min(trip.tare_mt, trip.gross_mt, trip.net_mt) < 0 or abs(trip.gross_mt - trip.tare_mt - trip.net_mt) > Decimal(".025"):
        db.rollback()
        raise HTTPException(422, "Gross - tare must equal net")
    trip.updated_at = datetime.now(timezone.utc)
    db.add(DespatchAudit(actor=user.login_id, action="CORRECT_TRIP",
                         entity_key=epass_no, reason=p.reason.strip(), before_json=json.dumps(before),
                         after_json=json.dumps(_as_dict(trip))))
    db.commit()
    return {"ok": True, "trip": _as_dict(trip)}


@router.get("/trips")
def trips(request: Request, from_date: date = date(2026, 10, 1),
          to_date: date = date(2026, 10, 31), q: str = "", page: int = 1,
          show_cancelled: bool = False, db: Session = Depends(get_db)):
    _user(request, db)
    if (to_date - from_date).days > 370 or page < 1:
        raise HTTPException(422, "Invalid date range or page")
    query = select(DespatchTrip).where(DespatchTrip.report_date.between(from_date, to_date))
    if not show_cancelled:
        query = query.where(DespatchTrip.status == "ACTIVE")
    if q.strip():
        from sqlalchemy import or_
        term = "%" + q.strip()[:90] + "%"
        query = query.where(or_(DespatchTrip.epass_no.ilike(term),
                                DespatchTrip.vehicle_no.ilike(term),
                                DespatchTrip.permit_no.ilike(term),
                                DespatchTrip.transporter.ilike(term)))
    count = db.scalar(select(func.count()).select_from(query.subquery()))
    result = db.scalars(query.order_by(DespatchTrip.report_date.desc(), DespatchTrip.epass_no)
                        .limit(100).offset((page-1)*100)).all()
    return {"count": count, "page": page, "rows": [_as_dict(r) for r in result]}


@router.get("/dashboard")
def dashboard(request: Request, from_date: date = date(2026, 10, 1),
              to_date: date = date(2026, 10, 31), db: Session = Depends(get_db)):
    _user(request, db)
    if to_date < from_date or (to_date-from_date).days > 370:
        raise HTTPException(422, "Invalid period")
    base = (DespatchTrip.report_date.between(from_date, to_date))
    active = (base & (DespatchTrip.status == "ACTIVE"))
    totals = db.execute(select(func.count(), func.coalesce(func.sum(DespatchTrip.net_mt), 0))
                        .where(active)).one()
    def breakdown(col, top=None):
        stmt = (select(col.label("name"), func.count().label("trips"),
                       func.sum(DespatchTrip.net_mt).label("tonnes"))
                .where(active).group_by(col).order_by(func.sum(DespatchTrip.net_mt).desc()))
        if top:
            stmt = stmt.limit(top)
        return [{"name": k or "(blank)", "trips": int(t), "tonnes": float(mt)}
                for k, t, mt in db.execute(stmt)]
    daily = [{"date": d.isoformat(), "trips": int(t), "tonnes": float(mt)}
             for d, t, mt in db.execute(
                 select(DespatchTrip.report_date, func.count(), func.sum(DespatchTrip.net_mt))
                 .where(active).group_by(DespatchTrip.report_date).order_by(DespatchTrip.report_date)
             )]
    unreviewed = db.scalar(select(func.count()).select_from(DespatchDocument)
                            .where(DespatchDocument.review_status == "NEEDS_REVIEW"))
    cancelled = db.scalar(select(func.count()).select_from(DespatchTrip)
                          .where(base & (DespatchTrip.status == "CANCELLED")))
    return {"period": {"from": from_date.isoformat(), "to": to_date.isoformat()},
            "totals": {"trips": int(totals[0]), "tonnes": float(totals[1]),
                       "cancelled": int(cancelled or 0), "imagesReview": int(unreviewed or 0)},
            "daily": daily, "destination": breakdown(DespatchTrip.destination, 15),
            "material": breakdown(DespatchTrip.material_grade, 15),
            "transporter": breakdown(DespatchTrip.transporter, 15),
            "buyer": breakdown(DespatchTrip.buyer, 15)}


@router.post("/image")
async def upload_image(request: Request, file: UploadFile = File(...),
                       report_date: date | None = Form(default=None),
                       db: Session = Depends(get_db)):
    csrf(request)
    user = _user(request, db, edit=True)
    contents = await file.read(8_000_001)
    ctype = (file.content_type or "").lower()
    if not (ctype in {"image/jpeg", "image/png", "image/webp"} and len(contents) <= 8_000_000):
        raise HTTPException(422, "Upload a JPEG, PNG or WEBP image up to 8 MB")
    digest = hashlib.sha256(contents).hexdigest()
    old = db.scalar(select(DespatchDocument).where(DespatchDocument.sha256 == digest))
    if old:
        return {"ok": True, "id": old.id, "duplicate": True}
    doc = DespatchDocument(id=uuid4().hex, report_date=report_date,
                           filename=(file.filename or "AMNS image")[:255],
                           media_type=ctype, sha256=digest, image_data=contents)
    db.add(doc)
    db.add(DespatchAudit(actor=user.login_id, action="IMAGE_UPLOAD", entity_key=doc.id,
                         reason="Image retained for comparison; needs manual review",
                         before_json=None, after_json=json.dumps({"filename": doc.filename, "date": str(report_date)})))
    db.commit()
    return {"ok": True, "id": doc.id, "reviewStatus": "NEEDS_REVIEW"}


@router.get("/images")
def images(request: Request, db: Session = Depends(get_db)):
    _user(request, db)
    rows = db.scalars(select(DespatchDocument).order_by(DespatchDocument.created_at.desc()).limit(100)).all()
    return [_as_dict(x) for x in rows]


@router.get("/image/{document_id}")
def image(document_id: str, request: Request, db: Session = Depends(get_db)):
    _user(request, db)
    row = db.get(DespatchDocument, document_id)
    if not row:
        raise HTTPException(404, "Image not found")
    return Response(row.image_data, media_type=row.media_type, headers={"Cache-Control": "private, no-store"})


class PermitReview(BaseModel):
    report_date: date
    permit_no: str
    buyer: str = ""
    destination: str = ""
    material: str = ""
    grade_percent: str = ""
    permit_qty: Decimal | None = None
    reported_trips: int | None = None
    reported_mt: Decimal | None = None
    balance_mt: Decimal | None = None
    document_id: str | None = None


@router.post("/permit-review")
def save_permit_review(p: PermitReview, request: Request, db: Session = Depends(get_db)):
    csrf(request)
    user = _user(request, db, edit=True)
    if not p.permit_no.strip():
        raise HTTPException(422, "Permit number required")
    if p.grade_percent and (not re.fullmatch(r"\d{1,2}(?:\.\d{1,3})?%?", p.grade_percent) or float(p.grade_percent.rstrip("%")) > 100):
        raise HTTPException(422, "Grade requires correction; cannot exceed 100%")
    if p.document_id and not db.get(DespatchDocument, p.document_id):
        raise HTTPException(422, "Image reference not found")
    old = db.scalar(select(DespatchPermitDay).where(
        DespatchPermitDay.report_date == p.report_date,
        DespatchPermitDay.permit_no == p.permit_no))
    before = _as_dict(old) if old else None
    if not old:
        old = DespatchPermitDay(report_date=p.report_date, permit_no=p.permit_no)
        db.add(old)
    for key, value in p.model_dump(exclude={"report_date", "permit_no"}).items():
        setattr(old, key, value)
    old.review_status = "CONFIRMED"
    db.add(DespatchAudit(actor=user.login_id, action="AMNS_PERMIT_REVIEW",
                         entity_key=p.permit_no, reason="Image verified entry",
                         before_json=json.dumps(before), after_json=p.model_dump_json()))
    db.commit()
    return {"ok": True, "permit": _as_dict(old)}


@router.get("/permit-reviews")
def permit_reviews(request: Request, from_date: date = date(2026, 10, 1),
                   to_date: date = date(2026, 10, 31), db: Session = Depends(get_db)):
    _user(request, db)
    rows = db.scalars(select(DespatchPermitDay)
                      .where(DespatchPermitDay.report_date.between(from_date, to_date))
                      .order_by(DespatchPermitDay.report_date.desc())).all()
    return [_as_dict(x) for x in rows]


@router.get("/export-check/{report_type}")
def export_check(report_type: str, request: Request,
                 from_date: date, to_date: date,
                 db: Session = Depends(get_db)):
    """Fast report readiness check, before triggering a browser-native download."""
    _user(request, db)
    if report_type not in {"trip", "transporter", "material", "destination", "despatch"}:
        raise HTTPException(404, "Unknown client report")
    if to_date < from_date or (to_date-from_date).days > 370:
        raise HTTPException(422, "Invalid reporting date range")
    if (from_date.year, from_date.month) != (to_date.year, to_date.month):
        raise HTTPException(422, "Choose one calendar month per client report")
    begin=from_date.replace(day=1)
    rows=db.scalars(select(DespatchTrip).where(
        DespatchTrip.report_date.between(begin,to_date),
        DespatchTrip.status=="ACTIVE")).all()
    day_count=sum(x.report_date==to_date for x in rows)
    total_mt=sum((Decimal(str(x.net_mt)) for x in rows),Decimal(0))
    last_date=max((x.report_date for x in rows),default=None)
    if not rows or not day_count:
        last=last_date.isoformat() if last_date else "none"
        return {"ready":False,"activeTrips":len(rows),"activeMt":float(total_mt),
                "reportDate":to_date.isoformat(),
                "message":f"No verified E-Pass trips for {to_date.isoformat()}. Latest available date: {last}. Check MIS Gmail import and the selected report date."}
    from app.services.tiom_despatch_client_reports import _keymap, _material_counts
    errors=[]
    if report_type in {"destination","transporter"}:
        _, missing = _keymap(rows)
        if missing:
            errors.append(f"{len(missing)} unmatched buyer/destination categories (e.g. {missing[0].epass_no})")
    if report_type in {"material","despatch","transporter"}:
        _, missing = _material_counts(rows)
        if missing:
            errors.append(f"{len(missing)} unmatched material grades (e.g. {missing[0].epass_no})")
    if errors:
        return {"ready":False,"activeTrips":len(rows),"activeMt":float(total_mt),
                "reportDate":to_date.isoformat(),
                "message":"Export needs source mapping corrections: "+"; ".join(errors)}
    return {"ready":True,"activeTrips":len(rows),"activeMt":float(total_mt),
            "reportDate":to_date.isoformat(),"lastImportedDay":last_date.isoformat()}


@router.get("/export/{report_type}")
def export_report(report_type: str, request: Request,
                  from_date: date = date(2026, 10, 1),
                  to_date: date = date(2026, 10, 31),
                  db: Session = Depends(get_db)):
    """Original client workbook layouts; cumulative data starts from month day 1."""
    _user(request, db)
    if report_type not in {"trip", "transporter", "material", "destination", "despatch"}:
        raise HTTPException(404, "Unknown report")
    if to_date < from_date or (to_date - from_date).days > 370:
        raise HTTPException(422, "Invalid date range")
    if (from_date.year, from_date.month) != (to_date.year, to_date.month):
        raise HTTPException(422, "Export one calendar month at a time")
    beginning = from_date.replace(day=1)
    records = db.scalars(
        select(DespatchTrip).where(
            DespatchTrip.report_date.between(beginning, to_date),
            DespatchTrip.status.in_(["ACTIVE", "CANCELLED"]),
        ).order_by(DespatchTrip.report_date, DespatchTrip.epass_no)
    ).all()
    from app.services.tiom_despatch_client_reports import build_client_workbook
    try:
        buffer, summary = build_client_workbook(report_type, records, from_date, to_date)
    except ValueError as exc:
        raise HTTPException(409, str(exc)) from exc
    filename = f"AMIPL_{report_type}_{to_date:%Y-%m-%d}_Client_Format.xlsx"
    return StreamingResponse(
        buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"',
                 "X-Despatch-Trip-Count": str(summary["trips"]),
                 "X-Despatch-Quantity-MT": str(summary["quantityMt"])}
    )


def _local_gmail_setup_allowed(request: Request, user) -> bool:
    """Never accept OAuth secrets through an unencrypted office LAN connection."""
    from ipaddress import ip_address
    try:
        host = request.headers.get("host", "").split(":")[0].lower()
        return bool(user.admin and request.client
                    and ip_address(request.client.host).is_loopback
                    and host in {"127.0.0.1", "localhost", "[::1]"})
    except ValueError:
        return False


@router.post("/gmail-connect-local")
async def gmail_connect_local(request: Request, file: UploadFile = File(...),
                              db: Session = Depends(get_db)):
    """Connect PC MIS Gmail via admin-only loopback JSON upload, never cloud/LAN."""
    csrf(request)
    user = _user(request, db, edit=True)
    if not _local_gmail_setup_allowed(request, user):
        raise HTTPException(403, "Gmail OAuth setup is permitted only on the PC at 127.0.0.1.")
    if not (file.filename or "").lower().endswith(".json"):
        raise HTTPException(422, "Select an authorized Gmail OAuth token JSON file.")
    raw = await file.read(65537)
    if len(raw) > 65536:
        raise HTTPException(413, "Gmail OAuth token file is too large.")
    try:
        payload = json.loads(raw.decode("utf-8-sig"))
    except (UnicodeError, ValueError) as exc:
        raise HTTPException(422, "Invalid Gmail OAuth JSON file.") from exc
    if not isinstance(payload, dict) or not all(
        isinstance(payload.get(k), str) and payload[k].strip()
        for k in ("refresh_token", "client_id", "client_secret", "token_uri")
    ):
        raise HTTPException(422, "This must be an authorized OAuth token JSON, not a Google client-secret file.")
    # Never save credentials from an unverified account. Checking Google first
    # prevents an accidental change to the existing WB Gmail authorization.
    try:
        from google.oauth2.credentials import Credentials
        from google.auth.transport.requests import Request as GoogleRequest
        from googleapiclient.discovery import build
        creds = Credentials.from_authorized_user_info(
            payload, scopes=["https://www.googleapis.com/auth/gmail.modify"])
        if not creds.valid:
            creds.refresh(GoogleRequest())
        address = (build("gmail", "v1", credentials=creds, cache_discovery=False)
                   .users().getProfile(userId="me").execute().get("emailAddress") or "").lower().strip()
    except Exception as exc:
        raise HTTPException(422, "Google could not verify this OAuth token. Check authorization and Gmail permissions.") from exc
    if address != "mistiom@nmtpl.com":
        raise HTTPException(422, "The token authorizes a different Gmail account, not mistiom@nmtpl.com.")
    import os
    import tempfile
    from pathlib import Path
    target = Path(__file__).resolve().parents[2] / "config" / "gmail_token.json"
    target.parent.mkdir(parents=True, exist_ok=True)
    temp_name = None
    try:
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=target.parent,
                                         prefix=".gmail_token_", suffix=".tmp",
                                         delete=False) as tmp:
            temp_name = tmp.name
            tmp.write(creds.to_json())
        os.chmod(temp_name, 0o600)
        os.replace(temp_name, target)
    finally:
        if temp_name and os.path.exists(temp_name):
            os.unlink(temp_name)
    return {"ok": True, "mailbox": address, "tokenStored": True,
            "next": "Use Sync Approved Gmail Now to verify the mail importer."}


@router.get("/mail-attachments")
def mail_attachments(request: Request, db: Session = Depends(get_db)):
    """Only current unresolved exceptions—not imported or historical files."""
    _user(request, db)
    from app.services.tiom_despatch_mail import open_despatch_reviews
    records=sorted(open_despatch_reviews(db),
                   key=lambda x: x.received_at or datetime.min, reverse=True)[:16]
    return [{
        "id": x.id, "filename": x.filename, "sender": x.sender,
        "reportDate": x.report_date.isoformat() if x.report_date else None,
        "status": x.status, "reason": x.reason,
        "sizeBytes": len(x.file_data), "receivedAt": x.received_at.isoformat(),
    } for x in records]


@router.get("/mail-attachment/{attachment_id}")
def download_mail_attachment(attachment_id: str, request: Request, db: Session = Depends(get_db)):
    """Download an approved-sender original safely, never execute it on server."""
    _user(request, db)
    from app.services.tiom_despatch_mail import DespatchMailAttachment
    item=db.get(DespatchMailAttachment, attachment_id)
    if not item:
        raise HTTPException(404, "Despatch mail attachment not found")
    name=re.sub(r"[^a-zA-Z0-9._-]", "_", item.filename)[:150] or "despatch-attachment.xls"
    return Response(content=item.file_data, media_type="application/octet-stream", headers={
        "Content-Disposition": f'attachment; filename="{name}"',
        "Cache-Control": "private, no-store",
        "X-Content-Type-Options": "nosniff",
    })


@router.get("/mail-status")
def mail_status(request: Request, db: Session = Depends(get_db)):
    """Expose verified MIS Gmail importer health, never claim sync from configuration alone."""
    user = _user(request, db)
    from app.routers.automation import _gmail_token_raw
    try:
        oauth = json.loads(_gmail_token_raw() or "{}")
        token_ready = isinstance(oauth, dict) and bool(oauth.get("refresh_token"))
    except (ValueError, TypeError):
        token_ready = False
    from app.services.tiom_despatch_mail import DespatchMailPoll, DespatchEmailProcessed
    p=db.get(DespatchMailPoll, "MIS")
    imported=db.scalar(select(func.count()).select_from(DespatchEmailProcessed).where(
        DespatchEmailProcessed.status.like("IMPORTED%"))) or 0
    from app.services.tiom_despatch_mail import open_despatch_reviews
    active_issues = open_despatch_reviews(db)
    review_files = len(active_issues)
    review_emails = len({item.message_id for item in active_issues})
    rows=db.scalar(select(func.count()).select_from(DespatchEmailProcessed)) or 0
    return {
        "configuredMailbox": "mistiom@nmtpl.com",
        "gmailTokenConfigured": token_ready,
        "localSetupAvailable": _local_gmail_setup_allowed(request, user),
        "lastAttempt": p.last_attempt.isoformat() if p and p.last_attempt else None,
        "lastSuccess": p.last_success.isoformat() if p and p.last_success else None,
        "lastError": p.last_error if p else "",
        "lastResult": json.loads(p.last_result or "{}") if p else {},
        "messagesProcessed":rows,
        "messagesWithDespatchData":imported,
        "emailsNeedingReview":review_emails,
        "attachmentsNeedingReview":review_files,
        "activeReviewFrom":"2026-10-27",
        "automaticPolling":"enabled_if_oauth_token_available",
        "imageRule":"Emailed images are collected independently; permit figures require verification before reporting",
    }


class DespatchReviewReconcileIn(BaseModel):
    after: str = Field(default="", max_length=128)
    force: bool = False


@router.post("/mail-review-reconcile")
def mail_review_reconcile(p: DespatchReviewReconcileIn, request: Request,
                          db: Session = Depends(get_db)):
    """Re-check archived approved-sender Excel sources against the live ledger.

    Only an authenticated editor may initiate; the source files and current
    E-Pass records are never erased or overwritten.
    """
    csrf(request)
    _user(request, db, edit=True)
    from app.services.tiom_despatch_mail import recheck_archived_despatch
    return recheck_archived_despatch(limit_messages=8, after=p.after, force=p.force)


@router.post("/mail-sync")
def mail_sync(request: Request, db: Session = Depends(get_db)):
    """Management-triggered mailbox sync, separate from WB ingestion."""
    csrf(request)
    _user(request, db, edit=True)
    from app.services.tiom_despatch_mail import gmail_scan
    try:
        return gmail_scan(50)
    except RuntimeError as exc:
        raise HTTPException(503, str(exc)) from exc
