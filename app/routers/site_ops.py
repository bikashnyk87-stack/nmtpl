from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from hashlib import sha256
from uuid import uuid4
import json
from urllib.request import Request as UrlRequest, urlopen
from urllib.error import URLError, HTTPError

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import RedirectResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session

from app.auth import WebUser, csrf, get_user
from app.db import get_db
from app.models import (
    Equipment, Person, LoadTrip, WbMovement, LoadWbMatch, PersonAttendance,
    EquipmentAttendance, ShiftDeployment, Location, Product, ActivityMaster,
)
from app.site_auth import accessible_sites, has_permission, require_permission, require_site
from app.site_models import (
    EquipmentSiteAssignment,
    PersonSiteAssignment,
    Site,
    SiteAssetAttendance,
    SiteAssetMeter,
    SiteAuditLog,
    SiteBillingReconciliation,
    SiteDataQualityIssue,
    SiteDeployment,
    SiteHsdTransaction,
    SiteOperationalImportBatch,
    SiteOperationalImportRow,
    SiteLocation,
    SiteMaterial,
    SitePersonAttendance,
    SiteSatelliteObservation,
    SiteSurveyMeasurement,
    SiteShift,
    SiteTrip,
    SiteTripReconciliation,
    SiteWbImportBatch,
    SiteWbMovement,
    SiteWeightFactor,
    TiomMisReport,
    TiomMisTripRow,
    TiomMisReconciliation,
)
from app.services.site_context import resolve_site_context
from app.site_schemas import (
    SiteAssetAttendanceIn,
    SiteAssetMeterIn,
    SiteAttendanceIn,
    SiteDataQualityResolveIn,
    SiteDecisionIn,
    SiteDeploymentIn,
    SiteMapConfigIn,
    SiteSatelliteObservationIn,
    SiteSurveyMeasurementIn,
    TiomMisReportIn,
)

router = APIRouter(prefix="/api/site-ops", tags=["site-operations"])


def _audit(db: Session, user, site_id: str | None, action: str, entity: str, entity_id: str | None, before=None, after=None, reason=None):
    db.add(SiteAuditLog(
        site_id=site_id,
        actor=user.login_id,
        action=action,
        entity=entity,
        entity_id=entity_id,
        before_json=json.dumps(before, default=str) if before is not None else None,
        after_json=json.dumps(after, default=str) if after is not None else None,
        reason=reason,
    ))


def _site_shift_or_422(db: Session, site_id: str, shift: str):
    code=str(shift or "").strip().upper()
    row=db.get(SiteShift, {"site_id":site_id,"shift":code}) if code else None
    if not row or not row.active:
        raise HTTPException(422, f"Choose a valid active shift for {site_id}.")
    return row


def _site_equipment_or_422(db: Session, site_id: str, machine_id: str, operating_day: date, label: str="equipment"):
    eq=db.get(Equipment,machine_id) if machine_id else None
    assignment=db.scalar(select(EquipmentSiteAssignment).where(
        EquipmentSiteAssignment.site_id==site_id,
        EquipmentSiteAssignment.machine_id==machine_id,
        EquipmentSiteAssignment.active.is_(True),
        EquipmentSiteAssignment.effective_from<=operating_day,
        or_(EquipmentSiteAssignment.effective_to.is_(None),EquipmentSiteAssignment.effective_to>=operating_day),
    ).order_by(EquipmentSiteAssignment.effective_from.desc(),EquipmentSiteAssignment.id.desc()).limit(1)) if machine_id else None
    if not eq or not eq.active or not assignment:
        raise HTTPException(422, f"Choose a valid active {label} assigned to {site_id}.")
    return eq


def _site_person_or_422(db: Session, site_id: str, employee_id: str, operating_day: date):
    person=db.get(Person,employee_id) if employee_id else None
    assignment=db.scalar(select(PersonSiteAssignment).where(
        PersonSiteAssignment.site_id==site_id,
        PersonSiteAssignment.employee_id==employee_id,
        PersonSiteAssignment.active.is_(True),
        PersonSiteAssignment.effective_from<=operating_day,
        or_(PersonSiteAssignment.effective_to.is_(None),PersonSiteAssignment.effective_to>=operating_day),
    ).order_by(PersonSiteAssignment.effective_from.desc(),PersonSiteAssignment.id.desc()).limit(1)) if employee_id else None
    if not person or not person.active or not assignment:
        raise HTTPException(422, f"Choose a valid active employee assigned to {site_id}.")
    return person


def _site_location_or_422(db: Session, site_id: str, location_id: str):
    row=db.get(SiteLocation,location_id) if location_id else None
    if not row or row.site_id!=site_id or not row.active:
        raise HTTPException(422, f"Choose a valid active location for {site_id}.")
    return row


def _decimal(value) -> Decimal:
    if value is None:
        return Decimal("0")
    return Decimal(str(value))


def _fmt_decimal(value):
    return float(_decimal(value))


def _require_any(db, user, site_id: str, module: str, actions=("VIEW", "CREATE", "EDIT", "APPROVE")):
    if user.admin:
        return
    if not any(has_permission(db, user, site_id, module, a) for a in actions):
        raise HTTPException(403, f"{module} permission is required for {site_id}.")


def _period_bounds(from_date: date | None, to_date: date | None):
    today = date.today()
    start = from_date or today.replace(day=1)
    end = to_date or today
    if end < start:
        raise HTTPException(422, "to_date cannot be before from_date")
    return start, end




# ---------------------------------------------------------------------------
# RC2 SIMPLE FIELD UX
# Read-only operational lookups are intentionally separate from Master Control.
# Site users can select canonical records without receiving master-edit rights.
# TIOM falls back to the established legacy masters during the transition.
# ---------------------------------------------------------------------------

def _lookup_limit(value: int) -> int:
    return max(1, min(int(value or 30), 100))


def _contains_text(*values, q: str) -> bool:
    if not q:
        return True
    needle = q.strip().lower()
    return any(needle in str(v or '').lower() for v in values)


@router.get("/{site_id}/lookups")
def operational_lookup(site_id: str, request: Request, kind: str, q: str = "", limit: int = 30, db: Session = Depends(get_db)):
    user = get_user(db, request)
    site_id = require_site(db, user, site_id)
    kind = str(kind or '').strip().upper()
    limit = _lookup_limit(limit)
    q = str(q or '').strip()
    out = []

    if kind == "EMPLOYEE":
        if site_id == "TIOM":
            rows = list(db.scalars(select(Person).where(Person.active.is_(True)).order_by(Person.name).limit(2000)))
            for x in rows:
                if _contains_text(x.employee_id, x.name, x.role, x.department, q=q):
                    out.append({"id": x.employee_id, "label": x.name, "sub": " · ".join(v for v in [x.employee_id, x.role] if v), "type": x.role})
                    if len(out) >= limit: break
        else:
            rows = db.execute(select(Person, PersonSiteAssignment).join(PersonSiteAssignment, PersonSiteAssignment.employee_id == Person.employee_id).where(PersonSiteAssignment.site_id == site_id, Person.active.is_(True), PersonSiteAssignment.active.is_(True)).order_by(Person.name).limit(2000)).all()
            for x, a in rows:
                if _contains_text(x.employee_id, x.name, x.role, x.department, q=q):
                    out.append({"id": x.employee_id, "label": x.name, "sub": " · ".join(v for v in [x.employee_id, x.role] if v), "type": x.role})
                    if len(out) >= limit: break
        return out

    if kind == "ASSET":
        if site_id == "TIOM":
            rows = list(db.scalars(select(Equipment).where(Equipment.active.is_(True)).order_by(Equipment.machine_id).limit(2000)))
            pairs = [(x, None) for x in rows]
        else:
            pairs = db.execute(select(Equipment, EquipmentSiteAssignment).join(EquipmentSiteAssignment, EquipmentSiteAssignment.machine_id == Equipment.machine_id).where(EquipmentSiteAssignment.site_id == site_id, Equipment.active.is_(True), EquipmentSiteAssignment.active.is_(True)).order_by(Equipment.machine_id).limit(2000)).all()
        for x, a in pairs:
            if _contains_text(x.machine_id, x.vehicle_no, x.door_no, x.type, x.group, x.make_model, q=q):
                primary = x.vehicle_no or x.door_no or x.machine_id
                out.append({"id": x.machine_id, "label": primary, "sub": " · ".join(v for v in [x.machine_id, x.type, x.group] if v), "type": x.type, "group": x.group, "vehicleNo": x.vehicle_no, "doorNo": x.door_no})
                if len(out) >= limit: break
        return out

    if kind == "LOCATION":
        if site_id == "TIOM":
            rows = list(db.scalars(select(Location).where(Location.active.is_(True)).order_by(Location.location_name).limit(1000)))
            for x in rows:
                if _contains_text(x.location_id, x.location_name, x.location_type, q=q):
                    out.append({"id": x.location_id, "label": x.location_name, "sub": " · ".join(v for v in [x.location_id, x.location_type] if v), "type": x.location_type})
                    if len(out) >= limit: break
        else:
            rows = list(db.scalars(select(SiteLocation).where(SiteLocation.site_id == site_id, SiteLocation.active.is_(True)).order_by(SiteLocation.name).limit(1000)))
            for x in rows:
                if _contains_text(x.site_location_id, x.code, x.name, x.location_type, q=q):
                    out.append({"id": x.site_location_id, "label": x.name, "sub": " · ".join(v for v in [x.code, x.location_type] if v), "type": x.location_type})
                    if len(out) >= limit: break
        return out

    if kind == "MATERIAL":
        if site_id == "TIOM":
            rows = list(db.scalars(select(Product).where(Product.active.is_(True)).order_by(Product.name).limit(1000)))
            for x in rows:
                if _contains_text(x.product_id, x.name, x.size_spec, q=q):
                    out.append({"id": x.product_id, "label": x.name, "sub": x.product_id, "type": "MATERIAL"})
                    if len(out) >= limit: break
        else:
            rows = list(db.scalars(select(SiteMaterial).where(SiteMaterial.site_id == site_id, SiteMaterial.active.is_(True)).order_by(SiteMaterial.name).limit(1000)))
            for x in rows:
                if _contains_text(x.site_material_id, x.code, x.name, x.material_group, q=q):
                    unit = x.default_unit or x.billable_unit or ''
                    out.append({"id": x.site_material_id, "label": x.name, "sub": " · ".join(v for v in [x.code, unit] if v), "type": x.material_group or "MATERIAL", "unit": unit})
                    if len(out) >= limit: break
        return out

    if kind == "ACTIVITY":
        rows = list(db.scalars(select(ActivityMaster).where(ActivityMaster.active.is_(True)).order_by(ActivityMaster.activity).limit(200)))
        for x in rows:
            if _contains_text(x.activity, q=q):
                out.append({"id": x.activity, "label": x.activity.replace('_', ' ').title(), "sub": "Vehicle required" if x.vehicle_required else "", "type": "ACTIVITY"})
                if len(out) >= limit: break
        return out

    raise HTTPException(422, "kind must be EMPLOYEE, ASSET, LOCATION, MATERIAL or ACTIVITY")


def _parse_simple_dt(raw, field_name: str):
    if not raw:
        raise HTTPException(422, f"{field_name} is required")
    if isinstance(raw, datetime):
        return raw
    try:
        value = str(raw).strip().replace('Z', '+00:00')
        return datetime.fromisoformat(value)
    except Exception:
        raise HTTPException(422, f"Invalid {field_name}")


@router.post("/{site_id}/trips/simple")
def save_simple_site_trip(site_id: str, payload: dict, request: Request, db: Session = Depends(get_db)):
    csrf(request)
    user = get_user(db, request)
    site_id = require_site(db, user, site_id)
    if site_id not in {"SOCP", "KOCP"}:
        raise HTTPException(409, "Simple site trip entry is configured for SOCP/KOCP")
    kind = str(payload.get("kind") or "TRIP").strip().upper()
    needed_module = "OB" if site_id == "KOCP" and kind == "OB" else "TRIP"
    require_permission(db, user, site_id, needed_module, "CREATE")
    try:
        operating_day = date.fromisoformat(str(payload.get("operatingDate") or ''))
    except Exception:
        raise HTTPException(422, "Operating date is required")
    shift = str(payload.get("shift") or '').strip().upper()
    shift_row = db.get(SiteShift, {"site_id": site_id, "shift": shift})
    if not shift_row or not shift_row.active:
        raise HTTPException(422, "Choose a valid shift")
    request_id = str(payload.get("requestId") or '').strip()
    if not request_id:
        raise HTTPException(422, "Request ID is required")
    existing = db.scalar(select(SiteTrip).where(SiteTrip.site_id == site_id, SiteTrip.source_type == "PORTAL_SIMPLE", SiteTrip.source_record_uid == request_id))
    if existing:
        same = (
            existing.operating_date == operating_day and existing.shift == shift and
            (existing.vehicle_id or None) == (str(payload.get("vehicleId") or '').strip() or None) and
            (existing.loading_equipment_id or None) == (str(payload.get("loadingEquipmentId") or '').strip() or None) and
            (existing.destination_location_id or None) == (str(payload.get("destinationLocationId") or '').strip() or None) and
            (existing.material_id or None) == (str(payload.get("materialId") or '').strip() or None) and
            (existing.gp_no or None) == (str(payload.get("gpNo") or '').strip() or None)
        )
        qmt = payload.get("quantityMt"); qcum = payload.get("quantityCum")
        if qmt not in (None, ''): same = same and existing.quantity_mt == Decimal(str(qmt))
        if qcum not in (None, ''): same = same and existing.quantity_cum == Decimal(str(qcum))
        if not same:
            raise HTTPException(409, "This retry ID belongs to a different trip. The new row was NOT saved; refresh and save again.")
        return {"ok": True, "tripId": existing.trip_id, "idempotent": True}

    vehicle = str(payload.get("vehicleId") or '').strip() or None
    loader = str(payload.get("loadingEquipmentId") or '').strip() or None
    source = str(payload.get("sourceLocationId") or '').strip() or None
    destination = str(payload.get("destinationLocationId") or '').strip() or None
    material = str(payload.get("materialId") or '').strip() or None
    gp_no = str(payload.get("gpNo") or '').strip() or None
    def assigned_equipment(machine_id, label):
        if not machine_id:
            return None
        eq=db.get(Equipment,machine_id)
        assignment=db.scalar(select(EquipmentSiteAssignment).where(
            EquipmentSiteAssignment.site_id==site_id,
            EquipmentSiteAssignment.machine_id==machine_id,
            EquipmentSiteAssignment.active.is_(True),
        ).order_by(EquipmentSiteAssignment.effective_from.desc(),EquipmentSiteAssignment.id.desc()).limit(1))
        if not eq or not eq.active or not assignment:
            raise HTTPException(422, f"Choose a valid active {label} assigned to {site_id}")
        return eq
    if vehicle: assigned_equipment(vehicle,"vehicle")
    if loader: assigned_equipment(loader,"loader/excavator")
    for loc in (source, destination):
        if loc:
            x = db.get(SiteLocation, loc)
            if not x or x.site_id != site_id or not x.active: raise HTTPException(422, "Choose a valid site location")
    if material:
        x = db.get(SiteMaterial, material)
        if not x or x.site_id != site_id or not x.active: raise HTTPException(422, "Choose a valid material")

    mt = payload.get("quantityMt")
    cum = payload.get("quantityCum")
    mt = Decimal(str(mt)) if mt not in (None, '') else None
    cum = Decimal(str(cum)) if cum not in (None, '') else None
    factor = None
    basis = "DIRECT_ENTRY"
    if site_id == "KOCP":
        factor_type = "OB_CUM_PER_TRIP" if kind == "OB" else "COAL_AVG_MT_PER_TRIP"
        stmt = select(SiteWeightFactor).where(
            SiteWeightFactor.site_id == site_id, SiteWeightFactor.factor_type == factor_type,
            SiteWeightFactor.status == "APPROVED", SiteWeightFactor.effective_from <= operating_day,
            or_(SiteWeightFactor.effective_to.is_(None), SiteWeightFactor.effective_to >= operating_day),
            or_(SiteWeightFactor.operating_date.is_(None), SiteWeightFactor.operating_date == operating_day),
            or_(SiteWeightFactor.shift.is_(None), SiteWeightFactor.shift == shift),
        ).order_by(SiteWeightFactor.operating_date.desc(), SiteWeightFactor.effective_from.desc(), SiteWeightFactor.version.desc())
        factor = db.scalars(stmt).first()
        if not factor:
            friendly = "OB CuM/trip" if kind == "OB" else "Coal average MT/trip"
            raise HTTPException(409, f"No approved MCL {friendly} rule is configured for {operating_day} Shift {shift}. Ask the supervisor to set it first.")
        if kind == "OB":
            cum = factor.factor_value; basis = "MCL_DUMPER_FACTOR"
        else:
            mt = factor.factor_value; basis = "MCL_SHIFT_AVERAGE"

    if site_id == "SOCP" and mt is None:
        raise HTTPException(422, "Quantity MT is required")
    event_at = datetime.combine(operating_day, shift_row.start_time)
    trip_id = f"TRIP-{site_id}-{uuid4().hex[:14].upper()}"
    row = SiteTrip(
        trip_id=trip_id, site_id=site_id, operating_date=operating_day, shift=shift, event_at=event_at,
        vehicle_id=vehicle, loading_equipment_id=loader, source_location_id=source, destination_location_id=destination,
        material_id=material, gp_no=gp_no, quantity_mt=mt, quantity_cum=cum, weight_basis=basis,
        factor_id=factor.factor_id if factor else None, source_type="PORTAL_SIMPLE", source_record_uid=request_id,
        status="POSTED", entered_by=user.login_id,
    )
    db.add(row)
    _audit(db, user, site_id, "CREATE_SIMPLE_TRIP", "site_trip", trip_id, after={
        "kind": kind, "operatingDate": str(operating_day), "shift": shift, "vehicle": vehicle, "loader": loader,
        "destination": destination, "material": material, "quantityMt": str(mt) if mt is not None else None,
        "quantityCum": str(cum) if cum is not None else None, "weightBasis": basis,
    })
    db.commit()
    return {"ok": True, "tripId": trip_id, "operatingDate": operating_day, "shift": shift, "quantityMt": mt, "quantityCum": cum, "weightBasis": basis}


@router.get("/TIOM/production/simple")
def list_tiom_simple_production(request: Request, operating_date: date | None = None, shift: str | None = None, limit: int = 100, db: Session = Depends(get_db)):
    user = get_user(db, request)
    require_site(db, user, "TIOM")
    if not user.admin and not (has_permission(db, user, "TIOM", "PRODUCTION", "VIEW") or has_permission(db, user, "TIOM", "FIELD_ENTRY", "VIEW")):
        raise HTTPException(403, "TIOM production view permission is required")
    stmt = select(LoadTrip)
    if operating_date:
        stmt = stmt.where(LoadTrip.operating_date == operating_date)
    if shift:
        stmt = stmt.where(LoadTrip.shift == shift.upper())
    rows = list(db.scalars(stmt.order_by(LoadTrip.loading_start_at.desc()).limit(max(1, min(int(limit), 300)))))
    return [{
        "tripId": r.trip_id, "operatingDate": r.operating_date, "shift": r.shift,
        "sourceLocationId": r.source_location_id, "destinationLocationId": r.destination_location_id,
        "activity": r.activity, "machineId": r.machine_id, "vehicleId": r.vehicle_id,
        "materialId": r.material_id, "loadingAt": r.loading_start_at, "unloadingAt": r.unload_at,
        "status": r.status, "tripSeq": r.trip_seq,
    } for r in rows]


@router.post("/TIOM/production/simple")
def save_tiom_simple_production(payload: dict, request: Request, db: Session = Depends(get_db)):
    csrf(request)
    user = get_user(db, request)
    require_site(db, user, "TIOM")
    if not user.admin and not (has_permission(db, user, "TIOM", "PRODUCTION", "CREATE") or has_permission(db, user, "TIOM", "FIELD_ENTRY", "CREATE")):
        raise HTTPException(403, "TIOM production create permission is required")

    try:
        operating_day = date.fromisoformat(str(payload.get("operatingDate") or ''))
    except Exception:
        raise HTTPException(422, "Operating date is required")
    shift = str(payload.get("shift") or '').strip().upper()
    if shift not in {"A", "B", "C"}:
        raise HTTPException(422, "Choose Shift A, B or C")
    loading_at = _parse_simple_dt(payload.get("loadingAt"), "Loading time")
    unloading_at = _parse_simple_dt(payload.get("unloadingAt"), "Unloading time")
    if unloading_at < loading_at:
        raise HTTPException(422, "Unloading time cannot be before loading time")

    source = str(payload.get("sourceLocationId") or '').strip()
    destination = str(payload.get("destinationLocationId") or '').strip()
    machine = str(payload.get("machineId") or '').strip()
    vehicle = str(payload.get("vehicleId") or '').strip()
    material = str(payload.get("materialId") or '').strip()
    activity = str(payload.get("activity") or 'LOADING').strip().upper()
    request_id = str(payload.get("requestId") or '').strip()
    if not request_id:
        raise HTTPException(422, "Request ID is required")

    existing = db.scalar(select(LoadTrip).where(LoadTrip.request_id == request_id))
    if existing:
        same = (
            existing.operating_date == operating_day and existing.shift == shift and
            existing.source_location_id == source and (existing.destination_location_id or '') == destination and
            existing.machine_id == machine and (existing.vehicle_id or '') == vehicle and
            (existing.material_id or '') == material and existing.activity == activity and
            existing.loading_start_at == loading_at and existing.unload_at == unloading_at
        )
        if not same:
            raise HTTPException(409, "This retry ID belongs to a different production entry. The new entry was NOT saved; refresh and save again.")
        return {"ok": True, "tripId": existing.trip_id, "idempotent": True, "status": existing.status}

    loc_source = db.get(Location, source) if source else None
    loc_dest = db.get(Location, destination) if destination else None
    loader = db.get(Equipment, machine) if machine else None
    truck = db.get(Equipment, vehicle) if vehicle else None
    product = db.get(Product, material) if material else None
    activity_row = db.get(ActivityMaster, activity)
    if not loc_source or not loc_source.active: raise HTTPException(422, "Choose a valid source location")
    if destination and (not loc_dest or not loc_dest.active): raise HTTPException(422, "Choose a valid destination")
    if not loader or not loader.active: raise HTTPException(422, "Choose a valid loader/excavator")
    if vehicle and (not truck or not truck.active): raise HTTPException(422, "Choose a valid vehicle")
    if material and (not product or not product.active): raise HTTPException(422, "Choose a valid material")
    if not activity_row or not activity_row.active: raise HTTPException(422, "Choose a valid activity")
    if activity_row.vehicle_required and not vehicle: raise HTTPException(422, "Vehicle is required for this activity")

    seq = None
    if vehicle:
        seq = (db.scalar(select(func.count()).select_from(LoadTrip).where(LoadTrip.operating_date == operating_day, LoadTrip.shift == shift, LoadTrip.vehicle_id == vehicle)) or 0) + 1
    now = datetime.now(timezone.utc)
    trip_id = f"TIOM-{uuid4().hex[:18].upper()}"
    row = LoadTrip(
        trip_id=trip_id, operating_date=operating_day, shift=shift, source_location_id=source,
        destination_location_id=destination or None, activity=activity, machine_id=machine,
        machine_operator_id=None, vehicle_id=vehicle or None, vehicle_driver_id=None, trip_seq=seq,
        material_id=material or None, loading_start_at=loading_at, loading_end_at=loading_at,
        unload_at=unloading_at, status="CLOSED", request_id=request_id,
        created_by=user.login_id, created_at=now, updated_at=now,
    )
    db.add(row)
    _audit(db, user, "TIOM", "CREATE_SIMPLE_PRODUCTION", "load_trip", trip_id, after={
        "operatingDate": str(operating_day), "shift": shift, "machine": machine, "vehicle": vehicle or None,
        "source": source, "destination": destination or None, "material": material or None,
        "loadingAt": str(loading_at), "unloadingAt": str(unloading_at),
        "attendanceDependency": False, "deploymentDependency": False,
    })
    db.commit()
    return {"ok": True, "tripId": trip_id, "operatingDate": operating_day, "shift": shift, "status": row.status}


@router.get("/central/overview")
def central_overview(request: Request, operating_date: date | None = None, db: Session = Depends(get_db)):
    user = get_user(db, request)
    day = operating_date or date.today()
    sites = []
    totals = {"trips": 0, "quantityMt": 0.0, "quantityCum": 0.0, "hsdIssuedL": 0.0, "wbRows": 0, "present": 0, "openIssues": 0}
    for site_id in sorted(accessible_sites(db, user)):
        if site_id == "TIOM":
            trips = db.scalar(select(func.count()).select_from(LoadTrip).where(LoadTrip.operating_date == day)) or 0
            wb_rows = db.scalar(select(func.count()).select_from(WbMovement).where(WbMovement.operating_date == day)) or 0
            net_kg = db.scalar(select(func.coalesce(func.sum(WbMovement.net_kg), 0)).where(WbMovement.operating_date == day)) or 0
            present = db.scalar(select(func.count()).select_from(PersonAttendance).where(PersonAttendance.operating_date == day, PersonAttendance.status == "PRESENT")) or 0
            item = {"siteId": site_id, "trips": trips, "quantityMt": float(_decimal(net_kg) / Decimal("1000")), "quantityCum": 0.0, "hsdIssuedL": 0.0, "wbRows": wb_rows, "present": present, "openIssues": 0}
        else:
            trips = db.scalar(select(func.count()).select_from(SiteTrip).where(SiteTrip.site_id == site_id, SiteTrip.operating_date == day, SiteTrip.status != "VOID")) or 0
            qty_mt = db.scalar(select(func.coalesce(func.sum(SiteTrip.quantity_mt), 0)).where(SiteTrip.site_id == site_id, SiteTrip.operating_date == day, SiteTrip.status != "VOID")) or 0
            qty_cum = db.scalar(select(func.coalesce(func.sum(SiteTrip.quantity_cum), 0)).where(SiteTrip.site_id == site_id, SiteTrip.operating_date == day, SiteTrip.status != "VOID")) or 0
            hsd = db.scalar(select(func.coalesce(func.sum(SiteHsdTransaction.litres), 0)).where(SiteHsdTransaction.site_id == site_id, SiteHsdTransaction.operating_date == day, SiteHsdTransaction.transaction_type == "ISSUE", SiteHsdTransaction.status == "POSTED")) or 0
            wb_rows = db.scalar(select(func.count()).select_from(SiteWbMovement).where(SiteWbMovement.site_id == site_id, SiteWbMovement.operating_date == day, SiteWbMovement.row_status == "VALID")) or 0
            present = db.scalar(select(func.count()).select_from(SitePersonAttendance).where(SitePersonAttendance.site_id == site_id, SitePersonAttendance.operating_date == day, SitePersonAttendance.status == "PRESENT")) or 0
            issues = db.scalar(select(func.count()).select_from(SiteDataQualityIssue).where(SiteDataQualityIssue.site_id == site_id, SiteDataQualityIssue.status == "OPEN")) or 0
            item = {"siteId": site_id, "trips": trips, "quantityMt": _fmt_decimal(qty_mt), "quantityCum": _fmt_decimal(qty_cum), "hsdIssuedL": _fmt_decimal(hsd), "wbRows": wb_rows, "present": present, "openIssues": issues}
        sites.append(item)
        for k in totals:
            totals[k] += item[k]
    return {"date": day, "totals": totals, "sites": sites}


@router.get("/{site_id}/attendance")
def list_attendance(site_id: str, request: Request, operating_date: date | None = None, shift: str | None = None, db: Session = Depends(get_db)):
    user = get_user(db, request); site_id = require_site(db, user, site_id); _require_any(db, user, site_id, "ATTENDANCE")
    if site_id == "TIOM":
        stmt = select(PersonAttendance).where(PersonAttendance.operating_date == (operating_date or date.today()))
        if shift: stmt = stmt.where(PersonAttendance.shift == shift)
        rows = db.scalars(stmt.order_by(PersonAttendance.shift, PersonAttendance.employee_id).limit(5000)).all()
        return [{"id": r.id, "operatingDate": r.operating_date, "shift": r.shift, "employeeId": r.employee_id, "status": r.status, "inAt": r.in_at, "outAt": r.out_at, "workedHours": r.worked_hours, "remarks": r.remarks, "sourceType": "LEGACY_TIOM"} for r in rows]
    stmt = select(SitePersonAttendance).where(SitePersonAttendance.site_id == site_id)
    if operating_date: stmt = stmt.where(SitePersonAttendance.operating_date == operating_date)
    if shift: stmt = stmt.where(SitePersonAttendance.shift == shift)
    rows = db.scalars(stmt.order_by(SitePersonAttendance.operating_date.desc(), SitePersonAttendance.shift, SitePersonAttendance.employee_id).limit(5000)).all()
    return [{"id": r.id, "operatingDate": r.operating_date, "shift": r.shift, "employeeId": r.employee_id, "status": r.status, "inAt": r.in_at, "outAt": r.out_at, "workedHours": r.worked_hours, "remarks": r.remarks, "sourceType": r.source_type} for r in rows]


@router.post("/{site_id}/attendance")
def save_attendance(site_id: str, p: SiteAttendanceIn, request: Request, db: Session = Depends(get_db)):
    csrf(request); user = get_user(db, request); site_id = require_site(db, user, site_id); require_permission(db, user, site_id, "ATTENDANCE", "CREATE")
    if site_id == "TIOM": raise HTTPException(409, "Use the existing TIOM attendance workflow during migration.")
    _site_shift_or_422(db,site_id,p.shift)
    _site_person_or_422(db,site_id,p.employeeId,p.operatingDate)
    row = db.scalar(select(SitePersonAttendance).where(SitePersonAttendance.site_id == site_id, SitePersonAttendance.operating_date == p.operatingDate, SitePersonAttendance.shift == p.shift.upper(), SitePersonAttendance.employee_id == p.employeeId))
    created = row is None
    if row is None:
        row = SitePersonAttendance(site_id=site_id, operating_date=p.operatingDate, shift=p.shift.upper(), employee_id=p.employeeId, entered_by=user.login_id)
        db.add(row)
    before = None if created else {"status": row.status, "inAt": row.in_at, "outAt": row.out_at}
    row.status=p.status.upper(); row.in_at=p.inAt; row.out_at=p.outAt; row.worked_hours=p.workedHours; row.source_type=p.sourceType; row.source_record_uid=p.sourceRecordUid; row.remarks=p.remarks; row.updated_at=datetime.now(timezone.utc)
    _audit(db,user,site_id,"CREATE_ATTENDANCE" if created else "UPDATE_ATTENDANCE","site_person_attendance",f"{p.operatingDate}:{p.shift}:{p.employeeId}",before=before,after=p.model_dump())
    db.commit(); return {"ok": True, "created": created}


@router.get("/{site_id}/asset-attendance")
def list_asset_attendance(site_id: str, request: Request, operating_date: date | None = None, shift: str | None = None, db: Session = Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); _require_any(db,user,site_id,"FLEET")
    if site_id == "TIOM":
        stmt=select(EquipmentAttendance).where(EquipmentAttendance.operating_date==(operating_date or date.today()))
        if shift: stmt=stmt.where(EquipmentAttendance.shift==shift)
        rows=db.scalars(stmt.order_by(EquipmentAttendance.shift,EquipmentAttendance.machine_id).limit(5000)).all()
        return [{"id":r.id,"operatingDate":r.operating_date,"shift":r.shift,"assetId":r.machine_id,"status":r.status,"condition":r.condition,"remarks":r.remarks,"sourceType":"LEGACY_TIOM"} for r in rows]
    stmt=select(SiteAssetAttendance).where(SiteAssetAttendance.site_id==site_id)
    if operating_date: stmt=stmt.where(SiteAssetAttendance.operating_date==operating_date)
    if shift: stmt=stmt.where(SiteAssetAttendance.shift==shift)
    rows=db.scalars(stmt.order_by(SiteAssetAttendance.operating_date.desc(),SiteAssetAttendance.shift,SiteAssetAttendance.asset_id).limit(5000)).all()
    return [{"id":r.id,"operatingDate":r.operating_date,"shift":r.shift,"assetId":r.asset_id,"status":r.status,"condition":r.condition,"remarks":r.remarks} for r in rows]


@router.post("/{site_id}/asset-attendance")
def save_asset_attendance(site_id: str, p: SiteAssetAttendanceIn, request: Request, db: Session = Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id); require_permission(db,user,site_id,"FLEET","CREATE")
    if site_id=="TIOM": raise HTTPException(409,"Use the existing TIOM equipment attendance workflow during migration.")
    _site_shift_or_422(db,site_id,p.shift)
    _site_equipment_or_422(db,site_id,p.assetId,p.operatingDate,"equipment")
    row=db.scalar(select(SiteAssetAttendance).where(SiteAssetAttendance.site_id==site_id,SiteAssetAttendance.operating_date==p.operatingDate,SiteAssetAttendance.shift==p.shift.upper(),SiteAssetAttendance.asset_id==p.assetId))
    created=row is None
    if not row:
        row=SiteAssetAttendance(site_id=site_id,operating_date=p.operatingDate,shift=p.shift.upper(),asset_id=p.assetId,entered_by=user.login_id); db.add(row)
    row.status=p.status.upper(); row.condition=p.condition; row.remarks=p.remarks; row.updated_at=datetime.now(timezone.utc)
    _audit(db,user,site_id,"UPSERT_ASSET_ATTENDANCE","site_asset_attendance",f"{p.operatingDate}:{p.shift}:{p.assetId}",after=p.model_dump()); db.commit(); return {"ok":True,"created":created}


@router.get("/{site_id}/deployments")
def list_deployments(site_id:str,request:Request,operating_date:date|None=None,shift:str|None=None,db:Session=Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); _require_any(db,user,site_id,"SHIFT_CONTROL")
    if site_id=="TIOM":
        stmt=select(ShiftDeployment).where(ShiftDeployment.operating_date==(operating_date or date.today()))
        if shift: stmt=stmt.where(ShiftDeployment.shift==shift)
        rows=db.scalars(stmt.order_by(ShiftDeployment.shift,ShiftDeployment.machine_id).limit(5000)).all()
        return [{"deploymentId":str(r.id),"operatingDate":r.operating_date,"shift":r.shift,"assetId":r.machine_id,"employeeId":None,"locationId":r.location_id,"activity":None,"fromAt":r.from_at,"toAt":r.to_at,"status":r.status,"remarks":r.reason,"sourceType":"LEGACY_TIOM"} for r in rows]
    stmt=select(SiteDeployment).where(SiteDeployment.site_id==site_id)
    if operating_date: stmt=stmt.where(SiteDeployment.operating_date==operating_date)
    if shift: stmt=stmt.where(SiteDeployment.shift==shift)
    rows=db.scalars(stmt.order_by(SiteDeployment.operating_date.desc(),SiteDeployment.shift,SiteDeployment.entered_at.desc()).limit(5000)).all()
    return [{"deploymentId":r.deployment_id,"operatingDate":r.operating_date,"shift":r.shift,"assetId":r.asset_id,"employeeId":r.employee_id,"locationId":r.location_id,"activity":r.activity,"fromAt":r.from_at,"toAt":r.to_at,"status":r.status,"remarks":r.remarks} for r in rows]


@router.post("/{site_id}/deployments")
def save_deployment(site_id:str,p:SiteDeploymentIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id); require_permission(db,user,site_id,"SHIFT_CONTROL","CREATE")
    if site_id=="TIOM": raise HTTPException(409,"Use the existing TIOM shift deployment workflow during migration.")
    if not p.assetId and not p.employeeId: raise HTTPException(422,"assetId or employeeId is required.")
    _site_shift_or_422(db,site_id,p.shift)
    if p.assetId: _site_equipment_or_422(db,site_id,p.assetId,p.operatingDate,"equipment / vehicle")
    if p.employeeId: _site_person_or_422(db,site_id,p.employeeId,p.operatingDate)
    if p.locationId: _site_location_or_422(db,site_id,p.locationId)
    did=f"DEP-{site_id}-{uuid4().hex[:14].upper()}"; row=SiteDeployment(deployment_id=did,site_id=site_id,operating_date=p.operatingDate,shift=p.shift.upper(),asset_id=p.assetId,employee_id=p.employeeId,location_id=p.locationId,activity=p.activity,from_at=p.fromAt,to_at=p.toAt,status=p.status.upper(),remarks=p.remarks,entered_by=user.login_id); db.add(row); _audit(db,user,site_id,"CREATE_DEPLOYMENT","site_deployment",did,after=p.model_dump()); db.commit(); return {"ok":True,"deploymentId":did}


@router.get("/{site_id}/meters")
def list_meters(site_id:str,request:Request,operating_date:date|None=None,shift:str|None=None,meter_type:str|None=None,db:Session=Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); _require_any(db,user,site_id,"HMR_KMR" if site_id=="KOCP" else "FLEET")
    stmt=select(SiteAssetMeter).where(SiteAssetMeter.site_id==site_id)
    if operating_date: stmt=stmt.where(SiteAssetMeter.operating_date==operating_date)
    if shift: stmt=stmt.where(SiteAssetMeter.shift==shift)
    if meter_type: stmt=stmt.where(SiteAssetMeter.meter_type==meter_type.upper())
    rows=db.scalars(stmt.order_by(SiteAssetMeter.operating_date.desc(),SiteAssetMeter.shift,SiteAssetMeter.asset_id).limit(5000)).all()
    return [{"readingId":r.reading_id,"operatingDate":r.operating_date,"shift":r.shift,"assetId":r.asset_id,"meterType":r.meter_type,"opening":r.opening_reading,"closing":r.closing_reading,"usage":r.usage,"remarks":r.remarks} for r in rows]


@router.post("/{site_id}/meters")
def save_meter(site_id:str,p:SiteAssetMeterIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id); module="HMR_KMR" if site_id=="KOCP" else "FLEET"; require_permission(db,user,site_id,module,"CREATE")
    _site_shift_or_422(db,site_id,p.shift)
    _site_equipment_or_422(db,site_id,p.assetId,p.operatingDate,"machine / vehicle")
    meter_type=str(p.meterType or "").strip().upper()
    if meter_type not in {"HMR","KMR","OTHER"}: raise HTTPException(422,"Choose HMR, KMR or OTHER.")
    if p.openingReading is not None and p.closingReading is not None and p.closingReading < p.openingReading:
        raise HTTPException(422,"Closing reading cannot be less than opening reading.")
    existing=db.scalar(select(SiteAssetMeter).where(SiteAssetMeter.site_id==site_id,SiteAssetMeter.operating_date==p.operatingDate,SiteAssetMeter.shift==p.shift.upper(),SiteAssetMeter.asset_id==p.assetId,SiteAssetMeter.meter_type==meter_type))
    usage=(p.closingReading-p.openingReading) if p.openingReading is not None and p.closingReading is not None else None
    if existing:
        existing.opening_reading=p.openingReading; existing.closing_reading=p.closingReading; existing.usage=usage; existing.remarks=p.remarks; rid=existing.reading_id
    else:
        rid=f"MTR-{site_id}-{uuid4().hex[:14].upper()}"; db.add(SiteAssetMeter(reading_id=rid,site_id=site_id,operating_date=p.operatingDate,shift=p.shift.upper(),asset_id=p.assetId,meter_type=meter_type,opening_reading=p.openingReading,closing_reading=p.closingReading,usage=usage,source_type=p.sourceType,remarks=p.remarks,entered_by=user.login_id))
    _audit(db,user,site_id,"UPSERT_METER","site_asset_meter",rid,after=p.model_dump()); db.commit(); return {"ok":True,"readingId":rid,"usage":usage}


def _norm_vehicle(v:str|None)->str:
    return "".join(ch for ch in (v or "").upper() if ch.isalnum())


@router.post("/{site_id}/reconciliation/run")
def run_reconciliation(site_id:str,request:Request,operating_date:date,shift:str|None=None,window_minutes:int=180,db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id); require_permission(db,user,site_id,"RECONCILIATION","CREATE")
    if site_id=="TIOM": raise HTTPException(409,"TIOM uses the existing field/WB reconciliation engine.")
    trips=list(db.scalars(select(SiteTrip).where(SiteTrip.site_id==site_id,SiteTrip.operating_date==operating_date,SiteTrip.status!="VOID")))
    wbs=list(db.scalars(select(SiteWbMovement).where(SiteWbMovement.site_id==site_id,SiteWbMovement.operating_date==operating_date,SiteWbMovement.row_status=="VALID")))
    if shift:
        trips=[x for x in trips if x.shift==shift]; wbs=[x for x in wbs if x.shift==shift]
    existing_trip=set(db.scalars(select(SiteTripReconciliation.trip_id).where(SiteTripReconciliation.site_id==site_id,SiteTripReconciliation.operating_date==operating_date,SiteTripReconciliation.trip_id.is_not(None))))
    existing_wb=set(db.scalars(select(SiteTripReconciliation.wb_movement_key).where(SiteTripReconciliation.site_id==site_id,SiteTripReconciliation.operating_date==operating_date,SiteTripReconciliation.wb_movement_key.is_not(None))))
    created=matched=0
    used_wb=set(existing_wb)
    for t in trips:
        if t.trip_id in existing_trip: continue
        tv=_norm_vehicle(t.vehicle_raw or (db.get(Equipment,t.vehicle_id).vehicle_no if t.vehicle_id and db.get(Equipment,t.vehicle_id) else None))
        candidates=[]
        for w in wbs:
            if w.movement_key in used_wb: continue
            if tv and _norm_vehicle(w.vehicle_raw)!=tv: continue
            if t.event_at and w.weigh_at:
                diff=abs((w.weigh_at.replace(tzinfo=None)-t.event_at.replace(tzinfo=None)).total_seconds())/60
                if diff>window_minutes: continue
            else: diff=999
            candidates.append((diff,w))
        rid=f"REC-{site_id}-{uuid4().hex[:14].upper()}"
        if len(candidates)==1 or candidates:
            candidates.sort(key=lambda x:x[0]); diff,w=candidates[0]; used_wb.add(w.movement_key)
            qty_diff=None
            if t.quantity_mt is not None: qty_diff=_decimal(w.net_kg)-(_decimal(t.quantity_mt)*Decimal("1000"))
            confidence=max(10,100-int(min(diff,90)))
            status="MATCHED" if len(candidates)==1 else "AMBIGUOUS"
            db.add(SiteTripReconciliation(reconciliation_id=rid,site_id=site_id,operating_date=operating_date,shift=t.shift,trip_id=t.trip_id,wb_movement_key=w.movement_key,status=status,method="VEHICLE_TIME",confidence=confidence,quantity_diff_kg=qty_diff,reason=("Multiple candidates; closest time selected" if status=="AMBIGUOUS" else None))); matched+=status=="MATCHED"
        else:
            db.add(SiteTripReconciliation(reconciliation_id=rid,site_id=site_id,operating_date=operating_date,shift=t.shift,trip_id=t.trip_id,status="TRIP_ONLY",method="AUTO",confidence=0,reason="No unmatched WB movement found"))
        created+=1
    for w in wbs:
        if w.movement_key in used_wb: continue
        rid=f"REC-{site_id}-{uuid4().hex[:14].upper()}"; db.add(SiteTripReconciliation(reconciliation_id=rid,site_id=site_id,operating_date=operating_date,shift=w.shift,wb_movement_key=w.movement_key,status="WB_ONLY",method="AUTO",confidence=0,reason="No unmatched trip found")); created+=1
    _audit(db,user,site_id,"RUN_RECONCILIATION","site_trip_reconciliation",str(operating_date),after={"created":created,"matched":matched,"windowMinutes":window_minutes}); db.commit(); return {"ok":True,"created":created,"matched":matched}


@router.get("/{site_id}/reconciliation")
def list_reconciliation(site_id:str,request:Request,operating_date:date|None=None,status:str|None=None,limit:int=1000,db:Session=Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); _require_any(db,user,site_id,"RECONCILIATION")
    stmt=select(SiteTripReconciliation).where(SiteTripReconciliation.site_id==site_id)
    if operating_date: stmt=stmt.where(SiteTripReconciliation.operating_date==operating_date)
    if status: stmt=stmt.where(SiteTripReconciliation.status==status.upper())
    rows=db.scalars(stmt.order_by(SiteTripReconciliation.operating_date.desc(),SiteTripReconciliation.shift,SiteTripReconciliation.created_at.desc()).limit(min(limit,5000))).all()
    return [{"reconciliationId":r.reconciliation_id,"operatingDate":r.operating_date,"shift":r.shift,"tripId":r.trip_id,"wbMovementKey":r.wb_movement_key,"status":r.status,"method":r.method,"confidence":r.confidence,"quantityDiffKg":r.quantity_diff_kg,"reason":r.reason} for r in rows]


@router.get("/{site_id}/performance")
def performance(site_id:str,request:Request,dimension:str="vehicle",from_date:date|None=None,to_date:date|None=None,db:Session=Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); start,end=_period_bounds(from_date,to_date)
    module={"vehicle":"FLEET","loader":"LOADER","excavator":"EXCAVATOR","driver":"DRIVER","destination":"GP_DESTINATION"}.get(dimension,"REPORTS"); _require_any(db,user,site_id,module)
    if site_id=="TIOM": return {"siteId":site_id,"dimension":dimension,"fromDate":start,"toDate":end,"source":"LEGACY_TIOM","rows":[]}
    field={"vehicle":func.coalesce(SiteTrip.vehicle_id,SiteTrip.vehicle_raw),"loader":SiteTrip.loading_equipment_id,"excavator":SiteTrip.loading_equipment_id,"driver":SiteTrip.driver_id,"destination":SiteTrip.destination_location_id}.get(dimension)
    if field is None: raise HTTPException(422,"Unsupported dimension")
    stmt=select(field.label("key"),func.count().label("trips"),func.coalesce(func.sum(SiteTrip.quantity_mt),0).label("mt"),func.coalesce(func.sum(SiteTrip.quantity_cum),0).label("cum")).where(SiteTrip.site_id==site_id,SiteTrip.operating_date>=start,SiteTrip.operating_date<=end,SiteTrip.status!="VOID").group_by(field).order_by(func.count().desc())
    rows=[]
    for key,trips,mt,cum in db.execute(stmt):
        if key is None: continue
        hsd=db.scalar(select(func.coalesce(func.sum(SiteHsdTransaction.litres),0)).where(SiteHsdTransaction.site_id==site_id,SiteHsdTransaction.operating_date>=start,SiteHsdTransaction.operating_date<=end,SiteHsdTransaction.transaction_type=="ISSUE",SiteHsdTransaction.asset_id==str(key))) or 0
        rows.append({"key":key,"trips":trips,"quantityMt":_fmt_decimal(mt),"quantityCum":_fmt_decimal(cum),"hsdL":_fmt_decimal(hsd),"lPerMt":(_fmt_decimal(hsd)/_fmt_decimal(mt) if _fmt_decimal(mt)>0 else None),"lPerTrip":(_fmt_decimal(hsd)/trips if trips else None)})
    return {"siteId":site_id,"dimension":dimension,"fromDate":start,"toDate":end,"rows":rows}


@router.get("/{site_id}/hsd-summary")
def hsd_summary(site_id:str,request:Request,from_date:date|None=None,to_date:date|None=None,db:Session=Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); _require_any(db,user,site_id,"HSD"); start,end=_period_bounds(from_date,to_date)
    if site_id=="TIOM": return {"siteId":site_id,"fromDate":start,"toDate":end,"source":"LEGACY_TIOM","opening":0,"receipts":0,"issues":0,"adjustments":0,"closing":0}
    sums={}
    for typ in ("OPENING","RECEIPT","ISSUE","ADJUSTMENT"):
        sums[typ.lower()+"s" if typ!="ISSUE" else "issues"]=_fmt_decimal(db.scalar(select(func.coalesce(func.sum(SiteHsdTransaction.litres),0)).where(SiteHsdTransaction.site_id==site_id,SiteHsdTransaction.operating_date>=start,SiteHsdTransaction.operating_date<=end,SiteHsdTransaction.transaction_type==typ,SiteHsdTransaction.status=="POSTED")) or 0)
    opening=sums.get("openings",0); receipts=sums.get("receipts",0); issues=sums.get("issues",0); adjustments=sums.get("adjustments",0)
    return {"siteId":site_id,"fromDate":start,"toDate":end,"opening":opening,"receipts":receipts,"issues":issues,"adjustments":adjustments,"closing":opening+receipts+adjustments-issues}


@router.post("/{site_id}/surveys")
def create_survey(site_id:str,p:SiteSurveyMeasurementIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id); require_permission(db,user,site_id,"MCL_SURVEY","CREATE")
    mid=f"SUR-{site_id}-{uuid4().hex[:14].upper()}"; db.add(SiteSurveyMeasurement(measurement_id=mid,site_id=site_id,period_start=p.periodStart,period_end=p.periodEnd,material_id=p.materialId,measurement_type=p.measurementType.upper(),measured_cum=p.measuredCum,measured_mt=p.measuredMt,source_org=p.sourceOrg,reference_no=p.referenceNo,reference_date=p.referenceDate,status="DRAFT",notes=p.notes,entered_by=user.login_id)); _audit(db,user,site_id,"CREATE_SURVEY","site_survey_measurement",mid,after=p.model_dump()); db.commit(); return {"ok":True,"measurementId":mid,"status":"DRAFT"}


@router.get("/{site_id}/surveys")
def list_surveys(site_id:str,request:Request,db:Session=Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); _require_any(db,user,site_id,"MCL_SURVEY")
    rows=db.scalars(select(SiteSurveyMeasurement).where(SiteSurveyMeasurement.site_id==site_id).order_by(SiteSurveyMeasurement.period_end.desc(),SiteSurveyMeasurement.entered_at.desc()).limit(1000)).all()
    return [{"measurementId":r.measurement_id,"periodStart":r.period_start,"periodEnd":r.period_end,"materialId":r.material_id,"measurementType":r.measurement_type,"measuredCum":r.measured_cum,"measuredMt":r.measured_mt,"sourceOrg":r.source_org,"referenceNo":r.reference_no,"status":r.status,"enteredBy":r.entered_by,"approvedBy":r.approved_by} for r in rows]


@router.post("/{site_id}/surveys/{measurement_id}/decision")
def decide_survey(site_id:str,measurement_id:str,p:SiteDecisionIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id); require_permission(db,user,site_id,"MCL_SURVEY","APPROVE")
    row=db.get(SiteSurveyMeasurement,measurement_id)
    if not row or row.site_id!=site_id: raise HTTPException(404,"Survey measurement not found")
    if p.action not in {"APPROVE","REJECT"}: raise HTTPException(422,"Only APPROVE or REJECT is supported")
    row.status="REJECTED" if p.action=="REJECT" else "APPROVED"; row.approved_by=user.login_id if p.action=="APPROVE" else None; row.approved_at=datetime.now(timezone.utc) if p.action=="APPROVE" else None; _audit(db,user,site_id,p.action+"_SURVEY","site_survey_measurement",measurement_id,reason=p.reason); db.commit(); return {"ok":True,"status":row.status}


@router.post("/{site_id}/billing/build")
def build_billing(site_id:str,request:Request,period_start:date,period_end:date,material_id:str|None=None,survey_id:str|None=None,db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id); require_permission(db,user,site_id,"BILLING","CREATE")
    if period_end<period_start: raise HTTPException(422,"Invalid billing period")
    trip_filters=[SiteTrip.site_id==site_id,SiteTrip.operating_date>=period_start,SiteTrip.operating_date<=period_end,SiteTrip.status!="VOID"]
    if material_id: trip_filters.append(SiteTrip.material_id==material_id)
    op_mt=db.scalar(select(func.coalesce(func.sum(SiteTrip.quantity_mt),0)).where(*trip_filters)) or 0; op_cum=db.scalar(select(func.coalesce(func.sum(SiteTrip.quantity_cum),0)).where(*trip_filters)) or 0
    survey=db.get(SiteSurveyMeasurement,survey_id) if survey_id else db.scalar(select(SiteSurveyMeasurement).where(SiteSurveyMeasurement.site_id==site_id,SiteSurveyMeasurement.period_start<=period_start,SiteSurveyMeasurement.period_end>=period_end,SiteSurveyMeasurement.status=="APPROVED").order_by(SiteSurveyMeasurement.approved_at.desc()).limit(1))
    cert_mt=_decimal(survey.measured_mt) if survey and survey.measured_mt is not None else None; cert_cum=_decimal(survey.measured_cum) if survey else None
    bid=f"BILL-{site_id}-{uuid4().hex[:14].upper()}"; row=SiteBillingReconciliation(billing_id=bid,site_id=site_id,period_start=period_start,period_end=period_end,material_id=material_id,operational_mt=op_mt,operational_cum=op_cum,certified_mt=cert_mt,certified_cum=cert_cum,billable_mt=cert_mt if cert_mt is not None else op_mt,billable_cum=cert_cum if cert_cum is not None else op_cum,variance_mt=(cert_mt-_decimal(op_mt)) if cert_mt is not None else None,variance_cum=(cert_cum-_decimal(op_cum)) if cert_cum is not None else None,source_measurement_id=survey.measurement_id if survey else None,status="DRAFT",created_by=user.login_id); db.add(row); _audit(db,user,site_id,"BUILD_BILLING","site_billing_reconciliation",bid,after={"periodStart":period_start,"periodEnd":period_end,"operationalMt":str(op_mt),"operationalCum":str(op_cum),"surveyId":survey.measurement_id if survey else None}); db.commit(); return {"ok":True,"billingId":bid}


@router.get("/{site_id}/billing")
def list_billing(site_id:str,request:Request,db:Session=Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); _require_any(db,user,site_id,"BILLING")
    rows=db.scalars(select(SiteBillingReconciliation).where(SiteBillingReconciliation.site_id==site_id).order_by(SiteBillingReconciliation.period_end.desc(),SiteBillingReconciliation.created_at.desc()).limit(1000)).all()
    return [{"billingId":r.billing_id,"periodStart":r.period_start,"periodEnd":r.period_end,"materialId":r.material_id,"operationalMt":r.operational_mt,"operationalCum":r.operational_cum,"certifiedMt":r.certified_mt,"certifiedCum":r.certified_cum,"billableMt":r.billable_mt,"billableCum":r.billable_cum,"varianceMt":r.variance_mt,"varianceCum":r.variance_cum,"surveyId":r.source_measurement_id,"status":r.status} for r in rows]


@router.post("/{site_id}/billing/{billing_id}/decision")
def decide_billing(site_id:str,billing_id:str,p:SiteDecisionIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id); require_permission(db,user,site_id,"BILLING","APPROVE")
    row=db.get(SiteBillingReconciliation,billing_id)
    if not row or row.site_id!=site_id: raise HTTPException(404,"Billing reconciliation not found")
    if p.action not in {"APPROVE","REJECT"}: raise HTTPException(422,"Only APPROVE or REJECT is supported")
    row.status="APPROVED" if p.action=="APPROVE" else "REJECTED"; row.approved_by=user.login_id if p.action=="APPROVE" else None; row.approved_at=datetime.now(timezone.utc) if p.action=="APPROVE" else None; _audit(db,user,site_id,p.action+"_BILLING","site_billing_reconciliation",billing_id,reason=p.reason); db.commit(); return {"ok":True,"status":row.status}


def _upsert_issue(db,site_id,operating_date,shift,issue_type,severity,entity,entity_id,description):
    existing=db.scalar(select(SiteDataQualityIssue).where(SiteDataQualityIssue.site_id==site_id,SiteDataQualityIssue.status=="OPEN",SiteDataQualityIssue.issue_type==issue_type,SiteDataQualityIssue.entity==entity,SiteDataQualityIssue.entity_id==entity_id))
    if existing: return False
    db.add(SiteDataQualityIssue(issue_id=f"DQ-{site_id}-{uuid4().hex[:14].upper()}",site_id=site_id,operating_date=operating_date,shift=shift,issue_type=issue_type,severity=severity,entity=entity,entity_id=entity_id,description=description,status="OPEN")); return True


@router.post("/{site_id}/data-quality/scan")
def scan_quality(site_id:str,request:Request,operating_date:date|None=None,db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id); require_permission(db,user,site_id,"DATA_QUALITY","CREATE"); day=operating_date or date.today(); created=0
    if site_id!="TIOM":
        for t in db.scalars(select(SiteTrip).where(SiteTrip.site_id==site_id,SiteTrip.operating_date==day,SiteTrip.status!="VOID")):
            if not (t.vehicle_id or t.vehicle_raw): created+=_upsert_issue(db,site_id,day,t.shift,"MISSING_VEHICLE","CRITICAL","TRIP",t.trip_id,"Trip has no vehicle mapping or vehicle number.")
            if not t.loading_equipment_id: created+=_upsert_issue(db,site_id,day,t.shift,"MISSING_LOADER","WARNING","TRIP",t.trip_id,"Trip has no loader/excavator mapping.")
            if not t.material_id: created+=_upsert_issue(db,site_id,day,t.shift,"MISSING_MATERIAL","WARNING","TRIP",t.trip_id,"Trip has no material mapping.")
            if not t.destination_location_id: created+=_upsert_issue(db,site_id,day,t.shift,"MISSING_DESTINATION","WARNING","TRIP",t.trip_id,"Trip has no destination mapping.")
        for r in db.scalars(select(SiteTripReconciliation).where(SiteTripReconciliation.site_id==site_id,SiteTripReconciliation.operating_date==day,SiteTripReconciliation.status.in_(["TRIP_ONLY","WB_ONLY","AMBIGUOUS","CONFLICT"]))):
            created+=_upsert_issue(db,site_id,day,r.shift,"RECONCILIATION_"+r.status,"CRITICAL" if r.status in {"TRIP_ONLY","WB_ONLY","CONFLICT"} else "WARNING","RECONCILIATION",r.reconciliation_id,r.reason or r.status)
    _audit(db,user,site_id,"DATA_QUALITY_SCAN","site_data_quality_issue",str(day),after={"created":created}); db.commit(); return {"ok":True,"created":created}


@router.get("/{site_id}/data-quality")
def list_quality(site_id:str,request:Request,status:str|None="OPEN",limit:int=1000,db:Session=Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); _require_any(db,user,site_id,"DATA_QUALITY")
    stmt=select(SiteDataQualityIssue).where(SiteDataQualityIssue.site_id==site_id)
    if status: stmt=stmt.where(SiteDataQualityIssue.status==status.upper())
    rows=db.scalars(stmt.order_by(SiteDataQualityIssue.detected_at.desc()).limit(min(limit,5000))).all()
    return [{"issueId":r.issue_id,"operatingDate":r.operating_date,"shift":r.shift,"issueType":r.issue_type,"severity":r.severity,"entity":r.entity,"entityId":r.entity_id,"description":r.description,"status":r.status,"owner":r.owner,"detectedAt":r.detected_at,"resolution":r.resolution} for r in rows]


@router.post("/{site_id}/data-quality/{issue_id}/resolve")
def resolve_quality(site_id:str,issue_id:str,p:SiteDataQualityResolveIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id); require_permission(db,user,site_id,"DATA_QUALITY","EDIT")
    row=db.get(SiteDataQualityIssue,issue_id)
    if not row or row.site_id!=site_id: raise HTTPException(404,"Issue not found")
    row.status="RESOLVED"; row.resolved_by=user.login_id; row.resolved_at=datetime.now(timezone.utc); row.resolution=p.resolution; _audit(db,user,site_id,"RESOLVE_DQ","site_data_quality_issue",issue_id,reason=p.resolution); db.commit(); return {"ok":True}


@router.get("/{site_id}/map")
def map_context(site_id:str,request:Request,db:Session=Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); _require_any(db,user,site_id,"MAP"); site=db.get(Site,site_id); locs=db.scalars(select(SiteLocation).where(SiteLocation.site_id==site_id,SiteLocation.active.is_(True)).order_by(SiteLocation.name)).all(); return {"siteId":site_id,"mapUrl":site.map_url if site else None,"locations":[{"id":x.site_location_id,"code":x.code,"name":x.name,"type":x.location_type,"distanceKm":x.distance_km,"latitude":x.latitude,"longitude":x.longitude} for x in locs]}


@router.post("/{site_id}/map/config")
def save_map_config(site_id:str,p:SiteMapConfigIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id); require_permission(db,user,site_id,"MAP","EDIT"); site=db.get(Site,site_id)
    if not site: raise HTTPException(404,"Site not found")
    old=site.map_url; site.map_url=p.mapUrl; _audit(db,user,site_id,"UPDATE_MAP","site",site_id,before={"mapUrl":old},after={"mapUrl":p.mapUrl}); db.commit(); return {"ok":True}


@router.get("/{site_id}/satellite")
def list_satellite(site_id:str,request:Request,db:Session=Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); _require_any(db,user,site_id,"SATELLITE"); rows=db.scalars(select(SiteSatelliteObservation).where(SiteSatelliteObservation.site_id==site_id).order_by(SiteSatelliteObservation.image_date.desc()).limit(200)).all(); return [{"observationId":r.observation_id,"provider":r.provider,"imageDate":r.image_date,"cloudPct":r.cloud_pct,"imageRef":r.image_ref,"previousObservationId":r.previous_observation_id,"changeAreaHa":r.change_area_ha,"status":r.status,"notes":r.notes} for r in rows]


@router.get("/{site_id}/satellite/{observation_id}/open")
def open_satellite_image(site_id:str, observation_id:str, request:Request, db:Session=Depends(get_db)):
    """Open a viewable Sentinel image, resolving old STAC metadata links on demand."""
    user=get_user(db,request)
    site_id=require_site(db,user,site_id)
    _require_any(db,user,site_id,"SATELLITE")
    row=db.scalar(select(SiteSatelliteObservation).where(
        SiteSatelliteObservation.site_id==site_id,
        SiteSatelliteObservation.observation_id==observation_id,
    ))
    if row is None:
        raise HTTPException(404,"Satellite observation not found.")

    target=str(row.image_ref or "").strip()
    if not target:
        raise HTTPException(404,"No satellite image reference is available.")

    if "stac.dataspace.copernicus.eu" in target and "/items/" in target:
        try:
            req=UrlRequest(target,headers={"User-Agent":"NMTPL-Central-Operations/1.0"})
            with urlopen(req,timeout=20) as resp:
                item=json.load(resp)
            assets=item.get("assets") or {}
            thumbnail=assets.get("thumbnail") or {}
            preview=thumbnail.get("href")
            if not preview:
                tci=assets.get("TCI_10m") or assets.get("TCI_20m") or assets.get("TCI_60m") or {}
                preview=((tci.get("alternate") or {}).get("https") or {}).get("href")
            if preview:
                target=preview
                row.image_ref=preview
                db.commit()
        except (HTTPError,URLError,TimeoutError,ValueError,TypeError):
            db.rollback()

    if not target.startswith(("http://","https://")):
        raise HTTPException(502,"A browser-viewable Sentinel image URL is not available.")
    return RedirectResponse(url=target,status_code=302)


@router.post("/{site_id}/satellite")
def save_satellite(site_id:str,p:SiteSatelliteObservationIn,request:Request,db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id); require_permission(db,user,site_id,"SATELLITE","CREATE"); oid=f"SAT-{site_id}-{uuid4().hex[:14].upper()}"; db.add(SiteSatelliteObservation(observation_id=oid,site_id=site_id,provider=p.provider,image_date=p.imageDate,cloud_pct=p.cloudPct,image_ref=p.imageRef,previous_observation_id=p.previousObservationId,change_area_ha=p.changeAreaHa,status=p.status,notes=p.notes,created_by=user.login_id)); _audit(db,user,site_id,"CREATE_SATELLITE_OBSERVATION","site_satellite_observation",oid,after=p.model_dump()); db.commit(); return {"ok":True,"observationId":oid}



@router.post("/{site_id}/satellite/discover")
def discover_satellite(
    site_id: str,
    request: Request,
    min_lon: float,
    min_lat: float,
    max_lon: float,
    max_lat: float,
    from_date: date | None = None,
    to_date: date | None = None,
    max_cloud: float = 20.0,
    db: Session = Depends(get_db),
):
    """Discover recent Sentinel-2 L2A products from the public CDSE STAC catalogue.

    This stores catalogue metadata only; it does not download large raster files.
    The public endpoint can be run later as a scheduled worker once the site's AOI
    is finalized.
    """
    csrf(request)
    user = get_user(db, request)
    site_id = require_site(db, user, site_id)
    require_permission(db, user, site_id, "SATELLITE", "CREATE")
    if not (-180 <= min_lon < max_lon <= 180 and -90 <= min_lat < max_lat <= 90):
        raise HTTPException(422, "Invalid bounding box coordinates.")
    if not (0 <= max_cloud <= 100):
        raise HTTPException(422, "max_cloud must be between 0 and 100.")
    end = to_date or date.today()
    start = from_date or (end - timedelta(days=45))
    body = {
        "collections": ["sentinel-2-l2a"],
        "bbox": [min_lon, min_lat, max_lon, max_lat],
        "datetime": f"{start.isoformat()}T00:00:00Z/{end.isoformat()}T23:59:59Z",
        "query": {"eo:cloud_cover": {"lte": max_cloud}},
        "limit": 25,
    }
    req = UrlRequest(
        "https://stac.dataspace.copernicus.eu/v1/search",
        data=json.dumps(body).encode("utf-8"),
        headers={"Content-Type": "application/json", "User-Agent": "NMTPL-Central-Operations/1.0"},
        method="POST",
    )
    try:
        with urlopen(req, timeout=20) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, ValueError) as exc:
        raise HTTPException(502, f"Copernicus catalogue unavailable: {type(exc).__name__}")
    features = payload.get("features") or []

    # A single Sentinel-2 acquisition date can return multiple STAC products/tiles.
    # The database intentionally stores one observation per site/provider/date, so
    # collapse the catalogue response to the lowest-cloud product for each date
    # before touching the database.
    best_by_day = {}
    for item in features:
        props = item.get("properties") or {}
        raw_dt = props.get("datetime") or props.get("start_datetime")
        if not raw_dt:
            continue
        try:
            image_day = datetime.fromisoformat(raw_dt.replace("Z", "+00:00")).date()
        except ValueError:
            continue
        cloud_raw = props.get("eo:cloud_cover")
        try:
            cloud_value = float(cloud_raw) if cloud_raw is not None else 101.0
        except (TypeError, ValueError):
            cloud_value = 101.0
        current = best_by_day.get(image_day)
        if current is None or cloud_value < current[0]:
            best_by_day[image_day] = (cloud_value, item)

    selected = sorted(best_by_day.items(), key=lambda x: x[0], reverse=True)[:10]
    stored = 0
    updated = 0
    output = []
    for image_day, (_, item) in selected:
        props = item.get("properties") or {}
        cloud = props.get("eo:cloud_cover")
        assets = item.get("assets") or {}
        thumbnail = assets.get("thumbnail") or {}
        item_url = thumbnail.get("href")
        if not item_url:
            tci = assets.get("TCI_10m") or assets.get("TCI_20m") or {}
            item_url = ((tci.get("alternate") or {}).get("https") or {}).get("href")
        if not item_url:
            links = item.get("links") or []
            item_url = next((x.get("href") for x in links if x.get("rel") in {"self", "alternate"} and x.get("href")), None)
        item_id = str(item.get("id") or "")
        existing = db.scalar(select(SiteSatelliteObservation).where(
            SiteSatelliteObservation.site_id == site_id,
            SiteSatelliteObservation.provider == "SENTINEL-2",
            SiteSatelliteObservation.image_date == image_day,
        ))
        if existing is None:
            existing = SiteSatelliteObservation(
                observation_id=f"SAT-{site_id}-{uuid4().hex[:14].upper()}",
                site_id=site_id,
                provider="SENTINEL-2",
                image_date=image_day,
                cloud_pct=cloud,
                image_ref=item_url,
                status="DISCOVERED",
                notes=f"CDSE STAC item {item_id}" if item_id else "CDSE STAC discovery",
                created_by=user.login_id,
            )
            db.add(existing)
            stored += 1
        else:
            old_cloud = float(existing.cloud_pct) if existing.cloud_pct is not None else 101.0
            new_cloud = float(cloud) if cloud is not None else 101.0
            old_ref = str(existing.image_ref or "")
            old_ref_is_metadata = "stac.dataspace.copernicus.eu" in old_ref and "/items/" in old_ref
            if new_cloud < old_cloud or not existing.image_ref or old_ref_is_metadata:
                existing.cloud_pct = cloud
                existing.image_ref = item_url
                existing.status = "DISCOVERED"
                existing.notes = f"CDSE STAC item {item_id}" if item_id else "CDSE STAC discovery"
                updated += 1
        output.append({"id": item_id, "imageDate": image_day, "cloudPct": cloud, "itemUrl": item_url})

    _audit(
        db, user, site_id, "SATELLITE_DISCOVER", "SiteSatelliteObservation", None,
        after={"bbox": body["bbox"], "found": len(features), "uniqueDates": len(best_by_day), "stored": stored, "updated": updated},
    )
    db.commit()
    return {
        "ok": True,
        "provider": "Copernicus Data Space STAC",
        "found": len(features),
        "uniqueDates": len(best_by_day),
        "stored": stored,
        "updated": updated,
        "items": output,
    }

@router.get("/{site_id}/reports/summary")
def report_summary(site_id:str,request:Request,from_date:date|None=None,to_date:date|None=None,db:Session=Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); _require_any(db,user,site_id,"REPORTS"); start,end=_period_bounds(from_date,to_date)
    if site_id=="TIOM":
        trips=db.scalar(select(func.count()).select_from(LoadTrip).where(LoadTrip.operating_date>=start,LoadTrip.operating_date<=end)) or 0; net=db.scalar(select(func.coalesce(func.sum(WbMovement.net_kg),0)).where(WbMovement.operating_date>=start,WbMovement.operating_date<=end)) or 0; return {"siteId":site_id,"fromDate":start,"toDate":end,"trips":trips,"quantityMt":float(_decimal(net)/1000),"quantityCum":0,"hsdIssuedL":0,"wbRows":db.scalar(select(func.count()).select_from(WbMovement).where(WbMovement.operating_date>=start,WbMovement.operating_date<=end)) or 0,"source":"LEGACY_TIOM"}
    filters=[SiteTrip.site_id==site_id,SiteTrip.operating_date>=start,SiteTrip.operating_date<=end,SiteTrip.status!="VOID"]
    trips=db.scalar(select(func.count()).select_from(SiteTrip).where(*filters)) or 0; mt=db.scalar(select(func.coalesce(func.sum(SiteTrip.quantity_mt),0)).where(*filters)) or 0; cum=db.scalar(select(func.coalesce(func.sum(SiteTrip.quantity_cum),0)).where(*filters)) or 0; hsd=db.scalar(select(func.coalesce(func.sum(SiteHsdTransaction.litres),0)).where(SiteHsdTransaction.site_id==site_id,SiteHsdTransaction.operating_date>=start,SiteHsdTransaction.operating_date<=end,SiteHsdTransaction.transaction_type=="ISSUE",SiteHsdTransaction.status=="POSTED")) or 0; wb=db.scalar(select(func.count()).select_from(SiteWbMovement).where(SiteWbMovement.site_id==site_id,SiteWbMovement.operating_date>=start,SiteWbMovement.operating_date<=end,SiteWbMovement.row_status=="VALID")) or 0
    shifts=[]
    for sh,tr,sm,sc in db.execute(select(SiteTrip.shift,func.count(),func.coalesce(func.sum(SiteTrip.quantity_mt),0),func.coalesce(func.sum(SiteTrip.quantity_cum),0)).where(*filters).group_by(SiteTrip.shift).order_by(SiteTrip.shift)): shifts.append({"shift":sh,"trips":tr,"quantityMt":_fmt_decimal(sm),"quantityCum":_fmt_decimal(sc)})
    return {"siteId":site_id,"fromDate":start,"toDate":end,"trips":trips,"quantityMt":_fmt_decimal(mt),"quantityCum":_fmt_decimal(cum),"hsdIssuedL":_fmt_decimal(hsd),"wbRows":wb,"shifts":shifts,"lPerMt":(_fmt_decimal(hsd)/_fmt_decimal(mt) if _fmt_decimal(mt)>0 else None)}


@router.get("/{site_id}/reports/export.xlsx")
def report_export(site_id:str,request:Request,from_date:date|None=None,to_date:date|None=None,db:Session=Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); _require_any(db,user,site_id,"REPORTS"); start,end=_period_bounds(from_date,to_date)
    wb=Workbook(); ws=wb.active; ws.title="Summary"; title=f"NMTPL {site_id} Operations Report"; ws.append([title]); ws.append(["Period",str(start),"to",str(end)]); ws.append([])
    if site_id=="TIOM":
        ws.append(["TIOM detailed exports remain available through the legacy report engine during migration."])
    else:
        headers=["Operating Date","Shift","Trip ID","Vehicle","Driver","Loader/Excavator","Source","Destination","Material","GP No","MT","CuM","Weight Basis","Status"] ; ws.append(headers)
        for cell in ws[4]: cell.font=Font(bold=True); cell.fill=PatternFill("solid",fgColor="D9EAF7")
        rows=db.scalars(select(SiteTrip).where(SiteTrip.site_id==site_id,SiteTrip.operating_date>=start,SiteTrip.operating_date<=end,SiteTrip.status!="VOID").order_by(SiteTrip.operating_date,SiteTrip.shift,SiteTrip.entered_at)).all()
        for r in rows: ws.append([r.operating_date,r.shift,r.trip_id,r.vehicle_id or r.vehicle_raw,r.driver_id,r.loading_equipment_id,r.source_location_id,r.destination_location_id,r.material_id,r.gp_no,float(r.quantity_mt) if r.quantity_mt is not None else None,float(r.quantity_cum) if r.quantity_cum is not None else None,r.weight_basis,r.status])
        h=wb.create_sheet("HSD"); h.append(["Date","Shift","Type","Asset Type","Asset","Litres","Supplier","Reference","Meter","Status"]); [setattr(c,"font",Font(bold=True)) for c in h[1]]
        for r in db.scalars(select(SiteHsdTransaction).where(SiteHsdTransaction.site_id==site_id,SiteHsdTransaction.operating_date>=start,SiteHsdTransaction.operating_date<=end).order_by(SiteHsdTransaction.operating_date,SiteHsdTransaction.shift)).all(): h.append([r.operating_date,r.shift,r.transaction_type,r.asset_type,r.asset_id,float(r.litres),r.supplier,r.reference_no,float(r.meter_reading) if r.meter_reading is not None else None,r.status])
        if site_id=="SOCP":
            s=wb.create_sheet("Weighbridge"); s.append(["WB ID","Date/Time","Operating Date","Shift","Vehicle","Party","Source","Destination","Gross kg","Tare kg","Net kg","Status"]); [setattr(c,"font",Font(bold=True)) for c in s[1]]
            for r in db.scalars(select(SiteWbMovement).where(SiteWbMovement.site_id==site_id,SiteWbMovement.operating_date>=start,SiteWbMovement.operating_date<=end).order_by(SiteWbMovement.weigh_at)).all(): s.append([r.wb_id,r.weigh_at,r.operating_date,r.shift,r.vehicle_raw,r.party,r.source_raw,r.destination_raw,float(r.gross_kg),float(r.tare_kg),float(r.net_kg),r.row_status])
    for sheet in wb.worksheets:
        for col in sheet.columns:
            width=min(35,max(10,max(len(str(c.value or "")) for c in col)+2)); sheet.column_dimensions[col[0].column_letter].width=width
        sheet.freeze_panes="A2"
    data=BytesIO(); wb.save(data); data.seek(0); fn=f"NMTPL_{site_id}_{start}_{end}.xlsx"; return StreamingResponse(data,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="{fn}"'})


def _xlsx_response(wb: Workbook, filename: str):
    data=BytesIO(); wb.save(data); data.seek(0)
    return StreamingResponse(data,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="{filename}"'})


def _header_map(headers):
    def norm(x): return "".join(ch for ch in str(x or "").upper().strip() if ch.isalnum())
    return {norm(h):i for i,h in enumerate(headers) if h not in (None,"")}


@router.get("/{site_id}/wb/template.xlsx")
def wb_template(site_id:str,request:Request,db:Session=Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); _require_any(db,user,site_id,"WB")
    wb=Workbook(); ws=wb.active; ws.title="WB Data"; headers=["ID","Date","REGDNO","PARTY","SOURCE","DESTINATION","GROSS","TARE","NET"]; ws.append(headers); ws.append(["740976","2026-09-28 06:39:08","OD23T8597","NMTPL","SURFACE MINER","RLS",49930,17240,32690]); ins=wb.create_sheet("Instructions"); ins.append(["NMTPL WB Bulk Upload"]); ins.append(["Site",site_id]); ins.append(["Rule","Gross - Tare must equal Net within 1 kg"]); ins.append(["Rule","ID must be unique within the site"]); ins.append(["Rule","Date/time is the actual WB timestamp; operating shift is derived automatically"]); return _xlsx_response(wb,f"NMTPL_{site_id}_WB_Template.xlsx")


@router.post("/{site_id}/wb/preview")
async def wb_preview(site_id:str,request:Request,file:UploadFile=File(...),db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id); require_permission(db,user,site_id,"WB","CREATE")
    payload=await file.read()
    if len(payload)>20*1024*1024: raise HTTPException(413,"WB upload exceeds 20 MB")
    try: wb=load_workbook(BytesIO(payload),read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]
    except Exception as exc: raise HTTPException(422,"Upload a valid .xlsx workbook") from exc
    headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]; hm=_header_map(headers); required=["ID","DATE","REGDNO","GROSS","TARE","NET"]; missing=[x for x in required if x not in hm]
    batch_id=f"WBIMP-{site_id}-{uuid4().hex[:12].upper()}"; batch=SiteWbImportBatch(batch_id=batch_id,site_id=site_id,source_type="WB_XLSX",file_name=file.filename,file_hash=sha256(payload).hexdigest(),status="INVALID_HEADERS" if missing else "STAGED",uploaded_by=user.login_id); db.add(batch)
    sample=[]; valid=review=0; seen=set()
    def val(cells,key):
        i=hm.get(key); return cells[i] if i is not None and i<len(cells) else None
    for row_no,cells in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
        if not any(x not in (None,"") for x in cells): continue
        errors=[]; wb_id=str(val(cells,"ID") or "").strip(); dt=val(cells,"DATE"); vehicle=str(val(cells,"REGDNO") or "").strip(); gross=val(cells,"GROSS"); tare=val(cells,"TARE"); net=val(cells,"NET")
        if not wb_id: errors.append("ID required")
        if wb_id in seen: errors.append("Duplicate ID inside upload")
        seen.add(wb_id)
        if db.scalar(select(SiteWbMovement.movement_key).where(SiteWbMovement.site_id==site_id,SiteWbMovement.wb_id==wb_id)): errors.append("WB ID already exists")
        if not isinstance(dt,datetime):
            try: dt=datetime.fromisoformat(str(dt))
            except Exception: errors.append("Invalid Date")
        try: gross=Decimal(str(gross)); tare=Decimal(str(tare)); net=Decimal(str(net))
        except Exception: errors.append("Gross/Tare/Net must be numeric")
        if not errors and gross <= tare: errors.append("Gross must be greater than Tare")
        if not errors and net <= 0: errors.append("Net must be greater than 0")
        if not errors and abs((gross-tare)-net)>Decimal("1"): errors.append("Gross - Tare != Net")
        if not vehicle: errors.append("REGDNO required")
        status="ERROR" if errors or missing else "STAGED"; review+=int(status=="ERROR"); valid+=int(status=="STAGED")
        if status=="STAGED":
            ctx=resolve_site_context(db,site_id,dt); key=f"{site_id}:WB:{wb_id}"; db.add(SiteWbMovement(movement_key=key,site_id=site_id,batch_id=batch_id,operating_date=ctx.operating_date,shift=ctx.shift,wb_id=wb_id,weigh_at=dt,vehicle_raw=_norm_vehicle(vehicle),party=str(val(cells,"PARTY") or "").strip() or None,source_raw=str(val(cells,"SOURCE") or "").strip() or None,destination_raw=str(val(cells,"DESTINATION") or "").strip() or None,gross_kg=gross,tare_kg=tare,net_kg=net,row_status="STAGED",issue=None,entered_by=user.login_id))
        if len(sample)<30: sample.append({"row":row_no,"wbId":wb_id,"date":str(dt),"vehicle":vehicle,"gross":str(gross),"tare":str(tare),"net":str(net),"status":status,"errors":errors})
    batch.total_rows=valid+review; batch.valid_rows=valid; batch.review_rows=review; _audit(db,user,site_id,"WB_IMPORT_PREVIEW","site_wb_import_batch",batch_id,after={"valid":valid,"errors":review,"missingHeaders":missing}); db.commit(); return {"ok":not missing and review==0,"batchId":batch_id,"validRows":valid,"errorRows":review,"missingHeaders":missing,"sample":sample}


@router.post("/{site_id}/wb/imports/{batch_id}/confirm")
def wb_confirm(site_id:str,batch_id:str,request:Request,db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id); require_permission(db,user,site_id,"WB","APPROVE")
    batch=db.get(SiteWbImportBatch,batch_id)
    if not batch or batch.site_id!=site_id: raise HTTPException(404,"WB batch not found")
    if batch.status!="STAGED": raise HTTPException(409,f"Batch cannot be confirmed from {batch.status}")
    if batch.review_rows: raise HTTPException(409,"Upload has validation errors; correct and upload again")
    rows=db.scalars(select(SiteWbMovement).where(SiteWbMovement.batch_id==batch_id,SiteWbMovement.row_status=="STAGED")).all()
    for r in rows: r.row_status="VALID"
    batch.status="CONFIRMED"; batch.confirmed_by=user.login_id; batch.confirmed_at=datetime.now(timezone.utc); _audit(db,user,site_id,"WB_IMPORT_CONFIRM","site_wb_import_batch",batch_id,after={"rows":len(rows)}); db.commit(); return {"ok":True,"rows":len(rows)}


@router.get("/{site_id}/trips/template.xlsx")
def trip_template(site_id:str,request:Request,db:Session=Depends(get_db)):
    user=get_user(db,request); site_id=require_site(db,user,site_id); _require_any(db,user,site_id,"TRIP")
    wb=Workbook(); ws=wb.active; ws.title="Trips"; headers=["Event At","Vehicle","Driver ID","Loading Equipment ID","Loading Operator ID","Source Code","Destination Code","Material Code","GP No","Quantity MT","Quantity CuM","Factor ID","Source UID"]; ws.append(headers); ws.append(["2026-09-30 06:15:00","OD-XX-0001","EMP001","L001","EMP010","FACE1","RLS","COAL","GP001",31.5,"","","SAMPLE-001"]); ins=wb.create_sheet("Instructions"); ins.append(["NMTPL Operational Trip Import"]); ins.append(["Site",site_id]); ins.append(["Note","Location and material codes are resolved against site masters."]); ins.append(["Note","For KOCP, approved Factor ID can populate MCL average MT/trip or OB CuM/trip."]); return _xlsx_response(wb,f"NMTPL_{site_id}_Trip_Template.xlsx")


@router.post("/{site_id}/trips/preview")
async def trip_preview(site_id:str,request:Request,file:UploadFile=File(...),db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id)
    if not user.admin and not (has_permission(db,user,site_id,"TRIP","CREATE") or has_permission(db,user,site_id,"OB","CREATE")): raise HTTPException(403,"Trip/OB create permission required")
    payload=await file.read()
    if len(payload)>25*1024*1024: raise HTTPException(413,"Trip upload exceeds 25 MB")
    try: wb=load_workbook(BytesIO(payload),read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]
    except Exception as exc: raise HTTPException(422,"Upload a valid .xlsx workbook") from exc
    headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]; hm=_header_map(headers); required=["EVENTAT","VEHICLE"]; missing=[x for x in required if x not in hm]; batch_id=f"TRIMP-{site_id}-{uuid4().hex[:12].upper()}"; batch=SiteOperationalImportBatch(batch_id=batch_id,site_id=site_id,import_type="TRIP",file_name=file.filename or "trips.xlsx",file_hash=sha256(payload).hexdigest(),status="INVALID_HEADERS" if missing else "PREVIEW",uploaded_by=user.login_id); db.add(batch)
    valid=errors_count=0; sample=[]
    def val(cells,key):
        i=hm.get(key); return cells[i] if i is not None and i<len(cells) else None
    for row_no,cells in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
        if not any(x not in (None,"") for x in cells): continue
        errors=[]; dt=val(cells,"EVENTAT"); vehicle=str(val(cells,"VEHICLE") or "").strip(); uid=str(val(cells,"SOURCEUID") or f"{batch_id}-R{row_no}")
        if not isinstance(dt,datetime):
            try: dt=datetime.fromisoformat(str(dt))
            except Exception: errors.append("Invalid Event At")
        if not vehicle: errors.append("Vehicle required")
        if db.scalar(select(SiteTrip.trip_id).where(SiteTrip.site_id==site_id,SiteTrip.source_type=="XLSX_IMPORT",SiteTrip.source_record_uid==uid)): errors.append("Source UID already imported")
        src=str(val(cells,"SOURCECODE") or "").strip().upper(); dst=str(val(cells,"DESTINATIONCODE") or "").strip().upper(); mat=str(val(cells,"MATERIALCODE") or "").strip().upper(); factor_id=str(val(cells,"FACTORID") or "").strip() or None
        if src and not db.get(SiteLocation,f"{site_id}:{src}"): errors.append(f"Unknown source code {src}")
        if dst and not db.get(SiteLocation,f"{site_id}:{dst}"): errors.append(f"Unknown destination code {dst}")
        if mat and not db.get(SiteMaterial,f"{site_id}:{mat}"): errors.append(f"Unknown material code {mat}")
        if factor_id:
            f=db.get(SiteWeightFactor,factor_id)
            if not f or f.site_id!=site_id or f.status!="APPROVED": errors.append("Factor is not approved for this site")
        record={"eventAt":dt.isoformat() if isinstance(dt,datetime) else str(dt),"vehicle":vehicle,"driverId":str(val(cells,"DRIVERID") or "").strip() or None,"loadingEquipmentId":str(val(cells,"LOADINGEQUIPMENTID") or "").strip() or None,"loadingOperatorId":str(val(cells,"LOADINGOPERATORID") or "").strip() or None,"sourceCode":src or None,"destinationCode":dst or None,"materialCode":mat or None,"gpNo":str(val(cells,"GPNO") or "").strip() or None,"quantityMt":str(val(cells,"QUANTITYMT")) if val(cells,"QUANTITYMT") not in (None,"") else None,"quantityCum":str(val(cells,"QUANTITYCUM")) if val(cells,"QUANTITYCUM") not in (None,"") else None,"factorId":factor_id,"sourceUid":uid}
        status="ERROR" if errors or missing else "VALID"; db.add(SiteOperationalImportRow(batch_id=batch_id,row_no=row_no,row_json=json.dumps(record),status=status,error_text="; ".join(errors) if errors else ("Missing headers: "+", ".join(missing) if missing else None))); valid+=int(status=="VALID"); errors_count+=int(status=="ERROR");
        if len(sample)<30: sample.append({"row":row_no,"status":status,"errors":errors,**record})
    batch.total_rows=valid+errors_count; batch.valid_rows=valid; batch.error_rows=errors_count; _audit(db,user,site_id,"TRIP_IMPORT_PREVIEW","site_operational_import_batch",batch_id,after={"valid":valid,"errors":errors_count,"missingHeaders":missing}); db.commit(); return {"ok":not missing and errors_count==0,"batchId":batch_id,"validRows":valid,"errorRows":errors_count,"missingHeaders":missing,"sample":sample}


@router.post("/{site_id}/trips/imports/{batch_id}/confirm")
def trip_confirm(site_id:str,batch_id:str,request:Request,db:Session=Depends(get_db)):
    csrf(request); user=get_user(db,request); site_id=require_site(db,user,site_id)
    if not user.admin and not (has_permission(db,user,site_id,"TRIP","APPROVE") or has_permission(db,user,site_id,"OB","APPROVE")): raise HTTPException(403,"Trip/OB approve permission required")
    batch=db.get(SiteOperationalImportBatch,batch_id)
    if not batch or batch.site_id!=site_id or batch.import_type!="TRIP": raise HTTPException(404,"Trip import batch not found")
    if batch.status!="PREVIEW" or batch.error_rows: raise HTTPException(409,"Batch must be a clean PREVIEW before confirmation")
    rows=db.scalars(select(SiteOperationalImportRow).where(SiteOperationalImportRow.batch_id==batch_id,SiteOperationalImportRow.status=="VALID").order_by(SiteOperationalImportRow.row_no)).all(); created=0
    for sr in rows:
        r=json.loads(sr.row_json); dt=datetime.fromisoformat(r["eventAt"]); ctx=resolve_site_context(db,site_id,dt); factor=db.get(SiteWeightFactor,r.get("factorId")) if r.get("factorId") else None; mt=Decimal(r["quantityMt"]) if r.get("quantityMt") else None; cum=Decimal(r["quantityCum"]) if r.get("quantityCum") else None; basis=None
        if factor:
            if ctx.operating_date<factor.effective_from or (factor.effective_to and ctx.operating_date>factor.effective_to) or (factor.operating_date and factor.operating_date!=ctx.operating_date) or (factor.shift and factor.shift!=ctx.shift): raise HTTPException(409,f"Factor {factor.factor_id} is not valid for row {sr.row_no}")
            if factor.factor_type=="COAL_AVG_MT_PER_TRIP": mt=factor.factor_value; basis="MCL_SHIFT_AVERAGE"
            elif factor.factor_type=="OB_CUM_PER_TRIP": cum=factor.factor_value; basis="MCL_DUMPER_FACTOR"
        vehicle_match=db.scalar(select(Equipment).where(or_(Equipment.machine_id==r["vehicle"],func.replace(func.replace(Equipment.vehicle_no,"-","")," ","")==_norm_vehicle(r["vehicle"]))).limit(1))
        tid=f"TRIP-{site_id}-{uuid4().hex[:14].upper()}"; db.add(SiteTrip(trip_id=tid,site_id=site_id,operating_date=ctx.operating_date,shift=ctx.shift,event_at=dt,vehicle_id=vehicle_match.machine_id if vehicle_match else None,vehicle_raw=None if vehicle_match else r["vehicle"],driver_id=r.get("driverId"),loading_equipment_id=r.get("loadingEquipmentId"),loading_operator_id=r.get("loadingOperatorId"),source_location_id=f"{site_id}:{r['sourceCode']}" if r.get("sourceCode") else None,destination_location_id=f"{site_id}:{r['destinationCode']}" if r.get("destinationCode") else None,material_id=f"{site_id}:{r['materialCode']}" if r.get("materialCode") else None,gp_no=r.get("gpNo"),quantity_mt=mt,quantity_cum=cum,weight_basis=basis,factor_id=r.get("factorId"),source_type="XLSX_IMPORT",source_record_uid=r["sourceUid"],source_batch_id=batch_id,status="POSTED",entered_by=user.login_id)); sr.status="IMPORTED"; created+=1
    batch.status="IMPORTED"; batch.confirmed_by=user.login_id; batch.confirmed_at=datetime.now(timezone.utc); _audit(db,user,site_id,"TRIP_IMPORT_CONFIRM","site_operational_import_batch",batch_id,after={"created":created}); db.commit(); return {"ok":True,"created":created}


# ---------------------------------------------------------------------------
# TIOM MIS Office Entry / paper-field-WB reconciliation
# ---------------------------------------------------------------------------

def _tiom_mis_permission(db, user, action: str):
    if user.admin:
        return
    require_permission(db, user, "TIOM", "MIS", action)


def _time_delta_seconds(a: datetime | None, b: datetime | None) -> float | None:
    if a is None or b is None:
        return None
    try:
        return abs((a - b).total_seconds())
    except TypeError:
        aa = a.replace(tzinfo=None) if getattr(a, "tzinfo", None) else a
        bb = b.replace(tzinfo=None) if getattr(b, "tzinfo", None) else b
        return abs((aa - bb).total_seconds())


def _mis_report_json(db: Session, report: TiomMisReport, include_rows: bool = False):
    rows = list(db.scalars(select(TiomMisTripRow).where(TiomMisTripRow.report_id == report.report_id).order_by(TiomMisTripRow.row_no)))
    recon = {x.row_id: x for x in db.scalars(select(TiomMisReconciliation).where(TiomMisReconciliation.row_id.in_([r.row_id for r in rows] if rows else ["__none__"]))) }
    status_counts = {}
    for x in recon.values():
        status_counts[x.match_status] = status_counts.get(x.match_status, 0) + 1
    out = {
        "reportId": report.report_id, "operatingDate": report.operating_date, "shift": report.shift,
        "vehicleId": report.vehicle_id, "operatorId": report.operator_id,
        "openingKmr": _fmt_decimal(report.opening_kmr) if report.opening_kmr is not None else None,
        "closingKmr": _fmt_decimal(report.closing_kmr) if report.closing_kmr is not None else None,
        "openingHmr": _fmt_decimal(report.opening_hmr) if report.opening_hmr is not None else None,
        "closingHmr": _fmt_decimal(report.closing_hmr) if report.closing_hmr is not None else None,
        "paperRef": report.paper_ref, "sourceDocument": report.source_document, "status": report.status,
        "enteredBy": report.entered_by, "enteredAt": report.entered_at, "submittedBy": report.submitted_by,
        "submittedAt": report.submitted_at, "approvedBy": report.approved_by, "approvedAt": report.approved_at,
        "version": report.version, "notes": report.notes, "rowCount": len(rows), "matchSummary": status_counts,
        "authoritativeTonnesChanged": False,
    }
    if include_rows:
        out["rows"] = [{
            "rowId": r.row_id, "rowNo": r.row_no, "loadingAt": r.loading_at, "unloadingAt": r.unloading_at,
            "material": r.material_raw, "source": r.source_raw, "destination": r.destination_raw, "remarks": r.remarks,
            "reconciliation": None if r.row_id not in recon else {
                "status": recon[r.row_id].match_status, "confidence": recon[r.row_id].confidence,
                "fieldTripId": recon[r.row_id].field_trip_id, "wbMovementKey": recon[r.row_id].wb_movement_key,
                "reason": recon[r.row_id].reason,
            }
        } for r in rows]
    return out


@router.get("/TIOM/mis/reports")
def tiom_mis_reports(request: Request, operating_date: date | None = None, shift: str | None = None, status: str | None = None, db: Session = Depends(get_db)):
    user = get_user(db, request); require_site(db, user, "TIOM"); _tiom_mis_permission(db, user, "VIEW")
    q = select(TiomMisReport)
    if operating_date: q = q.where(TiomMisReport.operating_date == operating_date)
    if shift: q = q.where(TiomMisReport.shift == shift.upper())
    if status: q = q.where(TiomMisReport.status == status.upper())
    rows = list(db.scalars(q.order_by(TiomMisReport.operating_date.desc(), TiomMisReport.entered_at.desc()).limit(500)))
    return [_mis_report_json(db, r, False) for r in rows]


@router.get("/TIOM/mis/reports/{report_id}")
def tiom_mis_report_detail(report_id: str, request: Request, db: Session = Depends(get_db)):
    user = get_user(db, request); require_site(db, user, "TIOM"); _tiom_mis_permission(db, user, "VIEW")
    report = db.get(TiomMisReport, report_id)
    if not report: raise HTTPException(404, "MIS report not found")
    return _mis_report_json(db, report, True)


@router.post("/TIOM/mis/reports")
def create_tiom_mis_report(p: TiomMisReportIn, request: Request, db: Session = Depends(get_db)):
    csrf(request); user = get_user(db, request); require_site(db, user, "TIOM"); _tiom_mis_permission(db, user, "CREATE")
    if not db.get(Equipment, p.vehicleId): raise HTTPException(422, "Unknown TIOM vehicle/equipment ID")
    if p.operatorId and not db.get(Person, p.operatorId): raise HTTPException(422, "Unknown operator/driver employee ID")
    report_id = "MIS-" + uuid4().hex[:18].upper()
    report = TiomMisReport(
        report_id=report_id, operating_date=p.operatingDate, shift=p.shift.upper(), vehicle_id=p.vehicleId, operator_id=p.operatorId,
        opening_kmr=p.openingKmr, closing_kmr=p.closingKmr, opening_hmr=p.openingHmr, closing_hmr=p.closingHmr,
        paper_ref=p.paperRef, source_document=p.sourceDocument, notes=p.notes, status="DRAFT", entered_by=user.login_id,
    )
    db.add(report)
    for idx, item in enumerate(p.rows, start=1):
        db.add(TiomMisTripRow(
            row_id=f"{report_id}-R{idx:03d}", report_id=report_id, row_no=idx, loading_at=item.loadingAt, unloading_at=item.unloadingAt,
            material_raw=item.material, source_raw=item.source, destination_raw=item.destination, remarks=item.remarks,
        ))
    _audit(db, user, "TIOM", "TIOM_MIS_CREATE", "tiom_mis_report", report_id, after={"date": str(p.operatingDate), "shift": p.shift.upper(), "vehicle": p.vehicleId, "rows": len(p.rows)})
    db.commit(); db.refresh(report)
    return _mis_report_json(db, report, True)


@router.post("/TIOM/mis/reports/{report_id}/submit")
def submit_tiom_mis_report(report_id: str, request: Request, db: Session = Depends(get_db)):
    csrf(request); user = get_user(db, request); require_site(db, user, "TIOM"); _tiom_mis_permission(db, user, "CREATE")
    report = db.get(TiomMisReport, report_id)
    if not report: raise HTTPException(404, "MIS report not found")
    if report.status not in {"DRAFT", "REJECTED"}: raise HTTPException(409, f"Cannot submit report in {report.status} status")
    report.status = "SUBMITTED"; report.submitted_by = user.login_id; report.submitted_at = datetime.now(timezone.utc); report.version += 1
    _audit(db, user, "TIOM", "TIOM_MIS_SUBMIT", "tiom_mis_report", report_id, after={"status": report.status})
    db.commit(); return {"ok": True, "reportId": report_id, "status": report.status}


def _best_tiom_field_match(db: Session, report: TiomMisReport, row: TiomMisTripRow):
    if not row.loading_at and not row.unloading_at:
        return None, False, 0, "No loading/unloading time to match safely"
    candidates = list(db.scalars(select(LoadTrip).where(
        LoadTrip.operating_date == report.operating_date, LoadTrip.shift == report.shift, LoadTrip.vehicle_id == report.vehicle_id
    )))
    scored = []
    for t in candidates:
        score = 0.0; reasons=[]
        dl = _time_delta_seconds(row.loading_at, t.loading_start_at)
        du = _time_delta_seconds(row.unloading_at, t.unload_at)
        if dl is not None:
            if dl > 2*3600: continue
            score += max(0, 60 - dl/120)
            reasons.append(f"load Δ {int(dl//60)}m")
        if du is not None and t.unload_at is not None:
            if du > 3*3600: continue
            score += max(0, 35 - du/300)
            reasons.append(f"unload Δ {int(du//60)}m")
        if row.material_raw and t.material_id and row.material_raw.strip().upper() == str(t.material_id).strip().upper():
            score += 10; reasons.append("material id exact")
        scored.append((score, t, ", ".join(reasons)))
    scored.sort(key=lambda x: x[0], reverse=True)
    if not scored or scored[0][0] < 20:
        return None, False, 0, "No sufficiently close Field trip"
    ambiguous = len(scored) > 1 and (scored[0][0] - scored[1][0]) < 7
    return scored[0][1], ambiguous, min(95, int(scored[0][0])), scored[0][2]


def _best_tiom_wb_match(db: Session, report: TiomMisReport, row: TiomMisTripRow, field_trip: LoadTrip | None):
    if field_trip:
        link = db.scalar(select(LoadWbMatch).where(LoadWbMatch.trip_id == field_trip.trip_id))
        if link:
            movement = db.get(WbMovement, link.movement_key)
            if movement:
                return movement, min(100, int(link.confidence or 0) + 10), "Existing Field↔WB reconciliation"
    ref = row.unloading_at or row.loading_at
    if not ref:
        return None, 0, "No trip time to match WB safely"
    movements = list(db.scalars(select(WbMovement).where(
        WbMovement.operating_date == report.operating_date, WbMovement.shift == report.shift, WbMovement.vehicle_id == report.vehicle_id
    )))
    ranked=[]
    for w in movements:
        d=_time_delta_seconds(ref,w.weigh_at)
        if d is not None and d <= 4*3600:
            ranked.append((d,w))
    ranked.sort(key=lambda x:x[0])
    if not ranked:
        return None,0,"No close WB movement"
    d,w=ranked[0]
    return w,max(20,90-int(d/180)),f"nearest WB Δ {int(d//60)}m"


@router.post("/TIOM/mis/reports/{report_id}/reconcile")
def reconcile_tiom_mis_report(report_id: str, request: Request, db: Session = Depends(get_db)):
    csrf(request); user=get_user(db,request); require_site(db,user,"TIOM"); _tiom_mis_permission(db,user,"EDIT")
    report=db.get(TiomMisReport,report_id)
    if not report: raise HTTPException(404,"MIS report not found")
    rows=list(db.scalars(select(TiomMisTripRow).where(TiomMisTripRow.report_id==report_id).order_by(TiomMisTripRow.row_no)))
    summary={}
    for row in rows:
        field, ambiguous, field_conf, field_reason = _best_tiom_field_match(db, report, row)
        wb, wb_conf, wb_reason = _best_tiom_wb_match(db, report, row, None if ambiguous else field)
        if ambiguous:
            status="AMBIGUOUS"; confidence=min(field_conf,60); reason=f"Field candidates too close; {field_reason}"
            field_id=None
        elif field and wb:
            status="MATCHED_FIELD_WB"; confidence=min(100,int((field_conf+wb_conf)/2)+5); reason=f"{field_reason}; {wb_reason}"; field_id=field.trip_id
        elif field:
            status="MATCHED_FIELD"; confidence=field_conf; reason=f"{field_reason}; WB not matched"; field_id=field.trip_id
        elif wb:
            status="MATCHED_WB_ONLY"; confidence=wb_conf; reason=wb_reason; field_id=None
        else:
            status="MIS_ONLY"; confidence=0; reason=f"{field_reason}; {wb_reason}"; field_id=None
        rec=db.scalar(select(TiomMisReconciliation).where(TiomMisReconciliation.row_id==row.row_id))
        if not rec:
            rec=TiomMisReconciliation(reconciliation_id="TMR-"+uuid4().hex[:18].upper(),row_id=row.row_id,reconciled_by=user.login_id)
            db.add(rec)
        rec.field_trip_id=field_id; rec.wb_movement_key=wb.movement_key if wb else None; rec.match_status=status; rec.confidence=confidence; rec.reason=reason; rec.reconciled_by=user.login_id; rec.reconciled_at=datetime.now(timezone.utc)
        summary[status]=summary.get(status,0)+1
    _audit(db,user,"TIOM","TIOM_MIS_RECONCILE","tiom_mis_report",report_id,after={"summary":summary,"authoritativeTonnesChanged":False})
    db.commit()
    return {"ok":True,"reportId":report_id,"summary":summary,"authoritativeTonnesChanged":False,"rows":_mis_report_json(db,report,True)["rows"]}


@router.post("/TIOM/mis/reports/{report_id}/decision")
def decide_tiom_mis_report(report_id: str, p: SiteDecisionIn, request: Request, db: Session = Depends(get_db)):
    csrf(request); user=get_user(db,request); require_site(db,user,"TIOM"); _tiom_mis_permission(db,user,"APPROVE")
    report=db.get(TiomMisReport,report_id)
    if not report: raise HTTPException(404,"MIS report not found")
    if p.action not in {"APPROVE","REJECT"}: raise HTTPException(422,"Only APPROVE/REJECT are valid")
    if p.action=="APPROVE":
        bad=db.scalar(select(func.count()).select_from(TiomMisReconciliation).join(TiomMisTripRow,TiomMisTripRow.row_id==TiomMisReconciliation.row_id).where(TiomMisTripRow.report_id==report_id,TiomMisReconciliation.match_status=="AMBIGUOUS")) or 0
        if bad: raise HTTPException(409,"Resolve ambiguous MIS rows before approval")
        report.status="APPROVED"; report.approved_by=user.login_id; report.approved_at=datetime.now(timezone.utc)
    else:
        report.status="REJECTED"; report.approved_by=user.login_id; report.approved_at=datetime.now(timezone.utc)
    report.version += 1
    _audit(db,user,"TIOM","TIOM_MIS_DECISION","tiom_mis_report",report_id,after={"status":report.status},reason=p.reason)
    db.commit(); return {"ok":True,"reportId":report_id,"status":report.status,"authoritativeTonnesChanged":False}
