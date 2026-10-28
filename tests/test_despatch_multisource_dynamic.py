"""Dynamic multi-source E-Pass import regressions for the MIS Date Wise and
Transporter Wise sheets received October 27, 2026.

Synthetic test values only; no confidential original MIS exports committed.
"""
from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
import sys

from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.db import Base
from app.routers import tiom_despatch as dispatch
from app.services.tiom_despatch_source_adapters import parse_despatch_sources, normalize_header


def sample_datewise():
    book=Workbook()
    ws=book.active
    ws.title="Rpt_Minrealwise"
    ws["F2"]="Date Wise Report"
    ws["K4"]="Report On :27-Oct-2026"
    for col,val in {
        1:"Sl#", 3:"Pass#", 4:"Permit#", 7:"Stack#",
        8:"Vehicle Reg. No.", 10:"Consignee Name",
        13:"Tare Wt.", 14:"Gross Wt.", 15:"Net Wt.", 16:"Date"
    }.items():
        ws.cell(6,col,val)
    for row,epass,vehicle,wt in (
        (7,"L102602644/100","OD09Z6765",31.91),
        (8,"L102602644/101","OD09Z5233",31.92),
    ):
        values={1:row-6,3:epass,4:"L102602644",7:"J26085113340",8:vehicle,
                10:"123013638611/DNA MINERALS PVT LTD",13:15.57,
                14:round(15.57+wt,2),15:wt,
                16:"Oct 27 2026 11:15:31:910AM"}
        for col,value in values.items():ws.cell(row,col,value)
    ws["A9"]="Total :";ws["C9"]=2;ws["O9"]=63.83
    data=BytesIO();book.save(data);return data.getvalue()


def sample_transporter():
    book=Workbook()
    ws=book.active;ws.title="TransportwiseRpt"
    for col,val in {2:"Sl#",3:"Pass No",7:"Vehicle RegNo",
                    11:"Mineral Name",13:"Qty(Tons)",
                    16:"Transporting Date"}.items():
        ws.cell(8,col,val)
    ws["B9"]="Transporter Name : RIDHICON LOGISTICS PRIVATE LIMITED"
    for row,epass,vehicle,wt in (
        (10,"L102602644/100","OD09Z6765",31.91),
        (11,"L102602644/101","OD09Z5233",31.92),
    ):
        for col,val in {2:row-9,3:epass,7:vehicle,11:"Iron",13:wt,16:"27-Oct-2026"}.items():
            ws.cell(row,col,val)
    ws["B12"]="Total Qty(Tons):63.830"
    data=BytesIO();book.save(data);return data.getvalue()


def test_provider_header_detection_for_both_independent_source_layouts():
    datewise=parse_despatch_sources(sample_datewise(), "Unfamiliar Vendor Report.xlsx")
    transporter=parse_despatch_sources(sample_transporter(), "Unknown File.xlsx")
    assert len(datewise)==len(transporter)==2
    assert datewise[0]["_role"]=="detailed"
    assert transporter[0]["_role"]=="enrichment"
    assert datewise[0]["epass_no"]==transporter[0]["epass_no"]=="L102602644/100"
    assert datewise[0]["report_date"]==transporter[0]["report_date"]=="2026-10-27"
    assert datewise[0]["net_mt"]=="31.91"
    assert datewise[0]["tare_mt"]=="15.57"
    assert transporter[0]["transporter"]=="RIDHICON LOGISTICS PRIVATE LIMITED"
    assert normalize_header("  E-Pass No. ")=="epassno"


def test_import_both_sources_enriches_without_duplicate_trips():
    engine=create_engine("sqlite://",poolclass=StaticPool,connect_args={"check_same_thread":False})
    Base.metadata.create_all(engine,tables=[dispatch.DespatchTrip.__table__,dispatch.DespatchAudit.__table__])
    with Session(engine) as db:
        a=dispatch._import_data(db,sample_datewise(),"Date WIse 27-10-2026.xlsx","TEST")
        b=dispatch._import_data(db,sample_transporter(),"Transporter WIse 27-10-2026.xlsx","TEST")
        assert a["ok"] and a["inserted"]==2
        assert b["ok"] and b["inserted"]==0 and b["enriched"]==2
        rows=list(db.scalars(select(dispatch.DespatchTrip)))
        assert len(rows)==2
        assert all(x.transporter=="RIDHICON LOGISTICS PRIVATE LIMITED" for x in rows)
        assert sum(float(x.net_mt) for x in rows)==63.83
        assert dispatch._import_data(db,sample_datewise(),"Date WIse 27-10-2026.xlsx","TEST")["skippedExisting"]==2
        assert dispatch._import_data(db,sample_transporter(),"Transporter WIse 27-10-2026.xlsx","TEST")["enriched"]==0


def test_transport_manifest_does_not_create_unverified_epass():
    engine=create_engine("sqlite://",poolclass=StaticPool)
    Base.metadata.create_all(engine,tables=[dispatch.DespatchTrip.__table__,dispatch.DespatchAudit.__table__])
    with Session(engine) as db:
        output=dispatch._import_data(db,sample_transporter(),"unknown.xls","TEST")
        assert output["ok"] is False and output["inserted"]==0
        assert output["conflictCount"]==2
        assert list(db.scalars(select(dispatch.DespatchTrip)))==[]


def test_biff_legacy_xls_maps_headers_without_filename_hardcoding(monkeypatch):
    import app.services.tiom_despatch_source_adapters as adapter
    from openpyxl import load_workbook
    data=sample_datewise()
    book=load_workbook(BytesIO(data),read_only=True,data_only=True)
    rows=[list(row) for row in book.active.iter_rows(values_only=True)]
    book.close()
    class FakeWorksheet:
        name="Rpt_Minrealwise"
        nrows=len(rows)
        def row_values(self,i):return rows[i]
    class FakeWorkbook:
        def sheets(self):return [FakeWorksheet()]
        def release_resources(self):pass
    monkeypatch.setitem(sys.modules,"xlrd",SimpleNamespace(
        open_workbook=lambda **kwargs:FakeWorkbook()))
    parsed=adapter.parse_despatch_sources(bytes.fromhex("D0CF11E0A1B11AE1")+b"fake", "source_export.xls")
    assert len(parsed)==2
    assert parsed[0]["report_date"]=="2026-10-27"


def test_transport_name_html_entities_are_normalized():
    book=Workbook();ws=book.active
    ws["B8"]="Pass No";ws["G8"]="Vehicle RegNo";ws["M8"]="Qty(Tons)"
    ws["P8"]="Transporting Date";ws["B9"]="Transporter Name : M/S MAA KUANRI LOADING &amp; TRANSPORT"
    ws["C8"]="Pass No";ws["B8"]="Sl#"
    ws["C10"]="L102602644/100";ws["G10"]="OD09Z6765";ws["M10"]=31.91;ws["P10"]="27-Oct-2026"
    x=BytesIO();book.save(x)
    assert parse_despatch_sources(x.getvalue(),"other supplier.xlsx")[0]["transporter"]=="M/S MAA KUANRI LOADING & TRANSPORT"


def test_gmail_recovery_query_includes_previously_reviewed_messages():
    from app.services.tiom_despatch_mail import _gmail_query
    query=_gmail_query({"rajugupta97093@gmail.com"},recovery=True)
    assert "label:NMTPL_DESPATCH_REVIEW_REQUIRED" in query
    assert "label:NMTPL_DESPATCH_PROCESSED" in query
