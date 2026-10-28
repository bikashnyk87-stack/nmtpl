"""Keep accepted 1–26 Oct Despatch untouched and process 27 Oct forward."""
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock
import uuid

from openpyxl import Workbook
from sqlalchemy import select

from app.services import tiom_despatch_mail as mail
from app.services.tiom_despatch_source_adapters import parse_despatch_sources
from tests.test_despatch_archive_recheck import _isolated_db, _review
from tests.test_despatch_multisource_dynamic import sample_datewise


def test_review_queue_filters_closed_period_but_keeps_archives(monkeypatch):
    make = _isolated_db(monkeypatch)
    with make() as db:
        db.add_all([
            _review("old", "DATEWISE 26-10-2026.xls", b"old", date(2026,10,26)),
            _review("old-undated", "DATEWISE 25-10-2026.xls", b"old", None),
            _review("current", "Date Wise 27-10-2026.xls", b"new", date(2026,10,27)),
            _review("current-undated", "Transporter Wise 28-10-2026.xls", b"new", None),
        ])
        db.commit()
        rows=mail.open_despatch_reviews(db)
        assert {x.message_id for x in rows} == {"current", "current-undated"}
        assert len(list(db.scalars(select(mail.DespatchMailAttachment)))) == 4


def test_gmail_search_starts_from_unfinished_reporting_period():
    q=mail._gmail_query({"rajugupta97093@gmail.com"})
    assert "after:2026/10/26" in q
    assert "after:2026/09/30" not in q


def test_gmail_scan_does_not_touch_26_oct_record(monkeypatch):
    msg={
        "payload": {"headers": [
            {"name":"From","value":"Bansilal Gupta <rajugupta97093@gmail.com>"},
            {"name":"Subject","value":"Dispatch report on dt. 26.10.26"},
        ],"parts":[]},
    }
    svc=Mock()
    svc.users.return_value.messages.return_value.get.return_value.execute.return_value=msg
    db=Mock()
    db.get.return_value=None
    result={"skipped":0}
    assert mail._scan_message(
        svc,db,"old-msg",{"rajugupta97093@gmail.com"},result
    ) is None
    assert result["skipped"]==1
    db.commit.assert_not_called()


def test_datewise_numeric_count_footer_not_treated_as_epass():
    source=sample_datewise()
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(source))
    ws=wb.active
    # Real MIS "Pass#" subtotal cell is often 532 without a date. This
    # must not abort import of the actual E-Pass rows.
    ws["C9"]=532
    dst=BytesIO()
    wb.save(dst)
    rows=parse_despatch_sources(dst.getvalue(), "Date Wise 27-10-2026.xlsx")
    assert len(rows)==2
    assert all(x["epass_no"].startswith("L102602644/") for x in rows)


def test_only_current_issues_visible_on_frontend():
    from pathlib import Path
    js=(Path(__file__).parents[1]/"app"/"static"/"tiom_despatch.js").read_text(encoding="utf-8")
    assert "New Despatch exceptions · 27 Oct onward" in js
    assert "through 26 Oct are already accepted" in js
    assert "Reconcile saved Gmail files" not in js
