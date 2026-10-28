"""Regression coverage for formerly marked-processed legacy .xls MIS Despatch emails."""
import io
import sys
from types import SimpleNamespace
from unittest.mock import Mock
from openpyxl import Workbook
import pytest

from app.routers.tiom_despatch import _extract
from app.services import tiom_despatch_mail as mail


def _wb_bytes(sheet_name="27.10.26", *, detailed=True):
    wb=Workbook()
    ws=wb.active
    ws.title=sheet_name
    ws.append(["Trip Summary"])
    ws.append(["Sl.No", "Permit No.", "Permit Qty.", "Buyer", "E-Pass No",
               "Vehicle", "Transporter", "Destination", "Material", "Stack", "Consignee",
               "Tare Wt.", "Gross Wt.", "Net Wt.", "Dispatch Date", "Time"] if detailed
              else ["Date", "Transporter", "Trips", "Net MT"])
    ws.append([1, "PERMIT-100", 500, "AMIPL", "EP-2710-1", "OD09AB1234",
               "TR-1", "PLANT", "FINES", "S1", "AMIPL", 10, 55, 45, "27.10.26", "06:15"]
              if detailed else ["27.10.26", "TR-1", 100, 4500])
    buffer=io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_xlsx_named_xls_reads_real_epass_rows():
    rows=_extract(_wb_bytes(), "Trip Summary 27-10-2026.xls")
    assert len(rows)==1
    assert rows[0]["epass_no"]=="EP-2710-1"
    assert rows[0]["report_date"]=="2026-10-27"


def test_legacy_biff_xls_reader_accepts_dated_epass_sheets(monkeypatch):
    src=_wb_bytes()
    from openpyxl import load_workbook
    original=load_workbook(io.BytesIO(src),read_only=True,data_only=True).active
    rows=[list(r) for r in original.iter_rows(values_only=True)]
    class FakeSheet:
        name="27.10.26"
        nrows=len(rows)
        def row_values(self,i): return rows[i]
    class FakeBook:
        def sheets(self): return [FakeSheet()]
        def release_resources(self): pass
    dummy=SimpleNamespace(open_workbook=lambda **kwargs: FakeBook())
    monkeypatch.setitem(sys.modules,"xlrd",dummy)
    result=_extract(bytes.fromhex("D0CF11E0A1B11AE1")+b"TEST", "Trip Summary.xls")
    assert len(result)==1
    assert result[0]["epass_no"]=="EP-2710-1"


def test_aggregate_datewise_report_is_never_invented_as_epass_trips():
    with pytest.raises(ValueError, match="aggregate"):
        _extract(_wb_bytes(detailed=False), "Date Wise Report 27-10-2026.xls")


def test_processed_label_query_excludes_review_and_has_recovery_mode():
    senders={"rajugupta97093@gmail.com"}
    new=mail._gmail_query(senders)
    backfill=mail._gmail_query(senders,recovery=True)
    assert "-label:NMTPL_DESPATCH_PROCESSED" in new
    assert "-label:NMTPL_DESPATCH_REVIEW_REQUIRED" in new
    assert "label:NMTPL_DESPATCH_PROCESSED" in backfill
    assert "from:rajugupta97093@gmail.com" in new


def test_legacy_no_matching_data_status_becomes_reviewable_not_imported(monkeypatch):
    prior=SimpleNamespace(status="NO_MATCHING_DATA",details="",checked_at=None)
    db=Mock()
    db.get.return_value=prior
    msg={"payload":{"headers":[{"name":"From","value":"Bansilal Gupta <rajugupta97093@gmail.com>"},
                              {"name":"Subject","value":"Dispatch report on dt. 27.10.26"}],
                    "parts":[{"filename":"Date Wise Report.xls","mimeType":"application/vnd.ms-excel",
                              "body":{"data":"AA=="}}]}}
    service=Mock()
    service.users.return_value.messages.return_value.get.return_value.execute.return_value=msg
    monkeypatch.setattr(mail,"_attachment",lambda *args: _wb_bytes(detailed=False))
    monkeypatch.setattr(mail,"_archive_for_review",lambda *args: True)
    monkeypatch.setattr(mail,"_import_data",lambda *args: (_ for _ in ()).throw(ValueError("No E-Pass lines")))
    result={"tripsInserted":0,"imagesSaved":0,"attachmentsForReview":0,
            "messagesImported":0,"recoveredEmails":0,"skipped":0}
    status=mail._scan_message(service,db,"msg-27",{"rajugupta97093@gmail.com"},result,recovery=True)
    assert status=="REVIEW_REQUIRED"
    assert prior.status=="REVIEW_REQUIRED"
    assert result["tripsInserted"]==0
    assert result["attachmentsForReview"]==1
    assert result["recoveredEmails"]==1
    assert result["messagesImported"]==0
    db.commit.assert_called_once()


def test_review_ui_keeps_original_xls_download_separate_from_trip_ledger():
    from pathlib import Path
    source=(Path(__file__).resolve().parents[1]/"app"/"static"/"tiom_despatch.js").read_text(encoding="utf-8")
    assert "api('mail-attachments')" in source
    assert "mail-attachment/" in source
    assert "Only verified E-Pass rows are imported" in source
