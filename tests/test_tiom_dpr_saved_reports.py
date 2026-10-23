"""End-user DPR checks: Saved Reports UI, shifts, MTD and downloadable workbook."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.db import Base
from app.models import ShiftMaster, WbImportBatch, WbMovement, Equipment
from app.routers import webapp
from app.services.tiom_dpr import build_dpr_workbook
from app.site_models import TiomDrillSet, TiomDrillingShift


def rec(day,sh,rom,final=0):
    return {"date":day,"shift":sh,"metrics":{
        "ROM":{"qty":rom,"trips":2,"source":"WB"},
        "WASTE":{"qty":0,"trips":None,"source":"—"},
        "TOTAL_EXCAVATION":{"qty":rom,"trips":None,"source":"CALC"},
        "TOTAL_PRODUCTION":{"qty":final,"trips":None,"source":"CALC"},
        "SCREEN_FINES":{"qty":final,"trips":1,"source":"WB"}},
        "wbValid":2,"wbReview":0,"productionStatus":"SUBMITTED",
        "drilling":[{"label":"ELGI-1","holes":5,"meterage":32,
            "drillHours":2,"metersPerHour":16,"remarks":"Checked"}],
        "drillHoles":5,"drillMetres":32,"drillHours":2,
        "plantHours":{"MSP-3":1.5},"plantProduction":{"MSP-3":final},
        "quality":[],"warnings":[]}


def sample():
    a=rec("2026-10-11","A",100.0,70.0)
    b=rec("2026-10-11","B",120.0,50.0)
    c=rec("2026-10-11","C",198.16,80.0)
    return {"dates":["2026-10-11"],"shifts":["A","B","C"],"shift":"ALL",
        "fromDate":"2026-10-11","toDate":"2026-10-11",
        "entries":[a,b,c],
        "mtdEntries":[rec("2026-10-01","A",40.0,25.0),a,b,c]}


def test_dpr_excel_dates_sections_formulas_and_mtd():
    wb=load_workbook(build_dpr_workbook(sample()))
    assert wb.sheetnames==["Summary","11.10.26","Stack Summary","Stacking",
                           "Drilling Detail","Source Data","Data Quality"]
    ws=wb["11.10.26"]
    assert ws["A5"].value=="ROM"
    assert ws["E5"].value==100.0
    assert ws["R5"].value==120.0
    assert ws["AE5"].value==198.16
    assert ws["AR5"].data_type=="f"
    assert ws["AR5"].value=="=SUM(E5,R5,AE5)"
    assert ws["BA5"].value=="Fresh ROM"
    assert ws["BF5"].value==pytest.approx(418.16)
    assert ws["BI5"].value is None or isinstance(ws["BI5"].value,str)  # no false hardcoded opening stock
    assert ws["AG32"].value=="ELGI-1"
    assert ws["AM32"].value==2
    assert wb["Summary"]["G4"].value==100.0
    assert wb["Drilling Detail"]["C4"].value=="ELGI-1"
    assert wb["Stack Summary"]["G4"].value is None  # no invented dispatch MT
    assert wb["Stack Summary"]["H4"].value is None  # no invented closing stock
    assert "#REF!" not in str(ws["AR5"].value)
    assert ws.page_setup.orientation=="landscape"


def test_missing_shift_data_remains_blank_not_zero():
    payload=sample()
    missing=rec("2026-10-11","B",0)
    missing["metrics"]["ROM"]["qty"]=None
    missing["metrics"]["TOTAL_EXCAVATION"]["qty"]=None
    missing["quality"]=["No saved production data"]
    missing["drilling"]=[]
    payload["entries"][1]=missing
    payload["mtdEntries"][2]=missing
    wb=load_workbook(build_dpr_workbook(payload))
    assert wb["11.10.26"]["R5"].value is None
    assert wb["Data Quality"]["D4"].value=="No saved production data"


def test_restricted_shift_and_invalid_date_range():
    engine=create_engine("sqlite://",connect_args={"check_same_thread":False},
                         poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        db.add_all([
            ShiftMaster(shift="A",start_time=datetime.strptime("05:00","%H:%M").time(),
                end_time=datetime.strptime("13:00","%H:%M").time(),active=True,scheduled_hours=8),
            ShiftMaster(shift="B",start_time=datetime.strptime("13:00","%H:%M").time(),
                end_time=datetime.strptime("21:00","%H:%M").time(),active=True,scheduled_hours=8)])
        db.flush()
        u=SimpleNamespace(admin=False,modules="PRODUCTION",shifts="A")
        _,_,allowed,requested=webapp._tiom_dpr_inputs(db,u,
            {"fromDate":"2026-10-11","toDate":"2026-10-11","shift":"ALL"})
        assert allowed==["A"] and requested=="ALL"
        from fastapi import HTTPException
        with pytest.raises(HTTPException) as err:
            webapp._tiom_dpr_inputs(db,u,
                {"fromDate":"2026-10-11","toDate":"2026-10-11","shift":"B"})
        assert err.value.status_code==403
        with pytest.raises(HTTPException) as err:
            webapp._tiom_dpr_inputs(db,u,
                {"fromDate":"2026-10-13","toDate":"2026-10-11","shift":"ALL"})
        assert err.value.status_code==422
    engine.dispose()


def test_http_end_user_flow_exports_xlsx(monkeypatch):
    engine=create_engine("sqlite://",connect_args={"check_same_thread":False},
                         poolclass=StaticPool)
    Base.metadata.create_all(engine)
    data=sample()
    app=FastAPI()
    app.include_router(webapp.router)
    def session():
        with Session(engine) as db:yield db
    app.dependency_overrides[webapp.get_db]=session
    monkeypatch.setattr(webapp,"get_user",lambda db,request:
                        SimpleNamespace(admin=True,modules="PRODUCTION",shifts="ALL"))
    monkeypatch.setattr(webapp,"_tiom_dpr_data",lambda db,start,end,shifts,request:data)
    with TestClient(app) as client:
        response=client.get("/api/web/tiom/dpr.xlsx",params={
            "fromDate":"2026-10-11","toDate":"2026-10-11","shift":"ALL"})
        assert response.status_code==200,response.text[:500]
        assert response.content.startswith(b"PK")
        assert "TIOM_DPR_2026-10-11_to_2026-10-11_ALL.xlsx" in response.headers["content-disposition"]
        assert load_workbook(BytesIO(response.content)).sheetnames[1]=="11.10.26"
        assert client.get("/api/web/tiom/dpr.xlsx",params={
            "fromDate":"2026-10-15","toDate":"2026-10-11","shift":"ALL"}).status_code==422
    engine.dispose()


def test_non_destructive_saved_reports_menu_and_mtd_action():
    root=Path(__file__).parents[1]
    js=(root/"app/static/tiom_phase1.js").read_text("utf-8")
    api=(root/"app/routers/webapp.py").read_text("utf-8")
    for term in ["Daily Production Report (DPR)","SAVED DRIVER SHIFT REPORTS",
                 "Download Detailed Entry Report","setDprMtd","loadSavedDprPreview",
                 "downloadSavedDpr","/api/web/tiom/dpr.xlsx",
                 "getTiomDprPreview","getUiState('SAVED_SECTION'"]:
        assert term in js
    assert "if method == 'getTiomDprPreview'" in api
    assert "'getTiomDprPreview'" in api
    assert "/tiom/mis-entry-export.xlsx" in api
