"""TIOM despatch end-user checks: import, cancel/restore, reports and sender maintenance."""
from __future__ import annotations

from datetime import date
from io import BytesIO
from types import SimpleNamespace

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.db import Base
from app.routers import tiom_despatch as mod
from app.services.tiom_despatch_mail import DespatchEmailProcessed


def client():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine, tables=[
        mod.DespatchAudit.__table__, mod.DespatchDocument.__table__,
        mod.DespatchPermitDay.__table__, mod.DespatchSender.__table__,
        mod.DespatchTrip.__table__, DespatchEmailProcessed.__table__,
    ])
    app = FastAPI()
    app.include_router(mod.router)
    def session():
        with Session(engine) as db:
            yield db
    app.dependency_overrides[mod.get_db] = session
    orig_get_user, orig_require, orig_csrf = mod.get_user, mod.require, mod.csrf
    mod.get_user = lambda db, request: SimpleNamespace(login_id="TEST_ADMIN", admin=True, modules="PRODUCTION")
    mod.require = lambda *args, **kwargs: None
    mod.csrf = lambda request: None
    return TestClient(app), engine, (orig_get_user, orig_require, orig_csrf)


def csv_data():
    rows = [
        "report_date,permit_no,permit_qty,buyer,epass_no,vehicle_no,transporter,destination,material_grade,stack_no,consignee,tare_mt,gross_mt,net_mt,dispatch_date,dispatch_time,status,source",
        "2026-10-01,L1,100,Self,L1/1,OD1,AAA,Dabuna,S.Fines(60-62),ST1,AMNS,15,45,30,2026-10-01,10:00,ACTIVE,TEST",
        "2026-10-26,L1,100,Self,L1/2,OD2,BBB,Dabuna,S.Fines(60-62),ST1,AMNS,15,47,32,2026-10-26,10:00,ACTIVE,TEST"
    ]
    return ("\n".join(rows)+"\n").encode()


def test_end_to_end_despatch_management():
    api, engine, old = client()
    try:
        r=api.get("/api/tiom/despatch/senders")
        assert r.status_code == 200 and len(r.json()) == 9
        r=api.post("/api/tiom/despatch/senders",json={
            "email": "new@example.com", "name": "New", "active": True})
        assert r.status_code == 200
        assert len(api.get("/api/tiom/despatch/senders").json()) == 10
        r=api.post("/api/tiom/despatch/upload",files={"file":("oct.csv",csv_data(),"text/csv")})
        assert r.status_code == 200 and r.json()["inserted"] == 2
        r=api.post("/api/tiom/despatch/upload",files={"file":("oct.csv",csv_data(),"text/csv")})
        assert r.status_code == 200 and r.json()["skippedExisting"] == 2
        dash=api.get("/api/tiom/despatch/dashboard").json()
        assert dash["totals"]["trips"] == 2 and dash["totals"]["tonnes"] == 62
        r=api.get("/api/tiom/despatch/trips?show_cancelled=true").json()
        assert r["count"] == 2
        r=api.post("/api/tiom/despatch/trip/L1%2F2/edit",json={
            "status":"CANCELLED","reason":"Verified permit cancellation"
        })
        assert r.status_code == 200, r.text
        dash=api.get("/api/tiom/despatch/dashboard").json()
        assert dash["totals"]["trips"] == 1 and dash["totals"]["tonnes"] == 30
        assert dash["totals"]["cancelled"] == 1
        r=api.post("/api/tiom/despatch/trip/L1%2F2/edit",json={
            "status":"ACTIVE","reason":"Cancellation reversed by AMNS"
        })
        assert r.status_code == 200, r.text
        dash=api.get("/api/tiom/despatch/dashboard").json()
        assert dash["totals"]["trips"] == 2 and dash["totals"]["tonnes"] == 62
        for kind in ("trip","transporter","material","destination","despatch"):
            resp=api.get("/api/tiom/despatch/export/"+kind)
            assert resp.status_code == 200, (kind, resp.text)
            wb=load_workbook(BytesIO(resp.content),read_only=True)
            ws=wb.active
            assert ws.max_row >= 6
            if kind == "trip":
                assert ws["N1"].value == 2
                assert ws["P1"].value == 62
                assert ws["P6"].value == "10:00"
            else:
                assert ws.cell(ws.max_row,ws.max_column).value == 62
            wb.close()
        with Session(engine) as db:
            actions=[x.action for x in db.scalars(select(mod.DespatchAudit)).all()]
            assert actions.count("CORRECT_TRIP") == 2
            assert db.scalar(select(func.count()).select_from(mod.DespatchTrip)) == 2
    finally:
        mod.get_user, mod.require, mod.csrf = old
        engine.dispose()


def test_non_admin_rejected_for_trip_mutation():
    api, engine, old = client()
    try:
        mod.get_user=lambda db, request: SimpleNamespace(login_id="TEST_READ", admin=False, modules="PRODUCTION")
        r=api.post("/api/tiom/despatch/trip/doesnotexist/edit",json={
            "status":"CANCELLED","reason":"No unauthorized cancellation"
        })
        assert r.status_code == 403
    finally:
        mod.get_user, mod.require, mod.csrf = old
        engine.dispose()
