"""26 October 2026 acceptance figures cross-checked against client source workbooks."""
from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

from app.services.tiom_despatch_client_layout import destination_bucket_index
from app.services.tiom_despatch_client_reports import build_client_workbook


DAY=date(2026,10,26)


def group(count, weight, buyer, destination, grade, prefix):
    millitons=int(Decimal(weight)*1000)
    per,extra=divmod(millitons,count)
    trips=[]
    for i in range(count):
        net=Decimal(per+(i<extra))/1000
        trips.append(SimpleNamespace(
            report_date=DAY, permit_no="L"+prefix, permit_qty=Decimal("15000"),
            buyer=buyer, epass_no=f"{prefix}/{i+1}", vehicle_no=f"OD09X{i:04d}",
            transporter="RIDHICON LOGISTICS PRIVATE LIMITED",
            destination=destination, material_grade=grade, stack_no="ST1",
            consignee=buyer, tare_mt=Decimal("15"), gross_mt=Decimal("15")+net,
            net_mt=net, dispatch_time="13:00", status="ACTIVE"))
    return trips


@pytest.fixture
def amns_day():
    return (
        group(190,"5791.460","Self","Dabuna","S.Fines(60-62)","SELF")
       +group(12,"373.280","DNA MINERALS PVT LTD","Chandil","5-18mm(60-62)","DNA")
       +group(3,"97.950","S N N MINERALS PRIVATE LIMITED.","Chandil","5-18mm(60-62)","SNN")
       +group(19,"606.580","PH ASSOCIATES PRIVATE LIMITED","ULIBURU","5-18mm(60-62)","PH")
    )


def test_26_oct_source_totals_and_bucket_mapping(amns_day):
    assert len(amns_day)==224
    assert sum((x.net_mt for x in amns_day),Decimal(0))==Decimal("6869.270")
    from collections import Counter
    slots=Counter(destination_bucket_index(x) for x in amns_day)
    assert slots=={6:190,41:12,23:3,30:19}


@pytest.mark.parametrize("type",["trip","transporter","material","destination","despatch"])
def test_client_report_against_26_oct_reference(amns_day,type):
    output, totals=build_client_workbook(type,amns_day,DAY,DAY)
    assert totals["trips"]==224
    assert totals["quantityMt"]==6869.27
    wb=load_workbook(BytesIO(output.getvalue()),read_only=True)
    if type=="trip":
        assert wb.sheetnames==["26.10.26","Cummulative","Cancellation"]
        assert wb["26.10.26"]["N1"].value==224
        assert wb["26.10.26"]["P1"].value==6869.27
        assert wb["Cummulative"]["N1"].value==224
    elif type=="transporter":
        ws=wb["26.10.26"]
        assert ws["C9"].value==224
        assert ws["D9"].value==6869.27
        assert ws["C10"].value==224
        assert ws["D10"].value==6869.27
        assert ws["O35"].value == "Trips"  # same-sheet cumulative header
    elif type=="material":
        ws=wb["Oct-26"]
        assert ws["H28"].value==5791.46
        assert ws["M28"].value==1077.81
        assert ws["O28"].value=="=SUM(B28:N28)"
    elif type=="destination":
        ws=wb["Oct-26"]
        assert ws["N31"].value==190
        assert ws["O31"].value==5791.46
        assert ws["AV31"].value==3
        assert ws["AW31"].value==97.95
        assert ws["BJ31"].value==19
        assert ws["BK31"].value==606.58
        assert ws["CF31"].value==12
        assert ws["CG31"].value==373.28
    elif type=="despatch":
        ws=wb["DESTINATION WISE"]
        assert ws["F29"].value==5791.46
        assert ws["T29"].value==1077.81
        assert ws["V29"].value.startswith("=SUM(B29:U29)")
    wb.close()


def test_unmapped_source_blocks_client_exports(amns_day):
    amns_day[0].buyer="UNVERIFIED CLIENT"
    with pytest.raises(ValueError,match="unmapped"):
        build_client_workbook("destination",amns_day,DAY,DAY)


def test_transporter_html_amp_alias_is_one_row_without_double_counting():
    """Regression: 14 trips/446.030 MT appeared twice as '&' and '&amp;'."""
    first=group(7,"223.015","Self","Dabuna","S.Fines(60-62)","MAA-A")
    second=group(7,"223.015","Self","Dabuna","S.Fines(60-62)","MAA-B")
    for trip in first:
        trip.transporter="M/S MAA KUANRI LOADING & TRANSPORT"
    for trip in second:
        trip.transporter="M/S MAA KUANRI LOADING &amp; TRANSPORT"
    records=first+second
    assert len({trip.epass_no for trip in records})==14
    result, totals=build_client_workbook("transporter",records,DAY,DAY)
    assert totals["trips"]==14
    assert totals["quantityMt"]==446.03
    wb=load_workbook(BytesIO(result.getvalue()),read_only=True)
    ws=wb["26.10.26"]
    assert ws["B28"].value=="M/S MAA KUANRI LOADING & TRANSPORT"
    assert ws["B29"].value=="Grand Total"  # no second transporter row
    assert ws["O28"].value==14   # Self / Dabuna / S.Fines(60-62) trips
    assert ws["P28"].value==446.03
    assert ws["C9"].value==14
    assert ws["D9"].value==446.03
    wb.close()


@pytest.mark.parametrize("report_type,cell,sheet",[
    ("trip","N1","Cancellation"),
    ("transporter","C28","26.10.26"),
    ("material","B28","Oct-26"),
    ("destination","B31","Oct-26"),
    ("despatch","B29","DESTINATION WISE"),
])
def test_numeric_zeros_are_visually_blank_but_still_zero(report_type,cell,sheet,amns_day):
    """Custom Excel number format hides displayed zeros without changing formulas or totals."""
    output,_=build_client_workbook(report_type,amns_day,DAY,DAY)
    wb=load_workbook(BytesIO(output.getvalue()),read_only=True)
    c=wb[sheet][cell]
    assert c.value==0, (report_type,cell,c.value)
    assert c.number_format.split(";")[2]=='""', (report_type,cell,c.number_format)
    wb.close()
