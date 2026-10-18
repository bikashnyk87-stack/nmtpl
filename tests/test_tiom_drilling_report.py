"""Regression checks for the supervisor's drilling report and Excel download."""
from pathlib import Path

from openpyxl import load_workbook

from app.services.tiom_drilling_report import drilling_workbook

ROOT = Path(__file__).resolve().parents[1]


def _summary():
    day={"hmr":7.5, "holes":31, "meterage":276.0, "metersPerHour":36.8}
    month={"hmr":67.9, "holes":228, "meterage":1688.0, "metersPerHour":24.86}
    return {
        "date":"2026-10-17","monthFrom":"2026-10-01",
        "rows":[{"drillSetId":"DM-01","label":"DM-01",
                 "day":day,"month":month,
                 "remarks":[{"shift":"B","text":"Diesel supply 1 hr"}]}],
        "dayTotals":day, "monthTotals":month,
    }


def _details():
    return [{
        "date":"2026-10-17","shift":"B","drillSetId":"DM-01",
        "label":"DM-01","drillMachine":"DM-01","compressor":"ELGI-1",
        "source":"17 Ha","drillOpenHmr":100,"drillCloseHmr":107.5,
        "drillHours":7.5,"compressorOpenHmr":200,"compressorCloseHmr":207,
        "compressorHours":7,"holes":31,"romObHoles":20,"bhjBhqHoles":11,
        "romObMeterage":180,"bhjBhqMeterage":96,"totalMeterage":276,
        "metersPerHour":36.8,"breakdownHours":1,
        "remarks":"Diesel supply 1 hr","enteredBy":"test","enteredAt":"2026-10-17",
    }]


def test_drilling_excel_layout_and_all_report_sheets():
    stream=drilling_workbook(_summary(),_details(),
        {"fromDate":"2026-10-01","toDate":"2026-10-17","shift":"ALL"})
    wb=load_workbook(stream)
    assert wb.sheetnames==["Drilling Report","Shift Details","ROM-OB Summary","BHJ-BHQ Summary"]
    ws=wb.worksheets[0]
    assert ws["A1"].value=="Drilling Report"
    assert str(ws["C3"].value)=="For The Day"
    assert str(ws["G3"].value)=="For The Month"
    assert ws["B5"].value=="DM-01"
    assert ws["C5"].value==7.5
    assert ws["D5"].value==31
    assert ws["E5"].value==276.0
    assert ws["F5"].data_type=="f"
    assert ws["G5"].value==67.9
    assert ws["I5"].value==1688.0
    assert "Diesel supply 1 hr" in ws["C7"].value
    assert ws.page_setup.orientation=="landscape"
    assert wb["Shift Details"]["U2"].value=="Diesel supply 1 hr"
    assert wb["ROM-OB Summary"]["C2"].value==20
    assert wb["BHJ-BHQ Summary"]["D2"].value==96


def test_missing_hmr_is_not_zero_and_never_divides_partial_hours():
    report=_summary()
    report["rows"][0]["day"]={"hmr":None,"holes":31,
                              "meterage":276.0,"metersPerHour":None}
    report["dayTotals"]=report["rows"][0]["day"]
    wb=load_workbook(drilling_workbook(report,_details(),
        {"fromDate":"2026-10-01","toDate":"2026-10-17","shift":"ALL"}))
    ws=wb["Drilling Report"]
    assert ws["C5"].value is None
    assert ws["F5"].value is None
    assert ws["F6"].value is None
    assert ws["D5"].value==31


def test_drilling_screen_and_export_are_connected():
    api=(ROOT/"app/routers/webapp.py").read_text(encoding="utf-8")
    js=(ROOT/"app/static/tiom_phase1.js").read_text(encoding="utf-8")
    assert "def _tiom_drilling_report_table" in api
    assert "/tiom/drilling-report.xlsx" in api
    assert "'reportTable':_tiom_drilling_report_table(db,day)" in api
    assert "downloadAllDrillingReports" in js
    assert "renderDrillingPublishedReport(d.reportTable)" in js
    assert "Shift / stoppage details" in js
