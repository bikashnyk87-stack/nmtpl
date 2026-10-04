from __future__ import annotations

import base64
import hashlib
import json
import logging
from logging.handlers import RotatingFileHandler
import os
from pathlib import Path
import re
import sys
import time
from datetime import date, datetime, time as dtime, timedelta, timezone
from decimal import Decimal
from email.utils import parseaddr
from io import BytesIO
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from openpyxl import load_workbook
from sqlalchemy import delete, select

from app.db import SessionLocal
from app.models import (
    AuditLog, Equipment, LoadWbMatch, ShiftMaster, VehicleAlias,
    WbImportBatch, WbMovement,
)
from app.services.reconcile import auto_reconcile
from app.services.time_context import TZ, now_local
from app.services.wb_mapping import (
    find_wb_sheet, row_value, build_location_resolver, ensure_wb_header_mapping,
    expected_movement_date, operating_date_from_movement,
)
from app.services.tiom_location_erp import canonicalize_wb_rows

CONFIG_PATH = ROOT / "config" / "wb_gmail_config.json"
TOKEN_PATH = ROOT / "config" / "gmail_token.json"
LOG_DIR = ROOT / "logs" / "wb_gmail"
LOG_PATH = LOG_DIR / "worker.log"
STATUS_PATH = LOG_DIR / "status.json"
ARCHIVE_DIR = ROOT / "data" / "wb_email_archive"
QUARANTINE_DIR = ROOT / "data" / "wb_quarantine"
SCOPES = ["https://www.googleapis.com/auth/gmail.modify"]
ENGINE_ACTOR = "WB_GMAIL_ENGINE"

LOG_DIR.mkdir(parents=True, exist_ok=True)
ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)
QUARANTINE_DIR.mkdir(parents=True, exist_ok=True)

logger = logging.getLogger("wb_gmail")
logger.setLevel(logging.INFO)
if not logger.handlers:
    fh = RotatingFileHandler(LOG_PATH, maxBytes=5_000_000, backupCount=5, encoding="utf-8")
    fh.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    logger.addHandler(fh)
    sh = logging.StreamHandler()
    sh.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    logger.addHandler(sh)


def write_status(**kwargs):
    payload = {"updated_at": datetime.now(timezone.utc).isoformat(), **kwargs}
    tmp = STATUS_PATH.with_suffix(".tmp")
    tmp.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")
    tmp.replace(STATUS_PATH)


def load_config():
    if not CONFIG_PATH.exists():
        raise RuntimeError(f"Missing configuration: {CONFIG_PATH}")
    cfg = json.loads(CONFIG_PATH.read_text(encoding="utf-8-sig"))
    required = [
        "enabled", "poll_seconds", "gmail_query", "allowed_senders",
        "processed_label", "duplicate_label", "error_label",
        "review_required_label", "review_rows_label",
        "max_review_percent", "max_file_mb", "start_after_utc",
    ]
    missing = [k for k in required if k not in cfg]
    if missing:
        raise RuntimeError("Missing config keys: " + ", ".join(missing))
    return cfg


def save_token(creds):
    TOKEN_PATH.write_text(creds.to_json(), encoding="utf-8")


def gmail_service():
    if not TOKEN_PATH.exists():
        raise RuntimeError("Gmail is not authorized. Run authorize_wb_gmail.bat once.")
    creds = Credentials.from_authorized_user_file(str(TOKEN_PATH), SCOPES)
    if creds.expired and creds.refresh_token:
        creds.refresh(Request())
        save_token(creds)
    if not creds.valid:
        raise RuntimeError("Gmail token is invalid. Run authorize_wb_gmail.bat again.")
    return build("gmail", "v1", credentials=creds, cache_discovery=False)


def ensure_labels(service, names):
    existing = service.users().labels().list(userId="me").execute().get("labels", [])
    by_name = {x["name"]: x["id"] for x in existing}
    result = {}
    for name in names:
        if name not in by_name:
            created = service.users().labels().create(
                userId="me",
                body={"name": name, "labelListVisibility": "labelShow", "messageListVisibility": "show"},
            ).execute()
            by_name[name] = created["id"]
        result[name] = by_name[name]
    return result


def sender_allowed(address, rules):
    address = (address or "").strip().lower()
    if not rules:
        return False
    for raw in rules:
        rule = str(raw or "").strip().lower()
        if not rule:
            continue
        if rule.startswith("*@"):
            if address.endswith(rule[1:]):
                return True
        elif rule.startswith("@"):
            if address.endswith(rule):
                return True
        elif address == rule:
            return True
    return False


def subject_allowed(subject, keywords):
    words = [str(x).strip().lower() for x in (keywords or []) if str(x).strip()]
    if not words:
        return True
    s = (subject or "").lower()
    return any(w in s for w in words)


def header_value(message, name):
    for h in message.get("payload", {}).get("headers", []):
        if h.get("name", "").lower() == name.lower():
            return h.get("value", "")
    return ""


def walk_parts(part):
    yield part
    for child in part.get("parts", []) or []:
        yield from walk_parts(child)


def xlsx_attachments(service, message):
    found = []
    for part in walk_parts(message.get("payload", {})):
        filename = (part.get("filename") or "").strip()
        if not filename.lower().endswith(".xlsx"):
            continue
        body = part.get("body", {}) or {}
        data = body.get("data")
        if not data and body.get("attachmentId"):
            data = service.users().messages().attachments().get(
                userId="me", messageId=message["id"], id=body["attachmentId"]
            ).execute().get("data")
        if not data:
            continue
        raw = base64.urlsafe_b64decode(data.encode("ascii"))
        found.append((filename, raw))
    return found


def safe_name(name):
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", name or "WB.xlsx").strip("._")
    return (cleaned or "WB.xlsx")[:180]


def norm_header(value):
    return re.sub(r"[^A-Z0-9]+", "", str(value or "").upper())


def norm_vehicle(value):
    return re.sub(r"[^A-Z0-9]+", "", str(value or "").upper())


def decimal_value(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        return Decimal(str(value).replace(",", "").strip())
    except Exception:
        return None


def clock_value(value):
    if isinstance(value, datetime):
        return value.time().replace(tzinfo=None)
    if isinstance(value, dtime):
        return value.replace(tzinfo=None)
    if isinstance(value, (int, float)) and 0 <= float(value) < 1:
        sec = int(round(float(value) * 86400)) % 86400
        return dtime(sec // 3600, (sec % 3600) // 60, sec % 60)
    txt = str(value or "").strip()
    for fmt in ("%H:%M:%S", "%H:%M", "%I:%M:%S %p", "%I:%M %p"):
        try:
            return datetime.strptime(txt, fmt).time()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(txt).time()
    except Exception:
        return None


def date_value(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    txt = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(txt, fmt).date()
        except ValueError:
            pass
    return None


def shift_bounds(day, definition):
    start = datetime.combine(day, definition.start_time, TZ)
    end = datetime.combine(day, definition.end_time, TZ)
    if end <= start:
        end += timedelta(days=1)
    return start, end


def wb_weigh_at(day, definition, clock):
    if not clock:
        return None
    start, end = shift_bounds(day, definition)
    stamp = datetime.combine(day, clock, TZ)
    if end.date() > day and stamp < start:
        stamp += timedelta(days=1)
    return stamp if start <= stamp < end else None


def find_header(ws):
    for rix, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=True), 1):
        cand = {norm_header(v): i for i, v in enumerate(row) if str(v or "").strip()}
        if "VEHICLENUMBER" in cand and ("GROSSWEIGHT" in cand or "NETWEIGHT" in cand):
            return rix, cand
    raise ValueError("WB headers not found. Expected Vehicle Number / Gross Weight / Net Weight.")


def col(headers, *names):
    for name in names:
        key = norm_header(name)
        if key in headers:
            return headers[key]
    return None


def infer_file_context(ws, header_row, idx, db):
    """Resolve operating date + shift using the proven PC-server C-shift rules."""
    shifts = set()
    rows_for_context = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(v not in (None, "") for v in row):
            continue
        vehicle = row_value(row, idx, "vehicle")
        if not str(vehicle or "").strip():
            continue
        raw_shift = str(row_value(row, idx, "shift") or "").strip().upper()[:1]
        if raw_shift in {"A", "B", "C"}:
            shifts.add(raw_shift)
        rows_for_context.append(row)

    if len(shifts) != 1:
        raise ValueError(
            f"Expected exactly one Shift Code A/B/C in the workbook; "
            f"found {sorted(shifts) or 'none'}."
        )
    shift = next(iter(shifts))
    definition = db.get(ShiftMaster, shift)
    if not definition or not definition.active:
        raise ValueError(f"Shift {shift} is not active in Shift Master.")

    operating_dates = set()
    operating_date_counts = {}
    raw_dates = set()
    bad_rows = 0
    for row in rows_for_context:
        movement_day = date_value(row_value(row, idx, "date"))
        clock = clock_value(row_value(row, idx, "time"))
        if movement_day is not None:
            raw_dates.add(movement_day)
        op_day = operating_date_from_movement(movement_day, clock, definition)
        if op_day:
            operating_dates.add(op_day)
            operating_date_counts[op_day] = operating_date_counts.get(op_day, 0) + 1
        else:
            bad_rows += 1

    # AMNS can use one operating date for the full overnight C shift.
    if (
        shift == "C"
        and len(operating_dates) != 1
        and len(raw_dates) == 1
        and bad_rows == 0
    ):
        operating_dates = set(raw_dates)

    # Or use calendar dates with a tiny isolated anomaly. Preserve the PC >=95% guard.
    if (
        shift == "C"
        and len(operating_dates) != 1
        and operating_date_counts
        and bad_rows == 0
    ):
        dominant_day, dominant_count = max(
            operating_date_counts.items(), key=lambda item: item[1]
        )
        total_resolved = sum(operating_date_counts.values())
        if total_resolved and (dominant_count / total_resolved) >= 0.95:
            operating_dates = {dominant_day}

    if len(operating_dates) != 1:
        raise ValueError(
            "Could not resolve one operating date from Movement Date + WB time. "
            f"Resolved dates: {sorted(map(str, operating_dates)) or 'none'}; "
            f"invalid rows: {bad_rows}."
        )
    return next(iter(operating_dates)), shift

def audit(db, action, entity, entity_id, detail):
    db.add(AuditLog(
        created_at=now_local(),
        actor=ENGINE_ACTOR,
        action=action,
        entity=entity,
        entity_id=str(entity_id) if entity_id is not None else None,
        detail=json.dumps(detail, default=str),
    ))


def archive_bytes(content, filename, day, shift, digest):
    folder = ARCHIVE_DIR / day.isoformat() / shift
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"{digest[:10]}_{safe_name(filename)}"
    if not target.exists():
        target.write_bytes(content)
    return target


def quarantine_bytes(content, filename, reason):
    QUARANTINE_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    target = QUARANTINE_DIR / f"{stamp}_{safe_name(filename)}"
    target.write_bytes(content)
    meta = target.with_suffix(target.suffix + ".error.txt")
    meta.write_text(reason, encoding="utf-8")
    return target


def _ingest_single_shift_xlsx(content, filename, cfg, message_id):
    max_bytes = int(float(cfg.get("max_file_mb", 15)) * 1024 * 1024)
    if len(content) > max_bytes:
        raise ValueError(f"Attachment exceeds {cfg['max_file_mb']} MB.")
    digest = hashlib.sha256(content).hexdigest()
    try:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError("Unable to read XLSX file.") from exc
    with SessionLocal() as db:
        ensure_wb_header_mapping(db)
        # Use the same DB-backed WB header mapping/sheet detection as the
        # working manual/base-server WB upload. This supports aliases,
        # duplicate headers and WB data living on a non-active worksheet.
        ws, header_row, idx, header_diag = find_wb_sheet(wb, db)
        operating_date, shift = infer_file_context(ws, header_row, idx, db)

        definition = db.get(ShiftMaster, shift)
        if not definition or not definition.active:
            raise ValueError(f"Shift {shift} is not active in Shift Master.")

        existing_same = list(db.scalars(select(WbImportBatch).where(
            WbImportBatch.operating_date == operating_date,
            WbImportBatch.shift == shift,
            WbImportBatch.file_hash == digest,
            WbImportBatch.status.in_(["PREVIEW", "CONFIRMED"]),
        )))
        confirmed_same = next(
            (batch for batch in existing_same if batch.status == "CONFIRMED"),
            None,
        )
        if confirmed_same:
            archive = archive_bytes(content, filename, operating_date, shift, digest)
            return {
                "outcome": "DUPLICATE",
                "batch_id": confirmed_same.batch_id,
                "date": operating_date.isoformat(),
                "shift": shift,
                "valid": confirmed_same.valid_rows,
                "review": confirmed_same.review_rows,
                "archive": str(archive),
                "message": "Exact WB file is already confirmed in SQL; no rows were inserted.",
            }
        for old_preview in existing_same:
            if old_preview.status == "PREVIEW":
                old_preview.status = "REPLACED"

        eq = list(db.scalars(select(Equipment)))
        vehicle_map = {norm_vehicle(e.machine_id): e.machine_id for e in eq}
        vehicle_map.update({norm_vehicle(e.vehicle_no): e.machine_id for e in eq if e.vehicle_no})
        for alias in db.scalars(select(VehicleAlias).where(VehicleAlias.active)):
            vehicle_map[norm_vehicle(alias.alias)] = alias.machine_id
        location_resolver = build_location_resolver(db)

        batch = WbImportBatch(
            batch_id=str(uuid4()),
            operating_date=operating_date,
            shift=shift,
            file_name=filename or "WB.xlsx",
            file_hash=digest,
            status="PREVIEW",
            valid_rows=0,
            review_rows=0,
        )
        db.add(batch)
        db.flush()

        valid = review = 0
        total_net = Decimal("0")
        unmapped_vehicles = set()
        tolerance = Decimal(str(cfg.get("weight_tolerance_kg", 5)))
        new_wb_rows = []

        for rno, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            if not any(v not in (None, "") for v in row):
                continue

            def value(key):
                return row_value(row, idx, key)

            vehicle_raw = str(value("vehicle") or "").strip()
            if not vehicle_raw:
                continue

            issues = []
            warnings = []
            move = str(value("move") or rno).strip()
            row_day = date_value(value("date"))
            raw_shift = str(value("shift") or "").strip().upper()[:1]

            if raw_shift != shift:
                issues.append(f"Shift {raw_shift or '?'} != file shift {shift}")

            clock = clock_value(value("time"))
            weigh_at = wb_weigh_at(operating_date, definition, clock)
            if not weigh_at:
                issues.append("Invalid or out-of-shift GW Created Time")
            else:
                expected_day = expected_movement_date(operating_date, clock, definition)
                midnight_date_exception = (
                    shift == "C"
                    and clock is not None
                    and definition.end_time <= definition.start_time
                    and clock < definition.end_time
                    and row_day == operating_date
                )
                if (
                    row_day
                    and expected_day
                    and row_day != expected_day
                    and not midnight_date_exception
                ):
                    issues.append(
                        f"Movement Date {row_day} != expected calendar date {expected_day} "
                        f"for operating date {operating_date} / Shift {shift}"
                    )

            tare = decimal_value(value("tare"))
            gross = decimal_value(value("gross"))
            net = decimal_value(value("net"))
            tare = tare if tare is not None else Decimal("0")
            if gross is None and net is not None:
                gross = tare + net
            if net is None and gross is not None and gross >= tare:
                net = gross - tare
            gross = gross if gross is not None else Decimal("0")
            if gross <= 0 or net is None or net <= 0 or gross < tare:
                issues.append("Invalid tare/gross/net weight")
            elif abs((gross - tare) - net) > tolerance:
                issues.append(f"Gross-Tare differs from Net by more than {tolerance} kg")

            source = location_resolver.resolve(
                value("source_code"), value("source_name"), "SOURCE"
            )
            dest = location_resolver.resolve(
                value("dest_code"), value("dest_name"), "DESTINATION"
            )
            if not source:
                issues.append("Source location missing")
            if not dest:
                issues.append("Destination location missing")

            vehicle_id = vehicle_map.get(norm_vehicle(vehicle_raw))
            if not vehicle_id:
                unmapped_vehicles.add(vehicle_raw)
                warnings.append("Vehicle not mapped to equipment master")

            status = "REVIEW" if issues else "VALID"
            issue_text = "; ".join(issues + warnings) or None
            key = f"{batch.batch_id}:{move}:{rno}"
            source_code_raw=str(value("source_code") or "").strip()
            source_name_raw=str(value("source_name") or "").strip()
            dest_code_raw=str(value("dest_code") or "").strip()
            dest_name_raw=str(value("dest_name") or "").strip()
            movement=WbMovement(
                movement_key=key,
                batch_id=batch.batch_id,
                operating_date=operating_date,
                shift=shift,
                movement_no=move,
                vehicle_raw=vehicle_raw,
                vehicle_id=vehicle_id,
                material_code=str(value("matcode") or "").strip() or None,
                material_name=str(value("matname") or "").strip() or None,
                source_raw=source_code_raw or source_name_raw or None,
                destination_raw=dest_code_raw or dest_name_raw or None,
                tare_kg=tare,
                gross_kg=gross,
                net_kg=net or Decimal("0"),
                weigh_at=weigh_at or datetime.combine(operating_date, definition.start_time, TZ),
                row_status=status,
                issue=issue_text,
            )
            db.add(movement); new_wb_rows.append(movement)
            if status == "VALID":
                valid += 1
                total_net += net or Decimal("0")
            else:
                review += 1

        if valid + review == 0:
            db.rollback()
            raise ValueError("No WB movement rows were found in the attachment.")
        db.flush()
        canonicalize_wb_rows(db,new_wb_rows)

        batch.valid_rows = valid
        batch.review_rows = review
        review_pct = (review * 100.0 / (valid + review)) if (valid + review) else 100.0

        audit(db, "WB_EMAIL_PREVIEW", "wb_import_batch", batch.batch_id, {
            "gmail_message_id": message_id,
            "file": filename,
            "file_hash": digest,
            "date": operating_date,
            "shift": shift,
            "valid": valid,
            "review": review,
            "review_percent": round(review_pct, 2),
            "unmapped_vehicles": sorted(unmapped_vehicles),
            "sheet": ws.title,
            "header_mappings": header_diag.get("matched", {}),
            "duplicate_headers": header_diag.get("duplicates", {}),
        })

        archive = archive_bytes(content, filename, operating_date, shift, digest)

        max_review = float(cfg.get("max_review_percent", 5.0))
        auto_confirm = bool(cfg.get("auto_confirm", True))
        if valid <= 0 or review_pct > max_review or not auto_confirm:
            db.commit()
            return {
                "outcome": "REVIEW_REQUIRED",
                "batch_id": batch.batch_id,
                "date": operating_date.isoformat(),
                "shift": shift,
                "valid": valid,
                "review": review,
                "review_percent": round(review_pct, 2),
                "tonnes": float(total_net / Decimal("1000")),
                "unmapped_vehicles": sorted(unmapped_vehicles),
                "archive": str(archive),
                "message": f"WB loaded as PREVIEW only: {valid} VALID / {review} REVIEW.",
            }

        # Atomic replacement: only once the incoming file has passed validation.
        prior = list(db.scalars(select(WbImportBatch).where(
            WbImportBatch.operating_date == operating_date,
            WbImportBatch.shift == shift,
            WbImportBatch.status == "CONFIRMED",
        )))
        for old in prior:
            keys = list(db.scalars(select(WbMovement.movement_key).where(WbMovement.batch_id == old.batch_id)))
            if keys:
                db.execute(delete(LoadWbMatch).where(LoadWbMatch.movement_key.in_(keys)))
            old.status = "REPLACED"

        batch.status = "CONFIRMED"
        batch.confirmed_at = now_local()
        db.flush()
        reconciliation = auto_reconcile(db, operating_date, shift)
        audit(db, "WB_EMAIL_CONFIRM", "wb_import_batch", batch.batch_id, {
            "gmail_message_id": message_id,
            "valid": valid,
            "review": review,
            "tonnes": float(total_net / Decimal("1000")),
            "reconciliation": reconciliation,
        })
        db.commit()
        return {
            "outcome": "CONFIRMED",
            "batch_id": batch.batch_id,
            "date": operating_date.isoformat(),
            "shift": shift,
            "valid": valid,
            "review": review,
            "review_percent": round(review_pct, 2),
            "tonnes": float(total_net / Decimal("1000")),
            "unmapped_vehicles": sorted(unmapped_vehicles),
            "archive": str(archive),
            "reconciliation": reconciliation,
            "message": (
                f"WB auto-confirmed: {valid} VALID / {review} REVIEW; "
                f"{reconciliation['matched']} matched, "
                f"{reconciliation['wb_unmatched']} WB unmatched."
            ),
        }



# NMTPL_WB_TIME_FIRST_SHIFT_ROUTING_V4
def _mixed_shift_row_value(row, idx, key):
    i = idx.get(key)
    return row[i] if i is not None and i < len(row) else None


def _mixed_shift_code(row, idx):
    raw = str(_mixed_shift_row_value(row, idx, "shift") or "").strip().upper()
    return raw[:1] if raw[:1] in {"A", "B", "C"} else ""


def _mixed_shift_make_subset(ws, header_row, selected_rows):
    from openpyxl import Workbook
    out = Workbook()
    dst = out.active
    dst.title = ws.title or "Sheet1"
    for row in ws.iter_rows(min_row=1, max_row=header_row, values_only=True):
        dst.append(list(row))
    for _, row in selected_rows:
        dst.append(list(row))
    buf = BytesIO()
    out.save(buf)
    return buf.getvalue()


def _mixed_shift_existing_movement(db, operating_date, shift, movement_no):
    if operating_date is None or not shift or not movement_no:
        return None
    return db.scalar(
        select(WbMovement)
        .join(WbImportBatch, WbMovement.batch_id == WbImportBatch.batch_id)
        .where(
            WbMovement.operating_date == operating_date,
            WbMovement.shift == shift,
            WbMovement.movement_no == movement_no,
            WbImportBatch.status == "CONFIRMED",
        )
    )


def _mixed_shift_vehicle_map(db):
    eq = list(db.scalars(select(Equipment)))
    vehicle_map = {norm_vehicle(e.machine_id): e.machine_id for e in eq}
    vehicle_map.update({
        norm_vehicle(e.vehicle_no): e.machine_id
        for e in eq if getattr(e, "vehicle_no", None)
    })
    for alias in db.scalars(select(VehicleAlias).where(VehicleAlias.active)):
        vehicle_map[norm_vehicle(alias.alias)] = alias.machine_id
    return vehicle_map


def _mixed_shift_append_missing_rows(rows, idx, filename, message_id, cfg):
    from collections import defaultdict
    grouped = defaultdict(list)
    for rno, row, target_shift in rows:
        if target_shift in {"A", "B", "C"}:
            grouped[target_shift].append((rno, row))

    summary = {
        "detected": sum(len(v) for v in grouped.values()),
        "duplicates": 0,
        "added_valid": 0,
        "added_review": 0,
        "errors": [],
        "by_shift": {},
    }

    for target_shift, shift_rows in sorted(grouped.items()):
        shift_summary = {
            "detected": len(shift_rows),
            "duplicates": 0,
            "added_valid": 0,
            "added_review": 0,
        }
        try:
            with SessionLocal() as db:
                definition = db.get(ShiftMaster, target_shift)
                if not definition or not definition.active:
                    raise ValueError(f"Shift {target_shift} is not active in Shift Master.")

                vehicle_map = _mixed_shift_vehicle_map(db)
                location_resolver = build_location_resolver(db)
                tolerance = Decimal(str(cfg.get("weight_tolerance_kg", 5)))
                prepared = []

                for rno, row in shift_rows:
                    move = str(_mixed_shift_row_value(row, idx, "move") or rno).strip()
                    vehicle_raw = str(_mixed_shift_row_value(row, idx, "vehicle") or "").strip()
                    row_day = date_value(_mixed_shift_row_value(row, idx, "date"))
                    clock = clock_value(_mixed_shift_row_value(row, idx, "time"))
                    op_day = operating_date_from_movement(row_day, clock, definition)

                    if _mixed_shift_existing_movement(
                        db, op_day, target_shift, move
                    ) is not None:
                        shift_summary["duplicates"] += 1
                        summary["duplicates"] += 1
                        continue

                    issues = []
                    warnings = []
                    if op_day is None:
                        issues.append("Movement Date missing/invalid")

                    weigh_at = wb_weigh_at(op_day, definition, clock) if op_day else None
                    if not weigh_at:
                        issues.append(f"GW Created Time does not belong to Shift {target_shift}")

                    tare = decimal_value(_mixed_shift_row_value(row, idx, "tare"))
                    gross = decimal_value(_mixed_shift_row_value(row, idx, "gross"))
                    net = decimal_value(_mixed_shift_row_value(row, idx, "net"))
                    tare = tare if tare is not None else Decimal("0")
                    if gross is None and net is not None:
                        gross = tare + net
                    if net is None and gross is not None and gross >= tare:
                        net = gross - tare
                    gross = gross if gross is not None else Decimal("0")
                    if gross <= 0 or net is None or net <= 0 or gross < tare:
                        issues.append("Invalid tare/gross/net weight")
                    elif abs((gross - tare) - net) > tolerance:
                        issues.append(
                            f"Gross-Tare differs from Net by more than {tolerance} kg"
                        )

                    source_code_raw = str(
                        _mixed_shift_row_value(row, idx, "source_code") or ""
                    ).strip()
                    source_name_raw = str(
                        _mixed_shift_row_value(row, idx, "source_name") or ""
                    ).strip()
                    dest_code_raw = str(
                        _mixed_shift_row_value(row, idx, "dest_code") or ""
                    ).strip()
                    dest_name_raw = str(
                        _mixed_shift_row_value(row, idx, "dest_name") or ""
                    ).strip()
                    source = location_resolver.resolve(
                        source_code_raw, source_name_raw, "SOURCE"
                    )
                    dest = location_resolver.resolve(
                        dest_code_raw, dest_name_raw, "DESTINATION"
                    )
                    if not source:
                        issues.append("Source location missing")
                    if not dest:
                        issues.append("Destination location missing")

                    vehicle_id = vehicle_map.get(norm_vehicle(vehicle_raw))
                    if not vehicle_id:
                        warnings.append("Vehicle not mapped to equipment master")

                    prepared.append({
                        "rno": rno,
                        "row": row,
                        "move": move,
                        "vehicle_raw": vehicle_raw,
                        "vehicle_id": vehicle_id,
                        "operating_date": op_day,
                        "weigh_at": weigh_at,
                        "tare": tare,
                        "gross": gross,
                        "net": net or Decimal("0"),
                        "source_raw": source_code_raw or source_name_raw or None,
                        "dest_raw": dest_code_raw or dest_name_raw or None,
                        "status": "REVIEW" if issues else "VALID",
                        "issue": "; ".join(issues + warnings) or None,
                    })

                by_date = defaultdict(list)
                for item in prepared:
                    by_date[item["operating_date"]].append(item)

                for op_day, items in by_date.items():
                    if op_day is None:
                        shift_summary["added_review"] += len(items)
                        summary["added_review"] += len(items)
                        summary["errors"].append(
                            f"{target_shift}: {len(items)} row(s) have no valid operating date"
                        )
                        continue

                    target_batch = db.scalar(
                        select(WbImportBatch)
                        .where(
                            WbImportBatch.operating_date == op_day,
                            WbImportBatch.shift == target_shift,
                            WbImportBatch.status == "CONFIRMED",
                        )
                        .order_by(WbImportBatch.confirmed_at.desc())
                    )

                    if target_batch is None:
                        valid_count = sum(1 for x in items if x["status"] == "VALID")
                        review_count = sum(1 for x in items if x["status"] == "REVIEW")
                        batch_status = (
                            "CONFIRMED"
                            if valid_count > 0 and review_count == 0
                            else "PREVIEW"
                        )
                        supplemental_hash = hashlib.sha256(
                            (
                                f"{message_id}|{filename}|{op_day}|{target_shift}|"
                                + "|".join(x["move"] for x in items)
                            ).encode("utf-8")
                        ).hexdigest()
                        target_batch = WbImportBatch(
                            batch_id=str(uuid4()),
                            operating_date=op_day,
                            shift=target_shift,
                            file_name=f"{filename} [ROUTED-{target_shift}]",
                            file_hash=supplemental_hash,
                            status=batch_status,
                            valid_rows=0,
                            review_rows=0,
                        )
                        if batch_status == "CONFIRMED":
                            target_batch.confirmed_at = now_local()
                        db.add(target_batch)
                        db.flush()

                    new_rows = []
                    for item in items:
                        if _mixed_shift_existing_movement(
                            db, op_day, target_shift, item["move"]
                        ) is not None:
                            shift_summary["duplicates"] += 1
                            summary["duplicates"] += 1
                            continue

                        movement = WbMovement(
                            movement_key=(
                                f"{target_batch.batch_id}:ROUTED:"
                                f"{item['move']}:{item['rno']}"
                            ),
                            batch_id=target_batch.batch_id,
                            operating_date=op_day,
                            shift=target_shift,
                            movement_no=item["move"],
                            vehicle_raw=item["vehicle_raw"],
                            vehicle_id=item["vehicle_id"],
                            material_code=str(
                                _mixed_shift_row_value(item["row"], idx, "matcode") or ""
                            ).strip() or None,
                            material_name=str(
                                _mixed_shift_row_value(item["row"], idx, "matname") or ""
                            ).strip() or None,
                            source_raw=item["source_raw"],
                            destination_raw=item["dest_raw"],
                            tare_kg=item["tare"],
                            gross_kg=item["gross"],
                            net_kg=item["net"],
                            weigh_at=item["weigh_at"] or datetime.combine(
                                op_day, definition.start_time, TZ
                            ),
                            row_status=item["status"],
                            issue=item["issue"],
                        )
                        db.add(movement)
                        new_rows.append(movement)

                        if item["status"] == "VALID":
                            target_batch.valid_rows = int(target_batch.valid_rows or 0) + 1
                            shift_summary["added_valid"] += 1
                            summary["added_valid"] += 1
                        else:
                            target_batch.review_rows = int(target_batch.review_rows or 0) + 1
                            shift_summary["added_review"] += 1
                            summary["added_review"] += 1

                        audit(
                            db,
                            "WB_MIXED_SHIFT_ROUTE",
                            "wb_movement",
                            item["move"],
                            {
                                "gmail_message_id": message_id,
                                "source_file": filename,
                                "target_date": op_day,
                                "target_shift": target_shift,
                                "source_row": item["rno"],
                                "status": item["status"],
                                "issue": item["issue"],
                            },
                        )

                    db.flush()
                    if new_rows:
                        canonicalize_wb_rows(db, new_rows)
                    if target_batch.status == "CONFIRMED":
                        try:
                            auto_reconcile(db, op_day, target_shift)
                        except Exception:
                            logger.exception(
                                "Mixed-shift reconciliation failed: date=%s shift=%s",
                                op_day, target_shift,
                            )
                db.commit()
        except Exception as exc:
            logger.exception("Mixed-shift routing failed for target shift %s", target_shift)
            summary["errors"].append(
                f"{target_shift}: {type(exc).__name__}: {exc}"
            )
        summary["by_shift"][target_shift] = shift_summary
    return summary


def ingest_xlsx(content, filename, cfg, message_id):
    """Time-first WB Gmail ingestion restored from the working PC server."""
    from collections import Counter

    try:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception:
        return _ingest_single_shift_xlsx(content, filename, cfg, message_id)

    with SessionLocal() as db:
        ensure_wb_header_mapping(db)
        ws, header_row, idx, _ = find_wb_sheet(wb, db)
        definitions = {
            str(d.shift).strip().upper(): d
            for d in db.scalars(select(ShiftMaster).where(ShiftMaster.active))
            if str(d.shift or "").strip().upper() in {"A", "B", "C"}
        }

    def clock_belongs(definition, clock):
        if clock is None or definition is None:
            return False
        start, end = definition.start_time, definition.end_time
        if end > start:
            return start <= clock < end
        return clock >= start or clock < end

    def shift_from_clock(clock):
        matches = [
            code for code, definition in definitions.items()
            if clock_belongs(definition, clock)
        ]
        return matches[0] if len(matches) == 1 else ""

    rows = []
    raw_counts = Counter()
    resolved_counts = Counter()
    corrections = []
    shift_index = idx.get("shift")

    for rno, row in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True),
        header_row + 1,
    ):
        vehicle = str(_mixed_shift_row_value(row, idx, "vehicle") or "").strip()
        if not vehicle:
            continue
        raw_code = _mixed_shift_code(row, idx)
        if raw_code:
            raw_counts[raw_code] += 1
        clock = clock_value(_mixed_shift_row_value(row, idx, "time"))
        effective_code = shift_from_clock(clock) or raw_code

        row_list = list(row)
        if effective_code and shift_index is not None:
            old_code = str(row_list[shift_index] or "").strip().upper()[:1]
            if old_code != effective_code:
                row_list[shift_index] = effective_code
                corrections.append({
                    "row": rno,
                    "movement": str(_mixed_shift_row_value(row, idx, "move") or rno).strip(),
                    "vehicle": vehicle,
                    "raw_shift": raw_code or old_code or "?",
                    "time": str(clock) if clock is not None else "",
                    "corrected_shift": effective_code,
                })

        corrected_row = tuple(row_list)
        rows.append((rno, corrected_row, effective_code))
        if effective_code:
            resolved_counts[effective_code] += 1

    if not rows or not resolved_counts:
        return _ingest_single_shift_xlsx(content, filename, cfg, message_id)

    if len(resolved_counts) == 1:
        if not corrections:
            return _ingest_single_shift_xlsx(content, filename, cfg, message_id)
        only_shift = next(iter(resolved_counts))
        normalized_bytes = _mixed_shift_make_subset(
            ws, header_row, [(rno, row) for rno, row, _ in rows]
        )
        result = _ingest_single_shift_xlsx(
            normalized_bytes,
            f"{filename} [TIME-NORMALIZED-{only_shift}]",
            cfg,
            message_id,
        )
        result["shift_code_corrections"] = corrections
        result["raw_shift_counts"] = dict(raw_counts)
        result["resolved_shift_counts"] = dict(resolved_counts)
        result["message"] = (
            str(result.get("message", "")).rstrip()
            + f" Shift-code normalization: {len(corrections)} row(s) corrected by GW Created Time."
        )
        return result

    primary_shift, _ = resolved_counts.most_common(1)[0]
    primary_rows = [
        (rno, row) for rno, row, code in rows
        if code == primary_shift or not code
    ]
    other_rows = [
        (rno, row, code) for rno, row, code in rows
        if code and code != primary_shift
    ]
    primary_bytes = _mixed_shift_make_subset(ws, header_row, primary_rows)
    result = _ingest_single_shift_xlsx(
        primary_bytes,
        f"{filename} [TIME-PRIMARY-{primary_shift}]",
        cfg,
        message_id,
    )
    routed = _mixed_shift_append_missing_rows(
        other_rows, idx, filename, message_id, cfg
    )
    result["mixed_shift"] = {
        "raw_counts": dict(raw_counts),
        "resolved_counts": dict(resolved_counts),
        "primary_shift": primary_shift,
        "other_rows": len(other_rows),
        "shift_code_corrections": corrections,
        "routed": routed,
    }
    result["review"] = int(result.get("review", 0)) + int(
        routed.get("added_review", 0)
    ) + len(routed.get("errors", []))
    result["message"] = (
        str(result.get("message", "")).rstrip()
        + f" Time-first routing: {len(corrections)} Shift Code correction(s); "
        + f"{len(other_rows)} row(s) genuinely belonged to other shift(s); "
        + f"{routed.get('duplicates', 0)} already existed and were ignored; "
        + f"{routed.get('added_valid', 0)} added to their correct shift; "
        + f"{routed.get('added_review', 0)} routed for review."
    )
    return result

def apply_labels(service, message_id, add_ids, remove_ids=None):
    service.users().messages().modify(
        userId="me",
        id=message_id,
        body={"addLabelIds": list(dict.fromkeys(add_ids)), "removeLabelIds": list(dict.fromkeys(remove_ids or []))},
    ).execute()


def process_message(service, message, cfg, labels):
    sender_text = header_value(message, "From")
    sender = parseaddr(sender_text)[1].lower()
    subject = header_value(message, "Subject")

    if not sender_allowed(sender, cfg.get("allowed_senders", [])):
        logger.warning("Skipped message %s from unapproved sender %s", message["id"], sender or sender_text)
        return {"status": "SKIPPED_SENDER", "message_id": message["id"], "sender": sender}

    if not subject_allowed(subject, cfg.get("subject_keywords", [])):
        logger.info("Skipped message %s because subject did not match configured keywords", message["id"])
        return {"status": "SKIPPED_SUBJECT", "message_id": message["id"], "subject": subject}

    start_after = datetime.fromisoformat(str(cfg["start_after_utc"]).replace("Z", "+00:00"))
    internal = datetime.fromtimestamp(int(message.get("internalDate", "0")) / 1000, tz=timezone.utc)
    if internal < start_after:
        return {"status": "SKIPPED_OLD", "message_id": message["id"]}

    attachments = xlsx_attachments(service, message)
    if not attachments:
        return {"status": "SKIPPED_NO_XLSX", "message_id": message["id"]}

    results = []
    for filename, content in attachments:
        try:
            result = ingest_xlsx(content, filename, cfg, message["id"])
            logger.info("%s | %s | %s", message["id"], filename, result["message"])
            results.append(result)
        except Exception as exc:
            reason = f"{type(exc).__name__}: {exc}"
            q = quarantine_bytes(content, filename, reason)
            logger.exception("WB attachment failed: message=%s file=%s quarantine=%s", message["id"], filename, q)
            results.append({"outcome": "ERROR", "message": reason, "quarantine": str(q), "file": filename})

    outcomes = {r["outcome"] for r in results}
    add = []
    remove = []

    processed_id = labels[cfg["processed_label"]]
    duplicate_id = labels[cfg["duplicate_label"]]
    error_id = labels[cfg["error_label"]]
    review_required_id = labels[cfg["review_required_label"]]
    review_rows_id = labels[cfg["review_rows_label"]]

    if "ERROR" in outcomes:
        add.append(error_id)
        remove.extend([processed_id, duplicate_id, review_required_id, review_rows_id])
    elif "REVIEW_REQUIRED" in outcomes:
        add.append(review_required_id)
        remove.extend([error_id, processed_id, duplicate_id, review_rows_id])
    elif outcomes == {"DUPLICATE"}:
        add.append(duplicate_id)
        remove.extend([error_id, processed_id, review_required_id, review_rows_id])
        if cfg.get("archive_success", True):
            remove.extend(["INBOX", "UNREAD"])
    else:
        add.append(processed_id)
        remove.extend([error_id, duplicate_id, review_required_id])
        if any(int(r.get("review", 0)) > 0 for r in results):
            add.append(review_rows_id)
        else:
            remove.append(review_rows_id)
        if cfg.get("archive_success", True):
            remove.extend(["INBOX", "UNREAD"])

    apply_labels(service, message["id"], add, remove)
    return {
        "status": "DONE",
        "message_id": message["id"],
        "sender": sender,
        "subject": subject,
        "results": results,
    }


def poll_once(service, cfg, labels):
    label_names = [
        cfg["processed_label"], cfg["duplicate_label"], cfg["error_label"],
        cfg["review_required_label"],
    ]
    exclusions = " ".join(f"-label:{name}" for name in label_names)
    query = (str(cfg.get("gmail_query", "has:attachment filename:xlsx")).strip() + " " + exclusions).strip()
    response = service.users().messages().list(userId="me", q=query, maxResults=int(cfg.get("max_messages_per_poll", 25))).execute()
    ids = [x["id"] for x in response.get("messages", [])]
    processed = []
    for message_id in reversed(ids):
        message = service.users().messages().get(userId="me", id=message_id, format="full").execute()
        result = process_message(service, message, cfg, labels)
        processed.append(result)
    return processed


def main():
    logger.info("NMTPL WB Gmail worker starting. Project root=%s", ROOT)
    consecutive_failures = 0
    while True:
        try:
            cfg = load_config()
            if not cfg.get("enabled", False):
                write_status(state="DISABLED", message="Set enabled=true in config after authorization.")
                time.sleep(max(30, int(cfg.get("poll_seconds", 180))))
                continue

            service = gmail_service()
            label_names = [
                cfg["processed_label"], cfg["duplicate_label"], cfg["error_label"],
                cfg["review_required_label"], cfg["review_rows_label"],
            ]
            labels = ensure_labels(service, label_names)
            results = poll_once(service, cfg, labels)
            consecutive_failures = 0

            counts = {}
            for r in results:
                key = r.get("status", "UNKNOWN")
                counts[key] = counts.get(key, 0) + 1
            write_status(
                state="RUNNING",
                last_poll=datetime.now(timezone.utc).isoformat(),
                messages_seen=len(results),
                counts=counts,
            )
            if results:
                logger.info("Poll complete: %s", counts)
            time.sleep(max(30, int(cfg.get("poll_seconds", 180))))
        except KeyboardInterrupt:
            logger.info("Worker stopped by keyboard interrupt.")
            write_status(state="STOPPED")
            return
        except Exception as exc:
            consecutive_failures += 1
            logger.exception("Worker poll failed")
            write_status(
                state="ERROR",
                error=f"{type(exc).__name__}: {exc}",
                consecutive_failures=consecutive_failures,
            )
            # Never hammer Gmail/DB during an outage.
            time.sleep(min(900, max(60, consecutive_failures * 60)))


if __name__ == "__main__":
    main()